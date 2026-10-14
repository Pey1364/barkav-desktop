import os
import io
import hmac
import uuid
from datetime import datetime, timedelta
import jdatetime
import openpyxl
from PIL import Image, ImageOps
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Flask, render_template, request, jsonify, send_from_directory, send_file, session, redirect, url_for
from database import get_connection, init_db, BARKAV_32_PHASES

app = Flask(__name__)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # حداکثر ۵۰ مگابایت در هر ارسال

# کلید نشست: روی هاست باید ثابت باشد، در غیر این صورت هر بار اجرا نشست‌ها باطل می‌شوند
app.secret_key = os.environ.get("BARKAV_SECRET_KEY") or os.urandom(32)
app.permanent_session_lifetime = timedelta(days=30)

# رمز ورود: اگر تنظیم نشود، ورود غیرفعال است (حالت محلی روی لپ‌تاپ)
APP_PASSWORD = os.environ.get("BARKAV_APP_PASSWORD", "")
LOGIN_EXEMPT_ENDPOINTS = {"login", "static"}

init_db()

PHASE_DICT = dict(BARKAV_32_PHASES)

@app.before_request
def require_login():
    if not APP_PASSWORD or request.endpoint in LOGIN_EXEMPT_ENDPOINTS or session.get("logged_in"):
        return None
    # صفحات و فایل‌ها به صفحه ورود می‌روند؛ درخواست‌های API با 401 جواب داده می‌شوند
    if request.method == "GET" and not request.path.startswith("/api/"):
        return redirect(url_for("login"))
    return jsonify({"success": False, "error": "ابتدا وارد شوید."}), 401

@app.route("/login", methods=["GET", "POST"])
def login():
    if not APP_PASSWORD:
        return redirect(url_for("index"))
    error = ""
    if request.method == "POST":
        # مقایسه بایتی تا رمز فارسی هم خطا ندهد و زمان پاسخ وابسته به محتوا نباشد
        given = request.form.get("password", "").encode("utf-8")
        if hmac.compare_digest(given, APP_PASSWORD.encode("utf-8")):
            session.clear()
            session["logged_in"] = True
            session.permanent = True
            return redirect(url_for("index"))
        error = "رمز ورود اشتباه است."
    return render_template("login.html", error=error)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

def get_today_jalali():
    return jdatetime.date.today().strftime("%Y/%m/%d")

def get_period_start_jalali(period_type):
    today = jdatetime.date.today()
    if period_type == "daily":
        return today.strftime("%Y/%m/%d"), today.strftime("%Y/%m/%d")
    elif period_type == "weekly":
        start = today - timedelta(days=6)
        return start.strftime("%Y/%m/%d"), today.strftime("%Y/%m/%d")
    elif period_type == "monthly":
        start = jdatetime.date(today.year, today.month, 1)
        return start.strftime("%Y/%m/%d"), today.strftime("%Y/%m/%d")
    return "", ""

PHOTO_MAX_SIDE = 1600        # بیشترین طول یا عرض عکس بعد از کوچک‌سازی (پیکسل)
PHOTO_JPEG_QUALITY = 80      # کیفیت JPEG؛ برای گزارش کارگاه کافی و حجم کم است

def save_attachment(file, stem, ext):
    """عکس را کوچک و فشرده (JPEG) ذخیره می‌کند تا فضای هاست رایگان کم مصرف شود.
    اگر فایل عکس خوانده نشود (مثلاً PDF یا HEIC)، بدون تغییر ذخیره می‌شود."""
    folder = app.config["UPLOAD_FOLDER"]
    try:
        file.stream.seek(0)
        img = ImageOps.exif_transpose(Image.open(file.stream))
        img.thumbnail((PHOTO_MAX_SIDE, PHOTO_MAX_SIDE))
        if img.mode in ("RGBA", "LA", "P"):
            # شفافیت را روی پس‌زمینه سفید می‌گذاریم تا بخش‌های شفاف سیاه نشوند
            img = img.convert("RGBA")
            background = Image.new("RGB", img.size, (255, 255, 255))
            background.paste(img, mask=img.split()[-1])
            img = background
        elif img.mode != "RGB":
            img = img.convert("RGB")
        saved_filename = f"{stem}.jpg"
        img.save(os.path.join(folder, saved_filename), "JPEG", quality=PHOTO_JPEG_QUALITY, optimize=True)
        return saved_filename
    except Exception:
        file.stream.seek(0)
        saved_filename = f"{stem}{ext}"
        file.save(os.path.join(folder, saved_filename))
        return saved_filename

@app.route("/")
def index():
    return render_template(
        "new_report.html",
        phases=BARKAV_32_PHASES,
        today_jalali=get_today_jalali()
    )

@app.route("/api/reports", methods=["POST"])
def create_report():
    report_date = request.form.get("report_date", "").strip()
    doc_category = request.form.get("doc_category", "گزارش روزانه کارگاه").strip()
    phase_id = request.form.get("phase_id", "").strip()
    phase_title = PHASE_DICT.get(phase_id, "")
    weather_condition = request.form.get("weather_condition", "").strip()
    temperature = request.form.get("temperature", "").strip()
    crew_count = int(request.form.get("crew_count") or 0)
    contractor_name = request.form.get("contractor_name", "").strip()
    work_summary = request.form.get("work_summary", "").strip()
    materials_delivered = request.form.get("materials_delivered", "").strip()
    machinery_used = request.form.get("machinery_used", "").strip()
    delays_and_safety = request.form.get("delays_and_safety", "").strip()
    supervisor_notes = request.form.get("supervisor_notes", "").strip()

    if not report_date or not work_summary:
        return jsonify({"success": False, "error": "تاریخ و شرح عملیات روزانه الزامی است."}), 400

    report_id = f"NFD-{uuid.uuid4().hex[:8].upper()}"

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO non_financial_daily_reports (
                id, report_date, doc_category, phase_id, phase_title,
                weather_condition, temperature, crew_count, contractor_name,
                work_summary, materials_delivered, machinery_used,
                delays_and_safety, supervisor_notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            report_id, report_date, doc_category, phase_id, phase_title,
            weather_condition, temperature, crew_count, contractor_name,
            work_summary, materials_delivered, machinery_used,
            delays_and_safety, supervisor_notes
        ))

        uploaded_files = request.files.getlist("photos")
        captions = request.form.getlist("photo_captions")

        for idx, file in enumerate(uploaded_files):
            if file and file.filename:
                ext = os.path.splitext(file.filename)[1].lower()
                stem = f"{report_id}_{idx+1}_{uuid.uuid4().hex[:6]}"
                saved_filename = save_attachment(file, stem, ext)
                cap = captions[idx] if idx < len(captions) else ""

                cursor.execute('''
                    INSERT INTO non_financial_attachments (report_id, file_name, original_name, caption)
                    VALUES (?, ?, ?, ?)
                ''', (report_id, saved_filename, file.filename, cap))

        conn.commit()

    return jsonify({"success": True, "report_id": report_id})

@app.route("/reports")
def list_reports():
    period = request.args.get("period", "")
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    phase_filter = request.args.get("phase_id", "").strip()
    category_filter = request.args.get("doc_category", "").strip()

    if period in ("daily", "weekly", "monthly") and not (start_date or end_date):
        start_date, end_date = get_period_start_jalali(period)

    query = "SELECT * FROM non_financial_daily_reports WHERE is_deleted = 0"
    params = []

    if start_date:
        query += " AND report_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND report_date <= ?"
        params.append(end_date)
    if phase_filter:
        query += " AND phase_id = ?"
        params.append(phase_filter)
    if category_filter:
        query += " AND doc_category = ?"
        params.append(category_filter)

    query += " ORDER BY report_date DESC, created_at DESC"

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(query, params)
        reports = [dict(row) for row in cursor.fetchall()]

        total_crew = 0
        total_photos = 0
        for r in reports:
            total_crew += int(r.get("crew_count") or 0)
            cursor.execute("SELECT * FROM non_financial_attachments WHERE report_id = ?", (r["id"],))
            atts = [dict(a) for a in cursor.fetchall()]
            r["attachments"] = atts
            total_photos += len(atts)

    stats = {
        "report_count": len(reports),
        "total_crew_man_days": total_crew,
        "total_photos": total_photos
    }

    return render_template(
        "reports_list.html",
        reports=reports,
        stats=stats,
        phases=BARKAV_32_PHASES,
        start_date=start_date,
        end_date=end_date,
        phase_filter=phase_filter,
        category_filter=category_filter,
        period=period
    )

@app.route("/api/reports/<report_id>/delete", methods=["POST"])
def soft_delete_report(report_id):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE non_financial_daily_reports SET is_deleted = 1, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (report_id,)
        )
        conn.commit()
    return jsonify({"success": True})

@app.route("/export/excel")
def export_excel():
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    phase_filter = request.args.get("phase_id", "").strip()
    category_filter = request.args.get("doc_category", "").strip()

    query = "SELECT * FROM non_financial_daily_reports WHERE is_deleted = 0"
    params = []
    if start_date:
        query += " AND report_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND report_date <= ?"
        params.append(end_date)
    if phase_filter:
        query += " AND phase_id = ?"
        params.append(phase_filter)
    if category_filter:
        query += " AND doc_category = ?"
        params.append(category_filter)

    query += " ORDER BY report_date ASC"

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(query, params)
        rows = [dict(r) for r in cursor.fetchall()]
        for r in rows:
            cursor.execute("SELECT COUNT(*) as cnt FROM non_financial_attachments WHERE report_id = ?", (r["id"],))
            r["photo_count"] = cursor.fetchone()["cnt"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "اسناد غیرمالی و گزارشات"
    ws.views.sheetView[0].rightToLeft = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Tahoma", size=10, bold=True, color="38BDF8")
    body_font = Font(name="Tahoma", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    headers = [
        "شناسه سند", "تاریخ شمسی", "نوع سند غیرمالی", "فاز اجرایی",
        "وضعیت جوی", "دما", "نفرات (نفر-روز)", "پیمانکار / اکیپ",
        "شرح عملیات و گزارش اجرایی", "مصالح وارده", "ماشین‌آلات و تجهیزات",
        "موانع، تاخیرات و ایمنی", "توضیحات نظارت / مدیریت", "تعداد عکس پیوست"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        row_data = [
            r["id"], r["report_date"], r["doc_category"], r["phase_title"] or r["phase_id"],
            r["weather_condition"], r["temperature"], r["crew_count"], r["contractor_name"],
            r["work_summary"], r["materials_delivered"], r["machinery_used"],
            r["delays_and_safety"], r["supervisor_notes"], r["photo_count"]
        ]
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_widths = [16, 14, 20, 28, 14, 10, 14, 20, 45, 30, 25, 30, 30, 15]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = f"BARKAV_NonFinancial_Reports_{start_date or 'Start'}_to_{end_date or 'End'}.xlsx".replace("/", "-")
    return send_file(
        output,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route("/uploads/<filename>")
def serve_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)

if __name__ == "__main__":
    print("====================================================")
    print(" BARKAV Non-Financial & Daily Reports Web Server")
    print(" Local URL:   http://127.0.0.1:5050")
    print(" Network URL: http://0.0.0.0:5050")
    print(" Login:       " + ("فعال" if APP_PASSWORD else "غیرفعال (BARKAV_APP_PASSWORD تنظیم نشده)"))
    print("====================================================")
    # حالت debug فقط با BARKAV_DEBUG=1 فعال می‌شود تا روی هاست عمومی خاموش بماند
    app.run(host="0.0.0.0", port=5050, debug=os.environ.get("BARKAV_DEBUG") == "1")
