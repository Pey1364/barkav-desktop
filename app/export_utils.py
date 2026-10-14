from __future__ import annotations

import html
import os
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from typing import Any, Iterable
from xml.sax.saxutils import escape


def _col_letter(n:int)->str:
    s=''
    while n:
        n,rem=divmod(n-1,26); s=chr(65+rem)+s
    return s


def _xlsx_column_width(values: Iterable[Any]) -> float:
    longest = 0
    for value in values:
        for line in str(value if value is not None else '').splitlines() or ['']:
            longest = max(longest, len(line))
    return min(42.0, max(12.0, longest * 1.1 + 2.0))


def _xlsx_row_height(row: Iterable[Any], widths: list[float]) -> float:
    max_lines = 1
    for index, value in enumerate(row):
        text = str(value if value is not None else '')
        per_line = max(12, int(widths[index] * 1.35)) if index < len(widths) else 32
        lines = 0
        for line in text.splitlines() or ['']:
            lines += max(1, (len(line) + per_line - 1) // per_line)
        max_lines = max(max_lines, lines)
    return min(90.0, max(24.0, 15.0 * max_lines + 6.0))


def export_xlsx(path:str|Path, headers:list[str], rows:list[list[Any]], title:str='گزارش',
                summary_lines:list[str]|None=None) -> Path:
    """Create a readable RTL XLSX with wrapped text and visible cell borders."""
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    summary_rows=[[str(line or ''),] for line in (summary_lines or ())]
    all_rows=summary_rows+[list(headers or [])]+[list(row or []) for row in rows]
    column_count=max(1, max((len(row) for row in all_rows), default=0))
    all_rows=[row + [''] * (column_count - len(row)) for row in all_rows]
    widths = [
        _xlsx_column_width(row[index] for row in all_rows)
        for index in range(column_count)
    ]

    def cell_xml(r:int,c:int,v:Any,style_id:int)->str:
        ref=f'{_col_letter(c)}{r}'
        style=f' s="{style_id}"'
        if v is None: return f'<c r="{ref}"{style}/>'
        if isinstance(v,bool): return f'<c r="{ref}"{style} t="b"><v>{1 if v else 0}</v></c>'
        if isinstance(v,(int,float)) and not isinstance(v,bool): return f'<c r="{ref}"{style}><v>{v}</v></c>'
        text=escape(str(v))
        return f'<c r="{ref}"{style} t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'

    sheet_rows=[]
    for ri,row in enumerate(all_rows,1):
        header_row=len(summary_rows)+1
        cells=''.join(cell_xml(ri,ci,v,1 if ri == header_row else 0) for ci,v in enumerate(row,1))
        height=f' ht="{_xlsx_row_height(row, widths):.1f}" customHeight="1"'
        sheet_rows.append(f'<row r="{ri}"{height}>{cells}</row>')
    widths_xml=''.join(
        f'<col min="{index}" max="{index}" width="{width:.1f}" customWidth="1"/>'
        for index, width in enumerate(widths, 1)
    )
    sheet=f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetViews><sheetView workbookViewId="0" rightToLeft="1" showGridLines="0"/></sheetViews><cols>{widths_xml}</cols><sheetData>{''.join(sheet_rows)}</sheetData></worksheet>'''
    content='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>'''
    rels='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>'''
    wb=f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="{escape(title[:31] or 'Report')}" sheetId="1" r:id="rId1"/></sheets></workbook>'''
    wb_rels='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>'''
    styles='''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="2"><font><sz val="11"/><name val="Tahoma"/></font><font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Tahoma"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF153B5B"/><bgColor indexed="64"/></patternFill></fill></fills><borders count="2"><border/><border><left style="thin"><color rgb="FFB8C2CC"/></left><right style="thin"><color rgb="FFB8C2CC"/></right><top style="thin"><color rgb="FFB8C2CC"/></top><bottom style="thin"><color rgb="FFB8C2CC"/></bottom></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="1" applyAlignment="1"><alignment horizontal="right" vertical="center" wrapText="1"/></xf><xf numFmtId="0" fontId="1" fillId="2" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf></cellXfs></styleSheet>'''
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml',content); z.writestr('_rels/.rels',rels); z.writestr('xl/workbook.xml',wb); z.writestr('xl/_rels/workbook.xml.rels',wb_rels); z.writestr('xl/worksheets/sheet1.xml',sheet); z.writestr('xl/styles.xml',styles)
    return path


def style_xlsx_workbook(workbook) -> None:
    """Apply the shared RTL/font/layout contract to openpyxl workbooks."""

    from openpyxl.styles import Border, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style='thin', color='B8C2CC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type='solid', fgColor='153B5B')

    for sheet in workbook.worksheets:
        sheet.sheet_view.rightToLeft = True
        sheet.sheet_view.showGridLines = False
        rows = list(sheet.iter_rows())
        non_empty_rows = {
            row_index for row_index, row in enumerate(rows, 1)
            if any(cell.value not in (None, '') for cell in row)
        }
        header_rows = set()
        for row_index in non_empty_rows:
            values = [str(cell.value or '').strip() for cell in rows[row_index - 1]]
            first = values[0] if values else ''
            if row_index == 1 or first in {'کادر ویزارد', 'جدول پرداخت‌ها'} or 'ردیف' in values[:3]:
                header_rows.add(row_index)

        max_column = max((len(row) for row in rows), default=0)
        widths: list[float] = []
        for column_index in range(1, max_column + 1):
            values = [rows[row_index - 1][column_index - 1].value for row_index in non_empty_rows]
            width = _xlsx_column_width(values)
            widths.append(width)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

        image_rows = set()
        for image in getattr(sheet, '_images', ()):
            origin = getattr(getattr(image, 'anchor', None), '_from', None)
            if origin is not None and getattr(origin, 'row', None) is not None:
                image_rows.add(int(origin.row) + 1)

        for row_index, row in enumerate(rows, 1):
            if row_index not in non_empty_rows:
                continue
            is_header = row_index in header_rows
            row_height = _xlsx_row_height((cell.value for cell in row), widths)
            if row_index in image_rows:
                row_height = max(row_height, 90.0)
            sheet.row_dimensions[row_index].height = row_height
            for cell in row:
                font = copy(cell.font)
                font.name = 'Tahoma'
                if is_header:
                    font.bold = True
                    font.color = 'FFFFFF'
                    cell.fill = header_fill
                cell.font = font
                alignment = copy(cell.alignment)
                alignment.horizontal = 'right'
                alignment.vertical = 'center'
                alignment.wrap_text = True
                cell.alignment = alignment
                cell.border = border

def choose_pdf_orientation(headers: list[Any], rows: Iterable[Iterable[Any]] = ()) -> str:
    """Choose A4 orientation from the rendered table's approximate width."""

    columns = [str(value or '') for value in headers]
    for row in rows:
        for index, value in enumerate(row):
            if index < len(columns):
                columns[index] = max(columns[index], str(value or ''), key=len)
    estimated_width = sum(min(170, max(48, len(value) * 4.8 + 22)) for value in columns)
    # A4 portrait has about 547 points between the standard 24-point margins.
    return 'افقی' if len(columns) >= 6 or estimated_width > 547 else 'عمودی'


def _pdf_column_widths(headers: list[Any], rows: list[list[Any]], available_width: float) -> list[float]:
    """Return readable RTL table widths constrained to the selected A4 page."""

    widths = []
    for index, header in enumerate(headers):
        values = [str(header or '')]
        values.extend(str(row[index] or '') for row in rows if index < len(row))
        longest = max((len(value) for value in values), default=0)
        widths.append(min(170, max(48, longest * 4.8 + 22)))
    if not widths:
        return [available_width]
    total = sum(widths)
    if total > available_width:
        factor = available_width / total
        widths = [max(34, width * factor) for width in widths]
        total = sum(widths)
        if total > available_width:
            factor = available_width / total
            widths = [width * factor for width in widths]
    return widths


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates=[r'C:\Windows\Fonts\tahoma.ttf',r'C:\Windows\Fonts\arial.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    for p in candidates:
        if Path(p).exists():
            try: pdfmetrics.registerFont(TTFont('PeymaanFont',p)); return 'PeymaanFont'
            except Exception: pass
    return 'Helvetica'


def _rtl(s:Any)->str:
    text=str(s if s is not None else '')
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
    except ImportError as exc:
        raise RuntimeError(
            'برای خروجی PDF فارسی، وابستگی‌های arabic-reshaper و python-bidi لازم هستند. '
            'ابتدا requirements_stage9.txt را نصب کنید.'
        ) from exc
    return get_display(arabic_reshaper.reshape(text))


def export_pdf(path:str|Path,headers:list[str],rows:list[list[Any]],title:str='گزارش',summary_lines:list[str]|None=None, *, orientation: str | None = None)->Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True); font=_pdf_font()
    normalized_headers = list(headers or ['گزارش'])
    normalized_rows = [list(row[:len(normalized_headers)]) + [''] * max(0, len(normalized_headers) - len(row)) for row in rows]
    selected_orientation = orientation if orientation in {'افقی', 'عمودی'} else choose_pdf_orientation(normalized_headers, normalized_rows)
    page_size = A4 if selected_orientation == 'عمودی' else landscape(A4)
    doc=SimpleDocTemplate(str(path),pagesize=page_size,rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    story=[]; style=ParagraphStyle('rtl',fontName=font,fontSize=10,alignment=TA_RIGHT,leading=14)
    title_style=ParagraphStyle('title',fontName=font,fontSize=15,alignment=TA_RIGHT,leading=20)
    story.append(Paragraph(_rtl(title),title_style)); story.append(Spacer(1,8))
    for line in summary_lines or []: story.append(Paragraph(_rtl(line),style))
    if summary_lines: story.append(Spacer(1,10))
    # Treeviews and Excel sheets display their first logical field on the
    # physical right.  ReportLab has no sheet-level RTL flag, so reverse the
    # physical cell order explicitly while keeping values untouched.
    display_headers = list(reversed(normalized_headers))
    display_rows = [list(reversed(row)) for row in normalized_rows]
    data=[[_rtl(h) for h in display_headers]]+[[_rtl(v) for v in row] for row in display_rows]
    if not data: data=[[_rtl('بدون داده')]]
    table=Table(data, colWidths=_pdf_column_widths(display_headers, display_rows, page_size[0] - 48), repeatRows=1, hAlign='RIGHT')
    table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#153B5B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#B8C2CC')),('ALIGN',(0,0),(-1,-1),'RIGHT'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FA')])]))
    story.append(table); doc.build(story); return path


def print_file(path:str|Path)->None:
    p=str(Path(path))
    if os.name=='nt': os.startfile(p,'print')  # type: ignore[attr-defined]
    else:
        import subprocess; subprocess.Popen(['lp',p])
