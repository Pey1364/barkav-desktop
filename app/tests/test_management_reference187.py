"""Rules remain prose; land and formula-driven estimates survive round trips."""
import json
from pathlib import Path

import openpyxl
import pytest

from budget_rich_v163 import RichBudgetService
from finance_db import FinanceDB
from price_update_service import PriceUpdateService, RULES_SHEET_NAME
from project_management_rules187 import MANAGEMENT_RULES
from tools.update_management_reference187 import update_reference, rebase_formula


def test_rules_export_import_is_idempotent_and_does_not_change_prices(tmp_path):
    service = PriceUpdateService(FinanceDB(tmp_path / 'test.sqlite3'), tmp_path)
    pid = service.upsert_catalog_item(None, description='خرید زمین', catalog_type='LAND',
                                     unit='مترمربع', code='1', phase_codes='01', estimate_rate=100)
    before = service.get_catalog_item(pid)
    destination = service.export_reference_workbook(out_path=tmp_path / 'rules.xlsx')
    workbook = openpyxl.load_workbook(destination)
    assert 'قیمت زمین' in workbook.sheetnames
    topics = [r[0] for r in list(workbook[RULES_SHEET_NAME].values)[3:]]
    assert all(topic in topics for topic, _text in MANAGEMENT_RULES)
    assert workbook[RULES_SHEET_NAME].sheet_view.rightToLeft
    service.import_reference_rules(destination)
    service.import_reference_rules(destination)
    assert len(service.list_calc_rules()) == len(MANAGEMENT_RULES)
    assert service.apply_calc_rules() == {'rules': 0, 'items': 0, 'rules_without_method': 0}
    assert service.get_catalog_item(pid) == before

    rule = service.list_calc_rules()[0]
    params = json.loads(rule['params_json'])
    params['rule_text'] = 'قانون ویرایش‌شده با علامت + و −'
    service.upsert_calc_rule(rule['rule_id'], catalog_type='ALL', description=rule['description'],
                             params_json=json.dumps(params, ensure_ascii=False))
    edited = service.export_reference_workbook(out_path=tmp_path / 'edited.xlsx')
    service.import_reference_rules(edited)
    assert len(service.list_calc_rules()) == len(MANAGEMENT_RULES)
    assert any(json.loads(r['params_json'])['rule_text'] == params['rule_text']
               for r in service.list_calc_rules())


def test_land_from_project_is_not_misclassified_as_material(tmp_path):
    db = FinanceDB(tmp_path / 'test.sqlite3')
    rich = RichBudgetService(db, tmp_path)
    pid = rich.ensure_land_price()
    with db.tx() as con:
        con.execute('UPDATE rich_price_items SET estimate_rate=100,forecast_rate=0,actual_rate=33 '
                    'WHERE price_id=?', (pid,))
    service = PriceUpdateService(db, tmp_path)
    service.import_from_rich_price_items()
    land = [r for r in service.list_catalog_items() if r['rich_price_id'] == pid]
    assert len(land) == 1
    assert land[0]['catalog_type'] == 'LAND'
    assert land[0]['forecast_rate'] == 0
    assert land[0]['actual_rate'] == 33
    path = service.export_reference_workbook(out_path=tmp_path / 'land.xlsx')
    other = PriceUpdateService(FinanceDB(tmp_path / 'other.sqlite3'), tmp_path)
    assert other.import_reference_workbook(path)['LAND'] == 1
    assert other.list_catalog_items(catalog_type='LAND')[0]['calc_method'] == ''
    assert other.list_catalog_items(catalog_type='LAND')[0]['forecast_rate'] == 0
    assert other.list_catalog_items(catalog_type='LAND')[0]['actual_rate'] == 33


def test_rebase_only_changes_target_ranges_and_preserves_signs_and_text():
    formula = "SUM($L$10:L20)+'براورد کلی'!M10-'قیمت خدمات'!D10+IF(B2=\"L10\",-2,+3)"
    assert rebase_formula(formula, 'براورد کلی', 'براورد کلی', 10, 2) == (
        "SUM($L$12:L22)+'براورد کلی'!M12-'قیمت خدمات'!D10+IF(B2=\"L10\",-2,+3)")
    assert rebase_formula("'براورد کلی'!L10-L10", 'روکش مالی', 'براورد کلی', 10, 2) == (
        "'براورد کلی'!L12-L10")


def reference_fixture(path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet('براورد کلی')
    sheet['B4'] = 'فاز 01 — خرید زمین'
    sheet['B6'] = 'C. خدمات'
    for row in (7, 8, 9):
        sheet.cell(row, 2, f'01-{row-6}')
        sheet.cell(row, 3, 'بنگاه' if row == 7 else 'خدمات زمین')
        sheet.cell(row, 4, 1)
        sheet.cell(row, 5, 'مقطوع')
        sheet.cell(row, 6, 10)
        sheet.cell(row, 9, 1)
        sheet.cell(row, 12, 10)
        sheet.row_dimensions[row].outlineLevel = 2
    sheet['B10'] = 'جمع فاز 01'
    sheet.merge_cells('B10:K10')
    for col in 'LMN':
        sheet[f'{col}10'] = f'=SUM({col}7:{col}9)'
    sheet['B12'] = 'فاز 02 — مطالعات اولیه'
    sheet['L13'] = -125
    sheet['M13'] = '=L13+10-5'
    sheet.row_dimensions[13].outlineLevel = 1
    sheet.row_dimensions[13].hidden = True
    sheet.row_dimensions[12].collapsed = True
    sheet.sheet_properties.outlinePr.showOutlineSymbols = True
    sheet.sheet_view.showOutlineSymbols = True
    cover = workbook.create_sheet('روکش مالی')
    cover['D4'] = "='براورد کلی'!L10"
    for title in ('قیمت مصالح', 'قیمت پیمانکار', 'قیمت خدمات'):
        sheet = workbook.create_sheet(title)
        sheet.append(['قیمت'])
        sheet.append([])
        sheet.append(['ردیف', 'عنوان', 'واحد', 'برآورد اولیه', 'برآورد بروز', 'واقعی', 'فازهای مورد استفاده'])
        sheet.append([1, 'خدمات زمین', 'مقطوع', 10, 20, 30, '01'])
        sheet.merge_cells('A1:G1')
        sheet.merge_cells('D2:F2')
    sheet = workbook.create_sheet('01-خرید زمین')
    sheet['A1'] = 'فاز 01'
    for row in (5, 6, 7):
        sheet.cell(row, 1, row - 4)
        sheet.cell(row, 3, 'خدمات زمین')
    sheet['A8'] = 'جمع کل فاز'
    for col in 'MNO':
        sheet[f'{col}8'] = f'=SUM({col}5:{col}7)'
    sheet = workbook.create_sheet(RULES_SHEET_NAME)
    sheet['A3'] = 'موضوع'
    sheet['B3'] = 'قانون'
    sheet['A4'] = '■ بخش ۱'
    sheet['A5'] = 'اصل قبلی'
    sheet['B5'] = 'علامت + و − حفظ شود'
    workbook.save(path)


def test_reference_copy_preserves_outline_negative_values_and_land_import(tmp_path):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    reference_fixture(source)
    original = source.read_bytes()
    update_reference(source, output)
    assert source.read_bytes() == original
    workbook = openpyxl.load_workbook(output)
    assert workbook['براورد کلی']['L15'].value == -125
    assert workbook['براورد کلی']['M15'].value == '=L15+10-5'
    assert workbook['براورد کلی'].row_dimensions[15].hidden
    assert workbook['براورد کلی'].row_dimensions[14].collapsed
    assert workbook['براورد کلی'].sheet_view.showOutlineSymbols
    assert workbook['روکش مالی']['D4'].value == "='براورد کلی'!L12"
    assert str(workbook['براورد کلی'].merged_cells) == 'B12:K12'
    assert workbook['قیمت زمین']['C4'].value == 'مترمربع'
    db = FinanceDB(tmp_path / 'test.sqlite3')
    rich = RichBudgetService(db, tmp_path)
    rich.import_workbook(output, preserve_manual=False)
    land_rows = [r for r in rich.list_phase_items(1) if r['cost_type'] == 'خرید زمین']
    assert len(land_rows) == 1
    assert land_rows[0]['description'] == 'خرید زمین'
    assert land_rows[0]['plan_ref'] == '01-4'
    assert rich.get_price(land_rows[0]['price_id'])['category'] == 'LAND'
    with pytest.raises(ValueError):
        update_reference(source, source)
    with pytest.raises(ValueError):
        update_reference(output, tmp_path / 'duplicate.xlsx')


def test_rules_window_displays_management_rules_once_after_import(tmp_path, monkeypatch):
    from types import SimpleNamespace
    monkeypatch.setenv('TELEGRAM_BOT_TOKEN', 'test')
    monkeypatch.setenv('GEMINI_API_KEY', 'test')
    monkeypatch.setenv('ADMIN_USER_ID', '1')
    from price_update_stage18_7 import GeneralRulesWindow187, ALL_LABEL

    class Table:
        def set_rows(self, rows, keys):
            self.rows, self.keys = rows, keys

    service = PriceUpdateService(FinanceDB(tmp_path / 'test.sqlite3'), tmp_path)
    window = SimpleNamespace(service=service, category_var=SimpleNamespace(get=lambda: ALL_LABEL),
                             params_table=Table(), rules_table=Table(), _all_phases=lambda _c: set())
    GeneralRulesWindow187._refresh(window)
    assert sum(k.startswith('MANAGEMENT:') for k in window.rules_table.keys) == len(MANAGEMENT_RULES)
    exported = service.export_reference_workbook(out_path=tmp_path / 'rules.xlsx')
    service.import_reference_rules(exported)
    GeneralRulesWindow187._refresh(window)
    assert sum(row[2].startswith('PM-') for row in window.rules_table.rows) == len(MANAGEMENT_RULES)
    assert not any(k.startswith('MANAGEMENT:') for k in window.rules_table.keys)


def test_excel_namespace_bindings_and_new_sheet_relationships_are_preserved(tmp_path):
    """openpyxl accepts dangling compatibility prefixes; real Excel may reject them."""
    from lxml import etree as xml
    from zipfile import ZipFile
    from tools.update_management_reference187 import MAIN, REL, Q, validate_excel_package

    source, output = tmp_path / 'namespaces.xlsx', tmp_path / 'corrected.xlsx'
    reference_fixture(source)
    with ZipFile(source) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    mc = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    unused = 'http://schemas.microsoft.com/office/spreadsheetml/2010/11/main'
    for path in ('xl/workbook.xml', 'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet5.xml'):
        original = xml.fromstring(files[path])
        root = xml.Element(original.tag, nsmap={**original.nsmap, 'mc': mc, 'unused': unused})
        root.attrib.update(original.attrib)
        root.set(f'{{{mc}}}Ignorable', 'unused')
        root[:] = list(original)
        if path == 'xl/workbook.xml':
            alternate = xml.SubElement(root, f'{{{mc}}}AlternateContent')
            xml.SubElement(alternate, f'{{{mc}}}Choice', Requires='unused')
            xml.SubElement(alternate, f'{{{mc}}}Fallback')
        if path == 'xl/worksheets/sheet5.xml':  # Source services sheet owns a VML part.
            xml.SubElement(root, Q('legacyDrawing'), {f'{{{REL}}}id': 'rIdDrawing'})
        files[path] = xml.tostring(root)
    files['xl/worksheets/_rels/sheet5.xml.rels'] = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdDrawing" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" '
        'Target="../drawings/test.vml"/></Relationships>').encode()
    files['xl/drawings/test.vml'] = b'<xml/>'
    with ZipFile(source, 'w') as archive:
        for path, content in files.items():
            archive.writestr(path, content)
    update_reference(source, output)
    with ZipFile(output) as archive:
        updated = {name: archive.read(name) for name in archive.namelist()}
    validate_excel_package(updated)
    workbook = xml.fromstring(updated['xl/workbook.xml'])
    assert workbook.nsmap['unused'] == unused
    assert workbook.get(f'{{{mc}}}Ignorable') == 'unused'
    roots = [xml.fromstring(data) for name, data in updated.items() if name.startswith('xl/worksheets/')
             and '/_rels/' not in name and name.endswith('.xml')]
    land = next(root for root in roots if any(cell.text == 'کاتالوگ قیمت زمین (برآورد اولیه / برآورد بروز / واقعی)'
                                            for cell in root.iter(Q('t'))))
    assert not list(land.iter(Q('legacyDrawing')))
    assert updated['xl/worksheets/_rels/sheet5.xml.rels'] == files['xl/worksheets/_rels/sheet5.xml.rels']

    # Demonstrate that validation catches precisely the previous failure.
    broken = dict(updated)
    broken['xl/workbook.xml'] = updated['xl/workbook.xml'].replace(
        f' xmlns:unused="{unused}"'.encode(), b'')
    with pytest.raises(ValueError, match='undeclared Excel compatibility prefix'):
        validate_excel_package(broken)
    broken = dict(updated)
    root = xml.fromstring(updated['xl/workbook.xml'])
    xml.SubElement(root, Q('externalReferences'), {f'{{{REL}}}id': 'rIdMissing'})
    broken['xl/workbook.xml'] = xml.tostring(root)
    with pytest.raises(ValueError, match='missing relationship'):
        validate_excel_package(broken)
