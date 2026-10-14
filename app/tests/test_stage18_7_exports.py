from pathlib import Path
import sys
import zipfile
from types import SimpleNamespace


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


class _FormatVar:
    def __init__(self, value):
        self.value = value
        self.set_values = []

    def get(self):
        return self.value

    def set(self, value):
        self.value = value
        self.set_values.append(value)


class _ExportTable:
    def __init__(self):
        self.calls = []

    def export_excel(self):
        self.calls.append('Excel')

    def export_pdf(self):
        self.calls.append('PDF')

    def print_report(self):
        self.calls.append('پرینت')


def test_transaction_format_selection_dispatches_and_resets():
    from dashboard_stage18_7 import Stage187Dashboard

    dashboard = object()
    for selected, expected in (('Excel', 'Excel'), ('PDF', 'PDF'), ('پرینت', 'پرینت')):
        table = _ExportTable()
        variable = _FormatVar(selected)
        Stage187Dashboard._export_transaction_187(dashboard, table, variable)
        assert table.calls == [expected]
        assert variable.value == 'خروجی'
        assert variable.set_values == ['خروجی']


def test_transaction_context_menu_dispatches_all_formats():
    from dashboard_stage18_7 import Stage187Dashboard

    dashboard = object()
    table = _ExportTable()
    for selected in ('Excel', 'PDF', 'پرینت'):
        Stage187Dashboard._export_transaction_context_187(dashboard, table, selected)
    assert table.calls == ['Excel', 'PDF', 'پرینت']


def test_transaction_toolbar_binds_format_selection_to_export_callback():
    source = (ROOT / 'dashboard_stage18_7.py').read_text(encoding='utf-8')
    assert "output_combo.bind(\n                '<<ComboboxSelected>>'" in source
    assert 'table._barkav_output_combo187 = output_combo' in source


def test_shared_xlsx_writer_is_rtl_and_uses_tahoma(tmp_path):
    from export_utils import export_xlsx

    destination = export_xlsx(
        tmp_path / 'transaction.xlsx',
        ['عنوان تراکنش', 'مبلغ'],
        [['هزینه ساخت', '۱۲۳۴۵۶']],
        title='گزارش تراکنش',
    )
    assert destination.exists()
    with zipfile.ZipFile(destination) as archive:
        sheet = archive.read('xl/worksheets/sheet1.xml').decode('utf-8')
        styles = archive.read('xl/styles.xml').decode('utf-8')
    assert 'rightToLeft="1"' in sheet
    assert 'Tahoma' in styles
    assert 'horizontal="right"' in styles
    assert 'vertical="center"' in styles
    assert 'borderId="1"' in styles
    assert 'showGridLines="0"' in sheet
    assert 'customHeight="1"' in sheet
    assert 'هزینه ساخت' in sheet


def test_pdf_auto_orientation_uses_standard_a4_and_keeps_logical_first_column_right(tmp_path):
    import re
    from export_utils import choose_pdf_orientation, export_pdf

    narrow_headers = ['ردیف', 'عنوان', 'مبلغ']
    wide_headers = [f'ستون{i}' for i in range(8)]
    assert choose_pdf_orientation(narrow_headers) == 'عمودی'
    assert choose_pdf_orientation(wide_headers) == 'افقی'

    narrow = export_pdf(
        tmp_path / 'narrow.pdf', narrow_headers,
        [['۱', 'هزینه ساخت', '۱۲۳']], title='گزارش باریک',
    )
    wide = export_pdf(
        tmp_path / 'wide.pdf', wide_headers,
        [['x'] * len(wide_headers)], title='گزارش عریض',
    )
    narrow_page = re.search(rb'/MediaBox\s*\[\s*0\s+0\s+([0-9.]+)\s+([0-9.]+)', narrow.read_bytes())
    wide_page = re.search(rb'/MediaBox\s*\[\s*0\s+0\s+([0-9.]+)\s+([0-9.]+)', wide.read_bytes())
    assert narrow_page and float(narrow_page.group(1)) < float(narrow_page.group(2))
    assert wide_page and float(wide_page.group(1)) > float(wide_page.group(2))

    # ReportLab has no sheet-level RTL flag: the shared writer reverses the
    # physical cells so the first logical field remains on the right edge.
    source = (ROOT / 'export_utils.py').read_text(encoding='utf-8')
    assert 'display_headers = list(reversed(normalized_headers))' in source
    assert "hAlign='RIGHT'" in source


def test_openpyxl_exports_receive_shared_rtl_font_contract(tmp_path):
    from openpyxl import Workbook, load_workbook
    from export_utils import style_xlsx_workbook

    book = Workbook()
    sheet = book.active
    sheet.append(['عنوان', 'مبلغ'])
    sheet.append(['ورودی تسهیلات', '۱۲۳'])
    style_xlsx_workbook(book)
    destination = tmp_path / 'openpyxl.xlsx'
    book.save(destination)

    loaded = load_workbook(destination)
    output = loaded.active
    assert output.sheet_view.rightToLeft is True
    assert output['A1'].font.name == 'Tahoma'
    assert output['A1'].alignment.horizontal == 'right'
    assert output['A1'].alignment.vertical == 'center'
    assert output['A1'].border.left.style == 'thin'
    assert output['A1'].fill.fgColor.rgb.endswith('153B5B')
    assert output.column_dimensions['B'].width >= 12
    assert output.row_dimensions[2].height >= 24
    assert output['A2'].alignment.wrap_text is True
    assert output['A2'].value == 'ورودی تسهیلات'


def test_pdf_exports_use_available_rtl_font_and_create_files(tmp_path):
    from reportlab.pdfbase import pdfmetrics
    from export_utils import _pdf_font, export_pdf
    from sales_claims_stage18_7 import export_contract_pdf_187, export_receipts_pdf_187

    font_name = _pdf_font()
    font = pdfmetrics.getFont(font_name)
    filename = str(getattr(font.face, 'filename', '')).lower()
    assert font_name == 'PeymaanFont'
    assert filename.endswith(('tahoma.ttf', 'arial.ttf', 'dejavusans.ttf'))

    paths = [
        export_pdf(tmp_path / 'shared.pdf', ['عنوان', 'مبلغ'], [['هزینه', '۱۲۳']], title='گزارش فارسی'),
        export_contract_pdf_187(tmp_path / 'contract.pdf', ['عنوان', 'مبلغ'], [{'عنوان': 'قرارداد', 'مبلغ': '۱۲۳'}], 'قرارداد'),
        export_receipts_pdf_187(tmp_path / 'receipts.pdf', ['مرحله', 'مبلغ'], [['قسط اول', '۱۲۳']], 'اسناد واریزی', [], []),
    ]
    assert all(path.exists() and path.stat().st_size > 1000 for path in paths)
    assert all(path.read_bytes().startswith(b'%PDF-') for path in paths)


class _TreeForExport:
    columns = ('row', 'title', 'amount')
    labels = {'row': 'ردیف', 'title': 'عنوان', 'amount': 'مبلغ'}
    values = {'1': ('۱', 'هزینه ساخت', '۱۲۳')}

    def selection(self):
        return ()

    def get_children(self):
        return tuple(self.values)

    def heading(self, key, option=None):
        return self.labels[key] if option else {'text': self.labels[key]}

    def item(self, iid, option=None):
        return self.values[iid] if option == 'values' else {'values': self.values[iid]}

    def __getitem__(self, key):
        if key == 'columns':
            return self.columns
        raise KeyError(key)


def test_sales_tree_exports_all_file_formats_with_rtl_xlsx(tmp_path, monkeypatch):
    import sales_claims_stage18_7 as sales

    tree = _TreeForExport()
    panel = SimpleNamespace(finance_tree=tree)
    destinations = {fmt: tmp_path / f'tree-{fmt}.{"xlsx" if fmt == "xlsx" else fmt}'
                    for fmt in ('xlsx', 'pdf', 'csv', 'txt', 'json')}
    monkeypatch.setattr(
        sales.filedialog, 'asksaveasfilename',
        lambda **_kwargs: str(destinations[monkeypatch.current_format]),
    )
    monkeypatch.setattr(sales.messagebox, 'showinfo', lambda *_args, **_kwargs: None)
    monkeypatch.setattr(sales.messagebox, 'showerror', lambda *_args, **_kwargs: None)
    for fmt in destinations:
        monkeypatch.current_format = fmt
        sales.SalesClaimsPanel187._export_tree(panel, tree, 'جدول تراکنش', fmt)
        assert destinations[fmt].exists() and destinations[fmt].stat().st_size > 0

    from openpyxl import load_workbook
    workbook = load_workbook(destinations['xlsx'])
    assert workbook.active.sheet_view.rightToLeft is True
    assert workbook.active['A1'].font.name == 'Tahoma'


def test_openpyxl_export_branches_keep_shared_font_and_rtl(tmp_path, monkeypatch):
    import contractors_stage18_7 as contractors
    import sales_claims_stage18_7 as sales

    outputs = iter(tmp_path / name for name in (
        'receipts.xlsx', 'receipts.pdf', 'contract.xlsx', 'contract.pdf',
        'finance.xlsx', 'finance.pdf', 'contractors.xlsx', 'contractors.pdf',
    ))
    monkeypatch.setattr(sales.filedialog, 'asksaveasfilename', lambda **_kwargs: str(next(outputs)))
    monkeypatch.setattr(contractors.filedialog, 'asksaveasfilename', lambda **_kwargs: str(next(outputs)))
    monkeypatch.setattr(sales.messagebox, 'showinfo', lambda *_args, **_kwargs: None)
    monkeypatch.setattr(sales.messagebox, 'showerror', lambda *_args, **_kwargs: None)
    monkeypatch.setattr(contractors.messagebox, 'showinfo', lambda *_args, **_kwargs: None)
    monkeypatch.setattr(contractors.messagebox, 'showerror', lambda *_args, **_kwargs: None)

    receipt_owner = SimpleNamespace()
    assert sales.PaymentDocuments187._save_export(
        receipt_owner, 'Excel', ['عنوان'], [{'عنوان': 'رسید'}], [], [], 'افقی',
    ) is True
    assert sales.PaymentDocuments187._save_export(
        receipt_owner, 'PDF', ['عنوان'], [{'عنوان': 'رسید'}], [], [], 'افقی',
    ) is True

    contract_owner = SimpleNamespace(_wizard_export_fields=[('وام', 'نام', 'ساخت')])
    assert sales.ContractDetails187._save_export(
        contract_owner, 'Excel', ['عنوان'], [{'عنوان': 'قرارداد'}], [], [], 'افقی', [],
    ) is True
    assert sales.ContractDetails187._save_export(
        contract_owner, 'PDF', ['عنوان'], [{'عنوان': 'قرارداد'}], [], [], 'افقی', [],
    ) is True

    class _FinanceOwner:
        def _finance_export_documents(self, _item_ids):
            return []

    sales.SalesClaimsPanel187._save_finance_export(
        _FinanceOwner(), 'Excel', ['عنوان'], [{'عنوان': 'پرداخت'}], [],
    )
    sales.SalesClaimsPanel187._save_finance_export(
        _FinanceOwner(), 'PDF', ['عنوان'], [{'عنوان': 'پرداخت'}], [],
    )

    contractor_owner = SimpleNamespace()
    contractors.ContractorsPanel187._write_export(
        contractor_owner, ['عنوان'], [{'عنوان': 'پیمانکار'}], 'xlsx', 'پیمانکاران',
    )
    contractors.ContractorsPanel187._write_export(
        contractor_owner, ['عنوان'], [{'عنوان': 'پیمانکار'}], 'pdf', 'پیمانکاران',
    )

    from openpyxl import load_workbook
    for name in ('receipts.xlsx', 'contract.xlsx', 'finance.xlsx', 'contractors.xlsx'):
        workbook = load_workbook(tmp_path / name)
        assert workbook.active.sheet_view.rightToLeft is True
        assert workbook.active['A1'].font.name == 'Tahoma'
    for name in ('receipts.pdf', 'contract.pdf', 'finance.pdf', 'contractors.pdf'):
        assert (tmp_path / name).read_bytes().startswith(b'%PDF-')


def test_sales_saved_workbook_has_the_same_three_section_header(tmp_path):
    from openpyxl import Workbook, load_workbook
    from export_utils import style_xlsx_workbook
    from sales_claims_stage18_7 import apply_export_workbook_header_187

    destination = tmp_path / 'sales-header.xlsx'
    book = Workbook()
    sheet = book.active
    headers = ['مرحله', 'تاریخ واریز', 'مبلغ واریز (تومان)', 'فیش واریزی']
    specs = [('۱. مورد معامله', 'نام پروژه', 'بارکاو۲')]
    apply_export_workbook_header_187(
        sheet, headers=headers, specs=specs, output_title='خروجی جزئیات قرارداد',
    )
    sheet.append(headers)
    sheet.append(['قسط اول', '۱۴۰۵/۰۶/۰۱', '۱۰۰', ''])
    style_xlsx_workbook(book)
    apply_export_workbook_header_187(
        sheet, headers=headers, specs=specs,
        output_title='خروجی جزئیات قرارداد', add_logo=False,
    )
    book.save(destination)

    output = load_workbook(destination).active
    assert output.sheet_view.rightToLeft is True
    assert output['A1'].value == 'جزئیات قرارداد فروش'
    assert output['B1'].value == 'شرکت مهندسی ساخت بارکاو'
    assert output['B2'].value == 'پروژه بارکاو۲'
    assert output['D1'].value == 'BARKAV'
    assert output['A4'].value == 'مرحله'


class _TransactionTreeForPayload:
    columns = ('row', 'hidden_id', 'title', 'amount')
    labels = {
        'row': 'ردیف', 'hidden_id': 'ID',
        'title': 'عنوان تراکنش', 'amount': 'مبلغ (تومان)',
    }
    values = {'item-1': ('۱', 'TX-1', 'هزینه ساخت', '۱۲۳,۴۵۶')}

    def cget(self, option):
        if option == 'columns':
            return self.columns
        if option == 'displaycolumns':
            return '#all'
        raise KeyError(option)

    def heading(self, key, option=None):
        return self.labels[key] if option == 'text' else {'text': self.labels[key]}

    def get_children(self):
        return tuple(self.values)

    def item(self, iid, option=None):
        return self.values[iid] if option == 'values' else {'values': self.values[iid]}


def test_transaction_export_payload_uses_display_headers_for_preview_and_file_rows():
    from dashboard_stage18_7 import Stage187Dashboard
    from sales_claims_stage18_7 import contract_preview_payment_rows_187

    table = SimpleNamespace(
        tree=_TransactionTreeForPayload(),
        _barkav_hidden_id_columns187=('hidden_id',),
        _iid_to_row={'item-1': {'id': 'TX-1'}},
    )
    dashboard = SimpleNamespace(
        _transaction_document_paths_187=lambda _table, _row: ['C:/docs/receipt.png'],
    )

    headers, rows = Stage187Dashboard._transaction_export_payload_187(dashboard, table)

    assert headers == ['ردیف', 'عنوان تراکنش', 'مبلغ (تومان)']
    assert rows == [{
        'ردیف': '۱', 'عنوان تراکنش': 'هزینه ساخت', 'مبلغ (تومان)': '۱۲۳,۴۵۶',
        '__document_paths_187': ['C:/docs/receipt.png'],
    }]
    assert contract_preview_payment_rows_187(headers, rows)[0]['عنوان تراکنش'] == 'هزینه ساخت'
    assert [rows[0].get(header, '') for header in headers] == ['۱', 'هزینه ساخت', '۱۲۳,۴۵۶']
