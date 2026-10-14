"""مالی، صورت‌وضعیت and آرشیو: the three پیمانکاران sub-tabs docs/theme_18_7's
مشخصات-تب-قراردادهای-پیمانکاران spec describes beyond قراردادها itself.

Same conventions as test_stage18_7_ui_contract.py: mostly source-string
inspection for wizard/window structure, a few real-Tk-root tests for
StringVar-level behaviour, and pure-data tests against a temporary JSON
store for the store-layer functions.
"""

from __future__ import annotations

import contextlib
import os
import tempfile
from pathlib import Path

from test_transaction_lifecycle_matrix import runtime

ROOT = Path(__file__).resolve().parents[1]


def _source(path: Path) -> str:
    return path.read_text(encoding='utf-8')


def _until_next_def_or_class(body: str) -> str:
    """Slice a module-level function/class body up to whichever of the next
    top-level def/class comes first - either may legitimately be absent."""

    candidates = [i for i in (body.find('\ndef ', 10), body.find('\nclass ', 10)) if i != -1]
    return body[:min(candidates)] if candidates else body


@contextlib.contextmanager
def _temp_store():
    """Point contractor_store_stage18_7 at an empty, throwaway JSON folder."""

    folder = tempfile.mkdtemp()
    previous = os.environ.get('STAGE4_DB_PATH')
    os.environ['STAGE4_DB_PATH'] = str(Path(folder) / 'x.sqlite3')
    try:
        yield
    finally:
        if previous is None:
            os.environ.pop('STAGE4_DB_PATH', None)
        else:
            os.environ['STAGE4_DB_PATH'] = previous


# ---------------------------------------------------------------------------
# صورت‌وضعیت store layer
# ---------------------------------------------------------------------------

def test_statement_amount_is_its_rows_plus_bonus_minus_penalty():
    import contractor_store_stage18_7 as store

    with _temp_store():
        statement = {
            'rows': [{'line_total': '500,000,000'}, {'line_total': '10,000,000'}],
            'bonus_amount': '5,000,000', 'penalty_amount': '2,000,000',
        }
        assert store.statement_rows_total(statement['rows']) == 510_000_000
        assert store.statement_amount(statement) == 513_000_000


def test_cumulative_statement_amount_adds_up_a_contracts_own_history_in_order():
    """مبلغ صورت‌وضعیت تجمعی is every صورت‌وضعیت on one قرارداد, oldest
    first - not every صورت‌وضعیت in the file, and not out of date order.
    """

    import contractor_store_stage18_7 as store

    with _temp_store():
        store.upsert_statement({'statement_id': 'S1', 'contract_id': 'C1',
                                'statement_date': '1405-01-01', 'rows': [{'line_total': '100'}]})
        store.upsert_statement({'statement_id': 'S2', 'contract_id': 'C1',
                                'statement_date': '1405-02-01', 'rows': [{'line_total': '200'}]})
        store.upsert_statement({'statement_id': 'S3', 'contract_id': 'C2',
                                'statement_date': '1405-01-15', 'rows': [{'line_total': '999'}]})

        history = store.statements_for_contract('C1')
        assert [s['statement_id'] for s in history] == ['S1', 'S2']
        assert store.cumulative_statement_amount('C1') == 300
        assert store.cumulative_statement_amount('C1', 'S1') == 100
        assert store.cumulative_statement_amount('C2') == 999


def test_statement_crud_round_trips_through_its_own_json_file():
    import contractor_store_stage18_7 as store

    with _temp_store():
        assert store.load_statements() == []
        saved = store.upsert_statement({'contract_id': 'C1', 'statement_number': '1'})
        assert saved['statement_id']
        assert store.get_statement(saved['statement_id'])['statement_number'] == '1'

        saved['statement_number'] = '1-اصلاحی'
        store.upsert_statement(saved)
        assert len(store.load_statements()) == 1, 'a second upsert with the same id replaces, not appends'
        assert store.get_statement(saved['statement_id'])['statement_number'] == '1-اصلاحی'

        store.delete_statement(saved['statement_id'])
        assert store.get_statement(saved['statement_id']) is None


# ---------------------------------------------------------------------------
# آرشیو store layer
# ---------------------------------------------------------------------------

def test_profile_crud_round_trips_through_its_own_json_file():
    import contractor_store_stage18_7 as store

    with _temp_store():
        assert store.load_profiles() == []
        saved = store.upsert_profile({'family_name': 'علی رضایی', 'responsiveness': '8'})
        assert saved['profile_id']
        assert store.get_profile(saved['profile_id'])['responsiveness'] == '8'

        saved['responsiveness'] = '9'
        store.upsert_profile(saved)
        assert len(store.load_profiles()) == 1
        assert store.get_profile(saved['profile_id'])['responsiveness'] == '9'

        store.delete_profile(saved['profile_id'])
        assert store.get_profile(saved['profile_id']) is None


# ---------------------------------------------------------------------------
# صورت‌وضعیت wizard/record window structure
# ---------------------------------------------------------------------------

def test_statement_row_totals_itself_from_the_contracts_own_unit_price():
    import tkinter as tk
    import contractor_statements_stage18_7 as cs

    root = tk.Tk()
    root.withdraw()
    try:
        row = cs.blank_statement_row()
        row['unit_price'].set('500000')
        assert row['line_total'].get() == '', 'no مقدار کارکرد yet - nothing to multiply'
        row['worked_quantity'].set('1000')
        digits = row['line_total'].get().translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789'))
        assert digits.replace(',', '').replace('٬', '') == '500000000'

        # حجم کل قرارداد and شرح are read-only reference fields.
        assert 'description' in cs.STATEMENT_READONLY_ROW_KEYS
        assert 'unit_price' in cs.STATEMENT_READONLY_ROW_KEYS
        assert 'contract_quantity' in cs.STATEMENT_READONLY_ROW_KEYS
        assert 'worked_quantity' not in cs.STATEMENT_READONLY_ROW_KEYS
    finally:
        root.destroy()


def test_statement_wizard_discards_a_never_confirmed_new_statement_on_close():
    """Unlike ContractWizard187/ArchiveWizard187 (which still carry a
    brand-new record on as a draft through the record window), closing a
    brand-new صورت‌وضعیت before تأیید نهایی must create no row at all - the
    operator's explicit reversal of the older shared rule, since a
    صورت‌وضعیت the wizard never reached تأیید نهایی on should never appear
    in the table. Closing an EDIT of an existing صورت‌وضعیت must still
    notify to reopen it unchanged, matching the other wizards' own cancel
    policy.
    """

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    close_body = source[source.index('    def close(self) -> None:'):]
    close_body = close_body[:close_body.index('\n    def ', 10)]
    assert 'if self._saved_once:' in close_body
    assert "if self.record.get('statement_id')" in close_body
    assert 'self.on_saved(self.record, True)' in close_body
    assert 'self._save(' not in close_body


def test_statement_amount_column_reflects_the_real_rows_and_bonus_penalty():
    """The wizard's own مبلغ صورت‌وضعیت caption must total the rows plus
    پاداش minus جریمه, the same rule statement_amount() applies in the
    store - not the plan's estimate and not the rows alone.
    """

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    retotal = source[source.index('    def _retotal(self) -> None:'):]
    retotal = retotal[:retotal.index('\n    def ', 10)]
    assert "amount_of(self._var('bonus_amount')" in retotal
    assert "amount_of(self._var('penalty_amount')" in retotal
    assert 'total + bonus - penalty' in retotal


# ---------------------------------------------------------------------------
# آرشیو wizard/record window structure
# ---------------------------------------------------------------------------

def test_archive_criteria_share_پاسخگویی_and_خوش‌قولی_across_every_category():
    """پاسخگویی و خوش‌قولی are common to all three دسته‌بندی criteria sets,
    each rated ۱ تا ۱۰ - the survey pages' own score dropdown range, shown
    in Persian digits since it is what the user sees in the کرکره directly.
    """

    import contractor_archive_stage18_7 as ca
    from ui_stage18_3 import fa
    assert ca.SCORE_RANGE == tuple(fa(str(n)) for n in range(1, 11))
    for category in ca.CATEGORY_OPTIONS:
        criteria = dict(ca.criteria_for(category))
        assert criteria.get('responsiveness') == 'پاسخگویی'
        assert criteria.get('reliability') == 'خوش‌قولی'


def test_archive_wizard_mirrors_the_contract_wizards_cancel_on_close_rule():
    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    close_body = source[source.index('    def close(self) -> None:'):]
    close_body = close_body[:close_body.index('\n    def ', 10)]
    assert 'if self._saved_once:' in close_body
    assert "if not self.record.get('profile_id'):" in close_body
    assert 'self.on_saved(self.record, True)' in close_body


def test_archive_wizard_opens_on_its_real_default_category_not_the_first_combo_option():
    """render_step(0, ...) fires the instant super().__init__() runs - before
    __init__ used to call _seed_from_record(). category's own _var() call in
    _render_fields ran first in that ordering and seeded the StringVar with
    CATEGORY_OPTIONS[0] ('تامین‌کننده مصالح'); _seed_from_record()'s own
    _var('category', DEFAULT_CATEGORY) came too late to overwrite it, since
    _var() never resets a key that already exists. A brand-new archive
    profile silently opened on the wrong دسته‌بندی (and so the wrong
    معیارهای نظرسنجی) until the user noticed and changed it by hand.
    """

    import tkinter as tk
    import contractor_archive_stage18_7 as ca

    root = tk.Tk()
    root.withdraw()
    try:
        wiz = ca.ArchiveWizard187(root, rich=None)
        try:
            assert wiz.values['category'].get() == ca.DEFAULT_CATEGORY == 'پیمانکار'
        finally:
            wiz.destroy()
    finally:
        root.destroy()


def test_archive_record_window_reads_its_own_stored_history_list():
    """۴. تاریخ‌های همکاری در طول پروژه - ردیف/تاریخ/فاز/موضوع per entry,
    read from the profile's own 'history' list rather than re-matched
    against قرارداد every time (auto-populating that list from قراردادها/
    تراکنش‌ها is separate follow-up work, not this window's own job).
    """

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    history = source[source.index("    def _render_history(self, parent) -> None:"):]
    history = history[:history.index('\n    def ', 10)]
    assert "self.profile.get('history')" in history
    assert "entry.get('date')" in history
    assert "entry.get('phase')" in history
    assert "entry.get('subject')" in history


def test_archive_record_window_lays_survey_in_three_columns_beside_score_and_history():
    """کادر ۲ (نظرسنجی) is split into criteria_page's own three chunks, laid
    out side by side rather than one long list; ۳ (محاسبهٔ امتیاز) sits in
    its own column beside those three. ۴ (تاریخ‌های همکاری) got its own
    full-width box below all of that instead once the operator's own ask
    (more room than a narrow side column left it) moved it there - see
    test_archive_record_window_lists_collaboration_dates_in_a_full_width_box_below_the_survey.
    The whole record window is scrollable, for whichever دسته‌بندی has
    enough معیار to overflow it."""

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    init_body = source[source.index('    def __init__(self, parent, *, profile: dict,'):]
    init_body = init_body[:init_body.index('\n    def ', 10)]
    assert 'scrollable=True' in init_body

    build = source[source.index('    def _build(self) -> None:', source.index('class ArchiveRecordWindow187')):]
    build = build[:build.index('\n    def ', 10)]
    assert 'for column in range(CRITERIA_PAGE_COUNT):' in build
    assert 'criteria_page(category, column)' in build
    assert "side.grid(row=0, column=0" in build

    with _temp_store():
        import tkinter as tk
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from contractor_archive_stage18_7 import ArchiveRecordWindow187

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            win = ArchiveRecordWindow187(root, profile={'family_name': 'تست', 'category': 'پیمانکار'})
            root.update()
            # Every criterion still got its own score/weight StringVar, split
            # across three columns instead of dropped.
            import contractor_archive_stage18_7 as ca
            for key, _label in ca.CONTRACTOR_CRITERIA:
                assert f'{key}_score' in win.values
                assert f'{key}_weight' in win.values
        finally:
            root.destroy()


def test_archive_record_window_lists_collaboration_dates_in_a_full_width_box_below_the_survey():
    """کادر ۴ used to sit squeezed into the narrow side column next to
    محاسبهٔ امتیاز (a third of the window's width) - the operator's own ask
    was more room than that left it, so it now gets its own full-width box
    below all three نظرسنجی columns instead."""

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    build = source[source.index('    def _build(self) -> None:', source.index('class ArchiveRecordWindow187')):]
    build = build[:build.index('\n    def ', 10)]
    assert "history_heading = tk.Frame(self.body, bg=p['bg'])" in build
    assert "history = ttk.Frame(self.body, style='Card.TFrame')" in build
    assert 'self._render_history(history)' in build


def test_archive_record_window_identity_fields_are_one_shared_two_by_three_grid():
    """مشخصات کلی used to be four separate _fields() calls (۳+۱+۱+۱ فیلد),
    each its own grid frame - ۵ ردیف پیاپی که سه‌تای آخرشان فقط نیمی از عرض
    را پر می‌کردند. یک گرید مشترک ۲×۳ با هر شش فیلد به‌جای آن."""

    import contractor_archive_stage18_7 as ca

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    build = source[source.index('    def _build(self) -> None:', source.index('class ArchiveRecordWindow187')):]
    build = build[:build.index('\n    def ', 10)]
    assert build.count('self._fields(') == 0
    assert 'self._render_identity_fields(identity)' in build

    fields_method = source[source.index('    def _render_identity_fields(self, group)'):]
    fields_method = fields_method[:fields_method.index('\n    def ', 10)]
    for group_name in ('IDENTITY_FIELDS', 'CATEGORY_FIELDS', 'ACTIVITY_FIELDS', 'PAYMENT_TYPE_FIELDS'):
        assert group_name in fields_method
    assert fields_method.count("grid.columnconfigure((0, 1)") == 1

    with _temp_store():
        import tkinter as tk
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from contractor_archive_stage18_7 import ArchiveRecordWindow187

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            win = ArchiveRecordWindow187(root, profile={'family_name': 'تست', 'category': 'پیمانکار'})
            root.update()
            for key, _label in (*ca.IDENTITY_FIELDS, *ca.CATEGORY_FIELDS, *ca.ACTIVITY_FIELDS,
                                *ca.PAYMENT_TYPE_FIELDS):
                assert key in win.values
        finally:
            root.destroy()


def test_archive_criteria_rows_pin_score_and_weight_to_a_fixed_column():
    """A bare pack(side='right') put امتیاز/ضریب right after however much
    room each row's OWN معیار text needed - رفع نواقص کار (کوتاه) و انجام
    کارهای خارج از قرارداد (بلند) در همان کادر, کادرهای عددی‌شان در دو x
    مختلف. یک عرض ثابت از بلندترین معیار همان صفحه/ستون آن‌ها را در یک خط
    عمودی می‌چیند - هم در ویزارد, هم در پنجرهٔ رکورد."""

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    wizard_method = source[source.index('    def _render_criteria_page(self, parent, page_index: int) -> None:'):]
    wizard_method = wizard_method[:wizard_method.index('\n    def ', 10)]
    assert 'label_width = max(len(label) for _key, label in chunk) + 1' in wizard_method
    assert 'width=label_width' in wizard_method

    record_method = source[source.index('    def _render_criteria(self, parent, criteria) -> None:'):]
    record_method = record_method[:record_method.index('\n    def ', 10)]
    assert 'label_width = max(len(label) for _key, label in criteria) + 1' in record_method
    assert 'width=label_width' in record_method


def test_archive_overall_score_rounds_to_one_decimal_place_not_a_whole_number():
    """امتیاز کلی با یک رقم اعشار حساب شود - the operator's own ask. Two
    معیار with weight ۱ and ۲ at امتیاز ۸ and ۹ average to 86.667 از ۱۰۰:
    rounding to a whole number (۸۷) and to one decimal (۸۶.۷) disagree,
    so this scenario actually exercises the fix instead of coincidentally
    passing either way."""

    import contractor_archive_stage18_7 as ca

    values = {'quality_of_work_score': '8', 'quality_of_work_weight': '1',
              'work_speed_score': '9', 'work_speed_weight': '2'}
    overall = ca.compute_overall_score(values, 'پیمانکار')
    assert round(overall, 1) == 86.7
    assert round(overall, 1) != round(overall)


def test_archive_wizard_opens_the_record_window_after_افزودن_and_close_with_no_confirmation():
    """افزودن → نام‌گذاری → بستن (بدون تأیید نهایی‌ای که این ویزارد اصلاً
    ندارد - محاسبهٔ امتیاز خودش آخرین قدم پیش از رکورد است) باید رکورد تازه
    را به‌عنوان پیش‌نویس نگه دارد و پنجرهٔ رکورد را باز کند - همان قاعدهٔ
    ContractWizard187, دست‌نخورده."""

    with _temp_store():
        import tkinter as tk
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from contractors_stage18_7 import ContractorsPanel187, ARCHIVE
        import contractor_archive_stage18_7 as ca

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            panel = ContractorsPanel187(root, db=None, base_dir=None)
            root.update()
            panel.open_wizard(ARCHIVE)
            root.update()
            wiz = next(c for c in panel.winfo_children() if isinstance(c, ca.ArchiveWizard187))
            wiz._var('family_name').set('تست رکورد')
            root.update()
            wiz.close()
            root.update()
            record = next((c for c in panel.winfo_children()
                          if isinstance(c, ca.ArchiveRecordWindow187)), None)
            assert record is not None, 'پنجره رکورد باید بعد از بستن ویزارد باز شود'
            assert 'تست رکورد' in record.title()
        finally:
            root.destroy()


def test_archive_detail_window_export_sits_left_and_identity_is_one_horizontal_row():
    """گزینهٔ خروجی به سمت چپ - همان تقاضای StatementFinanceDetailWindow187.
    اطلاعات کلی هم دیگر یک رشتهٔ سه‌ردیفهٔ \\n-دار نیست - شش ویجت Label
    جداگانه، پهلوی هم روی یک ردیف افقی، هم برای نمایش افقی و هم چون یک
    رشتهٔ ترکیبی چندبخشی گاهی زیر Tk (بدون الگوریتم bidi واقعی) به‌هم
    می‌ریزد."""

    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    build = source[source.index('    def _build(self) -> None:', source.index('class ArchiveDetailWindow187')):]
    build = build[:build.index('\n    def ', 10)]
    assert "combo.pack(side='left')" in build
    assert 'for segment in self._identity_segments():' in build
    assert 'identity = (' not in build


def test_archive_detail_window_export_includes_the_survey_table_in_two_columns():
    """نظرسنجی table used to be missing from _export_profile entirely - only
    نام/دسته‌بندی/عنوان فعالیت/امتیاز/احتمال went out. Now the same
    identity+score+criteria(دو ستون)+history data the detail window itself
    shows goes into a dedicated export instead of a flat five-column row."""

    source = _source(ROOT / 'contractors_stage18_7.py')
    export_method = source[source.index("    def _export_profile(self, profile, fmt: str, addressee='سایر') -> None:"):]
    export_method = export_method[:export_method.index('\n    def ', 10)]
    assert 'export_archive_profile_pdf_187(' in export_method
    assert 'archive_criteria_export_rows(profile)' in export_method
    assert 'archive_history_export_rows(profile)' in export_method
    assert 'self._export_profile_xlsx(profile, destination, addressee)' in export_method

    import contractor_archive_stage18_7 as ca

    profile = {'category': 'پیمانکار'}
    for key, _label in ca.CONTRACTOR_CRITERIA:
        profile[f'{key}_score'] = '8'
        profile[f'{key}_weight'] = '1'
    headers, rows = ca.archive_criteria_export_rows(profile)
    assert headers == ['معیار', 'امتیاز (از ۱۰)', 'ضریب', 'معیار', 'امتیاز (از ۱۰)', 'ضریب']
    # ۱۷ معیار در دو ستون یعنی ۹ ردیف, نه ۱۷ ردیفِ یک‌ستونه.
    assert len(rows) == (len(ca.CONTRACTOR_CRITERIA) + 1) // 2
    assert rows[0][0] == ca.CONTRACTOR_CRITERIA[0][1]
    half = (len(ca.CONTRACTOR_CRITERIA) + 1) // 2
    assert rows[0][3] == ca.CONTRACTOR_CRITERIA[half][1]


# ---------------------------------------------------------------------------
# ContractorsPanel187 wiring: all three tabs share قراردادها's own actions
# ---------------------------------------------------------------------------

def test_the_three_new_tabs_read_from_and_write_to_their_own_store():
    source = _source(ROOT / 'contractors_stage18_7.py')
    assert 'StatementRecordWindow187' in source and 'StatementWizard187' in source
    assert 'from contractor_statements_stage18_7 import (' in source
    assert 'from contractor_archive_stage18_7 import (' in source
    assert 'ArchiveDetailWindow187, ArchiveRecordWindow187, ArchiveWizard187,' in source

    load_rows = source[source.index('    def load_rows(self, sub_tab: str)'):]
    load_rows = load_rows[:load_rows.index('\n    def ', 10)]
    assert 'self._load_statement_rows()' in load_rows
    assert 'self._load_finance_rows()' in load_rows
    assert 'self._load_archive_rows()' in load_rows

    build_page = source[source.index('    def _build_page(self, page: ttk.Frame, label: str)'):]
    build_page = build_page[:build_page.index('\n    def ', 10)]
    assert 'صورت‌وضعیت جدید +' in build_page
    # آرشیو and امور مالی are both افزودن/ویرایش/حذف/کپی toolbars now.
    assert "table.add_tool('افزودن +', lambda: self.open_wizard(ARCHIVE))" in build_page
    assert "'افزودن +', self._add_finance_payment" in build_page
    assert "'ویرایش', self._edit_finance_payment" in build_page
    assert "'حذف ردیف', self._delete_finance_payment" in build_page
    assert "'کپی', self._copy_finance_payment" in build_page


def test_finance_tab_reports_a_contracts_recorded_statement_amount():
    """صورت‌وضعیت مبلغ in امور مالی is what STATEMENTS has actually
    recorded for that قرارداد, not a separate figure someone has to keep
    in step by hand - summed directly from statement_amount() over every
    صورت‌وضعیت on the contract (the same total cumulative_statement_amount()
    would give with its default upto_statement_id=None), computed from one
    load_statements() call instead of one re-read per قرارداد.
    """

    source = _source(ROOT / 'contractors_stage18_7.py')
    finance_rows = source[source.index('    def _load_finance_rows(self)'):]
    finance_rows = finance_rows[:finance_rows.index('\n    def ', 10)]
    assert 'statement_amount(statement)' in finance_rows
    assert 'paid_total(contract)' in finance_rows
    assert 'balance(contract)' in finance_rows

    with _temp_store():
        from ui_stage18_3 import fa
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({'family_name': 'الف', 'total_amount': '1000'})
        store.upsert_statement({'contract_id': contract['contract_id'], 'rows': [
            {'line_total': '100'}]})
        store.upsert_statement({'contract_id': contract['contract_id'], 'rows': [
            {'line_total': '50'}]})

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.finance_filter_field = None
        rows, _status = panel._load_finance_rows()
        assert len(rows) == 0  # no واریز recorded yet - the row list is per-payment
        expected_total = store.cumulative_statement_amount(contract['contract_id'])
        assert expected_total == 150

        store.upsert_contract({**contract, 'payments': [{'date': '1404/01/01', 'amount': '10'}]})
        rows, _status = panel._load_finance_rows()
        assert len(rows) == 1
        assert rows[0][6] == fa(str(expected_total))  # مبلغ صورت‌وضعیت تجمیعی column


# ---------------------------------------------------------------------------
# صورت‌وضعیت tab restructuring - toolbar, filters, columns and its new
# نوع صورت‌وضعیت field, all reported together against a real screenshot.
# ---------------------------------------------------------------------------

def test_tabs_use_their_own_colour_roles_not_the_old_reverted_names():
    """TabSelected.TButton (green) / TabUnselected.TButton (grey) are the
    tab strip's own roles - never the old, since-reverted TabActive.TButton/
    TabInactive.TButton pair from an earlier rounded-rectangle experiment.
    """

    tabs_source = _source(ROOT / 'barkav_ui' / 'tabs.py')
    assert "style='TabUnselected.TButton'" in tabs_source
    assert "'TabSelected.TButton' if active else 'TabUnselected.TButton'" in tabs_source
    assert 'TabActive.TButton' not in tabs_source
    assert 'TabInactive.TButton' not in tabs_source

    skin_source = _source(ROOT / 'barkav_ui' / 'skin.py')
    assert 'TAB_ROLES' not in skin_source
    role_block = skin_source[skin_source.index('BUTTON_ROLES = {'):skin_source.index('\n}\n')]
    assert "'TabSelected.TButton':" in role_block and "('tab_indicator'," in role_block
    assert "'TabUnselected.TButton':" in role_block and "('tab_indicator_inactive'," in role_block


def test_tabs_are_rounded_rectangles_everything_else_stays_pill():
    """Tabs (green/selected and grey/unselected) render as a small-radius
    rounded rectangle now, not the full stadium every other button keeps -
    a shape choice scoped to just those two roles via radius_override, with
    no change to a tab's own height or width-stretch behaviour.
    """

    skin_source = _source(ROOT / 'barkav_ui' / 'skin.py')
    from barkav_ui.skin import ROUNDED_RECT_ROLES
    assert ROUNDED_RECT_ROLES == {
        'TabSelected.TButton', 'TabUnselected.TButton',
        'WindowTabSelected.TButton', 'WindowTabUnselected.TButton',
    }
    assert 'radius_override=radius if name in ROUNDED_RECT_ROLES else None' in skin_source

    element_body = skin_source[skin_source.index('    def element('):skin_source.index('\n    # --- buttons')]
    assert 'radius_override' in element_body
    # height/edges still come from pill_height regardless of radius_override -
    # only rounding (the corner radius) may differ.
    assert 'height = pill_height if pill else None' in element_body

    import tkinter as tk
    from tkinter import ttk
    from barkav_ui.skin import install_rounded_skin
    from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

    root = tk.Tk()
    root.withdraw()
    try:
        style = ttk.Style(root)
        style.theme_use('clam')  # same image-capable base as the active theme engine
        install_rounded_skin(style, kit_palette(), kit_metrics(), root)
        tab = ttk.Button(root, text='تب', style='TabSelected.TButton')
        regular = ttk.Button(root, text='دکمه', style='Primary.TButton')
        tab.pack()
        regular.pack()
        root.update()
        # Same height and comparable reqwidth-per-character - the shape
        # changed, the box did not.
        assert tab.winfo_reqheight() >= regular.winfo_reqheight()  # font metrics can set both minima
        assert style.layout('TabSelected.TButton')[0][0].startswith('TabSelectedTButton.plate')
        assert style.layout('Primary.TButton')[0][0].startswith('PrimaryTButton.plate')
    finally:
        root.destroy()


def test_table_toolbar_button_is_its_own_role_not_the_shared_toolbar_button():
    """DataTable's own toolbar (افزودن/ویرایش/حذف/کپی, نمایش‌براساس, خروجی)
    gets toolbar_fill - a colour between bg and an inactive tab's grey - via
    its own role and its own ttk styles, so recolouring it never touches
    Toolbar.TButton (the brand bar's Save/Undo/Redo and dozens of other
    screens' small utility buttons) or Ghost.TButton (every window's own
    footer buttons).
    """

    table_source = _source(ROOT / 'barkav_ui' / 'table.py')
    assert "style='TableToolbar.TButton'" in table_source
    assert "style='Toolbar.TCombobox'" in table_source
    assert "style='Toolbar.TEntry'" in table_source

    skin_source = _source(ROOT / 'barkav_ui' / 'skin.py')
    role_block = skin_source[skin_source.index('BUTTON_ROLES = {'):skin_source.index('\n}\n')]
    assert "'TableToolbar.TButton':" in role_block and "('toolbar_fill'," in role_block
    # The generic Ghost role is untouched; Toolbar.TButton's fill moved from
    # 'glass' to 'tab_inactive' (same near-invisible-on-'panel' bug
    # ProjectSummary.TButton was split off to fix, applied at the shared
    # role instead of one more one-off split).
    assert "'Toolbar.TButton':  ('tab_inactive'," in role_block
    assert "'Ghost.TButton':    ('panel'," in role_block

def test_statements_columns_are_the_nine_requested_fields_in_order():
    import contractors_stage18_7 as cst
    assert cst.COLUMNS[cst.STATEMENTS] == (
        ('row', 'ردیف'), ('name', 'پیمانکار'), ('title', 'عنوان قرارداد'),
        ('contract_date', 'تاریخ قرارداد'), ('number', 'شماره صورت‌وضعیت'),
        ('type', 'نوع صورت‌وضعیت'), ('amount', 'مبلغ صورت‌وضعیت (تومان)'),
        ('paid', 'مبلغ پرداختی (تومان)'), ('balance', 'بدهکار / بستانکار (تومان)'),
    )


def test_statements_toolbar_has_کپی_before_حذف_a_نمایش_براساس_filter_and_export_but_no_search():
    """نمایش براساس is one cascading filter here too - همه/نام پیمانکار/
    نوع صورت‌وضعیت (موقت-قطعی), not a free-text search box.
    """

    source = _source(ROOT / 'contractors_stage18_7.py')
    build_page = source[source.index('    def _build_page(self, page: ttk.Frame, label: str)'):]
    build_page = build_page[:build_page.index('\n    def ', 10)]
    statements_branch = build_page[build_page.index("if label == STATEMENTS:"):]
    statements_branch = statements_branch[:statements_branch.index('return table') + len('return table')]

    # One order in every tab (operator round 2026-09-27): افزودن, ویرایش,
    # کپی, حذف - so کپی comes before حذف ردیف in add_tool's right-to-left
    # packing, as in قراردادها and ثبت دیتا.
    assert statements_branch.index("add_tool('کپی'") < statements_branch.index("add_tool('حذف ردیف'")
    assert "'نمایش براساس'" in statements_branch
    assert 'statement_filter_field' in statements_branch
    assert 'statement_filter_value' in statements_branch
    assert 'statement_export_choice' in statements_branch
    assert 'add_search' not in statements_branch
    # The same four toolbar actions, on the row itself.
    assert "table.tree.bind('<Button-3>', self._statement_context_menu" in statements_branch

    import contractors_stage18_7 as cst
    assert cst.STATEMENT_FILTER_FIELDS == (cst.ANY, 'نام پیمانکار', cst.STATEMENT_TYPE_FILTER_FIELD)


def test_statements_context_menu_mirrors_its_own_toolbar():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت'})
        store.upsert_statement({'contract_id': contract['contract_id']})

        import tkinter as tk
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

        root = tk.Tk()
        # Not withdrawn: an unmapped Treeview reports an empty bbox for
        # every item, and this test needs a real one to simulate the click.
        root.geometry('900x400')
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            panel = cst.ContractorsPanel187(root, db=None, rich=None)
            panel.pack(fill='both', expand=True)
            panel.nb.select(1)
            root.update()

            table = panel.tables[cst.STATEMENTS]
            first_iid = table.tree.get_children()[0]

            captured = {}
            original_popup = tk.Menu.tk_popup
            tk.Menu.tk_popup = lambda self, *a, **k: captured.setdefault('labels', [
                self.entrycget(i, 'label') for i in range(self.index('end') + 1)
                if self.type(i) != 'separator'])
            try:
                class FakeEvent:
                    pass
                event = FakeEvent()
                event.x, event.y = 5, 5
                event.x_root = event.y_root = 0
                bbox = table.tree.bbox(first_iid)
                event.y = bbox[1] + 2
                panel._statement_context_menu(event)
            finally:
                tk.Menu.tk_popup = original_popup

            assert captured['labels'] == ['صورت‌وضعیت جدید', 'ویرایش', 'کپی', 'حذف ردیف']
            assert table.tree.selection() == (first_iid,)
        finally:
            root.destroy()


def test_statements_filter_offers_statement_type_and_narrows_by_it():
    """صورت‌وضعیت's own نمایش براساس was missing نوع صورت‌وضعیت (موقت/قطعی) -
    unlike قراردادها, a صورت‌وضعیت row already IS one specific صورت‌وضعیت, so
    this reads the field directly rather than through its own قرارداد."""

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        assert cst.STATEMENT_TYPE_FILTER_FIELD in cst.STATEMENT_FILTER_FIELDS

        c1 = store.upsert_contract({'family_name': 'ابراهیم ازبک'})
        c2 = store.upsert_contract({'family_name': 'محمود کریمی'})
        store.upsert_statement({'contract_id': c1['contract_id'], 'statement_type': 'موقت'})
        store.upsert_statement({'contract_id': c2['contract_id'], 'statement_type': 'قطعی'})

        import tkinter as tk
        # Filtering needs Tcl variables, not a graphical Tk window.
        root = tk.Tcl()
        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.statement_filter_field = tk.StringVar(root, value=cst.STATEMENT_TYPE_FILTER_FIELD)
        panel.statement_filter_value = tk.StringVar(root, value=cst.ANY)

        rows, _status = panel._load_statement_rows()
        assert {r[1] for r in rows} == {'ابراهیم ازبک', 'محمود کریمی'}

        panel.statement_filter_value.set('موقت')
        rows, _status = panel._load_statement_rows()
        assert [r[1] for r in rows] == ['ابراهیم ازبک']

        panel.statement_filter_value.set('قطعی')
        rows, _status = panel._load_statement_rows()
        assert [r[1] for r in rows] == ['محمود کریمی']


def test_statements_statement_type_filter_value_choices_are_the_fixed_two_types():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({'family_name': 'ابراهیم ازبک'})
        store.upsert_statement({'contract_id': contract['contract_id']})

        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)

            class FakeCombo:
                def __init__(self):
                    self.values = None

                def configure(self, *, values):
                    self.values = values

            panel.statement_filter_field = tk.StringVar(root, value=cst.STATEMENT_TYPE_FILTER_FIELD)
            panel.statement_filter_value = tk.StringVar(root, value=cst.ANY)
            panel.statement_filter_value_combo = FakeCombo()
            panel.refresh = lambda: None
            panel._statement_filter_field_changed()
            assert panel.statement_filter_value_combo.values == (cst.ANY, 'موقت', 'قطعی')
            assert panel.statement_filter_value.get() == cst.ANY
        finally:
            root.destroy()


def test_copying_a_statement_drops_its_own_identity_but_keeps_everything_else():
    with _temp_store():
        import contractor_store_stage18_7 as store
        contract = store.upsert_contract({'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک'})
        original = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_number': 'SW-1',
            'statement_type': 'قطعی', 'period_from': '1405/01/01', 'period_to': '1405/02/01',
            'rows': [{'description': 'x', 'line_total': '1000'}],
        })

        source = _source(ROOT / 'contractors_stage18_7.py')
        method = source[source.index('    def copy_selected_statement(self)'):]
        method = method[:method.index('\n    def ', 10)]
        assert "clone.pop('statement_id', None)" in method
        assert "clone.pop('statement_number', None)" in method

        clone = dict(original)
        clone.pop('statement_id', None)
        clone.pop('statement_number', None)
        assert clone['contract_id'] == contract['contract_id']
        assert clone['statement_type'] == 'قطعی'
        assert clone['rows'] == original['rows']


# ---------------------------------------------------------------------------
# تاریخ قرارداد (new contract field) and نوع صورت‌وضعیت (new statement
# field, default موقت) - both reported in the same screenshot batch.
# ---------------------------------------------------------------------------

def test_contract_date_sits_beside_its_own_number_in_wizard_and_record():
    source = _source(ROOT / 'contractor_contracts_stage18_7.py')
    documents = source[source.index('    def _render_documents(self, parent)'):]
    documents = documents[:documents.index('\n    def ', 10)]
    assert "self._var('contract_date')" in documents
    assert "text='تاریخ قرارداد'" in documents

    page_documents = source[source.index('    def _page_documents(self, page)'):]
    page_documents = page_documents[:page_documents.index('\n    def ', 10)]
    assert "('contract_date', 'تاریخ قرارداد')" in page_documents


def test_statement_record_window_is_scrollable_so_its_footer_always_shows():
    """A صورت‌وضعیت with enough ردیف‌ها made this window request ~880px of
    content on a 768px-tall screen; the footer is packed before the body so
    it never gets pushed off, but with no scroll region it was still
    squeezed to a sliver a couple pixels tall - confirmed live at 1366x768.
    scrollable=True keeps ذخیره/بستن full-height regardless of row count or
    screen size, by letting the rows scroll instead of forcing the window
    to fit them all at once."""

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    init_body = source[source.index("class StatementRecordWindow187"):]
    init_body = init_body[init_body.index('def __init__'):]
    init_body = init_body[:init_body.index('\n    def ', 10)]
    assert 'scrollable=True' in init_body


def test_statement_type_defaults_to_موقت_in_wizard_first_step_and_record_first_box(tmp_path):
    import tkinter as tk
    os.environ['STAGE4_DB_PATH'] = str(tmp_path / 'x.sqlite3')
    import contractor_statements_stage18_7 as cs
    import contractor_store_stage18_7 as store

    assert cs.DEFAULT_STATEMENT_TYPE == 'موقت'
    assert cs.STATEMENT_TYPES == ('موقت', 'قطعی')

    contract = store.upsert_contract({'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک'})
    root = tk.Tk()
    root.withdraw()
    try:
        wiz = cs.StatementWizard187(root, rich=None)
        try:
            assert wiz._var('statement_type', cs.DEFAULT_STATEMENT_TYPE).get() == 'موقت'
        finally:
            wiz.destroy()

        statement = store.upsert_statement({'contract_id': contract['contract_id'], 'statement_number': 'SW-9'})
        rec = cs.StatementRecordWindow187(root, statement=statement, rich=None)
        try:
            assert rec.values['statement_type'].get() == 'موقت'
        finally:
            rec.destroy()
    finally:
        root.destroy()


# ---------------------------------------------------------------------------
# The same alternating-green row colour, extended from صورت‌وضعیت to
# قراردادها, امور مالی and آرشیو - none of the four should still tint by
# their own real status/بدهکاری once one of them was reported to.
# ---------------------------------------------------------------------------

def test_row_colour_alternates_everywhere_except_a_finance_row_gone_into_debt():
    with _temp_store():
        import contractors_stage18_7 as cst

        # Rows that would have read differently under the old owed-based
        # rule (paid off, still owed, overpaid) must still just alternate -
        # except امور مالی's own بدهکار/بستانکار column, which now marks a
        # genuinely negative one (پیمانکار بدهکار شده) red instead of
        # whatever the plain alternating stripe would have been. هر قرارداد
        # گیرد یک صورت‌وضعیت هم‌ارزش total_amount خودش، تا owed دقیقاً همان
        # سه حالتِ نامِ خودش را بسنجد: صفر، مثبت، منفی.
        # امور مالی is a per-واریز log now, so each contract needs exactly
        # one deposit to keep the row count matching one row per contract.
        import contractor_store_stage18_7 as store
        c1 = store.upsert_contract({'family_name': 'الف', 'total_amount': '1000',
                                    'payments': [{'amount': '1000', 'date': '1404/01/01'}]})  # owed == 0
        store.upsert_statement({'contract_id': c1['contract_id'],
                                'rows': [{'description': 'x', 'line_total': '1000'}]})
        c2 = store.upsert_contract({'family_name': 'ب', 'total_amount': '1000',
                                    'payments': [{'amount': '500', 'date': '1404/01/01'}]})  # owed > 0
        store.upsert_statement({'contract_id': c2['contract_id'],
                                'rows': [{'description': 'x', 'line_total': '1000'}]})
        c3 = store.upsert_contract({'family_name': 'ج', 'total_amount': '1000',
                                    'payments': [{'amount': '2000', 'date': '1404/01/01'}]})  # owed < 0
        store.upsert_statement({'contract_id': c3['contract_id'],
                                'rows': [{'description': 'x', 'line_total': '1000'}]})
        store.upsert_profile({'family_name': 'د', 'activity': 'اجرا'})
        store.upsert_profile({'family_name': 'ه', 'activity': 'اجرا'})

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.filters = {cst.CONTRACTS: type('V', (), {'get': lambda self: ''})(),
                         cst.FINANCE: type('V', (), {'get': lambda self: ''})(),
                         cst.ARCHIVE: type('V', (), {'get': lambda self: ''})()}
        panel.statement_filter_field = None
        panel.finance_filter_field = None

        _rows, status = cst.ContractorsPanel187.load_rows(panel, cst.CONTRACTS)
        assert status == {}, 'no green tint: alternate rows use the shared stripe'

        _rows, status = panel._load_finance_rows()
        assert status == {2: 'debt'}

        _rows, status = panel._load_archive_rows()
        # آرشیو auto-syncs از قراردادها هم - الف/ب/ج هرکدام یک ردیف پیمانکار
        # تازه می‌سازند چون هنوز به این نام پروفایلی نبود، در کنار د/ه خودشان.
        assert status == {}


# ---------------------------------------------------------------------------
# صورت‌وضعیت's first window: شماره صورت‌وضعیت before نوع صورت‌وضعیت, real
# Shamsi calendars for the three dates, Persian digits everywhere, and the
# same font as the rest of the wizard - all reported against one screenshot.
# ---------------------------------------------------------------------------

def test_statement_number_comes_before_نوع_صورت‌وضعیت_before_the_three_dates():
    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    helper = source[source.index('def _render_statement_identity_fields('):]
    helper = _until_next_def_or_class(helper)
    assert helper.index("field_row(*NUMBER_FIELD)") < helper.index("'statement_type'")
    assert helper.index("'statement_type'") < helper.index('for key, label in DATE_FIELDS:')


def test_the_three_period_dates_use_the_apps_own_shamsi_calendar():
    import contractor_statements_stage18_7 as cs
    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    assert 'from ui_stage10 import DateField' in source
    helper = source[source.index('def _render_statement_identity_fields('):]
    helper = _until_next_def_or_class(helper)
    assert 'DateField(holder, variable' in helper
    assert cs.DATE_FIELDS == (
        ('statement_date', 'تاریخ صورت‌وضعیت'), ('period_from', 'دوره از'), ('period_to', 'دوره تا'),
    )


def test_jalali_calendar_lays_out_right_to_left_and_picks_persian_digit_dates():
    """شنبه (WEEKDAYS[0]) belongs at the grid's right edge with هفته
    advancing leftward - the calendar used to lay itself out left-to-right
    like a Gregorian one instead. Also: _select() used to hand back a plain
    ASCII-digit date string even though the calendar's own day buttons and
    title already showed Persian digits, so a تاریخ picked through it (as
    opposed to typed by hand) recorded in the wrong script."""

    import tkinter as tk
    from ui_stage10 import JalaliCalendar

    root = tk.Tk()
    root.withdraw()
    try:
        cal = JalaliCalendar(root, '1404/06/15')
        try:
            root.update()
            # شنبه (index 0) at column 6 (rightmost); جمعه (index 6) at column 0.
            header_columns = {}
            for child in cal.grid.winfo_children():
                if isinstance(child, __import__('tkinter').ttk.Label):
                    info = child.grid_info()
                    if int(info['row']) == 0:
                        header_columns[child.cget('text')] = int(info['column'])
            assert header_columns[JalaliCalendar.WEEKDAYS[0]] == 6
            assert header_columns[JalaliCalendar.WEEKDAYS[6]] == 0

            captured = {}
            cal.on_select = lambda value: captured.setdefault('value', value)
            cal._select(cal.day or 1)
            assert captured['value']
            assert all(ch in '۰۱۲۳۴۵۶۷۸۹/' for ch in captured['value']), captured['value']
        finally:
            if cal.winfo_exists():
                cal.destroy()
    finally:
        root.destroy()


def test_a_write_trace_is_attached_after_its_entry_exists_not_before():
    """A write trace registered on a bare StringVar fires ahead of the Entry
    widget's own sync trace; by the time a nested re-set (Persian digits, or
    grouping) settles, the Entry has already read a stale value and shows
    what was typed, not what the variable actually holds. Reproduced and
    fixed empirically - this pins every call site that used to get the
    order backwards, so it cannot regress by one of them being missed.
    """

    source = _source(ROOT / 'contractor_statements_stage18_7.py')

    identity = source[source.index('def _render_statement_identity_fields('):]
    identity = _until_next_def_or_class(identity)
    assert identity.index("ttk.Entry(holder, textvariable=variable") < identity.index('_persian_digits_while_typing(variable)')

    grid = source[source.index('def render_statement_row_grid('):]
    grid = grid[:grid.index('\ndef ', 10)]
    assert grid.index('ttk.Entry(grid, textvariable=row[key]') < grid.index('group_amount_while_typing_187(row[key])')

    wizard_bonus = source[source.index('    def _render_bonus(self, parent)'):]
    wizard_bonus = wizard_bonus[:wizard_bonus.index('\n    def ', 10)]
    assert wizard_bonus.index("ttk.Entry(row, textvariable=variable") < wizard_bonus.index('group_amount_while_typing_187(variable)')

    record_build = source[source.index('    def _build(self) -> None:'):]
    record_build = record_build[:record_build.index('\n    def ', 10)]
    bonus_section = record_build[record_build.index('bonus = ttk.Frame'):]
    assert bonus_section.index("ttk.Entry(row, textvariable=variable") < bonus_section.index('group_amount_while_typing_187(variable)')

    # blank_statement_row must no longer attach grouping itself - that moved
    # to render_statement_row_grid, after each cell's own Entry exists.
    blank_row = source[source.index('def blank_statement_row('):]
    blank_row = blank_row[:blank_row.index('\ndef ', 10)]
    assert 'group_amount_while_typing_187' not in blank_row


def test_the_trace_reordering_actually_fixes_the_display_live(tmp_path):
    """Source order alone does not prove the widget shows the right thing -
    build the real wizard, type a value the same way a click plus keystroke
    would, and read the Entry's own .get(), not just the StringVar's."""

    import tkinter as tk
    from tkinter import ttk
    os.environ['STAGE4_DB_PATH'] = str(tmp_path / 'x.sqlite3')
    import contractor_statements_stage18_7 as cs
    import contractor_store_stage18_7 as store

    store.upsert_contract({'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک',
                           'rows': [{'description': 'اجرت اجرای اسکلت', 'unit_price': '750000',
                                    'quantity': '1914'}]})

    def find_entries(w):
        found = []
        for child in w.winfo_children():
            if isinstance(child, ttk.Entry):
                found.append(child)
            found.extend(find_entries(child))
        return found

    root = tk.Tk()
    root.withdraw()
    try:
        wiz = cs.StatementWizard187(root, rich=None)
        try:
            number_var = wiz._var('statement_number')
            number_var.set('12')
            root.update()
            entry = next(e for e in find_entries(wiz) if e.cget('textvariable') == number_var._name)
            assert entry.get() == number_var.get() == '۱۲'

            wiz._var('contract_id').set(next(c for c in cs.contract_choices() if 'اجرت اجرای اسکلت' in c))
            root.update()
            wiz.show_step(1)
            root.update()
            row = wiz.row_values[0]
            assert row['unit_price'].get() == '۷۵۰,۰۰۰'
            assert row['contract_quantity'].get() == '۱,۹۱۴'
            unit_entry = next(e for e in find_entries(wiz) if e.cget('textvariable') == row['unit_price']._name)
            assert unit_entry.get() == row['unit_price'].get()

            row['worked_quantity'].set('500')
            root.update()
            assert row['line_total'].get() == '۳۷۵,۰۰۰,۰۰۰'
        finally:
            wiz.destroy()
    finally:
        root.destroy()


# ---------------------------------------------------------------------------
# امور مالی: redesigned from a per-قرارداد rollup into a per-واریز log, with
# its own افزودن/ویرایش/حذف/کپی routed through the shared ثبت دیتا wizard.
# ---------------------------------------------------------------------------

def test_finance_columns_are_the_ten_requested_fields_in_order():
    import contractors_stage18_7 as cst

    assert [label for _key, label in cst.COLUMNS[cst.FINANCE]] == [
        'ردیف', 'تاریخ واریز', 'عنوان قرارداد', 'نام پیمانکار',
        'مبلغ واریز (تومان)', 'مبلغ تجمعی واریز (تومان)',
        'مبلغ صورت‌وضعیت تجمیعی (تومان)', 'بدهکار / بستانکار (تومان)',
        'مبلغ کل قرارداد (تومان)', 'مبلغ باقیمانده قرارداد (تومان)',
    ]


def test_finance_is_checkable_like_statements():
    source = _source(ROOT / 'contractors_stage18_7.py')
    assert 'checkable=label in (CONTRACTS, STATEMENTS, FINANCE, ARCHIVE)' in source


def test_finance_filter_offers_تاریخ_alongside_the_contract_fields():
    import contractors_stage18_7 as cst

    assert cst.FINANCE_FILTER_FIELDS == (cst.ANY, 'نام پیمانکار', 'عنوان قرارداد', 'تاریخ')


def test_finance_rows_are_per_payment_grouped_by_contract_oldest_first():
    """Each واریز is its own row, in date order within its own قرارداد, with
    a per-contract running مبلغ تجمعی واریز - not one rollup row per contract.
    """

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({
            'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک', 'total_amount': '2000000000',
            'payments': [
                {'date': '1404/02/01', 'amount': '500000000', 'pending_id': 'PEN1'},
                {'date': '1404/01/01', 'amount': '300000000', 'pending_id': 'PEN2'},
            ],
        })

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.finance_filter_field = None
        rows, _status = panel._load_finance_rows()

        assert len(rows) == 2
        # Oldest واریز first, even though it was appended to the JSON second.
        assert rows[0][1] == '۱۴۰۴/۰۱/۰۱' and rows[0][4] == '۳۰۰,۰۰۰,۰۰۰'
        assert rows[1][1] == '۱۴۰۴/۰۲/۰۱' and rows[1][4] == '۵۰۰,۰۰۰,۰۰۰'
        # مبلغ تجمعی واریز runs per contract: 300M, then 300M+500M.
        assert rows[0][5] == '۳۰۰,۰۰۰,۰۰۰'
        assert rows[1][5] == '۸۰۰,۰۰۰,۰۰۰'
        # بدهکار/بستانکار is صورت‌وضعیت-تجمیعی minus paid (0 - 800M here).
        assert rows[1][7] == '-۸۰۰,۰۰۰,۰۰۰'
        # مبلغ باقیمانده قرارداد is total_amount minus paid (balance()).
        assert rows[1][9] == '۱,۲۰۰,۰۰۰,۰۰۰'

        assert panel._finance_payments == [
            {'contract_id': contract['contract_id'], 'payment_index': 1, 'pending_id': 'PEN2'},
            {'contract_id': contract['contract_id'], 'payment_index': 0, 'pending_id': 'PEN1'},
        ]


def test_deleting_a_finance_row_removes_only_that_payment(runtime,monkeypatch):
    from types import SimpleNamespace as NS,MethodType
    import contractor_store_stage18_7 as store
    import contractors_stage18_7 as cst
    from test_transaction_lifecycle_matrix import m,case_payload,CASES,register
    from test_table_action_repair import center_for,choice_reply
    import transaction_lifecycle_ui as actions
    store.upsert_contract({"contract_id":"C1","title":"fixture","total_amount":2000,"payments":[]})
    payload={**case_payload(runtime,CASES[0]),"contract_id":"C1","cost_group":"پیمانکار"}
    ids=[]
    for mid in (1,2):
        pid=register(runtime,payload,mid);m.v18.submit_v18(pid,payload,runtime.rich);ids.append(pid)
    center,history=center_for(runtime);prompts,errors=choice_reply(monkeypatch)
    root=NS(_delete_lifecycle_rows187=lambda rows:actions.delete_rows(center,rows,runtime.db,m.messagebox))
    panel=NS(db=runtime.db,_finance_dashboard=lambda:(root,center),_changed=lambda:None,
        tables={cst.FINANCE:NS(selected_keys=lambda:["C1:1"])},
        _finance_payments=[{"contract_id":"C1","payment_index":index} for index in (0,1)])
    panel._selected_finance_sources=MethodType(cst.ContractorsPanel187._selected_finance_sources,panel)
    cst.ContractorsPanel187._delete_finance_payment(panel)
    assert prompts==[1] and not errors and len(history)==1
    assert [p["pending_id"] for p in store.get_contract("C1")["payments"]]==[ids[0]]
    assert runtime.db.snapshot().cash_balance==-1000
    assert runtime.db.get_pending_item(ids[0])["status"]=="APPROVED"
    assert runtime.db.get_pending_item(ids[1])["status"]=="REJECTED"


def test_abandoned_finance_pending_is_discarded_but_a_settled_one_is_left_alone():
    """افزودن/کپی create a placeholder pending item before the wizard even
    opens; closing that wizard without approving or rejecting anything used
    to leave it در انتظار بررسی forever. It must be discarded - but only if
    it is still exactly that untouched placeholder; one the wizard's own
    flow already approved or rejected must be left exactly as it is."""

    reject_calls = []
    _fake_approval_module(reject_calls=reject_calls)
    try:
        import tkinter as tk
        import contractors_stage18_7 as cst

        root = tk.Tk()
        root.withdraw()
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)

            class FakeDB:
                def __init__(self):
                    self.items = {}

                def get_pending_item(self, pending_id):
                    return self.items.get(pending_id)

            panel.db = FakeDB()

            changed_calls = []
            panel._changed = lambda: changed_calls.append(True)

            # Still exactly what افزودن left it as - discard it.
            panel.db.items['PEN1'] = {'pending_id': 'PEN1', 'status': 'PENDING_REVIEW'}
            panel._discard_finance_pending_if_untouched('PEN1')
            assert reject_calls == [('PEN1', 'GUI_187_FINANCE_ABANDONED')]
            assert changed_calls == [True]

            # The wizard's own flow already approved this one - leave it be.
            panel.db.items['PEN2'] = {'pending_id': 'PEN2', 'status': 'APPROVED'}
            panel._discard_finance_pending_if_untouched('PEN2')
            assert reject_calls == [('PEN1', 'GUI_187_FINANCE_ABANDONED')]

            # Already rejected (through the real flow) - leave it be too.
            panel.db.items['PEN3'] = {'pending_id': 'PEN3', 'status': 'REJECTED'}
            panel._discard_finance_pending_if_untouched('PEN3')
            assert reject_calls == [('PEN1', 'GUI_187_FINANCE_ABANDONED')]

            # Row vanished entirely (already cleaned up some other way).
            panel._discard_finance_pending_if_untouched('PEN-MISSING')
            assert reject_calls == [('PEN1', 'GUI_187_FINANCE_ABANDONED')]
        finally:
            root.destroy()
    finally:
        _unfake_approval_module()


def test_open_finance_pending_watches_the_wizards_real_parent_not_the_dashboard():
    """center.open_wizard(row) builds TransactionWizard186(self, ...) - self
    being center, not dashboard - so the new Toplevel lands in
    center.winfo_children(), never in dashboard's. _open_finance_pending
    used to scan dashboard.winfo_children() for it, so the <Destroy> bind
    never attached and an untouched افزودن/کپی placeholder was never
    cleaned up no matter how its wizard closed. A real tk.Toplevel here
    (not a fake stand-in) is what actually exercises that widget lookup."""

    reject_calls = []
    _fake_approval_module(reject_calls=reject_calls)
    try:
        import tkinter as tk
        import pytest
        import contractors_stage18_7 as cst

        try:
            root = tk.Tk()
        except tk.TclError as exc:
            pytest.skip(f'Tcl/Tk unavailable: {exc}')
        root.withdraw()
        try:
            class FakeDB:
                def __init__(self):
                    self.items = {'PEN1': {'pending_id': 'PEN1', 'status': 'PENDING_REVIEW'}}

                def get_pending_item(self, pending_id):
                    return self.items.get(pending_id)

            class Center(tk.Frame):
                def __init__(self, master):
                    super().__init__(master)
                    self.rows = [{'pending_id': 'PEN1'}]
                    self.filter = tk.StringVar()

                def refresh(self):
                    pass

                def open_wizard(self, row):
                    # Mirrors dashboard_stage18_6.py's own
                    # TransactionWizard186(self, ...): the real wizard's
                    # master is center, not dashboard.
                    tk.Toplevel(self)

            dashboard = tk.Frame(root)
            center = Center(dashboard)
            dashboard.approval_center = center
            dashboard.notebook = None
            dashboard.approval_tab = None

            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
            panel.db = FakeDB()
            panel._changed = lambda: None

            panel._open_finance_pending(dashboard, center, 'PEN1', cleanup_if_abandoned=True)
            root.update()
            wizard = next(w for w in center.winfo_children() if isinstance(w, tk.Toplevel))
            wizard.destroy()
            root.update()

            assert reject_calls == [('PEN1', 'GUI_187_FINANCE_ABANDONED')]
        finally:
            root.destroy()
    finally:
        _unfake_approval_module()


def test_finance_add_and_copy_wire_cleanup_but_edit_does_not():
    """ویرایش reopens an existing واریز's own pending item - closing that
    wizard untouched must leave the واریز exactly as it already was, not
    discard it. Only افزودن and کپی create a brand-new placeholder that
    needs the cleanup wired up."""

    source = _source(ROOT / 'contractors_stage18_7.py')

    def method_body(name: str) -> str:
        body = source[source.index(f'    def {name}('):]
        return body[:body.index('\n    def ', 10)]

    assert 'cleanup_if_abandoned=True' in method_body('_add_finance_payment')
    assert 'cleanup_if_abandoned=True' in method_body('_copy_finance_payment')
    edit_body = method_body('_edit_finance_payment')
    assert '_open_finance_pending(' in edit_body
    assert 'cleanup_if_abandoned' not in edit_body


def test_finance_edit_reopens_the_same_pending_add_and_copy_open_new_ones():
    """ویرایش must reuse the واریز's own pending_id (so revise_v18 corrects
    it in place); افزودن and کپی must each open a brand-new pending item -
    نکته copying a PAID deposit must not silently re-post it, so its
    payment_state comes back UNPAID until the wizard confirms it again.
    """

    from contextlib import nullcontext
    from types import SimpleNamespace as NS
    with _temp_store():
        import tkinter as tk
        import tkinter.messagebox as mb
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({
            'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک', 'total_amount': '1',
            'payments': [{'date': '1404/01/01', 'amount': '500000000',
                         'receiver': 'الف', 'issuer': 'ب', 'pending_id': 'PEN1'}],
        })

        class FakeDB:
            def __init__(self):
                self.pending = {}
                self._next = 1

            def atomic(self): return nullcontext()
            def read(self): return nullcontext(NS(execute=lambda *args:[]))
            def get_pending_item(self,pid): return self.pending.get(pid)

            def register_telegram_message(self, **_kw):
                pass

            def create_or_replace_pending(self, *, payload, status, **_kw):
                pid = f'PEN{self._next}'
                self._next += 1
                self.pending[pid] = {'pending_id': pid, 'payload': payload, 'status': status}
                return pid

        class FakeCenter:
            def __init__(self, db):
                self.db = db
                self.filter = tk.StringVar(value='Pending')
                self.rows = []
                self.opened = []

            def refresh(self):
                self.rows = list(self.db.pending.values())

            def open_wizard(self, row):
                self.opened.append(row)

            def winfo_children(self):
                # This test only cares about what row reaches open_wizard();
                # _open_finance_pending's own cleanup-widget-detection is
                # covered separately (see
                # test_open_finance_pending_watches_the_wizards_real_parent_not_the_dashboard).
                return []

        root = tk.Tk()
        root.withdraw()
        try:
            db = FakeDB()
            db.pending['PEN1'] = {'pending_id': 'PEN1', 'payload': {'v18_title':'هزینه ساخت','v18_subtype':'نقدی','amount':500000000,'payment_state':'PAID','contract_id':contract['contract_id']}, 'status':'APPROVED','telegram_chat_id':-187,'telegram_message_id':1}
            db._next = 2  # PEN1 was seeded directly, not through create_or_replace_pending
            center = FakeCenter(db)
            dashboard = tk.Frame(root)
            dashboard.approval_center = center
            dashboard.notebook = None
            dashboard.approval_tab = None

            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
            panel.db = db
            panel.winfo_toplevel = lambda: dashboard
            panel.finance_filter_field = None
            panel._finance_payments = [
                {'contract_id': contract['contract_id'], 'payment_index': 0, 'pending_id': 'PEN1'},
            ]
            panel.tables = {cst.FINANCE:NS(selected_keys=lambda:[str(contract['contract_id'])+':0'])}
            panel._selected_index = lambda sub_tab: 0

            panel._edit_finance_payment()
            assert center.opened[-1]['pending_id'] == 'PEN1'

            original_showinfo = mb.showinfo
            mb.showinfo = lambda *a, **k: None
            try:
                panel._copy_finance_payment()
            finally:
                mb.showinfo = original_showinfo
            copied = center.opened[-1]
            assert copied['pending_id'] != 'PEN1'
            assert copied['payload']['amount'] == 500000000
            assert copied['payload']['payment_state'] == 'UNPAID'

            panel._finance_payments = []
            panel._add_finance_payment()
            added = center.opened[-1]
            assert added['pending_id'] not in ('PEN1', copied['pending_id'])
            assert added['payload']['amount'] == 0
            assert added['payload']['contract_id'] == ''
        finally:
            root.destroy()


# ---------------------------------------------------------------------------
# صورت‌وضعیت row grid: مبلغ کارکرد دوره rename, کارکرد کل / مبلغ کارکرد
# تجمیعی, the footer totals, and a real bug fix - StatementRecordWindow187
# never wired مقدار کارکرد این دوره to مبلغ کارکرد دوره at all.
# ---------------------------------------------------------------------------

def test_statement_row_columns_add_the_two_cumulative_fields_and_rename_line_total():
    import contractor_statements_stage18_7 as cs

    assert [label for _key, label in cs.STATEMENT_ROW_COLUMNS] == [
        'ردیف', 'شرح', 'مبلغ واحد قرارداد', 'حجم کل قرارداد',
        'مقدار کارکرد این دوره', 'مبلغ کارکرد دوره', 'کارکرد کل', 'مبلغ کارکرد تجمیعی',
    ]
    assert 'cumulative_quantity' in cs.STATEMENT_READONLY_ROW_KEYS
    assert 'cumulative_amount' in cs.STATEMENT_READONLY_ROW_KEYS


def test_previous_statement_row_totals_is_zero_with_nothing_before_it():
    with _temp_store():
        import contractor_store_stage18_7 as store

        contract = store.upsert_contract({'title': 'ت', 'family_name': 'پ'})
        assert store.previous_statement_row_totals(contract['contract_id'], 'اجرت') == (0.0, 0)


def test_previous_statement_row_totals_reads_the_matching_rows_own_saved_cumulative():
    with _temp_store():
        import contractor_store_stage18_7 as store

        contract = store.upsert_contract({'title': 'ت', 'family_name': 'پ'})
        first = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_date': '1404/01/01',
            'rows': [{'description': 'اجرت', 'worked_quantity': '500', 'line_total': '350000000',
                     'cumulative_quantity': '500', 'cumulative_amount': '350000000'}],
        })
        second = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_date': '1404/02/01',
            'rows': [{'description': 'اجرت', 'worked_quantity': '300', 'line_total': '210000000',
                     'cumulative_quantity': '800', 'cumulative_amount': '560000000'}],
        })

        # A brand-new (unsaved) صورت‌وضعیت reads the LAST one on file.
        assert store.previous_statement_row_totals(contract['contract_id'], 'اجرت') == (800.0, 560000000)
        # Editing #2 itself must read #1's own totals, not its own.
        assert store.previous_statement_row_totals(
            contract['contract_id'], 'اجرت', before_statement_id=second['statement_id']) == (500.0, 350000000)
        # Editing #1 itself has nothing before it.
        assert store.previous_statement_row_totals(
            contract['contract_id'], 'اجرت', before_statement_id=first['statement_id']) == (0.0, 0)


def test_previous_statement_row_totals_orders_by_date_regardless_of_digit_script():
    """تاریخ صورت‌وضعیت may be stored in ASCII or Persian digits depending on
    when it was saved (JalaliCalendar started writing Persian digits partway
    through this app's life) - a mixed contract must still sort oldest
    first, not put every ASCII-digit date ahead of every Persian one by
    code point. Getting this wrong silently reset a later صورت‌وضعیت's own
    کارکرد کل/مبلغ کارکرد تجمیعی the next time it was opened and saved.
    """

    with _temp_store():
        import contractor_store_stage18_7 as store

        contract = store.upsert_contract({'title': 'ت', 'family_name': 'پ'})
        earlier = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_date': '۱۴۰۴/۰۱/۱۰',
            'rows': [{'description': 'اجرت', 'worked_quantity': '500', 'line_total': '350000000',
                     'cumulative_quantity': '500', 'cumulative_amount': '350000000'}],
        })
        later = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_date': '1404/02/15',
            'rows': [{'description': 'اجرت', 'worked_quantity': '300', 'line_total': '210000000',
                     'cumulative_quantity': '800', 'cumulative_amount': '560000000'}],
        })

        ordered = store.statements_for_contract(contract['contract_id'])
        assert [s['statement_id'] for s in ordered] == [earlier['statement_id'], later['statement_id']]
        assert store.previous_statement_row_totals(
            contract['contract_id'], 'اجرت', before_statement_id=later['statement_id']) == (500.0, 350000000)


def test_cumulative_columns_carry_forward_across_two_statements_end_to_end(tmp_path):
    """صورت‌وضعیت شماره ۱'s کارکرد کل/مبلغ کارکرد تجمیعی equal its own period
    figures; صورت‌وضعیت شماره ۲ adds its own period onto #1's - live, in the
    real wizard, not just at the store layer."""

    import tkinter as tk
    os.environ['STAGE4_DB_PATH'] = str(tmp_path / 'x.sqlite3')
    import contractor_statements_stage18_7 as cs
    import contractor_store_stage18_7 as store

    contract = store.upsert_contract({'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک',
                                      'rows': [{'description': 'اجرت اجرای اسکلت',
                                               'unit_price': '700000', 'quantity': '1914'}]})

    root = tk.Tk()
    root.withdraw()
    try:
        wiz1 = cs.StatementWizard187(root, rich=None)
        try:
            wiz1._var('contract_id').set(next(c for c in cs.contract_choices() if 'اجرت اجرای اسکلت' in c))
            root.update()
            wiz1.show_step(1)
            root.update()
            row1 = wiz1.row_values[0]
            row1['worked_quantity'].set('500')
            root.update()
            assert row1['cumulative_quantity'].get() == row1['worked_quantity'].get() == '۵۰۰'
            assert row1['cumulative_amount'].get() == row1['line_total'].get() == '۳۵۰,۰۰۰,۰۰۰'
            assert wiz1._row_totals_caption.cget('text') == (
                'مبلغ کل صورت‌وضعیت دوره: ۳۵۰,۰۰۰,۰۰۰ تومان'
                '    |    مبلغ کل صورت‌وضعیت تجمیعی: ۳۵۰,۰۰۰,۰۰۰ تومان')
            wiz1._var('statement_date').set('1404/01/01')
            wiz1.finish()
        finally:
            if wiz1.winfo_exists():
                wiz1.destroy()

        wiz2 = cs.StatementWizard187(root, rich=None)
        try:
            wiz2._var('contract_id').set(next(c for c in cs.contract_choices() if 'اجرت اجرای اسکلت' in c))
            root.update()
            wiz2.show_step(1)
            root.update()
            row2 = wiz2.row_values[0]
            # Before this period's own value is even typed, it already shows
            # #1's totals - not blank.
            assert row2['cumulative_quantity'].get() == '۵۰۰'
            assert row2['cumulative_amount'].get() == '۳۵۰,۰۰۰,۰۰۰'
            row2['worked_quantity'].set('300')
            root.update()
            assert row2['line_total'].get() == '۲۱۰,۰۰۰,۰۰۰'
            assert row2['cumulative_quantity'].get() == '۸۰۰'
            assert row2['cumulative_amount'].get() == '۵۶۰,۰۰۰,۰۰۰'
        finally:
            if wiz2.winfo_exists():
                wiz2.destroy()
    finally:
        root.destroy()


def test_record_window_recomputes_line_total_when_worked_quantity_is_typed(tmp_path):
    """Regression: StatementRecordWindow187 used to build its rows as a bare
    dict, skipping blank_statement_row entirely - typing into مقدار کارکرد
    این دوره there never touched مبلغ کارکرد دوره at all."""

    import tkinter as tk
    os.environ['STAGE4_DB_PATH'] = str(tmp_path / 'x.sqlite3')
    import contractor_statements_stage18_7 as cs
    import contractor_store_stage18_7 as store

    contract = store.upsert_contract({'title': 'اجرت اجرای اسکلت', 'family_name': 'ابراهیم ازبک'})
    statement = store.upsert_statement({
        'contract_id': contract['contract_id'],
        'rows': [{'description': 'اجرت اجرای اسکلت', 'unit_price': '700000', 'contract_quantity': '1914',
                 'worked_quantity': '500', 'line_total': '350000000',
                 'cumulative_quantity': '500', 'cumulative_amount': '350000000'}],
    })

    root = tk.Tk()
    root.withdraw()
    try:
        rec = cs.StatementRecordWindow187(root, statement=dict(statement), rich=None)
        try:
            row = rec.row_values[0]
            row['worked_quantity'].set('900')
            root.update()
            assert row['line_total'].get() == '۶۳۰,۰۰۰,۰۰۰'
            assert row['cumulative_quantity'].get() == '۹۰۰'
            assert row['cumulative_amount'].get() == '۶۳۰,۰۰۰,۰۰۰'
        finally:
            rec.destroy()
    finally:
        root.destroy()


# ---------------------------------------------------------------------------
# آرشیو: from a per-پیمانکار rating form into a registry of every طرف حساب
# (پیمانکار/تامین‌کننده مصالح/خدمات‌دهنده), with its own افزودن/ویرایش/حذف/کپی
# toolbar and دسته‌بندی/نام/نوع‌پرداخت نمایش‌براساس filter.
# ---------------------------------------------------------------------------

def test_archive_columns_are_the_nine_requested_fields_in_order():
    import contractors_stage18_7 as cst

    assert [label for _key, label in cst.COLUMNS[cst.ARCHIVE]] == [
        'ردیف', 'نام', 'دسته‌بندی', 'عنوان فعالیت', 'شماره تماس', 'آدرس',
        'نوع پرداخت', 'امتیاز کلی (از ۱۰۰)', 'احتمال همکاری مجدد',
    ]


def test_archive_toolbar_is_افزودن_ویرایش_حذف_کپی_then_filter_then_خروجی():
    source = _source(ROOT / 'contractors_stage18_7.py')
    build_page = source[source.index('    def _build_page(self, page: ttk.Frame, label: str)'):]
    build_page = build_page[:build_page.index('\n    def ', 10)]
    archive_block = build_page[build_page.index('# آرشیو: افزودن'):]
    assert "table.add_tool('افزودن +'" in archive_block
    assert "table.add_tool('ویرایش'" in archive_block
    assert "table.add_tool('حذف ردیف'" in archive_block
    assert "table.add_tool('کپی', self.copy_selected_profile)" in archive_block
    assert "'نمایش براساس', self.archive_filter_field" in archive_block
    assert "table.add_export(self.archive_export_choice" in archive_block
    assert "table.tree.bind('<Button-3>', self._archive_context_menu" in archive_block


def test_archive_filter_offers_دسته‌بندی_نام_and_نوع_پرداخت():
    import contractors_stage18_7 as cst

    assert cst.ARCHIVE_FILTER_FIELDS == (cst.ANY, 'دسته‌بندی', 'نام', 'نوع پرداخت')
    assert cst.ARCHIVE_FILTER_KEYS == {
        'دسته‌بندی': 'category', 'نام': 'family_name', 'نوع پرداخت': 'payment_type',
    }


def test_archive_rows_read_the_new_fields_in_the_new_column_order():
    with _temp_store():
        from ui_stage18_3 import fa
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_profile({
            'family_name': 'ابراهیم ازبک', 'category': 'پیمانکار', 'activity': 'اجرای اسکلت',
            'phone': '09121234567', 'address': 'تهران', 'payment_type': 'نقد',
            'overall_score': '85', 'rehire_chance': 'اولویت دوم',
        })

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.archive_filter_field = None
        rows, _status = panel._load_archive_rows()

        assert len(rows) == 1
        row = rows[0]
        assert row[1] == 'ابراهیم ازبک'
        assert row[2] == 'پیمانکار'
        assert row[3] == 'اجرای اسکلت'
        assert row[4] == fa('09121234567')
        assert row[5] == 'تهران'
        assert row[6] == 'نقد'
        assert row[7] == fa('85')
        # rehire_chance is a text priority label now (اولویت اول/دوم/...),
        # never a number - it must not be Persian-digit-converted.
        assert row[8] == 'اولویت دوم'


def test_archive_filter_narrows_by_category():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_profile({'family_name': 'الف', 'category': 'پیمانکار'})
        store.upsert_profile({'family_name': 'ب', 'category': 'تامین‌کننده مصالح'})

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.archive_filter_field = type('V', (), {'get': lambda self: 'دسته‌بندی'})()
        panel.archive_filter_value = type('V', (), {'get': lambda self: 'پیمانکار'})()
        rows, _status = panel._load_archive_rows()

        assert len(rows) == 1
        assert rows[0][1] == 'الف'


def test_archive_wizard_window_one_asks_identity_category_activity_and_payment_type():
    source = _source(ROOT / 'contractor_archive_stage18_7.py')
    assert "CATEGORY_OPTIONS = ('تامین‌کننده مصالح', 'پیمانکار', 'خدماتی')" in source
    assert "PAYMENT_TYPES = ('نقدی-چکی', 'نقدی', 'چکی')" in source
    assert "CATEGORY_FIELDS = (('category', 'دسته‌بندی فعالیت'),)" in source
    assert "ACTIVITY_FIELDS = (('activity', 'عنوان فعالیت'),)" in source
    assert "PAYMENT_TYPE_FIELDS = (('payment_type', 'نوع پرداخت'),)" in source

    render_step = source[source.index('    def render_step(self, index: int'):]
    render_step = render_step[:render_step.index('\n    def ', 10)]
    assert 'if index == 0:' in render_step
    assert 'CATEGORY_FIELDS, combo=CATEGORY_OPTIONS' in render_step
    assert 'ACTIVITY_FIELDS)' in render_step
    assert 'PAYMENT_TYPE_FIELDS, combo=PAYMENT_TYPES' in render_step


def test_archive_payment_type_defaults_to_the_new_mixed_option():
    """نوع پرداخت's own combo[0]-as-default convention (the same one category
    used to fall afoul of) means the operator's ask - نقدی-چکی as the
    default - is just a matter of it being FIRST in PAYMENT_TYPES."""

    import contractor_archive_stage18_7 as ca

    assert ca.PAYMENT_TYPES[0] == 'نقدی-چکی'
    assert set(ca.PAYMENT_TYPES) == {'نقدی-چکی', 'نقدی', 'چکی'}

    with _temp_store():
        import tkinter as tk
        import pytest
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

        try:
            root = tk.Tk()
        except tk.TclError as exc:
            pytest.skip(f'Tcl/Tk unavailable: {exc}')
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            wiz = ca.ArchiveWizard187(root, rich=None)
            assert wiz.values['payment_type'].get() == 'نقدی-چکی'
            record = ca.ArchiveRecordWindow187(root, profile={'family_name': 'ت'})
            assert record.values['payment_type'].get() == 'نقدی-چکی'
        finally:
            root.destroy()


def test_archive_phone_field_shows_persian_digits_while_typing():
    """شمارهٔ تماس, مثل هر عدد دیگر در این برنامه, باید با رقم فارسی تایپ
    شود - هم در ویزارد, هم در پنجرهٔ رکورد."""

    with _temp_store():
        import tkinter as tk
        import pytest
        from tkinter import ttk
        import contractor_archive_stage18_7 as ca
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

        try:
            root = tk.Tk()
        except tk.TclError as exc:
            pytest.skip(f'Tcl/Tk unavailable: {exc}')
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            wiz = ca.ArchiveWizard187(root, rich=None)
            wiz.values['phone'].set('09121234567')
            assert wiz.values['phone'].get() == '۰۹۱۲۱۲۳۴۵۶۷'

            record = ca.ArchiveRecordWindow187(root, profile={'family_name': 'ت'})
            record.values['phone'].set('09121234567')
            assert record.values['phone'].get() == '۰۹۱۲۱۲۳۴۵۶۷'
        finally:
            root.destroy()


def test_archive_criteria_pages_are_six_and_six_and_whatever_is_left():
    import contractor_archive_stage18_7 as ca

    for category in ca.CATEGORY_OPTIONS:
        total = ca.criteria_for(category)
        pages = [ca.criteria_page(category, i) for i in range(3)]
        assert len(pages[0]) == min(6, len(total))
        assert sum(len(page) for page in pages) == len(total)
        # No معیار repeated and none dropped across the three pages.
        assert [c for page in pages for c in page] == list(total)


def test_overall_score_is_the_weighted_average_out_of_100():
    import contractor_archive_stage18_7 as ca

    values = {'quality_of_work_score': '10', 'quality_of_work_weight': '2',
             'work_speed_score': '5', 'work_speed_weight': '1'}
    overall = ca.compute_overall_score(values, 'پیمانکار')
    assert abs(overall - (10 * 2 + 5 * 1) / 3 * 10) < 1e-9
    # A weightless/blank survey scores zero rather than raising.
    assert ca.compute_overall_score({}, 'پیمانکار') == 0.0


def test_rehire_chance_bands_match_the_spec_exactly():
    import contractor_archive_stage18_7 as ca

    cases = [
        (100, 'اولویت اول'), (90, 'اولویت اول'), (89.9, 'اولویت دوم'),
        (80, 'اولویت دوم'), (79.9, 'اولویت سوم'), (70, 'اولویت سوم'),
        (69.9, 'اولویت چهارم'), (60, 'اولویت چهارم'), (59.9, 'احتمال همکاری ندارد'),
        (0, 'احتمال همکاری ندارد'),
    ]
    for score, expected in cases:
        assert ca.compute_rehire_chance(score) == expected, score


def test_archive_detail_window_is_not_the_editable_record_window():
    """دبل‌کلیک باید ArchiveDetailWindow187 باز کند - پنجرهٔ رکورد قابل‌ویرایش
    فقط از ویرایش/کپی/راست‌کلیک باز می‌شود، دقیقاً مثل قراردادها."""

    source = _source(ROOT / 'contractors_stage18_7.py')
    open_record = source[source.index('    def open_record(self, sub_tab: str, item: str | None) -> None:'):]
    open_record = open_record[:open_record.index('\n    def ', 10)]
    archive_branch = open_record[open_record.index('if sub_tab == ARCHIVE:'):]
    assert 'ArchiveDetailWindow187(self, profile=profile, on_export=self._export_profile)' in archive_branch
    assert 'ArchiveRecordWindow187' not in archive_branch


def test_copying_a_profile_drops_its_own_id_but_keeps_everything_else():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        profile = store.upsert_profile({'family_name': 'الف', 'category': 'پیمانکار',
                                        'overall_score': '80'})

        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.rich = None
        panel._profile_ids = [profile['profile_id']]
        panel._selected_index = lambda sub_tab: 0

        captured = {}

        class FakeWizard:
            def __init__(self, parent, *, rich, profile, on_saved):
                captured['profile'] = profile

        original_wizard = cst.ArchiveWizard187
        cst.ArchiveWizard187 = FakeWizard
        try:
            cst.ContractorsPanel187.copy_selected_profile(panel)
        finally:
            cst.ArchiveWizard187 = original_wizard

        clone = captured['profile']
        assert 'profile_id' not in clone
        assert clone['family_name'] == 'الف'
        assert clone['category'] == 'پیمانکار'
        assert clone['overall_score'] == '80'


# ---------------------------------------------------------------------------
# آرشیو auto-sync: قراردادها + تراکنش‌های تأییدشده + عوامل پروژه, بدون
# ساختن ردیف تکراری و بدون دست‌زدن به امتیاز/ضریب دستی کسی.
# ---------------------------------------------------------------------------

def _fake_approval_module(reviews=(), *, reject_calls=None):
    """A stand-in for dashboard_stage18 so sync_archive_from_sources's own
    lazy import resolves without pulling in the real (heavy, AI/Telegram-
    coupled) module - only APPROVAL.list_reviews() is ever called on it.
    ``reject_calls``, if given, also records every APPROVAL.reject() call
    for the finance-pending-abandon tests."""

    import sys
    import types

    class FakeApproval:
        def list_reviews(self, limit=3000):
            return reviews

        def reject(self, pending_id, *, actor_id=''):
            if reject_calls is not None:
                reject_calls.append((pending_id, actor_id))

    module = types.ModuleType('dashboard_stage18')
    module.APPROVAL = FakeApproval()
    sys.modules['dashboard_stage18'] = module
    return module


def _unfake_approval_module():
    import sys
    sys.modules.pop('dashboard_stage18', None)


def test_sync_pulls_contractors_suppliers_service_providers_and_project_team():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import v18_store

        store.upsert_contract({'family_name': 'پیمانکار الف', 'phone': '0912', 'title': 'اجرای سقف'})

        _fake_approval_module([
            {'status': 'APPROVED', 'payload': {
                'cost_group': 'مصالح', 'receiver_name': 'فروشگاه فولاد', 'v18_subtype': 'نقدی'}},
            {'status': 'APPROVED', 'payload': {
                'cost_group': 'خدمات', 'receiver_name': 'نقشه‌بردار حسینی', 'v18_subtype': 'چک'}},
            # Not APPROVED yet - must not sync.
            {'status': 'PENDING_REVIEW', 'payload': {'cost_group': 'مصالح', 'receiver_name': 'نباید بیاید'}},
            # No cost_group مصالح/خدمات match (e.g. پیمانکار cost_group itself) - not this sync's job.
            {'status': 'APPROVED', 'payload': {'cost_group': 'پیمانکار', 'receiver_name': 'نادیده'}},
        ])
        original_load_profile = v18_store.load_profile
        v18_store.load_profile = lambda: {'team': {
            'executor': {'name': 'مجری رضایی', 'phone': '0910'},
            'supervisor': {'name': ''}, 'designer': {'name': ''},
        }}
        try:
            touched = store.sync_archive_from_sources()
        finally:
            v18_store.load_profile = original_load_profile
            _unfake_approval_module()

        profiles = {p['family_name']: p for p in store.load_profiles()}
        assert touched == 4
        assert 'نباید بیاید' not in profiles
        assert 'نادیده' not in profiles
        assert profiles['پیمانکار الف']['category'] == 'پیمانکار'
        assert profiles['پیمانکار الف']['activity'] == 'اجرای سقف'
        assert profiles['فروشگاه فولاد']['category'] == 'تامین‌کننده مصالح'
        assert profiles['فروشگاه فولاد']['payment_type'] == 'نقدی'
        assert profiles['نقشه‌بردار حسینی']['category'] == 'خدماتی'
        assert profiles['نقشه‌بردار حسینی']['payment_type'] == 'چکی'
        assert profiles['مجری رضایی']['category'] == 'خدماتی'
        assert profiles['مجری رضایی']['activity'] == 'مجری'
        assert profiles['مجری رضایی']['phone'] == '0910'


def test_sync_never_overwrites_a_persons_own_score_or_already_set_fields():
    with _temp_store():
        import contractor_store_stage18_7 as store

        store.upsert_profile({
            'family_name': 'پیمانکار الف', 'category': 'پیمانکار', 'address': 'اصفهان',
            'quality_of_work_score': '9', 'quality_of_work_weight': '3',
            'overall_score': '90', 'rehire_chance': 'اولویت اول',
        })
        store.upsert_contract({'family_name': 'پیمانکار الف', 'phone': '0912', 'title': 'اجرای سقف'})

        store.sync_archive_from_sources()

        assert len(store.load_profiles()) == 1  # no duplicate row for the same name
        profile = store.load_profiles()[0]
        assert profile['category'] == 'پیمانکار'  # already set - untouched
        assert profile['address'] == 'اصفهان'  # already set - untouched
        assert profile['phone'] == '0912'  # was blank - filled from the قرارداد
        assert profile['activity'] == 'اجرای سقف'  # was blank - filled
        assert profile['quality_of_work_score'] == '9'  # never touched by sync
        assert profile['quality_of_work_weight'] == '3'  # never touched by sync
        assert profile['overall_score'] == '90'  # never touched by sync
        assert profile['rehire_chance'] == 'اولویت اول'  # never touched by sync


def test_sync_dedups_by_normalized_name_not_exact_string():
    with _temp_store():
        import contractor_store_stage18_7 as store

        store.upsert_profile({'family_name': 'علي  رضا'})  # Arabic ي, double space
        store.upsert_contract({'family_name': 'علی رضا'})  # Persian ی, single space

        store.sync_archive_from_sources()
        assert len(store.load_profiles()) == 1


def test_sync_is_idempotent_running_twice_creates_nothing_new():
    with _temp_store():
        import contractor_store_stage18_7 as store

        store.upsert_contract({'family_name': 'پیمانکار الف', 'phone': '0912'})
        first = store.sync_archive_from_sources()
        second = store.sync_archive_from_sources()
        assert first == 1
        assert second == 0
        assert len(store.load_profiles()) == 1


def test_archive_tab_load_calls_sync_and_survives_it_failing():
    """_load_archive_rows must call the sync every refresh (the whole point
    is that آرشیو stays live), but a broken/absent legacy dashboard stack
    must never take the whole tab down with it - sync_archive_from_sources
    already swallows that per-source; this is the outer safety net."""

    source = _source(ROOT / 'contractors_stage18_7.py')
    load_rows = source[source.index('    def _load_archive_rows(self)'):]
    load_rows = load_rows[:load_rows.index('\n    def ', 10)]
    assert 'sync_archive_from_sources()' in load_rows
    assert 'except Exception' in load_rows

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'پیمانکار الف'})
        panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
        panel.archive_filter_field = None
        rows, _status = panel._load_archive_rows()
        assert len(rows) == 1
        assert rows[0][1] == 'پیمانکار الف'


# ---------------------------------------------------------------------------
# قراردادها: selection survives a background refresh
# ---------------------------------------------------------------------------

def test_a_selected_contract_row_survives_a_background_refresh():
    """dashboard_stage18_7.py's own refresh() calls contractors.refresh()
    unconditionally, and that refresh() is reached by two independent
    periodic ticks (the bot-status check every 4s, the pending-approval
    check) that have nothing to do with قراردادها's own data. DataTable
    rebuilds every row with a fresh Treeview iid on every set_rows() call,
    so without a stable per-row key the selection just vanishes a few
    seconds after the user clicks a row - confirmed live by simulating a
    click inside the real Stage187Dashboard and draining its event loop.
    ``ContractorsPanel187.refresh()`` must pass its own row identities
    (contract_id/statement_id/pending_id/profile_id) through as
    ``keys=`` so DataTable.set_rows can put the same row's selection back.
    """

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت'})
        store.upsert_contract({'family_name': 'محمود کریمی', 'title': 'اجرای سفت‌کاری'})

        import tkinter as tk
        from tkinter import ttk
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            panel = cst.ContractorsPanel187(root, db=None, rich=None)
            panel.pack(fill='both', expand=True)
            root.update()

            table = panel.tables[cst.CONTRACTS]
            first_iid = table.tree.get_children()[0]
            selected_contract_id = panel._contract_ids[0]
            table.tree.selection_set(first_iid)
            root.update()
            assert table.tree.selection() == (first_iid,)

            # Mimic a background tick calling the panel's own refresh(), same
            # as dashboard_stage18_7.py's refresh() -> contractors.refresh().
            panel.refresh()
            root.update()

            new_selection = table.tree.selection()
            assert new_selection, 'selection must not be dropped by a background refresh'
            reselected_index = table.tree.index(new_selection[0])
            assert panel._contract_ids[reselected_index] == selected_contract_id
        finally:
            root.destroy()


def test_finance_filter_by_contractor_only_offers_names_that_have_a_deposit():
    """A قرارداد with no واریز never produces a row in امور مالی's own
    جدول - so it must not appear as a پیمانکار option either, since
    picking it would always narrow the table to nothing."""

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'با واریز', 'title': 'قرارداد یک',
                               'payments': [{'date': '1404/01/01', 'amount': '1000'}]})
        store.upsert_contract({'family_name': 'بدون واریز', 'title': 'قرارداد دو'})

        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)

            class FakeCombo:
                def __init__(self):
                    self.values = None

                def configure(self, *, values):
                    self.values = values

            panel.finance_filter_field = tk.StringVar(root, value='نام پیمانکار')
            panel.finance_filter_value = tk.StringVar(root, value=cst.ANY)
            panel.finance_filter_value_combo = FakeCombo()
            panel.refresh = lambda: None
            panel._finance_filter_field_changed()
            assert panel.finance_filter_value_combo.values == (cst.ANY, 'با واریز')
        finally:
            root.destroy()


def test_contracts_filter_offers_statement_type_and_narrows_by_it():
    """قراردادها's own نمایش براساس is missing نوع صورت‌وضعیت (موقت/قطعی) -
    a قرارداد itself has no such field, so it must read from its own
    صورت‌وضعیت‌ها instead, defaulting to همه like every other filter here."""

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        assert cst.STATEMENT_TYPE_FILTER_FIELD in cst.FILTER_FIELDS

        c1 = store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت'})
        c2 = store.upsert_contract({'family_name': 'محمود کریمی', 'title': 'اجرای سفت‌کاری'})
        store.upsert_statement({'contract_id': c1['contract_id'], 'statement_type': 'موقت'})
        store.upsert_statement({'contract_id': c2['contract_id'], 'statement_type': 'قطعی'})

        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
            panel.filter_field = tk.StringVar(root, value=cst.STATEMENT_TYPE_FILTER_FIELD)
            panel.filter_value = tk.StringVar(root, value=cst.ANY)
            panel.filters = {}

            rows, _status = panel.load_rows(cst.CONTRACTS)
            assert {r[3] for r in rows} == {'ابراهیم ازبک', 'محمود کریمی'}

            panel.filter_value.set('موقت')
            rows, _status = panel.load_rows(cst.CONTRACTS)
            assert [r[3] for r in rows] == ['ابراهیم ازبک']

            panel.filter_value.set('قطعی')
            rows, _status = panel.load_rows(cst.CONTRACTS)
            assert [r[3] for r in rows] == ['محمود کریمی']
        finally:
            root.destroy()


def test_contracts_statement_type_filter_value_choices_are_the_fixed_two_types():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'ابراهیم ازبک'})

        import tkinter as tk
        root = tk.Tk()
        root.withdraw()
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)

            class FakeCombo:
                def __init__(self):
                    self.values = None

                def configure(self, *, values):
                    self.values = values

            panel.filter_field = tk.StringVar(root, value=cst.STATEMENT_TYPE_FILTER_FIELD)
            panel.filter_value = tk.StringVar(root, value=cst.ANY)
            panel.filter_value_combo = FakeCombo()
            panel.refresh = lambda: None
            panel._filter_field_changed()
            assert panel.filter_value_combo.values == (cst.ANY, 'موقت', 'قطعی')
            assert panel.filter_value.get() == cst.ANY
        finally:
            root.destroy()


def test_data_table_set_rows_ignores_keyless_rows_when_restoring_selection():
    """Most callers never pass keys= - set_rows must keep behaving exactly
    as before for them (selection cleared on rebuild) rather than treating
    every None key as the same row and selecting everything."""

    import tkinter as tk
    from tkinter import ttk
    from barkav_ui.skin import install_rounded_skin
    from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
    from barkav_ui.table import DataTable

    root = tk.Tk()
    root.withdraw()
    try:
        install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
        table = DataTable(root, (('a', 'الف'), ('b', 'ب')))
        table.pack()
        table.set_rows([('1', '2'), ('3', '4')])
        root.update()
        table.tree.selection_set(table.tree.get_children()[0])
        root.update()

        table.set_rows([('1', '2'), ('3', '4')])
        root.update()
        assert table.tree.selection() == (), 'no keys were given - nothing to match back up'
    finally:
        root.destroy()


# ---------------------------------------------------------------------------
# قراردادها: table exports read right-to-left, like the table on screen
# ---------------------------------------------------------------------------

def test_readonly_ttk_entry_background_is_explicitly_mapped_in_source():
    """The rounded TEntry pill has fully transparent corners in its own
    image; ttk paints the style's own `background` behind them, so any
    state left unmapped falls back to the active ttk theme's own default -
    confirmed live (screenshot) as a stark white rectangle around
    صورت‌وضعیت's مقدار/مبلغ کارکرد entries in one theme and its فعالیت‌های
    قرارداد readonly cells in another, since which state's default leaks
    through flips with the active palette. A style.lookup()-based
    assertion here is unreliable inside this suite - ttk's style database
    is not perfectly isolated between the many tk.Tk() roots other tests in
    this same process create and destroy - so this checks the fix is still
    present in source instead.
    """

    source = _source(ROOT / 'barkav_ui' / 'skin.py')
    entry_block = source[source.index("style.layout('TEntry',"):]
    entry_block = entry_block[:entry_block.index("style.map('TCombobox'")]
    assert "background=palette['bg']" in entry_block
    assert "style.map('TEntry', background=[" in entry_block
    assert "('readonly', palette['bg'])" in entry_block
    assert "('disabled', palette['bg'])" in entry_block


# ---------------------------------------------------------------------------
# صورت‌وضعیت/امور مالی's double-click: a detail window, not the record window
# ---------------------------------------------------------------------------

def test_statements_and_finance_double_click_open_the_detail_window_not_a_record():
    """صورت‌وضعیت used to open StatementRecordWindow187 on double click, and
    امور مالی mistakenly opened ContractRecordWindow187 - both must instead
    open the same StatementFinanceDetailWindow187, matching a هدایت‌شده
    reading of the shared spec sent for both tables."""

    source = _source(ROOT / 'contractors_stage18_7.py')
    open_record = source[source.index('    def open_record(self, sub_tab: str, item: str | None) -> None:'):]
    open_record = open_record[:open_record.index('\n    def ', 10)]
    statements_branch = open_record[open_record.index('if sub_tab == STATEMENTS:'):
                                    open_record.index('if sub_tab == FINANCE:')]
    finance_branch = open_record[open_record.index('if sub_tab == FINANCE:'):
                                 open_record.index('if sub_tab == ARCHIVE:')]
    assert 'StatementFinanceDetailWindow187(' in statements_branch
    assert 'StatementRecordWindow187(' not in statements_branch
    assert 'StatementFinanceDetailWindow187(' in finance_branch
    assert 'ContractRecordWindow187(' not in finance_branch


def test_finance_detail_window_uses_the_contracts_latest_statement():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        contract = store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت'})
        store.upsert_statement({'contract_id': contract['contract_id'], 'statement_number': 'S1',
                                'statement_date': '1404/01/01', 'statement_type': 'موقت'})
        latest = store.upsert_statement({'contract_id': contract['contract_id'], 'statement_number': 'S2',
                                         'statement_date': '1404/03/01', 'statement_type': 'قطعی'})

        captured = {}

        class FakeDetail:
            def __init__(self, _parent, *, contract, statement, db):
                captured['statement_id'] = statement.get('statement_id') if statement else None

        import contractor_statements_stage18_7 as cs
        original = cs.StatementFinanceDetailWindow187
        cst.StatementFinanceDetailWindow187 = FakeDetail
        try:
            panel = cst.ContractorsPanel187.__new__(cst.ContractorsPanel187)
            panel.db = None
            panel._selected_finance_payment = lambda: (contract, None)
            panel.open_record(cst.FINANCE, None)
        finally:
            cst.StatementFinanceDetailWindow187 = original

        assert captured['statement_id'] == latest['statement_id']


def test_finance_detail_window_handles_a_contract_with_no_statements_yet():
    """امور مالی can have a واریز before any صورت‌وضعیت exists - the detail
    window must show an empty ردیف‌ها section, not crash."""

    with _temp_store():
        import tkinter as tk
        from tkinter import ttk
        import contractor_store_stage18_7 as store
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from contractor_statements_stage18_7 import StatementFinanceDetailWindow187, DEFAULT_STATEMENT_TYPE

        contract = store.upsert_contract({
            'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت',
            'payments': [{'date': '1404/01/01', 'amount': '1000', 'receiver': 'ر', 'issuer': 'ص'}],
        })

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            win = StatementFinanceDetailWindow187(root, contract=contract, statement=None, db=None)
            root.update()
            # No صورت‌وضعیت at all still names a نوع - the same fallback
            # _statement_matches already uses, not a bare dash.
            assert f'نوع صورت‌وضعیت: {DEFAULT_STATEMENT_TYPE}' in win._summary_segments()[0]
        finally:
            root.destroy()


def test_finance_detail_summary_uses_separate_widgets_not_one_bidi_prone_string():
    """«بدهکار/بستانکار: عدد» در همان رشتهٔ چهار-بخشیِ خلاصه، حتی با RLM
    جلوی رشته، توسط Tk (بدون الگوریتم bidi واقعی) به‌هم می‌ریخت - هر بخش
    باید ویجت Label خودش را بگیرد، نه یک رشتهٔ ترکیبی."""

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    build = source[source.index('class StatementFinanceDetailWindow187'):]
    build = build[build.index('def _build(self)'):]
    build = build[:build.index('\n    def ', 10)]
    assert 'for segment in self._summary_segments():' in build
    assert "text=self._summary_line()" not in build


def test_finance_detail_row_and_payment_headers_share_the_row_cells_own_font():
    """tk.Label's width=N is a character count measured in THAT label's own
    font - a header at kit_font('small') and its row cells at kit_font('body')
    render the same width= to different pixel widths, so the column edges
    drift out of alignment going right to left. Both hand-rolled tables
    (ردیف‌های صورت‌وضعیت and واریزی‌ها) must size their header the same
    font as the cells under it.
    """

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    rows_method = source[source.index('    def _render_statement_rows(self, parent)'):]
    rows_method = rows_method[:rows_method.index('\n    def ', 10)]
    assert rows_method.count("font=kit_font('body')") >= 2
    assert "font=kit_font('small')" not in rows_method

    payments_method = source[source.index('    def _render_payments(self, parent)'):]
    payments_method = payments_method[:payments_method.index('\n    def ', 10)]
    assert payments_method.count("font=kit_font('body')") >= 2
    assert "font=kit_font('small')" not in payments_method


def test_finance_detail_export_combo_sits_on_the_left():
    """the operator's explicit ask - انتخاب خروجی moved off the RTL window's
    own right edge."""

    source = _source(ROOT / 'contractor_statements_stage18_7.py')
    build = source[source.index('class StatementFinanceDetailWindow187'):]
    build = build[build.index('def _build(self)'):]
    build = build[:build.index('\n    def ', 10)]
    combo_statement = build[build.index('combo = ttk.Combobox('):]
    combo_statement = combo_statement[:combo_statement.index(')', combo_statement.index('.pack('))]
    assert "side='left'" in combo_statement


def test_finance_detail_statement_type_falls_back_to_the_default_not_a_dash():
    """A صورت‌وضعیت saved before نوع صورت‌وضعیت existed as a field has no
    statement_type key at all - the summary must still name a real نوع
    (موقت, the same DEFAULT_STATEMENT_TYPE the wizard itself seeds new
    records with and _statement_matches already falls back to), not show
    an empty value after the label."""

    with _temp_store():
        import tkinter as tk
        from tkinter import ttk
        import contractor_store_stage18_7 as store
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from contractor_statements_stage18_7 import StatementFinanceDetailWindow187, DEFAULT_STATEMENT_TYPE

        contract = store.upsert_contract({'family_name': 'ت', 'title': 'ق'})
        statement = store.upsert_statement({
            'contract_id': contract['contract_id'], 'statement_date': '1404/01/01',
            'rows': [],
        })
        assert 'statement_type' not in statement

        root = tk.Tk()
        root.withdraw()
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            win = StatementFinanceDetailWindow187(root, contract=contract, statement=statement, db=None)
            root.update()
            assert f'نوع صورت‌وضعیت: {DEFAULT_STATEMENT_TYPE}' in win._summary_segments()[0]
        finally:
            root.destroy()


def test_export_statement_finance_detail_pdf_reverses_both_tables_columns():
    """ReportLab always lays a Table's column 0 out at the sheet's physical
    left regardless of alignment styling - confirmed by rendering a probe
    PDF with unambiguous markers during this feature's own development, the
    same way export_contract_pdf_187's caller and export_receipts_pdf_187
    already reverse their own tables. _rtl is stubbed to identity here so
    the captured cell text is exactly what was passed in, not its bidi-
    reshaped/reordered form."""

    import sales_claims_stage18_7 as sales
    import export_utils

    captured = {}
    from reportlab import platypus
    original_table_init = platypus.Table.__init__
    original_rtl = export_utils._rtl

    def traced_init(self, data, *a, **k):
        captured.setdefault('tables', []).append(data)
        return original_table_init(self, data, *a, **k)

    platypus.Table.__init__ = traced_init
    export_utils._rtl = lambda text: text
    try:
        destination = Path(tempfile.mkdtemp()) / 'out.pdf'
        sales.export_statement_finance_detail_pdf_187(
            destination, title='t', identity_lines=['x'],
            statement_headers=['ردیف', 'شرح', 'آخر'],
            statement_rows=[['1', 'x', 'z']],
            statement_caption='c',
            payment_headers=['ردیف', 'تاریخ', 'آخر'],
            payment_rows=[['1', 'd', 'z']],
            receipt_paths=[''],
        )
    finally:
        platypus.Table.__init__ = original_table_init
        export_utils._rtl = original_rtl

    def cell_text(flowable):
        return flowable.getPlainText() if hasattr(flowable, 'getPlainText') else str(flowable)

    tables = captured['tables']
    # First table built is the ردیف‌های صورت‌وضعیت one; its header row's
    # first physical column (ReportLab always lays column 0 out on the
    # left) must be the *last* logical header (آخر), not ردیف.
    statement_header_row = [cell_text(c) for c in tables[0][0]]
    assert statement_header_row == ['آخر', 'شرح', 'ردیف']

    payment_header_row = [cell_text(c) for c in tables[1][0]]
    assert payment_header_row == ['آخر', 'تاریخ', 'ردیف']


def test_pdf_export_reverses_columns_to_match_the_on_screen_rtl_order():
    """export_pdf (export_utils.py) lays its own table out left-to-right, so
    a logical-order header list puts ردیف at the sheet's physical left edge -
    backwards from the on-screen DataTable, which puts that same column at
    the right. export_receipts_pdf_187 already reverses its own columns for
    exactly this reason; _write_export must do the same for its four
    پیمانکاران tabs (قراردادها/صورت‌وضعیت/امور مالی/آرشیو all share it)."""

    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت',
                               'contract_number': 'C-1'})

        import tkinter as tk
        from tkinter import ttk, filedialog
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette

        root = tk.Tk()
        root.withdraw()
        original_asksave = filedialog.asksaveasfilename
        original_pdf = cst.export_contract_pdf_187
        original_showinfo = cst.messagebox.showinfo
        original_showerror = cst.messagebox.showerror
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            panel = cst.ContractorsPanel187(root, db=None, rich=None)
            panel.pack(fill='both', expand=True)
            root.update()

            captured = {}
            errors = []

            def fake_pdf(path, headers, rows, title, *, orientation=None, summary_lines=None):
                # Same keyword-only parameters as the real export_contract_pdf_187:
                # _write_export always passes summary_lines (the addressee line)
                # and passes orientation once the preview has chosen one.
                captured.update(headers=headers, rows=rows)

            cst.export_contract_pdf_187 = fake_pdf
            cst.messagebox.showinfo = lambda *a, **k: None
            # _write_export reports every exception through messagebox.showerror,
            # and a real one waits for a click for ever - a stale stub here once
            # held back the whole suite behind an unanswerable box. Record the
            # call instead, so a failure is an assertion with the message in it.
            cst.messagebox.showerror = lambda *a, **k: errors.append(a)
            filedialog.asksaveasfilename = lambda **k: str(Path(tempfile.mkdtemp()) / 'out.pdf')

            on_screen_headers = [label for _key, label in cst.COLUMNS[cst.CONTRACTS]]
            panel.export(cst.CONTRACTS, 'pdf')

            assert not errors, f'the export reported an error instead of writing: {errors}'
            assert captured['headers'] == list(reversed(on_screen_headers))
            assert captured['rows'][0] == list(reversed(list(
                dict(zip(on_screen_headers, panel.load_rows(cst.CONTRACTS)[0][0])).values())))
        finally:
            filedialog.asksaveasfilename = original_asksave
            cst.export_contract_pdf_187 = original_pdf
            cst.messagebox.showinfo = original_showinfo
            cst.messagebox.showerror = original_showerror
            root.destroy()


def test_xlsx_export_sets_right_to_left_sheet_view_and_right_aligned_cells():
    with _temp_store():
        import contractor_store_stage18_7 as store
        import contractors_stage18_7 as cst

        store.upsert_contract({'family_name': 'ابراهیم ازبک', 'title': 'اجرای اسکلت'})

        import tkinter as tk
        from tkinter import ttk, filedialog
        from ui_stage18_7 import style_ttk_187
        from barkav_ui.skin import install_rounded_skin
        from barkav_ui.theme import metrics as kit_metrics, palette as kit_palette
        from openpyxl import load_workbook

        root = tk.Tk()
        root.withdraw()
        original_asksave = filedialog.asksaveasfilename
        original_showinfo = cst.messagebox.showinfo
        original_showerror = cst.messagebox.showerror
        try:
            style_ttk_187(root)
            install_rounded_skin(ttk.Style(root), kit_palette(), kit_metrics(), root)
            panel = cst.ContractorsPanel187(root, db=None, rich=None)
            panel.pack(fill='both', expand=True)
            root.update()

            errors = []
            cst.messagebox.showinfo = lambda *a, **k: None
            # Same as the PDF test above: a real showerror would block the run.
            cst.messagebox.showerror = lambda *a, **k: errors.append(a)
            destination = str(Path(tempfile.mkdtemp()) / 'out.xlsx')
            filedialog.asksaveasfilename = lambda **k: destination
            panel.export(cst.CONTRACTS, 'xlsx')

            assert not errors, f'the export reported an error instead of writing: {errors}'
            workbook = load_workbook(destination)
            sheet = workbook.active
            assert sheet.sheet_view.rightToLeft is True
            # The shared export_xlsx writer lays the sheet out as the addressee
            # line (جهت ارائه به), then its own centred header band, then the
            # rows. Everything that carries content besides the header band
            # stays right-aligned.
            assert [cell.value for cell in sheet[2]] == [
                label for _key, label in cst.COLUMNS[cst.CONTRACTS]]
            assert all(cell.alignment.horizontal == 'right' for cell in sheet[1])
            assert all(cell.alignment.horizontal == 'right' for cell in sheet[3])
        finally:
            filedialog.asksaveasfilename = original_asksave
            cst.messagebox.showinfo = original_showinfo
            cst.messagebox.showerror = original_showerror
            root.destroy()


# ---------------------------------------------------------------------------
# _find_payment_receipt_path: فیش lookup for both pending and posted واریز
# ---------------------------------------------------------------------------
#
# These use a real FinanceDB against a temp sqlite file and the real
# ApprovalCenterService.approve() flow - not mocks - because the bug this
# guards against is exactly a DB-schema/relink chain (attachments losing
# their pending_id once relink_pending_attachments() moves them onto the
# posted EXPENSE/TRANSACTION/CHEQUE entity), which a mock would hide.

def _cheque_free_expense_payload(amount: int) -> dict:
    return {
        'detected_type': 'EXPENSE', 'document_type': 'PAYMENT_RECEIPT',
        'transaction_type_fa': 'هزینه ساخت', 'amount': amount,
        'invoice_amount': amount, 'final_amount': amount, 'discount_amount': 0,
        'description': 'بتن', 'expense_category': 'مصالح', 'construction_stage': 'اسکلت',
        'payer_name': 'پیمان', 'receiver_name': 'بتن جنوب', 'payment_method': 'BANK_TRANSFER',
        'payment_state': 'PAID', 'paid_amount': amount, 'is_cheque': False,
        'manager_fee_eligible': True, 'event_date_shamsi': '1405/05/20', 'confidence': 0.9,
    }


def test_receipt_path_finds_a_still_pending_payments_own_fish():
    """یک واریز تازه/در انتظار بررسی - فیش هنوز entity_type='PENDING' است."""

    import tempfile as _tempfile
    from finance_db import FinanceDB
    from telegram_stage5_workflow import TelegramStage5Workflow
    from contractor_statements_stage18_7 import _find_payment_receipt_path

    with _tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        db = FinanceDB(root / 'db.sqlite3')
        wf = TelegramStage5Workflow(db)
        pending = wf.receive_message(
            chat_id=-1001, message_id=1, raw_text='#test', content_type='PHOTO',
            message_date='1405/05/20 10:30', analysis=_cheque_free_expense_payload(120_000_000),
        )
        receipt = root / 'receipt.jpg'
        receipt.write_bytes(b'fake')
        db.attach_file_to_pending(pending_id=pending.pending_id, local_path=str(receipt),
                                  mime_type='image/jpeg')

        found = _find_payment_receipt_path(db, {'pending_id': pending.pending_id})
        assert found == str(receipt)


def test_receipt_path_follows_relink_to_the_posted_expense():
    """بعد از approve()، فیش دیگر روی PENDING نیست - relink_pending_attachments
    آن را به EXPENSE برده و attachments دیگر pending_id اصلی را نگه نمی‌دارد.
    _find_payment_receipt_path باید از FinanceDB.get_financial_posting همان
    pending_id مسیر EXPENSE تازه را پیدا کند، نه 'بدون فیش'."""

    import tempfile as _tempfile
    from finance_db import FinanceDB
    from approval_service import ApprovalCenterService
    from telegram_stage5_workflow import TelegramStage5Workflow
    from contractor_statements_stage18_7 import _find_payment_receipt_path

    with _tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        db = FinanceDB(root / 'db.sqlite3')
        wf = TelegramStage5Workflow(db)
        pending = wf.receive_message(
            chat_id=-1001, message_id=1, raw_text='#test', content_type='PHOTO',
            message_date='1405/05/20 10:30', analysis=_cheque_free_expense_payload(120_000_000),
        )
        receipt = root / 'receipt.jpg'
        receipt.write_bytes(b'fake')
        db.attach_file_to_pending(pending_id=pending.pending_id, local_path=str(receipt),
                                  mime_type='image/jpeg')

        svc = ApprovalCenterService(db, root)
        approval = svc.approve(pending.pending_id, actor_id='TEST')
        assert approval.posting.expense_id is not None

        # The old best-effort lookup would find nothing here - the whole
        # point of this fix.
        assert db.list_pending_attachments(pending.pending_id) == []

        found = _find_payment_receipt_path(db, {'pending_id': pending.pending_id})
        assert found == str(receipt)


def test_receipt_path_follows_relink_to_a_posted_cheque():
    """یک صورت‌وضعیت/پرداخت با چک: posting فقط cheque_id دارد (نه
    expense_id/transaction_id) - شاخهٔ CHEQUE در _find_payment_receipt_path
    باید همان فیش را پیدا کند."""

    import tempfile as _tempfile
    from finance_db import FinanceDB
    from approval_service import ApprovalCenterService
    from telegram_stage5_workflow import TelegramStage5Workflow
    from contractor_statements_stage18_7 import _find_payment_receipt_path

    with _tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        db = FinanceDB(root / 'db.sqlite3')
        wf = TelegramStage5Workflow(db)
        pending = wf.receive_message(
            chat_id=-1001, message_id=1, raw_text='#test cheque', content_type='PHOTO',
            message_date='1405/05/20 10:30', analysis={
                'detected_type': 'CHEQUE', 'document_type': 'CHEQUE',
                'amount': 50_000_000, 'description': 'چک پیمانکار',
                'payer_name': 'پیمان', 'receiver_name': 'پیمانکار الف',
                'due_date_shamsi': '1405/07/01', 'issue_date_shamsi': '1405/05/20',
                'cheque_serial': '12345', 'issuer_bank': 'ملت', 'confidence': 0.9,
            },
        )
        receipt = root / 'cheque_receipt.jpg'
        receipt.write_bytes(b'fake-cheque')
        db.attach_file_to_pending(pending_id=pending.pending_id, local_path=str(receipt),
                                  mime_type='image/jpeg')

        svc = ApprovalCenterService(db, root)
        approval = svc.approve(pending.pending_id, actor_id='TEST')
        assert approval.posting.cheque_id is not None
        assert approval.posting.expense_id is None

        found = _find_payment_receipt_path(db, {'pending_id': pending.pending_id})
        assert found == str(receipt)


def test_receipt_path_is_empty_without_crashing_when_nothing_matches():
    """بدون pending_id، pending_id ناموجود، یا db=None - همیشه '', نه خطا."""

    import tempfile as _tempfile
    from finance_db import FinanceDB
    from contractor_statements_stage18_7 import _find_payment_receipt_path

    assert _find_payment_receipt_path(None, {'pending_id': 'PEN-X'}) == ''
    assert _find_payment_receipt_path(object(), {}) == ''

    with _tempfile.TemporaryDirectory() as tmp:
        db = FinanceDB(Path(tmp) / 'db.sqlite3')
        assert _find_payment_receipt_path(db, {'pending_id': 'PEN-DOES-NOT-EXIST'}) == ''


def test_contractor_payment_schedule_has_fixed_ends_and_preserves_typed_values():
    import sys
    sys.path.insert(0, str(ROOT))
    from contractor_contracts_stage18_7 import (
        normalise_payment_schedule_187,
        payment_stage_labels_187,
    )

    assert payment_stage_labels_187(4) == (
        'پیش پرداخت', 'مرحله اول', 'مرحله دوم', 'تحویل کار',
    )
    schedule = normalise_payment_schedule_187(4, [
        {'label': 'پیش پرداخت', 'due_date_shamsi': '1405/01/10', 'amount': '10000000'},
        {'label': 'تحویل کار', 'due_date_shamsi': '1405/05/20', 'amount': '50000000'},
    ])
    assert [row['label'] for row in schedule] == list(payment_stage_labels_187(4))
    assert schedule[0]['due_date_shamsi'] == '1405/01/10'
    assert schedule[-1]['amount'] == '50000000'
