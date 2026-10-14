"""Actual Tk dialog tests against a disposable project (Xvfb or Windows)."""
import os
import time
import json
import tkinter as tk
from tkinter import ttk
from pathlib import Path
import pytest
from test_advanced_search187 import search, document
from ai_section_models187 import AIModelPicker187, section_choice, MODES
from advanced_search_dialog187 import AdvancedSearchDialog187
from ui_stage18_7 import style_ttk_187

pytestmark=pytest.mark.skipif(os.environ.get('BARKAV_RUN_UI_TESTS')!='1',reason='Tk display opt-in')

@pytest.fixture
def root():
    root=tk.Tk();root.withdraw();style_ttk_187(root)
    yield root
    root.destroy()


def walk(widget):
    for child in widget.winfo_children():
        yield child
        yield from walk(child)


def test_search_entry_context_clipboard_filters_and_result_table(root,search):
    document(search,'D1','DAILY_REPORT','ناصر — کارگاه')
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    root.update()
    assert window.scope.get()=='همه' and window.date_basis.get()=='همه'
    entries=[widget for widget in walk(window) if isinstance(widget,ttk.Entry)]
    assert all(widget.cget('justify')=='right' for widget in entries)
    assert any('BarkavFieldClipboard' in widget.bindtags() for widget in entries)
    window.query.set('ناصر');window.run()
    deadline=time.monotonic()+5
    while window._running and time.monotonic()<deadline:root.update();time.sleep(.01)
    assert not window._running
    assert any('ناصر' in str(row.get('شرح')) for row in window.rows)
    window.ai.set(True);window._toggle_ai();root.update()
    assert window.models.winfo_ismapped()
    window.destroy()


def test_model_dropdown_colours_and_group_modes_disable_checkboxes(root,search):
    picker=AIModelPicker187(root,section='management',base_dir=search.base_dir)
    picker.pack();root.update();picker.open();root.update()
    checks=[widget for widget in walk(picker.popup) if isinstance(widget,tk.Checkbutton)]
    assert checks and all(widget.cget('state')=='disabled' for widget in checks)
    assert {widget.cget('fg') for widget in checks}<={'#20ad79','#e36363'}
    mode=next(widget for widget in walk(picker.popup) if isinstance(widget,ttk.Combobox))
    mode.set(MODES['SELECTED']);root.update()
    assert all(widget.cget('state')=='normal' for widget in checks)
    checks[0].invoke();checks[1].invoke()
    next(widget for widget in walk(picker.popup) if isinstance(widget,ttk.Button) and widget.cget('text')=='تأیید انتخاب').invoke()
    assert len(picker.choice['models'])==2
    assert section_choice('management')['models']==picker.choice['models']
    assert section_choice('data_entry')['mode']=='AUTO'
    assert section_choice('calendar')['mode']=='AUTO'


def test_export_formats_open_same_preview_and_retain_footer_rows(root,search,monkeypatch):
    import advanced_search_dialog187 as module
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window._display(['ردیف','شرح','مبلغ (تومان)'],[{'ردیف':1,'شرح':'بتن','مبلغ (تومان)':1000},
        {'ردیف':'جمع','شرح':'هزینه ثبت‌شده','مبلغ (تومان)':1000}])
    captures=[]
    monkeypatch.setattr(module,'show_export_preview_187',lambda *args,**kwargs:captures.append(kwargs))
    for choice in ('PDF','Excel','پرینت'):window.preview(choice)
    assert [value['export_format'] for value in captures]==['PDF','Excel','پرینت']
    assert all(value['rows'][-1]['ردیف']=='جمع' for value in captures)
    assert all(value['include_addressee'] for value in captures)
    window.destroy()


def test_period_report_renders_in_main_table_with_management_selection(root,search,monkeypatch):
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window.start.set('1405/07/01');window.end.set('1405/07/07');window.period.set('هفتگی')
    captured={}
    def table(label,first,last,**kwargs):
        captured.update(label=label,first=first,last=last,**kwargs)
        return ['ردیف','شرح گزارش'],[{'ردیف':1,'شرح گزارش':'گزارش ثبت‌شده'}]
    monkeypatch.setattr(window.search_service,'period_table',table)
    window.period_report()
    deadline=time.monotonic()+5
    while window._running and time.monotonic()<deadline:root.update();time.sleep(.01)
    assert captured['label']=='هفتگی' and window._ai_section187=='management'
    assert captured['first']=='1405/07/01' and window.rows[0]['شرح گزارش']=='گزارش ثبت‌شده'
    window.destroy()


def test_request_modes_and_faq_selection(root,search,monkeypatch):
    from advanced_search_requests187 import REQUEST_TYPES
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    root.update()
    assert window.request_type.get()==REQUEST_TYPES[0] and 'سایر' in window.scope_combo.cget('values')
    window.request_type.set(REQUEST_TYPES[2]);window._mode_changed();root.update()
    assert window.scope_label.cget('text')=='درخواست از' and window.faq_combo.winfo_ismapped()
    assert window.query_entry.winfo_ismapped() and window.models.winfo_ismapped()
    assert not window.filter_widgets[2][1].winfo_ismapped()
    window.scope.set('پیام‌ها');window._update_questions()
    assert len(window.faq_combo.cget('values'))==4
    calls=[];monkeypatch.setattr(window,'run',lambda:calls.append(window.query.get()))
    window.faq.set(window.faq_combo.cget('values')[0]);window._choose_question()
    assert calls and 'خوانده' in calls[0]
    window.request_type.set(REQUEST_TYPES[1]);window._mode_changed();root.update()
    assert window.period_host.winfo_ismapped() and not window.query_entry.winfo_ismapped()
    window.request_type.set(REQUEST_TYPES[0]);window._mode_changed();root.update()
    assert window.query_entry.winfo_ismapped() and window.query_caption.winfo_ismapped()
    window.destroy()


def test_ai_failure_is_inline_and_preserves_source_evidence(root,search,monkeypatch):
    from advanced_search_requests187 import REQUEST_TYPES
    document(search,'D1','DAILY_REPORT','حضور ناصر')
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window.request_type.set(REQUEST_TYPES[2]);window._mode_changed();window.scope.set('تقویم');window.query.set('ناصر')
    monkeypatch.setattr(window.search_service,'ask',lambda *a,**kw:(_ for _ in ()).throw(RuntimeError('private payload')))
    window.run()
    deadline=time.monotonic()+5
    while window._running and time.monotonic()<deadline:root.update();time.sleep(.01)
    assert not window._running and window.rows[-1]['ردیف']=='وضعیت'
    assert any('ناصر' in row.get('شرح','') for row in window.rows)
    assert 'private payload' not in str(window.rows)
    assert str(window.run_button.cget('state'))=='normal'
    window.destroy()


def test_complete_report_switch_restores_all_text_without_limiting_export(root,search):
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    content='متن گزارش '*100
    window._display(['ردیف','شرح'],[{'ردیف':1,'شرح':content}])
    assert window.rows[0]['شرح']==content
    assert window.page_body.winfo_children()[0].rows[0]['شرح'].endswith('…')
    window.full_report.set(True);window._full_changed()
    assert window.page_body.winfo_children()[0].rows[0]['شرح']==content
    window.destroy()


def test_export_uses_result_snapshot_and_saves_readable_excel_and_pdf(root,search,tmp_path,monkeypatch):
    import advanced_search_dialog187 as module
    from openpyxl import load_workbook
    import pypdfium2
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window.result_metadata={'query':'سیمان','start':'1405/07/01','end':'1405/07/03','scope':'تراکنشها',
        'date_basis':'همه','audience':'کارفرما','documents':False}
    window._display(['ردیف','شرح','مبلغ (تومان)'],[{'ردیف':1,'شرح':'سیمان','مبلغ (تومان)':1000},
        {'ردیف':'جمع','شرح':'جمع هزینه','مبلغ (تومان)':1000}])
    captures=[]
    monkeypatch.setattr(module,'show_export_preview_187',lambda *a,**kw:captures.append(kw))
    window.query.set('پرسش جدید بدون جستجو')
    window.preview('Excel')
    assert captures[-1]['specs'][0][2]=='سیمان'
    monkeypatch.setattr(module.filedialog,'asksaveasfilename',lambda **kw:str(tmp_path/('report'+kw['defaultextension'])))
    assert captures[-1]['on_confirm']('Excel','افقی',False)
    book=load_workbook(tmp_path/'report.xlsx')
    assert book.active.sheet_view.rightToLeft
    assert any('سیمان' in str(cell.value) for row in book.active for cell in row)
    assert len(book.active._images)>0
    book.close()
    assert captures[-1]['on_confirm']('PDF','افقی',False)
    pdf=pypdfium2.PdfDocument(str(tmp_path/'report.pdf'));assert len(pdf)>=1;pdf.close()
    window.destroy()


def test_period_ai_question_starts_existing_report_directly(root,search,monkeypatch):
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window.ai.set(True);window.query.set('گزارش هفتگی هفته اول مهرماه 1405')
    calls=[]
    monkeypatch.setattr(window,'period_report',lambda **kw:calls.append(kw))
    window.run()
    assert calls==[{'generate':True}]
    assert window.period.get()=='هفتگی' and window.start.get()=='۱۴۰۵/۰۷/۰۱'
    window.destroy()


def test_filters_and_ai_controls_fit_minimum_window_without_clipping(root,search):
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    window.ai.set(True);window._toggle_ai()
    window.geometry('950x560');root.update()
    left=window.winfo_rootx();right=left+window.winfo_width()
    for widget in walk(window):
        if isinstance(widget,(ttk.Entry,ttk.Combobox,ttk.Checkbutton)) and widget.winfo_ismapped():
            assert widget.winfo_rootx()>=left
            assert widget.winfo_rootx()+widget.winfo_width()<=right
    assert window.models.winfo_ismapped()
    window.destroy()


def test_paging_keeps_all_result_rows_for_exports(root,search,monkeypatch):
    import advanced_search_dialog187 as module
    from calendar_report_tables_187 import ReportTable187
    window=AdvancedSearchDialog187(root,service=search.service,rich=search.rich)
    rows=[{'ردیف':i+1,'تاریخ':'1405/07/01','شرح':f'رکورد {i+1}'} for i in range(205)]
    rows.append({'ردیف':'جمع','شرح':'تعداد','تعداد':205})
    window._display(['ردیف','تاریخ','شرح','تعداد'],rows);root.update()
    table=next(child for child in window.page_body.winfo_children() if isinstance(child,ReportTable187))
    assert len(table.rows)==100 and len(window.rows)==206
    window._page(2);root.update()
    table=next(child for child in window.page_body.winfo_children() if isinstance(child,ReportTable187))
    assert len(table.rows)==6 and table.rows[-1]['ردیف']=='جمع'
    captures=[]
    monkeypatch.setattr(module,'show_export_preview_187',lambda *a,**kw:captures.append(kw))
    window.preview('Excel')
    assert len(captures[0]['rows'])==206
    assert captures[0]['rows'][-1]['تعداد']=='۲۰۵'
    window.destroy()
