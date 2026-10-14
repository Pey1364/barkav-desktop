"""Service-level tests for the pieces the «بروزرسانی قیمت» tab grew later:
the real دسته‌بندی->فاز cascade, a scoped run, «اعمال قوانین», and the
named/scoped reference-workbook export.

Headless on purpose (no Tk): these are decisions about data, and the matching
GUI-visibility tests live in ``test_stage18_7_ui.py`` where a real dashboard
exists to measure geometry against.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

import price_calc_engine as calc
from finance_db import FinanceDB
from price_update_service import PriceUpdateService


@pytest.fixture()
def service(tmp_path: Path) -> PriceUpdateService:
    db = FinanceDB(tmp_path / 'test.sqlite3')
    return PriceUpdateService(db, tmp_path)


# --------------------------------------------------------- cascading filter

def test_available_phase_codes_is_scoped_to_the_chosen_catalog(service: PriceUpdateService):
    """دسته‌بندی -> فاز must be a real cascade: only phases that category has rows in."""

    service.upsert_catalog_item(None, description='سیمان', catalog_type='MATERIAL',
                                estimate_rate=100_000, phase_codes='08')
    service.upsert_catalog_item(None, description='لوازم مشترک', catalog_type='MATERIAL',
                                estimate_rate=50_000, phase_codes='15,16')
    service.upsert_catalog_item(None, description='اجرای اسکلت', catalog_type='CONTRACTOR',
                                estimate_rate=900_000, phase_codes='08')
    service.upsert_catalog_item(None, description='نقشه‌کشی', catalog_type='SERVICE',
                                estimate_rate=10_000, phase_codes='03')

    assert service.available_phase_codes('MATERIAL') == ['08', '15', '16']
    assert service.available_phase_codes('CONTRACTOR') == ['08']
    assert service.available_phase_codes('SERVICE') == ['03']
    assert service.available_phase_codes() == ['03', '08', '15', '16']


def test_available_phase_codes_drops_a_soft_deleted_row(service: PriceUpdateService):
    price_id = service.upsert_catalog_item(None, description='قلم موقت', catalog_type='MATERIAL',
                                            estimate_rate=1_000, phase_codes='24')
    assert '24' in service.available_phase_codes('MATERIAL')
    service.soft_delete_catalog_item(price_id)
    assert '24' not in service.available_phase_codes('MATERIAL')


# ------------------------------------------------------------- scoped run

def test_run_update_scope_leaves_rows_outside_it_untouched(service: PriceUpdateService):
    """A scoped run is a subset run: out-of-scope rows keep their price and stay out of the snapshot."""

    service.save_parameter('usd', base_value=50_000, current_value=100_000)

    inside = service.upsert_catalog_item(
        None, description='داخل دامنه', catalog_type='MATERIAL', estimate_rate=100_000,
        phase_codes='08', calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [150_000]}))
    outside = service.upsert_catalog_item(
        None, description='بیرون دامنه', catalog_type='MATERIAL', estimate_rate=100_000,
        phase_codes='24', calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [150_000]}))

    result = service.run_update(trigger='manual', estimate_date_shamsi='1405/06/01',
                                phase_code='08')

    assert result.item_count == 1
    assert service.get_catalog_item(inside)['forecast_rate'] == 150_000
    # untouched: still the seeded بروز = اولیه, and never written into the snapshot
    assert service.get_catalog_item(outside)['forecast_rate'] == 100_000
    assert {r['price_id'] for r in service.snapshot_items(result.snapshot_id)} == {inside}


def test_run_update_archives_its_scope_in_the_snapshot(service: PriceUpdateService):
    service.upsert_catalog_item(None, description='قلم', catalog_type='SERVICE',
                                estimate_rate=10_000, phase_codes='03')
    scoped = service.run_update(estimate_date_shamsi='1405/06/01',
                                catalog_type='SERVICE', phase_code='03')
    totals = json.loads(next(s for s in service.list_snapshots()
                             if s['snapshot_id'] == scoped.snapshot_id)['totals_json'])
    assert totals['scope_catalog_type'] == 'SERVICE'
    assert totals['scope_phase_code'] == '03'

    full = service.run_update(estimate_date_shamsi='1405/07/01')
    full_totals = json.loads(next(s for s in service.list_snapshots()
                                  if s['snapshot_id'] == full.snapshot_id)['totals_json'])
    assert full_totals['scope_catalog_type'] == ''
    assert full_totals['scope_phase_code'] == ''


# ------------------------------------------------------------- apply rules

def test_apply_calc_rules_classifies_matching_rows_only(service: PriceUpdateService):
    """«اعمال قوانین» writes روش + ضرایب onto the rows a rule covers, and nothing else."""

    covered = service.upsert_catalog_item(None, description='مصالح فاز ۸', catalog_type='MATERIAL',
                                           estimate_rate=100_000, phase_codes='08')
    other_phase = service.upsert_catalog_item(None, description='مصالح فاز ۲۴', catalog_type='MATERIAL',
                                               estimate_rate=100_000, phase_codes='24')
    other_type = service.upsert_catalog_item(None, description='پیمانکار فاز ۸', catalog_type='CONTRACTOR',
                                              estimate_rate=100_000, phase_codes='08')

    service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='مصالح فاز ۸ با روش وب',
        affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_WEB, 'carry_rate': 0.3}))

    assert service.apply_calc_rules() == {'rules': 1, 'items': 1, 'rules_without_method': 0}

    applied = service.get_catalog_item(covered)
    assert applied['calc_method'] == calc.METHOD_WEB
    assert applied['suggested_method'] == calc.METHOD_WEB
    params = json.loads(applied['calc_params_json'])
    assert params['carry_rate'] == 0.3
    # 'method' is the rule's own key; it is never copied into a row's params
    assert 'method' not in params
    for untouched in (other_phase, other_type):
        assert service.get_catalog_item(untouched)['calc_method'] == ''


def test_apply_calc_rules_with_no_phase_codes_covers_the_whole_catalog_type(service: PriceUpdateService):
    first = service.upsert_catalog_item(None, description='الف', catalog_type='SERVICE',
                                         estimate_rate=1_000, phase_codes='03')
    second = service.upsert_catalog_item(None, description='ب', catalog_type='SERVICE',
                                          estimate_rate=2_000, phase_codes='31')
    service.upsert_calc_rule(None, catalog_type='SERVICE', description='همه خدمات',
                              params_json=json.dumps({'method': calc.METHOD_SERVICE_TARIFF}))

    assert service.apply_calc_rules()['items'] == 2
    for price_id in (first, second):
        assert service.get_catalog_item(price_id)['calc_method'] == calc.METHOD_SERVICE_TARIFF


def test_apply_calc_rules_merges_params_instead_of_replacing_them(service: PriceUpdateService):
    """A weights-only rule must not wipe out a row's hand-entered web prices."""

    price_id = service.upsert_catalog_item(
        None, description='قلم', catalog_type='MATERIAL', estimate_rate=100_000, phase_codes='08',
        calc_params_json=json.dumps({'web_prices': [140_000], 'carry_rate': 0.1}))
    service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='وزن‌ها', affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_WEIGHTED_INDEX, 'carry_rate': 0.3}))

    service.apply_calc_rules()
    params = json.loads(service.get_catalog_item(price_id)['calc_params_json'])
    assert params['web_prices'] == [140_000]   # kept
    assert params['carry_rate'] == 0.3         # the rule wins where the two overlap


def test_apply_calc_rules_reports_a_rule_with_no_method(service: PriceUpdateService):
    """A method-less rule is a silent no-op for the next run, so it must be counted."""

    price_id = service.upsert_catalog_item(None, description='قلم', catalog_type='MATERIAL',
                                           estimate_rate=100_000, phase_codes='08')
    service.upsert_calc_rule(None, catalog_type='MATERIAL', description='فقط ضرایب',
                              affected_phase_codes='08', params_json=json.dumps({'tax_rate': 0.1}))

    result = service.apply_calc_rules()
    assert result['rules_without_method'] == 1
    assert result['items'] == 1
    row = service.get_catalog_item(price_id)
    assert row['calc_method'] == ''   # still unclassified; run_update will skip it
    assert json.loads(row['calc_params_json'])['tax_rate'] == 0.1


def test_apply_calc_rules_applies_only_the_selected_rules(service: PriceUpdateService):
    price_id = service.upsert_catalog_item(None, description='قلم', catalog_type='MATERIAL',
                                           estimate_rate=100_000, phase_codes='08')
    wanted = service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='انتخاب‌شده', affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_WEB}))
    service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='انتخاب‌نشده', affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_SERVICE_TARIFF}))

    assert service.apply_calc_rules([wanted])['rules'] == 1
    assert service.get_catalog_item(price_id)['calc_method'] == calc.METHOD_WEB


def test_apply_calc_rules_lets_the_most_recently_edited_rule_win(service: PriceUpdateService):
    """Two rules covering one row: the newest edit is the last writer, not the first."""

    price_id = service.upsert_catalog_item(None, description='قلم', catalog_type='MATERIAL',
                                           estimate_rate=100_000, phase_codes='08')
    older = service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='قدیمی', affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_WEB}))
    newer = service.upsert_calc_rule(
        None, catalog_type='MATERIAL', description='جدید', affected_phase_codes='08',
        params_json=json.dumps({'method': calc.METHOD_WEIGHTED_INDEX}))

    # _now() is second-precision, so two rules created in one test tie on
    # updated_at and re-saving one would not separate them. The timestamps are
    # set directly instead, which is what "edited a minute apart" looks like in
    # the real table.
    with service.db.tx() as con:
        con.execute('UPDATE price_calc_rules SET updated_at=? WHERE rule_id=?',
                    ('2026-09-01T10:00:00', older))
        con.execute('UPDATE price_calc_rules SET updated_at=? WHERE rule_id=?',
                    ('2026-09-01T10:01:00', newer))
    assert service.list_calc_rules()[0]['rule_id'] == newer

    service.apply_calc_rules()
    assert service.get_catalog_item(price_id)['calc_method'] == calc.METHOD_WEIGHTED_INDEX


def test_list_calc_rules_order_is_stable_when_updated_at_ties(service: PriceUpdateService):
    """Same-second edits must not reorder between refreshes -- «اعمال قوانین» depends on the order."""

    for index in range(5):
        service.upsert_calc_rule(None, catalog_type='MATERIAL', description=f'قانون {index}',
                                  params_json=json.dumps({'method': calc.METHOD_WEB}))
    with service.db.tx() as con:
        con.execute('UPDATE price_calc_rules SET updated_at=?', ('2026-09-01T10:00:00',))
    first = [r['rule_id'] for r in service.list_calc_rules()]
    assert first == [r['rule_id'] for r in service.list_calc_rules()]
    assert len(set(first)) == 5


def test_apply_calc_rules_then_run_update_actually_moves_a_price(service: PriceUpdateService):
    """End to end: a fresh unclassified row only escalates once a rule has classified it."""

    service.save_parameter('usd', base_value=50_000, current_value=100_000)
    price_id = service.upsert_catalog_item(
        None, description='سیمان', catalog_type='MATERIAL', estimate_rate=100_000,
        phase_codes='08', calc_params_json=json.dumps({'web_prices': [130_000]}))

    service.run_update(estimate_date_shamsi='1405/06/01')
    assert service.get_catalog_item(price_id)['forecast_rate'] == 100_000   # untouched

    service.upsert_calc_rule(None, catalog_type='MATERIAL', description='وب',
                              affected_phase_codes='08',
                              params_json=json.dumps({'method': calc.METHOD_WEB}))
    service.apply_calc_rules()
    service.run_update(estimate_date_shamsi='1405/07/01')
    assert service.get_catalog_item(price_id)['forecast_rate'] == 130_000


# ------------------------------------------------------ named/scoped export

def test_export_reference_workbook_honours_an_explicit_path_and_scope(service: PriceUpdateService):
    service.upsert_catalog_item(None, description='مصالح', catalog_type='MATERIAL',
                                estimate_rate=100_000, phase_codes='08')
    service.upsert_catalog_item(None, description='خدمات', catalog_type='SERVICE',
                                estimate_rate=10_000, phase_codes='03')

    destination = service.base_dir / 'sub' / 'chosen name.xlsx'
    assert service.export_reference_workbook(out_path=destination) == destination

    destination.unlink()
    service.export_reference_workbook(out_path=destination, catalog_type='MATERIAL')
    workbook = openpyxl.load_workbook(destination)
    assert 'قیمت مصالح' in workbook.sheetnames
    assert 'قیمت خدمات' not in workbook.sheetnames   # scoped out
    assert 'پارامترها' in workbook.sheetnames        # parameters always travel along


def test_scoped_export_still_round_trips_back_in(service: PriceUpdateService):
    """A narrowed workbook is a valid round trip: it updates its own rows and no others."""

    material = service.upsert_catalog_item(None, description='مصالح', catalog_type='MATERIAL',
                                            code='M-1', estimate_rate=100_000, phase_codes='08')
    service_row = service.upsert_catalog_item(None, description='خدمات', catalog_type='SERVICE',
                                               estimate_rate=10_000, phase_codes='03')
    destination = service.base_dir / 'material only.xlsx'
    service.export_reference_workbook(out_path=destination, catalog_type='MATERIAL')

    workbook = openpyxl.load_workbook(destination)
    workbook['قیمت مصالح'].cell(row=4, column=4).value = 125_000   # برآورد اولیه, edited in Excel
    workbook.save(destination)
    service.import_reference_workbook(destination)

    assert service.get_catalog_item(material)['estimate_rate'] == 125_000
    assert service.get_catalog_item(service_row)['estimate_rate'] == 10_000
    assert len(service.list_catalog_items()) == 2   # nothing duplicated


def test_run_update_for_one_item_writes_only_that_item_but_uses_its_siblings(service: PriceUpdateService):
    """«قلم» in the wizard: one row is updated, against the whole catalog's context."""

    service.save_parameter('usd', base_value=50_000, current_value=100_000)
    web_id = service.upsert_catalog_item(
        None, description='کاشی', catalog_type='MATERIAL', phase_codes='22', estimate_rate=100_000,
        forecast_rate=100_000, calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [130_000]}))
    target_id = service.upsert_catalog_item(
        None, description='چسب کاشی', catalog_type='MATERIAL', phase_codes='22', estimate_rate=50_000,
        forecast_rate=50_000, calc_method=calc.METHOD_PHASE_AVERAGE)

    result = service.run_update(price_id=target_id)

    assert result.item_count == 1 and result.recalculated_count == 1
    # Priced from its sibling's web ratio, which a one-row pool would not have had.
    assert service.get_catalog_item(target_id)['forecast_rate'] == 65_000
    assert service.get_catalog_item(web_id)['forecast_rate'] == 100_000  # not written
    assert [row['price_id'] for row in service.snapshot_items(result.snapshot_id)] == [target_id]
