from __future__ import annotations
import os

import json
import tempfile
from pathlib import Path

import openpyxl
import pytest

import price_calc_engine as calc
from budget_rich_v163 import RichBudgetService
from finance_db import FinanceDB
from price_update_service import PriceUpdateService


@pytest.fixture()
def service(tmp_path: Path) -> PriceUpdateService:
    db = FinanceDB(tmp_path / 'test.sqlite3')
    return PriceUpdateService(db, tmp_path)


def test_migrate_seeds_default_parameters(service: PriceUpdateService):
    keys = {p['key'] for p in service.list_parameters()}
    assert keys == {'usd', 'copper', 'wage_annual_growth', 'months_elapsed'}
    for p in service.list_parameters():
        assert p['current_value'] == 0


def test_catalog_item_crud(service: PriceUpdateService):
    price_id = service.upsert_catalog_item(
        None, description='سیمان شاهرود', catalog_type='MATERIAL', unit='کیسه',
        estimate_rate=100_000, phase_codes='08',
    )
    row = service.get_catalog_item(price_id)
    assert row['description'] == 'سیمان شاهرود'
    assert row['estimate_rate'] == 100_000

    service.upsert_catalog_item(price_id, estimate_rate=120_000)
    assert service.get_catalog_item(price_id)['estimate_rate'] == 120_000

    copy_id = service.copy_catalog_item(price_id)
    assert copy_id != price_id
    assert service.get_catalog_item(copy_id)['description'].endswith('(رونوشت)')

    service.soft_delete_catalog_item(price_id)
    assert service.get_catalog_item(price_id)['active'] == 0
    assert price_id not in [r['price_id'] for r in service.list_catalog_items()]


def test_upsert_catalog_item_rejects_unknown_field(service: PriceUpdateService):
    with pytest.raises(ValueError):
        service.upsert_catalog_item(None, description='x', catalog_type='MATERIAL', bogus_field=1)


def test_save_parameter_tracks_base_previous_current(service: PriceUpdateService):
    service.save_parameter('usd', base_value=150_000, current_value=150_000, source='manual')
    service.save_parameter('usd', current_value=232_500, source='tgju')
    p = service.get_parameter('usd')
    assert p['base_value'] == 150_000
    assert p['previous_value'] == 150_000  # carried from the value current_value replaced
    assert p['current_value'] == 232_500
    assert p['source'] == 'tgju'


def test_save_parameter_rejects_unknown_key(service: PriceUpdateService):
    with pytest.raises(ValueError):
        service.save_parameter('not_a_real_key', current_value=1)


def test_calc_rule_crud(service: PriceUpdateService):
    rule_id = service.upsert_calc_rule(
        None, catalog_type='CONTRACTOR', description='کارگری خالص',
        params_json=json.dumps({'wage_weight': 0.6, 'adjustment_weight': 0.25, 'inflation_weight': 0.15}),
    )
    rules = service.list_calc_rules('CONTRACTOR')
    assert len(rules) == 1 and rules[0]['rule_id'] == rule_id
    service.delete_calc_rule(rule_id)
    assert service.list_calc_rules('CONTRACTOR') == []


def _set_params(service: PriceUpdateService, *, usd_base=150_000, usd_current=232_500,
                 copper_base=2_003_030, copper_current=2_603_290,
                 wage_annual=0.60, months=5):
    service.save_parameter('usd', base_value=usd_base, current_value=usd_current)
    service.save_parameter('copper', base_value=copper_base, current_value=copper_current)
    service.save_parameter('wage_annual_growth', base_value=wage_annual, current_value=wage_annual)
    service.save_parameter('months_elapsed', base_value=0, current_value=months)


def test_run_update_web_row_uses_its_own_ratio(service: PriceUpdateService):
    _set_params(service)
    price_id = service.upsert_catalog_item(
        None, description='میلگرد', catalog_type='MATERIAL', unit='کیلوگرم',
        estimate_rate=100_000, phase_codes='08', calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [140_000, 142_000], 'carry_rate': 0.035, 'tax_rate': 0.10}),
    )
    result = service.run_update(trigger='manual', estimate_date_shamsi='1405/07/01')
    assert result.item_count == 1
    assert result.recalculated_count == 1
    item = service.get_catalog_item(price_id)
    avg = (140_000 + 142_000) / 2
    expected_rate = round(avg * 1.035 * 1.10)
    assert item['forecast_rate'] == expected_rate
    assert item['previous_rate'] == 100_000  # what forecast_rate was before this run


def test_run_update_leaves_unclassified_rows_unchanged(service: PriceUpdateService):
    _set_params(service)
    price_id = service.upsert_catalog_item(
        None, description='قلم دسته‌بندی‌نشده', catalog_type='MATERIAL',
        estimate_rate=50_000, forecast_rate=55_000,
    )
    result = service.run_update(trigger='manual')
    assert result.recalculated_count == 0
    item = service.get_catalog_item(price_id)
    assert item['forecast_rate'] == 55_000  # untouched, never zeroed


def test_run_update_phase_average_fallback_uses_positive_web_ratios_in_same_phase(service: PriceUpdateService):
    _set_params(service)
    service.upsert_catalog_item(
        None, description='قلم وب‌دار فاز ۰۶ (۱)', catalog_type='MATERIAL', phase_codes='06',
        estimate_rate=100_000, calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [130_000]}),
    )
    fallback_id = service.upsert_catalog_item(
        None, description='کانکس', catalog_type='MATERIAL', phase_codes='06',
        estimate_rate=200_000, calc_method=calc.METHOD_PHASE_AVERAGE,
    )
    service.run_update(trigger='manual')
    fallback = service.get_catalog_item(fallback_id)
    assert fallback['forecast_rate'] == round(200_000 * 1.30)
    # Priced by phase average by design, not because an own ratio was
    # rejected -- the reference workbook's «لیست بررسی مجدد» does not list
    # such rows, so neither does the app.
    assert fallback['needs_review'] == 0
    assert fallback['executed_method'] == calc.METHOD_PHASE_AVERAGE


def test_run_update_rejected_web_ratio_takes_the_phase_average_and_is_flagged(service: PriceUpdateService):
    """«ممنوعیت ضریب کاهشی»: a web price below اولیه is never applied."""

    _set_params(service)
    service.upsert_catalog_item(
        None, description='سرامیک ۶۰×۳۰', catalog_type='MATERIAL', phase_codes='22',
        estimate_rate=100_000, calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [120_000]}))
    falling_id = service.upsert_catalog_item(
        None, description='سرامیک ۱۲۰×۶۰', catalog_type='MATERIAL', phase_codes='22',
        estimate_rate=400_000, calc_method=calc.METHOD_WEB,
        calc_params_json=json.dumps({'web_prices': [285_000]}))
    service.run_update(trigger='manual')
    falling = service.get_catalog_item(falling_id)
    assert falling['forecast_rate'] == round(400_000 * 1.20)
    assert falling['needs_review'] == 1


def test_run_update_contractor_weighted_index_floors_at_wage_growth(service: PriceUpdateService):
    _set_params(service)
    price_id = service.upsert_catalog_item(
        None, description='گمانه آزمایش خاک', catalog_type='CONTRACTOR', phase_codes='02',
        estimate_rate=260_000, calc_method=calc.METHOD_WEIGHTED_INDEX,
        calc_params_json=json.dumps({
            'wage_weight': 0.60, 'adjustment_weight': 0.25, 'inflation_weight': 0.15,
            'dollar_or_metal_weight': 0.0, 'adjustment_index': 1.0, 'inflation_rate': 1.0,
        }),
    )
    service.run_update(trigger='manual')
    item = service.get_catalog_item(price_id)
    wage_growth = calc.wage_growth_rate(0.60, 5)
    assert item['forecast_rate'] == round(260_000 * wage_growth)  # floor wins: 0.6+0.25+0.15 index is below it
    assert item['needs_review'] == 1


def test_run_update_creates_a_never_overwritten_snapshot(service: PriceUpdateService):
    _set_params(service)
    service.upsert_catalog_item(
        None, description='میلگرد', catalog_type='MATERIAL', estimate_rate=100_000,
        calc_method=calc.METHOD_WEB, calc_params_json=json.dumps({'web_prices': [150_000]}),
    )
    first = service.run_update(trigger='manual', estimate_date_shamsi='1405/06/01')
    _set_params(service, usd_current=250_000, months=6)
    second = service.run_update(trigger='manual', estimate_date_shamsi='1405/07/01')
    assert first.snapshot_id != second.snapshot_id
    snapshots = service.list_snapshots()
    assert {s['snapshot_id'] for s in snapshots} == {first.snapshot_id, second.snapshot_id}
    first_items = service.snapshot_items(first.snapshot_id)
    second_items = service.snapshot_items(second.snapshot_id)
    assert len(first_items) == 1 and len(second_items) == 1
    # The first snapshot's own frozen numbers never change once the second run happens.
    assert first_items[0]['computed_rate'] == round(150_000)


def test_needs_review_items_lists_only_flagged_active_rows(service: PriceUpdateService):
    _set_params(service)
    # A web price below اولیه: rejected, replaced by the fallback, flagged.
    flagged_id = service.upsert_catalog_item(
        None, description='کانکس', catalog_type='MATERIAL', phase_codes='06', estimate_rate=100_000,
        calc_method=calc.METHOD_WEB, calc_params_json=json.dumps({'web_prices': [90_000]}),
    )
    clean_id = service.upsert_catalog_item(
        None, description='میلگرد', catalog_type='MATERIAL', estimate_rate=100_000,
        calc_method=calc.METHOD_WEB, calc_params_json=json.dumps({'web_prices': [101_000]}),
    )
    service.run_update(trigger='manual')
    flagged = {r['price_id'] for r in service.needs_review_items()}
    assert flagged_id in flagged
    assert clean_id not in flagged


def test_sync_forecast_to_rich_only_updates_linked_items_by_value(tmp_path: Path):
    rich_db = FinanceDB(tmp_path / 'rich.sqlite3')
    RichBudgetService(rich_db, tmp_path)  # only needed to create the rich_price_items table
    with rich_db.tx() as con:
        con.execute(
            "INSERT INTO rich_price_items(price_id,code,description,category,unit,estimate_rate,"
            "actual_rate,forecast_rate,status,source,visible,aggregate,group_key,active,created_at,updated_at) "
            "VALUES('RICH-1','','سیمان','MATERIAL','کیسه',100000,100000,100000,'','',1,0,'',1,'2026-01-01','2026-01-01')"
        )

    price_db = FinanceDB(tmp_path / 'priceupdate.sqlite3')
    service = PriceUpdateService(price_db, tmp_path)
    linked_id = service.upsert_catalog_item(
        None, description='سیمان', catalog_type='MATERIAL', estimate_rate=100_000,
        forecast_rate=140_000, rich_price_id='RICH-1',
    )
    service.upsert_catalog_item(None, description='قلم بدون لینک', catalog_type='MATERIAL',
                                 estimate_rate=50_000, forecast_rate=60_000)

    updated = service.sync_forecast_to_rich(rich_db)
    assert updated == 1
    with rich_db.read() as con:
        row = con.execute('SELECT forecast_rate,category,description FROM rich_price_items WHERE price_id=?',
                           ('RICH-1',)).fetchone()
    assert row['forecast_rate'] == 140_000
    # Schema/other columns of the row are untouched -- only the value changed.
    assert row['category'] == 'MATERIAL'
    assert row['description'] == 'سیمان'


def _build_reference_fixture(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def write_catalog(name: str, header_row3: list[str], data_rows: list[list]):
        ws = wb.create_sheet(name)
        ws.append([f'کاتالوگ {name}'])
        ws.append([None, None, None, 'قیمت واحد'])
        ws.append(header_row3)
        for r in data_rows:
            ws.append(r)

    write_catalog(
        'قیمت مصالح',
        ['ردیف', 'عنوان', 'واحد', 'برآورد اولیه', 'برآورد بروز', 'واقعی', 'فازهای مورد استفاده'],
        [[1, 'سیمان شاهرود', 'کیسه', 100_000, 130_000, None, '08']],
    )
    write_catalog(
        'قیمت پیمانکار',
        ['ردیف', 'عنوان', 'واحد', 'برآورد اولیه', 'برآورد بروز', 'واقعی', 'فازهای مورد استفاده'],
        [[1, 'گمانه آزمایش خاک', 'متر', 260_000, 316_244, None, '02']],
    )
    write_catalog(
        'قیمت خدمات',
        ['ردیف', 'عنوان', 'واحد', 'برآورد اولیه', 'برآورد بروز', 'واقعی', 'فازهای مورد استفاده'],
        [[1, 'بنگاه خرید زمین', 'مقطوع', 75_000_000, 87_220_650, None, '01']],
    )

    params = wb.create_sheet('پارامترها')
    params.append(['کلید', 'مقدار'])
    params.append(['نرخ دلار آزاد - مبنا (فروردین ۱۴۰۵)', 150_000])
    params.append(['نرخ دلار آزاد - جاری', 232_500])
    params.append(['قیمت مس کاتد - مبنا', 2_003_030])
    params.append(['قیمت مس کاتد - جاری', 2_603_290])
    wb.save(path)


def test_import_reference_workbook_seeds_catalog_and_parameters(tmp_path: Path):
    fixture = tmp_path / 'reference.xlsx'
    _build_reference_fixture(fixture)
    db = FinanceDB(tmp_path / 'test.sqlite3')
    service = PriceUpdateService(db, tmp_path)

    counts = service.import_reference_workbook(fixture)
    assert counts == {'MATERIAL': 1, 'CONTRACTOR': 1, 'SERVICE': 1}

    items = service.list_catalog_items()
    assert len(items) == 3
    material = next(r for r in items if r['catalog_type'] == 'MATERIAL')
    assert material['description'] == 'سیمان شاهرود'
    assert material['estimate_rate'] == 100_000
    assert material['forecast_rate'] == 130_000  # seeded from column E, not zeroed
    assert material['phase_codes'] == '08'
    # Unclassified on first import -- a later rule assignment decides the method.
    assert material['calc_method'] == ''

    usd = service.get_parameter('usd')
    assert usd['base_value'] == 150_000
    assert usd['current_value'] == 232_500
    copper = service.get_parameter('copper')
    assert copper['current_value'] == 2_603_290


def test_import_reference_workbook_is_idempotent(tmp_path: Path):
    fixture = tmp_path / 'reference.xlsx'
    _build_reference_fixture(fixture)
    db = FinanceDB(tmp_path / 'test.sqlite3')
    service = PriceUpdateService(db, tmp_path)

    service.import_reference_workbook(fixture)
    service.import_reference_workbook(fixture)
    assert len(service.list_catalog_items()) == 3


def test_export_reference_workbook_never_overwrites_a_previous_export(service: PriceUpdateService):
    service.upsert_catalog_item(None, description='سیمان', catalog_type='MATERIAL', estimate_rate=100_000)
    first = service.export_reference_workbook()
    second = service.export_reference_workbook()
    assert first != second
    assert first.exists() and second.exists()


def test_export_reference_workbook_writes_readable_catalog_and_parameter_sheets(service: PriceUpdateService):
    service.upsert_catalog_item(
        None, description='سیمان شاهرود', catalog_type='MATERIAL', unit='کیسه',
        estimate_rate=100_000, forecast_rate=130_000, phase_codes='08', code='7')
    service.save_parameter('usd', base_value=150_000, current_value=232_500)

    out_path = service.export_reference_workbook()

    import openpyxl as openpyxl_module
    wb = openpyxl_module.load_workbook(out_path)
    assert set(wb.sheetnames) >= {'قیمت مصالح', 'قیمت پیمانکار', 'قیمت خدمات', 'پارامترها'}
    rows = list(wb['قیمت مصالح'].iter_rows(values_only=True))
    data_row = rows[3]  # title, blank, header, then data
    assert data_row[1] == 'سیمان شاهرود'
    assert data_row[3] == 100_000
    assert data_row[4] == 130_000

    param_rows = {r[0]: r[1] for r in wb['پارامترها'].iter_rows(values_only=True) if r[0]}
    assert param_rows['نرخ دلار آزاد - مبنا'] == 150_000
    assert param_rows['نرخ دلار آزاد - جاری'] == 232_500


def test_export_then_reimport_updates_an_app_created_item_in_place_by_description(service: PriceUpdateService):
    """The exact bug this design has to avoid: a UI-created row carries a
    random id, not the deterministic hash a first-ever import would compute,
    so re-importing an unedited export must not create a duplicate row."""

    price_id = service.upsert_catalog_item(
        None, description='قلم ساخته‌شده در برنامه', catalog_type='MATERIAL',
        estimate_rate=100_000, calc_method=calc.METHOD_WEB, brand='برند تست',
    )
    out_path = service.export_reference_workbook()

    service.import_reference_workbook(out_path)

    items = service.list_catalog_items()
    assert len(items) == 1  # no duplicate
    reimported = service.get_catalog_item(price_id)
    assert reimported is not None  # same price_id, updated in place
    assert reimported['calc_method'] == calc.METHOD_WEB  # app-only config survives re-import
    assert reimported['brand'] == 'برند تست'


def test_import_reference_workbook_matches_by_code_before_description(service: PriceUpdateService):
    price_id = service.upsert_catalog_item(
        None, description='نام قدیمی', catalog_type='MATERIAL', code='M-7', estimate_rate=100_000)

    fixture = service.base_dir / 'edited.xlsx'
    _build_reference_fixture_material_row(fixture, code='M-7', description='نام جدید ویرایش‌شده در اکسل',
                                           estimate_rate=120_000, forecast_rate=150_000)
    service.import_reference_workbook(fixture)

    items = service.list_catalog_items()
    assert len(items) == 1
    updated = service.get_catalog_item(price_id)
    assert updated['description'] == 'نام جدید ویرایش‌شده در اکسل'
    assert updated['estimate_rate'] == 120_000


def _build_reference_fixture_material_row(path: Path, *, code: str, description: str,
                                           estimate_rate: int, forecast_rate: int) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet = wb.create_sheet('قیمت مصالح')
    sheet.append(['کاتالوگ قیمت مصالح'])
    sheet.append([None, None, None, 'قیمت واحد'])
    sheet.append(['ردیف', 'عنوان', 'واحد', 'برآورد اولیه', 'برآورد بروز', 'واقعی', 'فازهای مورد استفاده'])
    sheet.append([code, description, 'کیسه', estimate_rate, forecast_rate, None, '08'])
    wb.save(path)


def _insert_rich_price_item(con, *, price_id: str, description: str, category: str,
                             estimate_rate: int = 100_000, forecast_rate: int = 100_000,
                             visible: int = 1, aggregate: int = 0, active: int = 1) -> None:
    con.execute(
        "INSERT INTO rich_price_items(price_id,code,description,category,unit,estimate_rate,"
        "actual_rate,forecast_rate,status,source,visible,aggregate,group_key,active,created_at,updated_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (price_id, '', description, category, 'کیسه', estimate_rate, estimate_rate,
         forecast_rate, '', '', visible, aggregate, '', active, '2026-01-01', '2026-01-01'),
    )


def _link_rich_phase(con, *, price_id: str, phase_no: int) -> None:
    con.execute(
        "INSERT INTO rich_phase_items(phase_item_id,phase_no,phase_name,description,price_id,"
        "active,created_at,updated_at) VALUES(?,?,?,?,?,1,?,?)",
        (f'PHASE-{price_id}-{phase_no}', phase_no, 'فاز', 'قلم', price_id, '2026-01-01', '2026-01-01'),
    )


def test_import_from_rich_price_items_seeds_the_empty_catalog_from_the_projects_own_price_list(
        tmp_path: Path):
    """A real project's rich_price_items is what «شروع براورد» should have
    something to run against on a first-ever open, instead of an empty
    catalog with no way in short of a manual add or an Excel import.

    Ported from origin/master (HANDOFF entry 15) - the operator hit this on
    a real project (2026-09-24): "با بارگذاری براورد جدید اصلا هیچ دیتایی
    وارد تب بروزرسانی نمیشه".
    """
    db = FinanceDB(tmp_path / 'shared.sqlite3')
    RichBudgetService(db, tmp_path)
    with db.tx() as con:
        _insert_rich_price_item(con, price_id='RICH-MAT', description='سیمان', category='MATERIAL')
        _insert_rich_price_item(con, price_id='RICH-LABOR', description='دستمزد بنا', category='LABOR')
        _insert_rich_price_item(con, price_id='RICH-INSTALL', description='نصب کابینت',
                                 category='INSTALLATION_COMPONENT')
        _insert_rich_price_item(con, price_id='RICH-HIDDEN', description='قلم مخفی', category='MATERIAL',
                                 visible=0)
        _insert_rich_price_item(con, price_id='RICH-AGGREGATE', description='جمع کل', category='MATERIAL',
                                 aggregate=1)
        _insert_rich_price_item(con, price_id='RICH-INACTIVE', description='قلم حذف‌شده', category='MATERIAL',
                                 active=0)
        _link_rich_phase(con, price_id='RICH-MAT', phase_no=8)
        _link_rich_phase(con, price_id='RICH-MAT', phase_no=9)

    service = PriceUpdateService(db, tmp_path)
    counts = service.import_from_rich_price_items()

    # Only the real, visible, non-aggregate, active rows come in.
    assert counts == {'imported': 3, 'updated': 0, 'total_source_rows': 3}
    items = {r['description']: r for r in service.list_catalog_items()}
    assert set(items) == {'سیمان', 'دستمزد بنا', 'نصب کابینت'}

    # category -> catalog_type mapping.
    assert items['سیمان']['catalog_type'] == 'MATERIAL'
    assert items['دستمزد بنا']['catalog_type'] == 'CONTRACTOR'  # LABOR -> CONTRACTOR
    assert items['نصب کابینت']['catalog_type'] == 'SERVICE'  # INSTALLATION_COMPONENT -> SERVICE

    # rich_phase_items links become phase_codes, zero-padded and sorted.
    assert items['سیمان']['phase_codes'] == '08,09'
    assert items['دستمزد بنا']['phase_codes'] == ''

    # Traceable back to its rich_price_items row for future syncing.
    assert items['سیمان']['rich_price_id'] == 'RICH-MAT'


def test_import_from_rich_price_items_is_idempotent_and_updates_in_place(tmp_path: Path):
    db = FinanceDB(tmp_path / 'shared.sqlite3')
    RichBudgetService(db, tmp_path)
    with db.tx() as con:
        _insert_rich_price_item(con, price_id='RICH-1', description='سیمان', category='MATERIAL')
    service = PriceUpdateService(db, tmp_path)
    first = service.import_from_rich_price_items()
    assert first == {'imported': 1, 'updated': 0, 'total_source_rows': 1}

    # The project's own price list changes; re-running picks the edit up
    # instead of creating a duplicate catalog row.
    with db.tx() as con:
        con.execute("UPDATE rich_price_items SET description='سیمان تیپ ۲' WHERE price_id='RICH-1'")
    second = service.import_from_rich_price_items()
    assert second == {'imported': 0, 'updated': 1, 'total_source_rows': 1}

    items = service.list_catalog_items()
    assert len(items) == 1
    assert items[0]['description'] == 'سیمان تیپ ۲'


# --------------------------------------------------------------------------
# The operator's real reference workbook. It lives outside the repository
# (private project data), so these run only on a machine that has it.

REAL_REFERENCE = Path(os.environ.get('BARKAV_REAL_REFERENCE_WORKBOOK', ''))
needs_real_reference = pytest.mark.skipif(not REAL_REFERENCE.is_file(),
                                          reason='reference workbook not on this machine')


@needs_real_reference
def test_real_reference_workbook_fills_methods_rules_parameters_and_reproduces_its_forecasts(tmp_path: Path):
    """Loading the workbook fills جزئیات/قوانین/پارامترها, and a run reproduces its column E.

    Every catalog row's «برآورد بروز» in the workbook is the result of its
    own H-column formula; importing the formulas' روش and figures and then
    running the engine against the workbook's own پارامترها must land on the
    very same number for all ۴۵۶ rows.
    """

    db = FinanceDB(tmp_path / 'test.sqlite3')
    rich = RichBudgetService(db, tmp_path)
    rich.import_workbook(REAL_REFERENCE)
    service = PriceUpdateService(db, tmp_path)
    service.import_from_rich_price_items()
    counts = service.import_reference_workbook(REAL_REFERENCE)
    assert counts == {'MATERIAL': 282, 'CONTRACTOR': 110, 'SERVICE': 64}

    items = service.list_catalog_items()
    assert len(items) == 456  # matched onto the rows seeded from the price list, not duplicated
    assert all(item['calc_method'] for item in items)
    assert all(item['rich_price_id'] for item in items)
    assert '07,09' in {item['phase_codes'] for item in items}
    cement = next(i for i in items if i['description'] == 'سیمان')
    params = json.loads(cement['calc_params_json'])
    assert cement['calc_method'] == calc.METHOD_WEB and cement['brand'] == 'سیمان شاهرود'
    assert params['web_prices'] == [200_000, 203_000, 216_000] and params['carry_rate'] == 0.3
    assert {i['description'] for i in service.needs_review_items()} >= {'سرامیک 60×120', 'سرامیک 80×80'}
    assert len(service.list_calc_rules()) > 40

    assert service.get_parameter('usd')['base_value'] == 150_000
    assert service.get_parameter('usd')['current_value'] == 232_500
    assert service.get_parameter('months_elapsed')['current_value'] == 5
    assert abs(service.growth_context()['wage'] - 1.216326007054794) < 1e-12

    expected = {i['price_id']: i['forecast_rate'] for i in items}
    service.run_update(trigger='manual')
    for item in service.list_catalog_items():
        assert abs(item['forecast_rate'] - expected[item['price_id']]) <= 2, item['description']


def test_ensure_parameter_defaults_fills_only_empty_values(tmp_path: Path):
    fixture = tmp_path / 'reference.xlsx'
    _build_reference_fixture(fixture)
    service = PriceUpdateService(FinanceDB(tmp_path / 'test.sqlite3'), tmp_path)
    service.save_parameter('copper', base_value=10, current_value=12, source='wizard')

    changed = service.ensure_parameter_defaults(
        workbook=fixture, settings={'usd_rate_baseline': 1, 'usd_rate_current': 2})

    usd = service.get_parameter('usd')
    assert (usd['base_value'], usd['current_value']) == (150_000, 232_500)  # workbook before settings
    assert usd['previous_value'] == 232_500
    copper = service.get_parameter('copper')
    assert (copper['base_value'], copper['current_value']) == (10, 12)  # typed values are never replaced
    assert copper['previous_value'] == 12  # an empty قبلی starts at بروز, not at 0
    assert 'usd' in changed


def test_normalize_phase_codes_reads_the_workbooks_persian_comma():
    from price_update_service import normalize_phase_codes
    assert normalize_phase_codes('07، 09') == '07,09'
    assert normalize_phase_codes('۸') == '08'
    assert normalize_phase_codes('') == ''
