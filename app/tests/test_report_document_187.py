from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unicodedata
import zipfile
import pytest
from PIL import Image, ImageChops
from pypdf import PdfReader
from openpyxl import load_workbook
from report_document_stage18_7 import (build_report_document, selected_tree_items,
    selected_table_rows, resolve_report_document_paths, fa_number)


def logical_text(page):
    return unicodedata.normalize('NFKC', page.extract_text())


class Tree:
    def __init__(self, selected=()): self.selected = selected
    def get_children(self): return ('b', 'a', 'c')
    def selection(self): return self.selected
    def index(self, iid): return self.get_children().index(iid)
    def item(self, iid, option): return [iid, 'visible']


def test_selected_export_keeps_visual_order():
    assert selected_tree_items(Tree(('c', 'b'))) == ['b', 'c']
    assert selected_tree_items(Tree()) == ['b', 'a', 'c']
    table = SimpleNamespace(tree=Tree(('c', 'b')), sorted_rows=lambda: [[2], [1], [3]])
    assert selected_table_rows(table) == [[2], [3]]


def test_saved_pdf_and_excel_pages_are_the_preview_document(tmp_path):
    report = build_report_document(['ردیف', 'شرح'],
        [{'ردیف': fa_number(i + 1), 'شرح': 'هزینه ساخت ۱۲۳'} for i in range(80)],
        [('مشخصات', 'نام پروژه', 'پروژه آزمایشی')], 'گزارش آزمایشی', orientation='عمودی')
    output = report.save(tmp_path / 'report.pdf', 'PDF')
    assert output.read_bytes() == report.pdf
    reader = PdfReader(BytesIO(report.pdf))
    assert report.page_count == len(reader.pages) >= 3
    for index, page in enumerate(reader.pages, 1):
        text = logical_text(page)
        assert 'شرکت مهندسی ساخت بارکاو' in text
        assert 'جهت ارائه به:' in text
        assert 'تعداد صفحات گزارش' in text
        assert f'{fa_number(index)} از {fa_number(report.page_count)}' in text
    report.save(tmp_path / 'report.xlsx', 'Excel')
    book = load_workbook(tmp_path / 'report.xlsx')
    assert len(book.worksheets) == report.page_count + 1
    assert book['داده‌های گزارش'].max_row == 81
    assert book['داده‌های گزارش'].sheet_view.rightToLeft
    images = list(report.page_images(scale=2))
    with zipfile.ZipFile(tmp_path / 'report.xlsx') as archive:
        for index, expected in enumerate(images, 1):
            actual = Image.open(BytesIO(archive.read(f'xl/media/image{index}.png'))).convert('RGB')
            assert actual.size == expected.size
            assert ImageChops.difference(actual, expected.convert('RGB')).getbbox() is None
    assert book.worksheets[0].print_area.startswith("'صفحه ۱'!$A$1:$A$")
    assert all(row.height <= 409 for row in book.worksheets[0].row_dimensions.values())
    assert sum(row.height for row in book.worksheets[0].row_dimensions.values()) == pytest.approx(images[0].height / 2 * .75)


def test_documents_toggle_controls_actual_pdf_and_excel_content(tmp_path):
    receipt = tmp_path / 'receipt.png'
    Image.new('RGB', (80, 120), '#ac1223').save(receipt)
    def make(include):
        return build_report_document(['شرح', 'نمایش سند'],
            [{'شرح':'ردیف انتخاب‌شده', 'نمایش سند': str(receipt), '__document_paths_187':[str(receipt)]}],
            [], 'گزارش', include_documents=include)
    enabled, disabled = make(True), make(False)
    with_images = sum(len(page.images) for page in PdfReader(BytesIO(enabled.pdf)).pages)
    without_images = sum(len(page.images) for page in PdfReader(BytesIO(disabled.pdf)).pages)
    assert with_images > without_images
    assert 'اسناد ردیف' in ''.join(logical_text(page) for page in PdfReader(BytesIO(enabled.pdf)).pages)
    assert disabled.rows[0][1] == ''
    assert enabled.rows[0][1] == receipt.name


def test_pdf_attachment_is_embedded_as_document_pages(tmp_path):
    receipt = build_report_document(['شرح'], [{'شرح':'ATTACHMENT-ONLY'}], [], 'سند')
    path = receipt.save(tmp_path/'attachment.pdf', 'PDF')
    report = build_report_document(['شرح'], [{'شرح':'گزارش', '__document_paths_187':[str(path)]}], [], 'گزارش')
    assert sum(len(page.images) for page in PdfReader(BytesIO(report.pdf)).pages) >= 2


def test_data_entry_includes_approved_relinked_evidence(tmp_path, monkeypatch):
    from dashboard_stage18_7 import DataEntryCenter187
    import reports_stage18_7 as reports
    receipt = tmp_path/'posted.png'
    Image.new('RGB',(30,40),'white').save(receipt)
    db = SimpleNamespace(
        list_pending_attachments=lambda identity: [],
        get_financial_posting=lambda identity: {'transaction_id':'TX-'+identity},
        list_entity_attachments=lambda kind, identity: [{'local_path':str(receipt)}],
    )
    assert resolve_report_document_paths(db, {'pending_id':'P1'}) == [str(receipt)]
    captured = {}
    def preview(parent, choice, headers, rows, title):
        captured.update(headers=headers, rows=rows)
    monkeypatch.setattr(reports, '_export_through_preview_187', preview)
    variable = SimpleNamespace(get=lambda:'PDF', set=lambda value:None)
    table = SimpleNamespace(tree=Tree(('c','b')), COLUMNS=[('id','ردیف',60),('text','شرح',100)],
        iid_to_row={'b':{'pending_id':'P1'},'a':{'pending_id':'P2'},'c':{'pending_id':'P3'}})
    DataEntryCenter187._output_selected(SimpleNamespace(output_var=variable,table=table,rich=SimpleNamespace(db=db)))
    assert [row['ردیف'] for row in captured['rows']] == ['b','c']
    assert all(row['__document_paths_187'] == [str(receipt)] for row in captured['rows'])


def test_preview_confirmation_copies_cached_document(tmp_path, monkeypatch):
    import sales_claims_stage18_7 as sales
    report = build_report_document(['شرح'], [{'شرح':'PREVIEW-SNAPSHOT'}], [], 'گزارش')
    target=tmp_path/'saved.pdf'
    monkeypatch.setattr(sales.filedialog,'asksaveasfilename',lambda **kwargs:str(target))
    monkeypatch.setattr(sales.messagebox,'showinfo',lambda *args,**kwargs:None)
    harness=SimpleNamespace(format_choice=SimpleNamespace(get=lambda:'PDF'),output_title='گزارش',
        _report_document=report,_take_focus=lambda:None)
    assert sales.ContractExportPreview187._save_document_187(harness)
    assert target.read_bytes() == report.pdf


def test_duplicate_display_headers_keep_each_cell_value(tmp_path):
    report = build_report_document(['معیار', 'امتیاز', 'معیار', 'امتیاز'],
        [['کیفیت', 8, 'تعهد', 9]], [], 'نظرسنجی')
    assert report.rows == (('کیفیت', 8, 'تعهد', 9),)
    text = logical_text(PdfReader(BytesIO(report.pdf)).pages[0])
    assert 'کیفیت' in text and 'تعهد' in text


def test_print_uses_same_cached_pdf(tmp_path, monkeypatch):
    import sales_claims_stage18_7 as sales
    import export_utils
    report = build_report_document(['شرح'], [{'شرح':'PRINT-SNAPSHOT'}], [], 'گزارش')
    target=tmp_path/'print.pdf'
    sent=[]
    monkeypatch.setattr(sales.filedialog,'asksaveasfilename',lambda **kwargs:str(target))
    monkeypatch.setattr(sales.messagebox,'showinfo',lambda *args,**kwargs:None)
    monkeypatch.setattr(export_utils,'print_file',lambda path:sent.append(Path(path).read_bytes()))
    harness=SimpleNamespace(format_choice=SimpleNamespace(get=lambda:'پرینت'),output_title='گزارش',
        _report_document=report,_take_focus=lambda:None)
    assert sales.ContractExportPreview187._save_document_187(harness)
    assert sent == [report.pdf]


def test_cancel_keeps_preview_snapshot_and_restores_focus(monkeypatch):
    import sales_claims_stage18_7 as sales
    focus=[]
    report=object()
    monkeypatch.setattr(sales.filedialog,'asksaveasfilename',lambda **kwargs:'')
    harness=SimpleNamespace(format_choice=SimpleNamespace(get=lambda:'PDF'),output_title='گزارش',
        _report_document=report,_take_focus=lambda:focus.append(True))
    assert sales.ContractExportPreview187._save_document_187(harness) is False
    assert harness._report_document is report and focus == [True]


def test_large_specs_and_long_rows_paginate_without_losing_content():
    report = build_report_document(['شرح'], [{'شرح':'LONG-END ' + 'متن بلند ' * 300}],
        [('مشخصات', f'ردیف {i}', f'DETAIL-{i}') for i in range(80)], 'گزارش', orientation='عمودی')
    assert report.page_count > 1
    text=''.join(logical_text(page) for page in PdfReader(BytesIO(report.pdf)).pages)
    assert 'DETAIL-79' in text
    # pypdf can drop a Latin run beside Persian. Read the renderer's own
    # text layer to verify the mixed run, independently of that extractor.
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(report.pdf) as document:
        actual = ''
        for page in document:
            layer = page.get_textpage()
            actual += layer.get_text_range()
            layer.close()
            page.close()
        assert 'LONG-END' in actual


def test_actual_windows_preview_rebuilds_options_and_keeps_caption(tmp_path):
    import os
    if os.name != 'nt': pytest.skip('Windows caption acceptance')
    import tkinter as tk
    import ctypes
    from sales_claims_stage18_7 import ContractExportPreview187
    receipt=tmp_path/'doc.png'
    Image.new('RGB',(50,60),'red').save(receipt)
    root=tk.Tk()
    root.withdraw()
    preview=None
    try:
        preview=ContractExportPreview187(root, export_format='PDF',
            headers=['شرح'], payload=[{'شرح':'دادهٔ آزمایشی', '__document_paths_187':[str(receipt)]}],
            specs=[], on_confirm=lambda *args:False, include_addressee=True)
        root.update()
        assert preview.wm_transient() == ''
        assert preview.wm_attributes('-toolwindow') == 0
        hwnd=ctypes.windll.user32.GetParent(preview.winfo_id())
        bits=ctypes.windll.user32.GetWindowLongW(hwnd,-16)
        required=0x00020000 | 0x00010000 | 0x00080000
        assert bits & required == required
        previous=preview._report_document
        preview.orientation_choice.set('عمودی')
        preview.addressee_choice.set('کارفرما')
        preview.include_documents.set(False)
        preview._render()
        root.update()
        assert preview._report_document is not previous
        assert len(preview._report_pages) == preview._report_document.page_count
        assert fa_number(preview._report_document.page_count) in preview.page_count_label.cget('text')
        page=PdfReader(BytesIO(preview._report_document.pdf)).pages[0]
        assert float(page.mediabox.height) > float(page.mediabox.width)
        assert 'کارفرما' in logical_text(page)
        assert sum(len(page.images) for page in PdfReader(BytesIO(previous.pdf)).pages) > len(page.images)
    finally:
        if preview is not None: preview.destroy()
        root.destroy()
