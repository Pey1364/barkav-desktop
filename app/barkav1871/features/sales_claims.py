"""Canonical 18.7.1 sales, contracts and claims workspace.

This module is deliberately a presentation layer over ``RevenueService``.  A
contract does not create cash; only the existing receipt workflow may do that.
Extended contract details are stored as JSON in the existing sales ``notes``
field, so this UI needs no database migration.
"""
from __future__ import annotations

import csv
import json
from datetime import date, datetime
import os
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from typing import Any

try:
    import jdatetime
except ModuleNotFoundError:  # The launcher reports the missing runtime dependency.
    jdatetime = None

from barkav1871.services.documents import AttachmentService
from barkav1871.services.exports import export_pdf
from barkav1871.services.sales_input import parse_money, parse_sale_form
from barkav1871.services.sales import RevenueService, gregorian_to_shamsi, normalize_shamsi
from barkav1871.design.controls import RtlNotebookStrip, en, fa


def _require_jdatetime() -> None:
    if jdatetime is None:
        raise RuntimeError("برای تقویم شمسی، بستهٔ jdatetime را نصب کنید.")


CONTRACT_STEPS = (
    '۱. بارگذاری قرارداد و مورد معامله', '۲. تاریخ‌ها', '۳. خریدار',
    '۴. مالی', '۵. شرایط پرداخت', '۶. امضاکنندگان', '۷. تأیید نهایی',
)
PAYMENT_LABELS = ('پرداخت نقدی', 'قسط اول', 'قسط دوم', 'قسط سوم', 'قسط چهارم', 'قسط پنجم',
                      'قسط ششم', 'قسط هفتم', 'موعد تحویل', 'موعد سند')
INSTALLMENT_LABELS = ('قسط اول', 'قسط دوم', 'قسط سوم', 'قسط چهارم', 'قسط پنجم',
                          'قسط ششم', 'قسط هفتم', 'قسط هشتم', 'قسط نهم', 'قسط دهم')


def payment_labels_for_installments(installment_count: int) -> tuple[str, ...]:
    """Contract-wizard order: cash, chosen installments, delivery, then deed."""
    count = max(1, min(10, int(installment_count)))
    return ('پرداخت نقدی', *INSTALLMENT_LABELS[:count], 'موعد تحویل', 'موعد سند')

CONTRACT_TABLE_COLUMNS = (
    ('row', 'ردیف', 80), ('floor', 'طبقه', 80), ('unit', 'شماره واحد', 90),
    ('date', 'تاریخ قرارداد', 110), ('buyer', 'خریدار', 150), ('kind', 'نوع فروش', 100),
    ('area', 'متراژ فروخته‌شده', 125), ('price', 'قیمت هر مترمربع (تومان)', 170),
    ('total', 'مبلغ قرارداد (تومان)', 150), ('delivery', 'تاریخ تحویل', 110),
    ('deed', 'تاریخ سند', 110), ('received', 'واریزی (تومان)', 140),
    ('remain', 'مانده (تومان)', 140), ('select', 'انتخاب', 75),
)
CONTRACT_NON_MONEY_COLUMNS = ('date', 'buyer', 'kind')
EXPORT_PLACEHOLDER = 'انتخاب فرمت خروجی'
SALES_FINANCE_PENDING_SOURCE = 'MANUAL_SALES_FINANCE'
SALES_FINANCE_PENDING_ACTIONS = frozenset({'ADD', 'EDIT'})

PAYMENT_DETAIL_COLUMNS = (
    ('row', 'ردیف', 70), ('stage', 'مرحله', 125), ('contract_date', 'تاریخ قرارداد', 110),
    ('contract_amount', 'مبلغ قرارداد (تومان)', 150), ('contract_due_date', 'تاریخ واریز طبق قرارداد', 150),
    ('received_date', 'تاریخ واریز', 110),
    ('received_amount', 'مبلغ واریز (تومان)', 150), ('late', 'دیرکرد (روز)', 100),
    ('balance', 'بدهکار/بستانکار (تومان)', 170), ('document', 'نمایش سند', 120),
)


def toman(value: Any) -> str:
    """Render monetary values with Persian digits and three-digit grouping."""
    return fa(f"{int(value or 0):,}")


def persian_digits(value: Any) -> str:
    """Display ASCII/Arabic-Indic digits as Persian without changing stored values."""
    return fa(en(value))


def formatted_amount_input(value: Any) -> str:
    """Format wizard money input with Persian digits and thousands separators."""
    raw = en(value).replace(',', '').replace('٬', '').strip()
    if not raw:
        return ''
    try:
        return fa(f'{int(raw):,}')
    except (TypeError, ValueError):
        return persian_digits(value)


def selectable_row(number: int, checked: bool = False) -> str:
    """A Persian row label with a visible, keyboard-independent check mark."""
    return f"{'☑' if checked else '☐'} {fa(number)}"


def selection_cell(checked: bool = False) -> str:
    """Render the independent selection column without polluting row numbers."""
    return '☑' if checked else '☐'


def visible_rows_payload(headers: list[str], rows: list[tuple[Any, ...]]) -> list[dict[str, str]]:
    """Return only the already-visible table data for safe, local export."""
    return [dict(zip(headers, (str(value) for value in row))) for row in rows]


def finance_export_rows(headers: list[str], rows: list[tuple[Any, ...]]) -> list[dict[str, str]]:
    """Export finance rows without UI checkbox glyphs in the row-number column."""
    payload = visible_rows_payload(headers, rows)
    for row in payload:
        if 'ردیف' in row:
            row['ردیف'] = row['ردیف'].replace('☐', '').replace('☑', '').strip()
        if 'انتخاب' in row:
            row['انتخاب'] = 'انتخاب شده' if '☑' in row['انتخاب'] else 'انتخاب نشده'
    return payload


def sales_finance_pending_payload(*, action: str, sale: dict[str, Any], installment: dict[str, Any] | None,
                                      receipt: dict[str, Any] | None) -> dict[str, Any]:
    """Build an approval-only sales receipt draft; it never posts cash by itself."""
    current_action = str(action or '').upper()
    if current_action not in {'ADD', 'EDIT', 'DELETE'}:
        raise ValueError('عملیات مالی فروش نامعتبر است.')
    installment = installment or {}
    receipt = receipt or {}
    amount = int(receipt.get('amount') or installment.get('amount') or 0)
    operation = {'ADD': 'ADD_REQUEST', 'DELETE': 'DELETE_REQUEST', 'EDIT': 'EDIT_REQUEST'}[current_action]
    description = {
        'ADD': 'درخواست افزودن ورودی فروش',
        'DELETE': 'درخواست حذف/اصلاح ردیف مالی',
        'EDIT': 'درخواست ویرایش تراکنش ثبت‌شده فروش',
    }[current_action]
    return {
        'detected_type': 'TRANSACTION',
        # These are the actual title/type fields consumed by the Data Entry wizard.
        'title': 'ورودی', 'subtype': 'فروش',
        'event_date_shamsi': receipt.get('received_date_shamsi') or gregorian_to_shamsi(date.today()),
        'amount': amount, 'payment_state': 'UNPAID',
        'sale_id': sale.get('sale_id') or '', 'buyer_name': sale.get('buyer_name') or '',
        'unit_number': sale.get('unit_number') or '', 'contract_number': sale.get('contract_number') or '',
        'installment_id': installment.get('installment_id') or '',
        'received_payment_id': receipt.get('received_payment_id') or '',
        'transaction_id': receipt.get('transaction_id') or '',
        'operation': operation, 'sales_pending_action': current_action,
        'description': description, 'source': SALES_FINANCE_PENDING_SOURCE,
    }


def sales_finance_pending_message_date(payload: dict[str, Any], now: datetime | None = None) -> str:
    """Return the Persian/Shamsi display date used by the sales Pending row."""
    raw = str((payload or {}).get('event_date_shamsi') or '').strip()
    if not raw:
        raw = gregorian_to_shamsi((now or datetime.now()).date())
    return persian_digits(raw.split()[0])


def is_temporary_sales_finance_pending(payload: dict[str, Any] | None) -> bool:
    """Only Add/Edit receipt drafts may be discarded before approval."""
    payload = payload or {}
    return (
        payload.get('source') == SALES_FINANCE_PENDING_SOURCE
        and str(payload.get('sales_pending_action') or '').upper() in SALES_FINANCE_PENDING_ACTIONS
    )


def payment_detail_export_payload(headers: list[str], rows: list[tuple[Any, ...]], installment_ids: list[str],
                                      documents_by_installment: dict[str, list[dict[str, Any]]], include_documents: bool) -> list[dict[str, str]]:
    """Keep documents opt-in and expose every receipt as its own preview row."""
    document_header = 'نمایش سند'
    if not include_documents:
        index = headers.index(document_header)
        headers = headers[:index] + headers[index + 1:]
        rows = [row[:index] + row[index + 1:] for row in rows]
        return visible_rows_payload(headers, rows)
    payload: list[dict[str, str]] = []
    for installment_id, row in zip(installment_ids, rows):
        payment = dict(zip(headers, (str(value) for value in row)))
        paths = []
        for document in documents_by_installment.get(str(installment_id), []):
            path = str(document.get('local_path') or document.get('attachment_id') or '').strip()
            if path and path not in paths:
                paths.append(path)
        # Each payment document must have a visible row.  Extra supporting
        # receipts intentionally leave amount/date blank to avoid duplicating
        # accounting figures in the exported contract table.
        stage_header = next((header for header in headers if header == 'مرحله' or 'مرحله پرداخت' in header), '')
        received_date_header = next((header for header in headers if header == 'تاریخ واریز'), '')
        received_amount_header = next((header for header in headers if 'مبلغ واریز' in header and 'کل مبلغ' not in header), '')
        for document_number, path in enumerate(paths or [''], 1):
            entry = dict(payment)
            if len(paths) > 1 and stage_header:
                entry[stage_header] = f"{entry.get(stage_header) or 'مرحله'} · فیش {persian_digits(document_number)}"
            if document_number > 1:
                if received_date_header:
                    entry[received_date_header] = '—'
                if received_amount_header:
                    entry[received_amount_header] = '—'
            entry[document_header] = path
            payload.append(entry)
    return payload


def export_document_paths(value: Any) -> list[str]:
    """Split the newline-separated, local attachment paths used by PDF/Excel cells."""
    return [item.strip() for item in str(value or '').replace('\r', '').split('\n') if item.strip()]


def export_payload_document_paths(payload: list[dict[str, Any]], document_header: str) -> list[str]:
    """Return every unique document referenced by visible export rows."""
    paths: list[str] = []
    for row in payload:
        for path in export_document_paths(row.get(document_header, '')):
            if path not in paths:
                paths.append(path)
    return paths


def wizard_export_payload(values: dict[str, Any], document_path: str = '') -> dict[str, str]:
    """Serialize the current wizard draft only; it does not save a contract or payment."""
    payload = {str(key): str(value or '') for key, value in values.items()}
    payload['contract_document_path'] = str(document_path or '')
    return payload


def export_contract_pdf(path: str | Path, headers: list[str], rows: list[Any], title: str) -> Path:
    """Write Persian contract text with the project's embedded RTL PDF font."""
    from barkav1871.services.exports import export_pdf
    values = [[row.get(header, '') for header in headers] if isinstance(row, dict) else list(row) for row in rows]
    return export_pdf(path, headers, values, title=title)


def export_receipts_pdf(path: str | Path, headers: list[str], rows: list[list[Any]], title: str,
                            summary_lines: list[str], document_paths: list[str], *, orientation: str = 'افقی',
                            contract_document_paths: list[str] | None = None) -> Path:
    """Export the same grouped specifications, summary and receipt-in-cell table used by the preview."""
    from reportlab import platypus
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from export_utils import _pdf_font, _rtl
    font = _pdf_font(); path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    # Source tables store fields in logical order; a Persian PDF table must
    # display that order from right to left, with the receipt column on left.
    headers = list(reversed(headers))
    rows = [list(reversed(row)) for row in rows]
    document_headers = [header for header in headers if header in {'نمایش سند', 'فیش واریزی'}]
    has_receipts = bool(document_paths) or any(any(export_document_paths(row[index]) for index, header in enumerate(headers) if header in document_headers) for row in rows)
    page_size = A3 if has_receipts else A4
    if str(orientation) == 'افقی':
        page_size = landscape(page_size)
    doc = platypus.SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=18, leftMargin=18, topMargin=18, bottomMargin=18)
    compact_vertical = str(orientation) == 'عمودی'
    # Payment-table text must remain readable beside a compact receipt image.
    style = ParagraphStyle('receipt_rtl', fontName=font, fontSize=8 if compact_vertical else 9,
                           leading=10 if compact_vertical else 12, alignment=TA_RIGHT)
    story = [platypus.Paragraph(_rtl(title), ParagraphStyle('receipt_title', parent=style, fontSize=15, leading=18))]
    page_width = page_size[0] - doc.leftMargin - doc.rightMargin
    grouped: dict[str, list[tuple[str, str]]] = {}
    plain_lines: list[str] = []
    for line in summary_lines:
        text = str(line)
        if ' — ' in text and ':' in text:
            group, remainder = text.split(' — ', 1)
            label, value = remainder.split(':', 1)
            grouped.setdefault(group.strip(), []).append((label.strip(), value.strip()))
        else:
            plain_lines.append(text)
    if plain_lines:
        summary_cells = [platypus.Paragraph(_rtl(line), style) for line in plain_lines]
        summary = platypus.Table([summary_cells], hAlign='RIGHT')
        summary.setStyle(platypus.TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#E8F0F7')),
            ('BOX', (0, 0), (-1, -1), .6, colors.HexColor('#496176')),
            ('INNERGRID', (0, 0), (-1, -1), .35, colors.HexColor('#B8C2CC')),
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story.extend([summary, platypus.Spacer(1, 8)])
    # The export and its preview use the same RTL grid: ۱/۲, ۳/۴ and ۵/۶.
    # ReportLab's left physical cell is the second RTL box in each pair.
    spec_boxes = []
    paired_box_width = (page_width - 8 * mm) / 2
    for group, entries in grouped.items():
        box_rows = [[platypus.Paragraph(_rtl(group), ParagraphStyle(
            f'box_title{len(story)}', parent=style, fontSize=10, leading=13,
        ))]]
        # ReportLab's physical right column is the RTL label column.
        box_rows.extend([[platypus.Paragraph(_rtl(value), style), platypus.Paragraph(_rtl(label), style)]
                         for label, value in entries])
        box = platypus.Table(box_rows, colWidths=[paired_box_width / 2, paired_box_width / 2], hAlign='RIGHT')
        box.setStyle(platypus.TableStyle([
            ('SPAN', (0, 0), (-1, 0)), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DCE8F2')),
            ('BOX', (0, 0), (-1, -1), .8, colors.HexColor('#496176')),
            ('INNERGRID', (0, 1), (-1, -1), .35, colors.HexColor('#B8C2CC')),
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        spec_boxes.append(box)
    if spec_boxes:
        pair_rows = []
        for index in range(0, len(spec_boxes), 2):
            right_box = spec_boxes[index]
            left_box = spec_boxes[index + 1] if index + 1 < len(spec_boxes) else ''
            pair_rows.append([left_box, right_box])
        pair_table = platypus.Table(pair_rows, colWidths=[paired_box_width, paired_box_width], hAlign='RIGHT')
        pair_table.setStyle(platypus.TableStyle([
            # Keep all contract sections inside one continuous frame: the outer
            # border and shared dividers belong to the paired grid, not a gap
            # between each nested section box.
            ('BOX', (0, 0), (-1, -1), .8, colors.HexColor('#496176')),
            ('INNERGRID', (0, 0), (-1, -1), .8, colors.HexColor('#496176')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0), ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(pair_table)
    document_width = 48 * mm if has_receipts else max(28 * mm, page_width / max(3, len(headers)))
    remaining_width = (page_width - document_width) / max(1, len(headers) - len(document_headers)) if compact_vertical else max(24 * mm, (page_width - document_width) / max(1, len(headers) - len(document_headers)))
    col_widths = [document_width if header in document_headers else remaining_width for header in headers]
    receipt_width = 45 * mm if has_receipts else 0
    receipt_height = 64 * mm if has_receipts else 0

    def cell(header: str, value: Any):
        if header not in document_headers:
            return platypus.Paragraph(_rtl(str(value or '-')), style)
        paths = export_document_paths(value)
        documents = []
        for item in paths:
            candidate = Path(item)
            if candidate.exists() and candidate.suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
                try:
                    documents.append(platypus.Image(str(candidate), width=receipt_width, height=receipt_height, kind='proportional'))
                except (OSError, ValueError):
                    continue
        return documents or platypus.Paragraph(_rtl('سند ثبت نشده است' if not paths else '\n'.join(Path(item).name for item in paths)), style)

    data = [[platypus.Paragraph(_rtl(header), style) for header in headers]]
    data.extend([[cell(header, value) for header, value in zip(headers, row)] for row in rows])
    # Keep the table itself on the right edge too.  Cell paragraphs already use
    # RTL/right alignment, but ReportLab otherwise places the whole table left.
    table = platypus.Table(data, colWidths=col_widths, repeatRows=1, splitByRow=1, hAlign='RIGHT')
    table.setStyle(platypus.TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), font), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#153B5B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#B8C2CC')),
        ('ALIGN', (0, 0), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(table)
    contract_images = [Path(str(item)) for item in (contract_document_paths or [])]
    if contract_images:
        story.extend([platypus.PageBreak(), platypus.Paragraph(_rtl('تصویر قرارداد'), ParagraphStyle(
            'contract_image_title', parent=style, fontSize=14, leading=18,
        ))])
        for candidate in contract_images:
            if not candidate.exists() or candidate.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
                story.append(platypus.Paragraph(_rtl(candidate.name), style)); continue
            try:
                story.extend([platypus.Spacer(1, 8), platypus.Image(str(candidate), width=105 * mm, height=148 * mm, kind='proportional')])
            except (OSError, ValueError):
                story.append(platypus.Paragraph(_rtl(candidate.name), style))
    doc.build(story); return path


def contract_price_per_m2(row: dict[str, Any], meta: dict[str, Any]) -> str:
    area = en(str(meta.get('contract_area') or '')).replace(',', '').strip()
    try:
        numeric_area = float(area)
    except ValueError:
        numeric_area = 0
    return toman(int(row.get('contract_amount') or 0) / numeric_area) if numeric_area > 0 else '-'


def contract_table_values(number: int, row: dict[str, Any], meta: dict[str, Any], detail: Any) -> tuple[str, ...]:
    """Present a contract from the existing ledger; derived prices are never posted back."""
    return (
        selectable_row(number), persian_digits(row.get('floor') or '-'), persian_digits(row.get('unit_number') or '-'),
        persian_digits(row.get('contract_date_shamsi') or '-'), persian_digits(row.get('buyer_name') or '-'),
        str(meta.get('sale_kind') or '-'), persian_digits(meta.get('contract_area') or '-'),
        contract_price_per_m2(row, meta), toman(row.get('contract_amount')),
        persian_digits(meta.get('delivery_date') or '-'), persian_digits(meta.get('deed_date') or '-'),
        toman(detail.cash_received), toman(detail.receivable), selection_cell(),
    )


def filter_contract_rows(rows: list[dict[str, Any]], kind: str, value: str) -> list[dict[str, Any]]:
    """Two-stage contract filtering with a stable «همه» default."""
    if kind == 'همه' or value in ('', 'همه'):
        return rows
    key = {'نام خریدار': 'buyer_name', 'شماره واحد': 'unit_number', 'شماره قرارداد': 'contract_number'}.get(kind)
    return [row for row in rows if key and str(row.get(key) or '') == value]


def contract_note(values: dict[str, Any]) -> str:
    """Keep only non-ledger contract metadata in the pre-existing notes column."""
    return json.dumps({'contract_v': 1, **values}, ensure_ascii=False, sort_keys=True)


def payment_plan(values: dict[str, Any]) -> list[dict[str, Any]]:
    """Build the installment plan; never invent money or dates.

    The delivery/deed dates are copied to their payment rows when the user did
    not type another date.  Zero-amount stages are intentionally omitted.
    """
    count_text = en(str(values.get('payment_count') or '4'))
    try: count = max(1, min(10, int(count_text)))
    except ValueError: count = 4
    labels = payment_labels_for_installments(count)
    plan: list[dict[str, Any]] = []
    for index, label in enumerate(labels, 1):
        amount = parse_money(str(values.get(f'pay{index}_amount') or ''))
        due = normalize_shamsi(str(values.get(f'pay{index}_date') or ''))
        if index == count + 2 and not due:
            due = normalize_shamsi(str(values.get('delivery_date') or ''))
        if index == count + 3 and not due:
            due = normalize_shamsi(str(values.get('deed_date') or ''))
        if not amount:
            continue
        if not due:
            raise ValueError(f'تاریخ «{label}» را وارد کنید.')
        plan.append({'sequence_no': index, 'amount': amount, 'due_date_shamsi': due, 'description': label})
    if not plan:
        raise ValueError('حداقل یک مرحله پرداخت با مبلغ و تاریخ لازم است.')
    return plan


class ShamsiDatePicker(tk.Toplevel):
    """RTL Persian calendar picker used by contract and payment wizard dates."""
    MONTHS = ('', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند')
    DAYS = ('شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه')

    def __init__(self, parent, variable):
        _require_jdatetime()
        super().__init__(parent)
        self.variable = variable
        today = jdatetime.date.today()
        raw = normalize_shamsi(variable.get())
        try: parsed = [int(part) for part in raw.split('/')]
        except (ValueError, AttributeError): parsed = []
        self.year, self.month = (parsed[0], parsed[1]) if len(parsed) == 3 else (today.year, today.month)
        self.title('انتخاب تاریخ شمسی'); self.geometry('470x430'); self.minsize(400, 360); self.transient(parent)
        self._build(); self._render()

    @staticmethod
    def _month_days(year, month):
        if month <= 6: return 31
        if month <= 11: return 30
        return 30 if jdatetime.date(year, 1, 1).isleap() else 29

    def _build(self):
        self.configure(bg='#15202B')
        nav = tk.Frame(self, bg='#1E2C3A'); nav.pack(fill='x', padx=8, pady=8)
        self.year_var = tk.StringVar(); self.month_var = tk.StringVar()
        self.month_box = ttk.Combobox(nav, textvariable=self.month_var, values=self.MONTHS[1:], state='readonly', justify='right', width=12)
        self.month_box.pack(side='right', padx=3, pady=6); self.month_box.bind('<<ComboboxSelected>>', lambda _e: self._month_changed())
        years = tuple(persian_digits(year) for year in range(jdatetime.date.today().year - 30, jdatetime.date.today().year + 31))
        self.year_box = ttk.Combobox(nav, textvariable=self.year_var, values=years, state='readonly', justify='right', width=8)
        self.year_box.pack(side='right', padx=3, pady=6); self.year_box.bind('<<ComboboxSelected>>', lambda _e: self._year_changed())
        button_style = {
            'bg': '#324657', 'fg': '#FFFFFF', 'activebackground': '#41627A',
            'activeforeground': '#FFFFFF', 'font': ('Segoe UI', 10, 'bold'),
            'relief': 'raised', 'bd': 1, 'width': 10, 'anchor': 'center',
            'padx': 6, 'pady': 4,
        }
        tk.Button(nav, text='ماه قبل', command=lambda: self._move(-1), **button_style).pack(side='left', padx=4, pady=6)
        tk.Button(nav, text='ماه بعد', command=lambda: self._move(1), **button_style).pack(side='left', padx=4, pady=6)
        self.grid_host = tk.Frame(self, bg='#15202B'); self.grid_host.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        for column in range(7): self.grid_host.grid_columnconfigure(column, weight=1)
        for row in range(7): self.grid_host.grid_rowconfigure(row, weight=1)

    def _sync_selectors(self):
        self.year_var.set(persian_digits(self.year)); self.month_var.set(self.MONTHS[self.month])

    def _year_changed(self):
        digits = en(self.year_var.get())
        try: self.year = int(digits)
        except ValueError: return
        self._render()

    def _month_changed(self):
        try: self.month = self.MONTHS.index(self.month_var.get())
        except ValueError: return
        self._render()

    def _move(self, delta):
        value = self.year * 12 + self.month - 1 + delta; self.year, zero = divmod(value, 12); self.month = zero + 1; self._render()

    def _render(self):
        self._sync_selectors()
        for child in self.grid_host.winfo_children(): child.destroy()
        for index, label in enumerate(self.DAYS):
            tk.Label(self.grid_host, text=label, bg='#20303F', fg='#DDE7F0', anchor='center').grid(row=0, column=6 - index, sticky='nsew', padx=1, pady=1)
        first = jdatetime.date(self.year, self.month, 1).togregorian(); first_index = (first.weekday() + 2) % 7
        for day in range(1, self._month_days(self.year, self.month) + 1):
            slot = first_index + day - 1; row, day_index = divmod(slot, 7)
            tk.Button(self.grid_host, text=persian_digits(day), command=lambda value=day: self._select(value),
                      bg='#20303F', fg='#DDE7F0', activebackground='#315D86', activeforeground='white', relief='flat').grid(
                          row=row + 1, column=6 - day_index, sticky='nsew', padx=1, pady=1)

    def _select(self, day):
        self.variable.set(persian_digits(f'{self.year:04d}/{self.month:02d}/{day:02d}')); self.destroy()


class SalesTabStrip(RtlNotebookStrip):
    """Sales-only RTL tab strip with rounded title frames and green active text."""
    def rebuild(self):
        bg, line, text, muted, _active_bg, _active_fg = self.colors
        for child in self.bar.winfo_children():
            child.destroy()
        self.buttons = {}
        self.canvas = tk.Canvas(self.bar, bg=bg, highlightthickness=0, bd=0)
        self.canvas.pack(fill='both', expand=True)
        self.canvas.bind('<Configure>', lambda _event: self._draw_tabs(), add='+')
        self._draw_tabs()

    def _rounded_box(self, x1, y1, x2, y2, radius, fill, outline, tag):
        canvas = self.canvas
        canvas.create_rectangle(x1 + radius, y1, x2 - radius, y2, fill=fill, outline=outline, tags=(tag,))
        canvas.create_rectangle(x1, y1 + radius, x2, y2 - radius, fill=fill, outline=outline, tags=(tag,))
        for start, box in ((180, (x1, y1, x1 + 2 * radius, y1 + 2 * radius)),
                           (90, (x2 - 2 * radius, y1, x2, y1 + 2 * radius)),
                           (270, (x1, y2 - 2 * radius, x1 + 2 * radius, y2)),
                           (0, (x2 - 2 * radius, y2 - 2 * radius, x2, y2))):
            canvas.create_arc(*box, start=start, extent=90, fill=fill, outline=outline, tags=(tag,))

    def _draw_tabs(self):
        if not hasattr(self, 'canvas'):
            return
        self.canvas.delete('all')
        try:
            selected = str(self.notebook.select())
        except Exception:
            selected = ''
        width = max(self.canvas.winfo_width(), 320)
        x = width - 8
        self.buttons = {}
        for index, tid in enumerate(self.notebook.tabs()):
            label = str(self.notebook.tab(tid, 'text'))
            tab_width = max(112, min(190, 30 + len(label) * 14))
            x1, x2 = x - tab_width, x
            active = str(tid) == selected
            tag = f'sales_tab{index}'
            self.canvas.create_rectangle(x1, 5, x2, 43,
                                         fill=self.colors[4] if active else self.colors[0],
                                         outline=self.colors[1], width=1, tags=(tag,))
            self.canvas.create_text((x1 + x2) / 2, 24, text=label,
                                    fill=self.colors[5] if active else self.colors[3],
                                    font=self.font, tags=(tag,))
            self.buttons[str(tid)] = tag
            self.canvas.tag_bind(tag, '<Button-1>', lambda _event, target=tid: self.notebook.select(target))
            x = x1 - 6

    def sync(self):
        self._draw_tabs()


class PaymentDocuments(tk.Toplevel):
    """Scrollable, grouped receipt-document viewer for selected finance rows."""
    def __init__(self, parent, *, db, base_dir, payments, contract_total=0):
        super().__init__(parent)
        self.db, self.base_dir = db, Path(base_dir)
        self.payments = list(payments)
        self.contract_total = int(contract_total or 0)
        self.attachments = AttachmentService(db, self.base_dir)
        self.title('اسناد واریزی')
        self.geometry('980x650'); self.minsize(700, 440); self.resizable(True, True)
        self._build(); self._refresh()

    def _build(self):
        self.configure(bg='#15202B')
        top = tk.Frame(self, bg='#1E2C3A'); top.pack(fill='x', padx=8, pady=8)
        tk.Label(top, text='نمایش اسناد واریزی', bg='#1E2C3A', fg='#DDE7F0',
                 font=('Segoe UI', 12, 'bold'), anchor='e').pack(side='right', padx=8, pady=6)
        self.stage_filter = tk.StringVar(value='همه موارد')
        self.filter_box = ttk.Combobox(top, textvariable=self.stage_filter, state='readonly', justify='right', width=24)
        self.filter_box.pack(side='right', padx=8, pady=5); self.filter_box.bind('<<ComboboxSelected>>', lambda _e: self._refresh())
        self.export_choice = tk.StringVar(value=EXPORT_PLACEHOLDER)
        self.export_box = ttk.Combobox(top, textvariable=self.export_choice,
                                       values=(EXPORT_PLACEHOLDER, 'Excel', 'PDF', 'پرینت'),
                                       state='readonly', justify='right', width=20)
        self.export_box.pack(side='left', padx=8, pady=5)
        self.export_box.bind('<<ComboboxSelected>>', self._export_selected)
        area = ttk.Frame(self); area.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        self.tree = ttk.Treeview(area, columns=('stage', 'date', 'amount', 'document'), show='headings')
        for key, label, width in (('stage', 'مرحله پرداخت', 180), ('date', 'تاریخ واریز', 140), ('amount', 'مبلغ واریز (تومان)', 170), ('document', 'فیش واریزی', 420)):
            self.tree.heading(key, text=label, anchor='e'); self.tree.column(key, width=width, anchor='e', stretch=True)
        scroll = ttk.Scrollbar(area, orient='vertical', command=self.tree.yview); xscroll = ttk.Scrollbar(area, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll.set, xscrollcommand=xscroll.set)
        self._model_tree_scrollbars = (scroll, xscroll)
        self.cards_canvas = tk.Canvas(area, bg='#15202B', highlightthickness=0)
        self.cards_scroll = ttk.Scrollbar(area, orient='vertical', command=self.cards_canvas.yview)
        self.cards_host = tk.Frame(self.cards_canvas, bg='#15202B')
        self.cards_window = self.cards_canvas.create_window((0, 0), window=self.cards_host, anchor='nw')
        self.cards_host.bind('<Configure>', lambda _e: self.cards_canvas.configure(scrollregion=self.cards_canvas.bbox('all')))
        self.cards_canvas.bind('<Configure>', lambda event: self.cards_canvas.itemconfigure(self.cards_window, width=event.width))
        self.cards_canvas.configure(yscrollcommand=self.cards_scroll.set)
        self.cards_scroll.pack(side='right', fill='y'); self.cards_canvas.pack(side='left', fill='both', expand=True)
        self.tree.bind('<Double-1>', self._open_document)
        summary_bar = tk.Frame(self, bg='#1E2C3A'); summary_bar.pack(fill='x', padx=8, pady=(0, 8))
        self.summary_values = {}
        for key, label in (('contract', 'مبلغ کل قرارداد (تومان)'), ('paid', 'مبلغ واریزی (تومان)'), ('debt', 'مبلغ بدهکار/بستانکار (تومان)')):
            item = tk.Frame(summary_bar, bg='#1E2C3A'); item.pack(side='right', padx=12, pady=7)
            value = tk.StringVar(value='۰'); self.summary_values[key] = value
            tk.Label(item, textvariable=value, bg='#1E2C3A', fg='#FFFFFF', anchor='e', justify='right', width=16).pack(side='right')
            tk.Label(item, text=label, bg='#1E2C3A', fg='#DDE7F0', anchor='e', justify='right').pack(side='right', padx=(5, 0))

    def _rows_for_filter(self):
        selected = self.stage_filter.get()
        return [item for item in self.payments if selected == 'همه موارد' or item['stage'] == selected]

    def _refresh(self):
        stages = ['همه موارد'] + list(dict.fromkeys(item['stage'] for item in self.payments))
        self.filter_box.configure(values=stages)
        if self.stage_filter.get() not in stages: self.stage_filter.set('همه موارد')
        for iid in self.tree.get_children(): self.tree.delete(iid)
        for index, item in enumerate(self._rows_for_filter(), 1):
            parent = self.tree.insert('', 'end', values=(item['stage'], item['date'], toman(item['amount']), ''))
            documents = item.get('documents', [])
            for document in documents:
                self.tree.insert(parent, 'end', values=(f"↳ {item['stage']}", '', '', document.get('local_path') or document.get('attachment_id') or ''), iid=f'doc{index}_{document.get("attachment_id", index)}')
            if not documents:
                self.tree.insert(parent, 'end', values=(f"↳ {item['stage']}", '', '', 'سند ثبت نشده است'))
            self.tree.item(parent, open=True)
        self._render_cards()

    def _render_cards(self):
        for child in self.cards_host.winfo_children():
            child.destroy()
        visible = self._rows_for_filter()
        paid_total = sum(int(item.get('amount') or 0) for item in visible)
        debt = max(0, self.contract_total - paid_total)
        self.summary_values['contract'].set(toman(self.contract_total))
        self.summary_values['paid'].set(toman(paid_total))
        self.summary_values['debt'].set(toman(debt))
        for index in range(0, len(visible), 2):
            row = tk.Frame(self.cards_host, bg='#15202B'); row.pack(fill='x', pady=5)
            for item in visible[index:index + 2]:
                self._payment_card(row, item)

    def _payment_card(self, parent, item):
        card = tk.Frame(parent, bg='#1E2C3A', highlightbackground='#496176', highlightthickness=1)
        card.pack(side='right', fill='both', expand=True, padx=5)
        tk.Label(card, text=str(item.get('stage') or 'مرحله پرداخت'), bg='#21364B', fg='#F3F7FA',
                 anchor='e', justify='right', padx=8, pady=5).pack(fill='x')
        meta = tk.Frame(card, bg='#1E2C3A'); meta.pack(fill='x', padx=8, pady=5)
        tk.Label(meta, text=f"مبلغ پرداختی (تومان): {toman(item.get('amount') or 0)}",
                 bg='#1E2C3A', fg='#DDE7F0', anchor='e').pack(fill='x')
        tk.Label(meta, text=f"تاریخ پرداختی: {persian_digits(item.get('date') or '-')}",
                 bg='#1E2C3A', fg='#DDE7F0', anchor='e').pack(fill='x')
        docs = list(item.get('documents') or [])
        if not docs:
            tk.Label(card, text='سند ثبت نشده است', bg='#1E2C3A', fg='#F2B8B5', anchor='e', padx=8).pack(fill='x', pady=8)
        for document in docs:
            self._document_card(card, document)

    def _document_card(self, parent, document):
        path = str(document.get('local_path') or '')
        if not path and document.get('attachment_id'):
            path = str((self.attachments.get(str(document['attachment_id'])) or {}).get('local_path') or '')
        label = Path(path).name if path else str(document.get('attachment_id') or 'سند')
        mime = str(document.get('mime_type') or '')
        if mime.startswith('image/') or Path(path).suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
            try:
                from PIL import Image, ImageTk
                image = Image.open(path).convert('RGB')
                image_host = tk.Frame(parent, bg='#1E2C3A'); image_host.pack(side='left', padx=8, pady=6)
                canvas = tk.Canvas(image_host, width=190, height=210, bg='#15202B', highlightthickness=1,
                                   highlightbackground='#496176')
                canvas.pack()
                zoom = tk.DoubleVar(value=1.0)
                def draw(_event=None):
                    width = max(1, canvas.winfo_width() - 8); height = max(1, canvas.winfo_height() - 8)
                    fit = min(width / max(1, image.width), height / max(1, image.height), 1.0)
                    scale = max(0.05, fit * zoom.get())
                    size = (max(1, int(image.width * scale)), max(1, int(image.height * scale)))
                    photo = ImageTk.PhotoImage(image.resize(size)); canvas.delete('all')
                    canvas.create_image(width // 2, height // 2, image=photo, anchor='center'); canvas.image = photo
                    canvas.configure(scrollregion=canvas.bbox('all'))
                controls = tk.Frame(image_host, bg='#1E2C3A'); controls.pack(fill='x')
                tk.Button(controls, text='+', command=lambda: (zoom.set(min(4.0, zoom.get() + .2)), draw()),
                          bg='#315D86', fg='white', relief='flat', width=2).pack(side='left', padx=1)
                tk.Button(controls, text='−', command=lambda: (zoom.set(max(.2, zoom.get() - .2)), draw()),
                          bg='#315D86', fg='white', relief='flat', width=2).pack(side='left', padx=1)
                canvas.bind('<Configure>', draw, add='+')
                canvas.bind('<MouseWheel>', lambda event: (zoom.set(max(.2, min(4.0, zoom.get() + (.2 if event.delta > 0 else -.2)))), draw()), add='+')
                # Keep the receipt in this card: left-click and drag pans a zoomed image.
                canvas.bind('<ButtonPress-1>', lambda event: canvas.scan_mark(event.x, event.y), add='+')
                canvas.bind('<B1-Motion>', lambda event: canvas.scan_dragto(event.x, event.y, gain=1), add='+')
                draw()
                return
            except Exception:
                pass
        tk.Button(parent, text=f'نمایش سند: {label}', command=lambda d=document: self._open_document_data(d),
                  bg='#315D86', fg='white', relief='flat', anchor='e').pack(fill='x', padx=8, pady=6)

    def _export_selected(self, _event=None):
        export_format = {'Excel': 'xlsx', 'PDF': 'pdf', 'پرینت': 'print'}.get(self.export_choice.get())
        if export_format:
            self._export(export_format)

    def _open_image_zoom(self, document, initial_zoom=1.0):
        path = str(document.get('local_path') or '')
        if not path and document.get('attachment_id'):
            path = str((self.attachments.get(str(document['attachment_id'])) or {}).get('local_path') or '')
        try:
            from PIL import Image, ImageTk
            image = Image.open(path)
        except Exception as exc:
            return messagebox.showerror('فیش واریزی', str(exc), parent=self)
        viewer = tk.Toplevel(self); viewer.title('پیش‌نمایش فیش واریزی'); viewer.geometry('800x600'); viewer.minsize(480, 360); viewer.resizable(True, True)
        zoom = tk.DoubleVar(value=max(0.2, min(4.0, float(initial_zoom or 1.0))))
        canvas = tk.Canvas(viewer, bg='#101820'); canvas.pack(fill='both', expand=True)
        def draw(_event=None):
            scale = max(0.2, min(4.0, zoom.get())); size = (max(1, int(image.width * scale)), max(1, int(image.height * scale)))
            photo = ImageTk.PhotoImage(image.resize(size)); canvas.delete('all'); canvas.create_image(10, 10, image=photo, anchor='nw'); canvas.image = photo
            canvas.configure(scrollregion=canvas.bbox('all'))
        controls = tk.Frame(viewer, bg='#1E2C3A'); controls.pack(fill='x')
        tk.Button(controls, text='+', command=lambda: (zoom.set(zoom.get() + .2), draw()), bg='#315D86', fg='white').pack(side='left')
        tk.Button(controls, text='−', command=lambda: (zoom.set(zoom.get() - .2), draw()), bg='#315D86', fg='white').pack(side='left')
        canvas.bind('<MouseWheel>', lambda event: (zoom.set(max(0.2, min(4.0, zoom.get() + (0.2 if event.delta > 0 else -0.2)))), draw()))
        viewer.bind('<Configure>', draw); draw()

    def _open_document_data(self, document):
        local_path = Path(str(document.get('local_path') or ''))
        if str(document.get('mime_type') or '').startswith('text/') or local_path.suffix.lower() in {'.txt', '.csv', '.log'}:
            try:
                text = local_path.read_text(encoding='utf-8', errors='replace')
                viewer = tk.Toplevel(self); viewer.title('متن سند مالی'); viewer.geometry('760x520'); viewer.minsize(480, 320)
                box = tk.Text(viewer, wrap='word', padx=10, pady=10); scroll = ttk.Scrollbar(viewer, orient='vertical', command=box.yview)
                box.configure(yscrollcommand=scroll.set); scroll.pack(side='right', fill='y'); box.pack(side='left', fill='both', expand=True); box.insert('1.0', text); box.configure(state='disabled')
            except OSError as exc:
                messagebox.showerror('متن سند مالی', str(exc), parent=self)
        else:
            try: self.attachments.open(str(document.get('attachment_id')))
            except Exception as exc: messagebox.showerror('فیش واریزی', str(exc), parent=self)

    def _open_document(self, _event=None):
        selected = self.tree.selection()
        if not selected: return
        values = self.tree.item(selected[0], 'values')
        path = str(values[3] or '')
        if path == 'سند ثبت نشده است':
            return messagebox.showinfo('اسناد واریزی', 'سند ثبت نشده است.', parent=self)
        document = next((d for item in self.payments for d in item.get('documents', []) if str(d.get('local_path') or d.get('attachment_id') or '') == path), None)
        if document:
            local_path = Path(str(document.get('local_path') or ''))
            if str(document.get('mime_type') or '').startswith('text/') or local_path.suffix.lower() in {'.txt', '.csv', '.log'}:
                try:
                    text = local_path.read_text(encoding='utf-8', errors='replace')
                    viewer = tk.Toplevel(self); viewer.title('متن سند مالی'); viewer.geometry('760x520'); viewer.minsize(480, 320)
                    box = tk.Text(viewer, wrap='word', padx=10, pady=10); scroll = ttk.Scrollbar(viewer, orient='vertical', command=box.yview)
                    box.configure(yscrollcommand=scroll.set); scroll.pack(side='right', fill='y'); box.pack(side='left', fill='both', expand=True)
                    box.insert('1.0', text); box.configure(state='disabled')
                except OSError as exc:
                    messagebox.showerror('متن سند مالی', str(exc), parent=self)
            else:
                try: self.attachments.open(str(document.get('attachment_id')))
                except Exception as exc: messagebox.showerror('فیش واریزی', str(exc), parent=self)

    def _visible_rows(self):
        rows = []
        for parent in self.tree.get_children():
            rows.append(tuple(self.tree.item(parent, 'values')))
            rows.extend(tuple(self.tree.item(child, 'values')) for child in self.tree.get_children(parent))
        return rows

    def _export(self, fmt):
        headers = ('مرحله پرداخت', 'تاریخ واریز', 'مبلغ واریز (تومان)', 'فیش واریزی')
        visible = self._rows_for_filter()
        if not visible: return messagebox.showinfo('خروجی اسناد', 'سندی برای خروجی وجود ندارد.', parent=self)
        paid_total = sum(int(item.get('amount') or 0) for item in visible)
        debt = max(0, self.contract_total - paid_total)
        summary_lines = [
            f'مبلغ کل قرارداد (تومان): {toman(self.contract_total)}',
            f'مبلغ واریزی (تومان): {toman(paid_total)}',
            f'مبلغ بدهکار/بستانکار (تومان): {toman(debt)}',
        ]
        payload = []
        document_paths = []
        for item in visible:
            paths = []
            for document in item.get('documents', []):
                path = str(document.get('local_path') or '').strip()
                if not path and document.get('attachment_id'):
                    path = str((self.attachments.get(str(document['attachment_id'])) or {}).get('local_path') or '').strip()
                if path and path not in paths:
                    paths.append(path)
                if path and path not in document_paths:
                    document_paths.append(path)
            # One visible row per receipt makes every selected voucher visible
            # in the preview.  Only the first supporting document carries the
            # payment amount/date, so summary totals are never duplicated.
            document_rows = paths or ['']
            for document_number, path in enumerate(document_rows, 1):
                row_stage = str(item.get('stage') or 'مرحله پرداخت')
                if len(document_rows) > 1:
                    row_stage = f'{row_stage} · فیش {persian_digits(document_number)}'
                payload.append({
                    'مرحله پرداخت': row_stage,
                    'تاریخ واریز': persian_digits(item.get('date') or '-') if document_number == 1 else '—',
                    'مبلغ واریز (تومان)': toman(item.get('amount') or 0) if document_number == 1 else '—',
                    'فیش واریزی': path,
                })
        specs = [('خلاصه', line.split(':', 1)[0], line.split(':', 1)[1].strip()) for line in summary_lines]
        choices = {'xlsx': 'Excel', 'pdf': 'PDF', 'print': 'پرینت'}
        return ContractExportPreview(
            self, export_format=choices.get(fmt, 'PDF'), headers=list(headers), payload=payload, specs=specs,
            window_title='پیش‌نمایش اسناد واریزی',
            on_confirm=lambda choice, orientation: self._save_export(choice, list(headers), payload, specs, document_paths, orientation),
        )

    def _save_export(self, choice, headers, payload, specs, document_paths, orientation='افقی'):
        export_format = {'Excel': 'xlsx', 'PDF': 'pdf', 'پرینت': 'print'}.get(choice, 'pdf')
        actual_format = 'pdf' if export_format == 'print' else export_format
        try:
            path = filedialog.asksaveasfilename(
                parent=self, title=f'خروجی {actual_format.upper()} اسناد واریزی',
                defaultextension=f'.{actual_format}',
                filetypes=[('Excel', '*.xlsx')] if actual_format == 'xlsx' else [('PDF', '*.pdf')],
            )
            if not path:
                return False
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            if actual_format == 'xlsx':
                from openpyxl import Workbook
                book = Workbook(); sheet = book.active; sheet.append(headers)
                for row in payload:
                    sheet.append([row.get(header, '') for header in headers])
                    document_path = str(row.get('فیش واریزی') or '')
                    if document_path and Path(document_path).exists() and Path(document_path).suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
                        try:
                            from openpyxl.drawing.image import Image as ExcelImage
                            image = ExcelImage(document_path); image.width, image.height = 160, 110
                            image.anchor = f'{chr(65 + max(0, len(headers) - 1))}{sheet.max_row}'
                            sheet.add_image(image)
                        except (OSError, ImportError, ValueError):
                            pass
                sheet.append(())
                for group, label, value in specs: sheet.append((label, value))
                book.save(path)
            else:
                export_receipts_pdf(
                    path, headers, [[row.get(header, '') for header in headers] for row in payload],
                    'اسناد واریزی', [f'{label}: {value}' for _group, label, value in specs], document_paths,
                    orientation=orientation,
                )
                if export_format == 'print' and hasattr(os, 'startfile'):
                    os.startfile(path, 'print')
        except Exception as exc:
            messagebox.showerror('خروجی اسناد', str(exc), parent=self)
            return False
        messagebox.showinfo('خروجی اسناد', 'خروجی آماده شد.', parent=self)
        return True


class ContractWizard(tk.Toplevel):
    def __init__(self, parent, *, db, base_dir: Path, on_saved, record_opener=None, initial: dict[str, Any] | None = None, document_path: str = '', sale_id: str = ''):
        super().__init__(parent)
        self.db, self.base_dir, self.on_saved, self.sale_id = db, Path(base_dir), on_saved, str(sale_id or '')
        self.record_opener = record_opener
        self.revenue = RevenueService(db)
        self.step, self.document_path = 0, str(document_path or '')
        self.values = {key: tk.StringVar() for key in (
            'project_name', 'block_name', 'deal_mode', 'sale_kind', 'unit_number', 'floor', 'map_area', 'contract_area',
            'contract_number', 'contract_date', 'delivery_date', 'deed_date', 'buyer_name', 'father_name',
            'national_id', 'buyer_phone', 'total_price', 'price_per_m2', 'broker_fee', 'payment_count',
            'notary_name', 'witness_1', 'witness_2', 'buyer_representative', 'seller_representative',
        )}
        for number in range(1, 14):
            self.values[f'pay{number}_amount'] = tk.StringVar()
            self.values[f'pay{number}_date'] = tk.StringVar()
        self.values['sale_kind'].set('پیش‌فروش')
        self.values['deal_mode'].set('واحد کامل')
        self.values['payment_count'].set('۴')
        self._payment_visible_installments = 4
        self._unit_combo = None
        self._unit_records: dict[str, dict[str, Any]] = {}
        for key, value in (initial or {}).items():
            if key in self.values and value not in (None, ''):
                self.values[key].set(str(value))
        self._digit_normalizing = False
        for key, variable in self.values.items():
            variable.trace_add('write', lambda *_args, current=variable, field=key: self._format_wizard_variable(field, current))
            self._persianize_variable(variable)
        self.values['project_name'].trace_add('write', lambda *_: self._refresh_unit_choices())
        self.values['block_name'].trace_add('write', lambda *_: self._refresh_unit_choices())
        self.title('ثبت قرارداد فروش — Wizard — BARKAV 18.7.1')
        self.geometry('900x650')
        self.minsize(740, 560)
        self.transient(parent)
        self.protocol('WM_DELETE_WINDOW', self._close_to_record)
        self._build_shell()
        self._show_step()

    def _persianize_variable(self, variable):
        if self._digit_normalizing:
            return
        value = variable.get(); display = persian_digits(value)
        if value != display:
            self._digit_normalizing = True
            try: variable.set(display)
            finally: self._digit_normalizing = False

    def _format_wizard_variable(self, key, variable):
        self._persianize_variable(variable)
        if key not in {'total_price', 'broker_fee', 'price_per_m2'} and not key.startswith('pay'):
            return
        value = formatted_amount_input(variable.get())
        if value != variable.get() and not self._digit_normalizing:
            self._digit_normalizing = True
            try: variable.set(value)
            finally: self._digit_normalizing = False

    def _build_shell(self):
        bg, surface, panel, line, text, muted = '#071B2C', '#0D2234', '#10283C', '#324657', '#DDE7F0', '#9FB3C8'
        self.configure(bg=bg)
        header = tk.Frame(self, bg=surface, highlightbackground=line, highlightthickness=1)
        header.pack(fill='x', padx=8, pady=8)
        header.grid_columnconfigure(1, weight=1)
        self.title_var = tk.StringVar()
        tk.Label(header, text='ثبت مرحله‌ای قرارداد', bg=surface, fg=text, anchor='e', justify='right',
                 font=('Segoe UI', 12, 'bold')).grid(row=0, column=2, sticky='e', padx=12, pady=10)
        self.progress_status = tk.Label(header, text='هر مرحله فقط اطلاعات مرتبط با قرارداد را نشان می‌دهد.', bg=surface,
                                        fg=muted, anchor='e', justify='right', font=('Segoe UI', 9))
        self.progress_status.grid(row=0, column=1, sticky='ew', padx=12)
        progress = tk.Frame(header, bg=surface)
        progress.grid(row=0, column=0, sticky='w', padx=12)
        self.step_label = tk.Label(progress, text='', bg=surface, fg=muted, anchor='e', font=('Segoe UI', 9))
        self.step_label.pack(side='right', padx=(0, 8))
        self.progress_canvas = tk.Canvas(progress, width=128, height=16, bg=surface, highlightthickness=0)
        self.progress_canvas.pack(side='left')
        self.progress_canvas.bind('<Configure>', lambda _event: self._render_progress(), add='+')
        body_host = tk.Frame(self, bg=panel, highlightbackground=line, highlightthickness=1)
        body_host.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        self.body_canvas = tk.Canvas(body_host, bg=panel, highlightthickness=0)
        body_scroll = ttk.Scrollbar(body_host, orient='vertical', command=self.body_canvas.yview)
        self.body_canvas.configure(yscrollcommand=body_scroll.set)
        body_scroll.pack(side='right', fill='y'); self.body_canvas.pack(side='left', fill='both', expand=True)
        self.body = tk.Frame(self.body_canvas, bg=panel)
        body_window = self.body_canvas.create_window((0, 0), window=self.body, anchor='nw')
        self.body.bind('<Configure>', lambda _e: self.body_canvas.configure(scrollregion=self.body_canvas.bbox('all')))
        self.body_canvas.bind('<Configure>', lambda event: self.body_canvas.itemconfigure(body_window, width=event.width))
        nav = tk.Frame(self, bg=bg)
        nav.pack(fill='x', padx=8, pady=(0, 8))
        self.back = tk.Button(nav, text='→ قبلی', command=self._back, bg=surface, fg=text,
                              relief='flat', padx=14, pady=8)
        self.back.pack(side='right', padx=3)
        self.next = tk.Button(nav, text='بعدی ←', command=self._next, bg='#2563A6', fg='white',
                              relief='flat', padx=18, pady=8, font=('Segoe UI', 9, 'bold'))
        self.next.pack(side='right', padx=3)
        tk.Button(nav, text='پنجره رکورد فروش', command=self._close_to_record, bg=surface, fg=text,
                  relief='flat', padx=12, pady=8).pack(side='left', padx=3)

    def _render_progress(self):
        """Draw the same compact stage-progress treatment used by the data wizard."""
        canvas = getattr(self, 'progress_canvas', None)
        if canvas is None:
            return
        canvas.delete('all')
        width, height = max(int(canvas.winfo_width()), 128), max(int(canvas.winfo_height()), 16)
        canvas.create_rectangle(0, 0, width - 1, height - 1, fill='#0A1A2A', outline='#324657')
        filled = max(2, int((width - 4) * (self.step + 1) / len(CONTRACT_STEPS)))
        canvas.create_rectangle(2, 2, 2 + filled, height - 3, fill='#31B9D0', outline='')
        self.step_label.configure(text=f"مرحله {fa(self.step + 1)} از {fa(len(CONTRACT_STEPS))}")

    def _export_wizard(self, export_format: str):
        payload = wizard_export_payload({key: value.get().strip() for key, value in self.values.items()}, self.document_path)
        extension = f'.{export_format}'
        path = filedialog.asksaveasfilename(parent=self, title=f'خروجی {export_format.upper()} ویزارد قرارداد', defaultextension=extension, filetypes=[(export_format.upper(), f'*{extension}')])
        if not path:
            return
        try:
            destination = Path(path)
            if export_format == 'csv':
                with destination.open('w', encoding='utf-8-sig', newline='') as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(payload)); writer.writeheader(); writer.writerow(payload)
            elif export_format == 'txt':
                destination.write_text('\n'.join(f'{key}: {value}' for key, value in payload.items()), encoding='utf-8')
            else:
                destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
        except OSError as exc:
            return messagebox.showerror('خروجی ویزارد', str(exc), parent=self)
        messagebox.showinfo('خروجی ویزارد', 'پیش‌نویس فعلی ویزارد صادر شد.', parent=self)

    def _field(self, parent, key, label, row, column=0, *, readonly=False, values=None):
        box = tk.Frame(parent, bg='#1E2C3A')
        box.grid(row=row, column=column, sticky='ew', padx=6, pady=5)
        tk.Label(box, text=label, bg='#1E2C3A', fg='#DDE7F0', anchor='e').pack(fill='x', padx=6, pady=(5, 1))
        # An empty project-unit list must still render as a dropdown; the
        # choices are populated as soon as project/block is selected.
        if values is not None:
            widget = ttk.Combobox(box, textvariable=self.values[key], values=values, state='readonly', justify='right')
            widget.pack(fill='x', padx=6, pady=(0, 6))
        else:
            date_field = key in {'contract_date', 'delivery_date', 'deed_date'} or (key.startswith('pay') and key.endswith('_date')) or key.endswith('_date')
            if date_field:
                line = tk.Frame(box, bg='#1E2C3A'); line.pack(fill='x', padx=6, pady=(0, 6))
                widget = ttk.Entry(line, textvariable=self.values[key], justify='right')
                widget.pack(side='right', fill='x', expand=True)
                tk.Button(line, text='تقویم', command=lambda current=self.values[key]: ShamsiDatePicker(self, current),
                          bg='#324657', fg='#DDE7F0', relief='flat', padx=7).pack(side='left', padx=(5, 0))
            else:
                widget = ttk.Entry(box, textvariable=self.values[key], justify='right')
            if readonly:
                widget.configure(state='readonly')
                widget.pack(fill='x', padx=6, pady=(0, 6))
            elif not (values is None and (key in {'contract_date', 'delivery_date', 'deed_date'} or key.endswith('_date'))):
                widget.pack(fill='x', padx=6, pady=(0, 6))
        return widget

    def _available_units(self) -> dict[str, dict[str, Any]]:
        """Read project units for selection; never creates or mutates units."""
        project = self.values['project_name'].get().strip()
        block = self.values['block_name'].get().strip()
        try:
            with self.db.read() as con:
                rows = con.execute(
                    """SELECT unit_number, floor, area FROM units
                       WHERE status!='CANCELLED'
                         AND (?='' OR COALESCE(project_name,'')=?)
                         AND (?='' OR COALESCE(block_name,'')=?)
                       ORDER BY unit_number""",
                    (project, project, block, block),
                ).fetchall()
        except Exception:
            rows = []
        return {str(row['unit_number']): {'floor': row['floor'], 'area': row['area']} for row in rows if row['unit_number'] not in (None, '')}

    def _refresh_unit_choices(self):
        combo = getattr(self, '_unit_combo', None)
        if combo is None:
            return
        self._unit_records = self._available_units()
        current = self.values['unit_number'].get().strip()
        choices = list(self._unit_records)
        if current and current not in choices:
            choices.insert(0, current)
        combo.configure(values=tuple(persian_digits(value) for value in choices))

    def _unit_selected(self, _event=None):
        selected = self.values['unit_number'].get().strip()
        key = en(selected)
        record = self._unit_records.get(selected) or self._unit_records.get(key)
        if not record:
            return
        self.values['floor'].set(persian_digits(record.get('floor') or ''))
        self.values['map_area'].set(persian_digits(record.get('area') or ''))

    def _show_step(self):
        for child in self.body.winfo_children():
            child.destroy()
        self.title_var.set(CONTRACT_STEPS[self.step])
        self._render_progress()
        self.back.configure(state='normal' if self.step else 'disabled')
        self.next.configure(text='ثبت قرارداد' if self.step == len(CONTRACT_STEPS) - 1 else 'بعدی ←')
        page = tk.LabelFrame(self.body, text=CONTRACT_STEPS[self.step], labelanchor='ne', bg='#15202B', fg='#DDE7F0',
                             font=('Segoe UI', 10, 'bold'), padx=8, pady=8)
        page.pack(fill='both', expand=True)
        page.grid_columnconfigure((0, 1), weight=1)
        if self.step == 0:
            ttk.Button(page, text='بارگذاری قرارداد', command=self._load_document).grid(row=0, column=1, sticky='e', padx=8)
            self.document_label = tk.Label(page, text=self.document_path or 'فایلی انتخاب نشده است', bg='#15202B', fg='#9FB3C8', anchor='e', justify='right')
            self.document_label.grid(row=1, column=0, columnspan=2, sticky='ew', padx=8, pady=10)
            self.deal_mode_combo = self._field(page, 'deal_mode', 'روش معامله', 2, 1, values=('واحد کامل', 'فروش متری'))
            self.sale_kind_combo = self._field(page, 'sale_kind', 'نوع فروش', 2, 0, values=('پیش‌فروش', 'آماده'))
            # Keep both step-one selectors explicit/read-only so the dropdown
            # affordance remains available with every active ttk theme.
            self.deal_mode_combo.configure(values=('واحد کامل', 'فروش متری'), state='readonly')
            self.sale_kind_combo.configure(values=('پیش‌فروش', 'آماده'), state='readonly')
            self._field(page, 'project_name', 'نام پروژه', 3, 1)
            self._field(page, 'block_name', 'بلوک', 3, 0)
            self._unit_records = self._available_units()
            self._deal_widgets = [
                self._field(page, 'unit_number', 'شماره واحد', 4, 1, values=tuple(persian_digits(value) for value in self._unit_records)),
                self._field(page, 'floor', 'طبقه', 4, 0),
                self._field(page, 'map_area', 'متراژ طبق نقشه', 5, 1),
            ]
            self._unit_combo = self._deal_widgets[0]
            self._unit_combo.bind('<<ComboboxSelected>>', self._unit_selected)
            self._field(page, 'contract_area', 'متراژ قرارداد', 5, 0)
            self.values['deal_mode'].trace_add('write', lambda *_: self._toggle_deal_mode())
            self._toggle_deal_mode()
        elif self.step == 1:
            self._field(page, 'contract_number', 'شماره قرارداد', 0, 1)
            self._field(page, 'contract_date', 'تاریخ قرارداد', 1, 1)
            self._field(page, 'delivery_date', 'تاریخ تحویل', 2, 1)
            self._field(page, 'deed_date', 'تاریخ سند', 3, 1)
        elif self.step == 2:
            self._field(page, 'buyer_name', 'نام و نام خانوادگی', 0, 1)
            self._field(page, 'father_name', 'نام پدر', 1, 1)
            self._field(page, 'national_id', 'شماره ملی', 2, 1)
            self._field(page, 'buyer_phone', 'شماره تماس', 3, 1)
        elif self.step == 3:
            self._field(page, 'total_price', 'قیمت کل (تومان)', 0, 1)
            price = self._field(page, 'price_per_m2', 'قیمت هر متر مربع (تومان)', 1, 1, readonly=True)
            self._field(page, 'broker_fee', 'کمیسیون بنگاه (تومان)', 2, 1)
            self.values['total_price'].trace_add('write', lambda *_: self._price_per_m2())
            self.values['contract_area'].trace_add('write', lambda *_: self._price_per_m2())
            self._price_per_m2()
        elif self.step == 4:
            count_box = self._field(page, 'payment_count', 'تعداد اقساط', 0, 1)
            count_box.bind('<FocusOut>', self._payment_count_changed)
            count_box.bind('<Return>', self._payment_count_changed)
            count = self._payment_installment_count()
            self._payment_visible_installments = count
            labels = payment_labels_for_installments(count)
            self._payment_widgets = []
            for index, label in enumerate(labels, 1):
                amount = self._field(page, f'pay{index}_amount', f'{label} (تومان)', index, 1)
                due = self._field(page, f'pay{index}_date', f'تاریخ {label}', index, 0)
                self._payment_widgets.append((self._payment_widget_container(amount), self._payment_widget_container(due)))
            self._toggle_payment_rows()
        elif self.step == 5:
            self._field(page, 'notary_name', 'نام دفترخانه', 0, 1)
            self._field(page, 'witness_1', 'شاهد ۱', 1, 1)
            self._field(page, 'witness_2', 'شاهد ۲', 2, 1)
            self._field(page, 'buyer_representative', 'نماینده خریدار', 3, 1)
            self._field(page, 'seller_representative', 'نماینده فروشنده', 4, 1)
        else:
            summary = self._summary_text()
            label = tk.Label(page, text=summary, bg='#1E2C3A', fg='#DDE7F0', justify='right', anchor='ne',
                             font=('Segoe UI', 11), padx=16, pady=16)
            label.pack(fill='both', expand=True)

    def _load_document(self):
        path = filedialog.askopenfilename(parent=self, title='قرارداد', filetypes=[('Documents', '*.pdf *.jpg *.jpeg *.png *.doc *.docx *.txt'), ('All files', '*.*')])
        if not path:
            return
        self.document_path = path
        text = ''
        if Path(path).suffix.lower() == '.txt':
            try: text = Path(path).read_text(encoding='utf-8')
            except UnicodeDecodeError: text = Path(path).read_text(encoding='utf-8-sig')
        else:
            try:
                from receipt_intelligence import LocalOCREngine
                text = LocalOCREngine().extract(Path(path)).text
            except Exception:
                text = ''
        parsed = parse_sale_form(text)
        mapping = {'project_name': 'project_name', 'block_name': 'block_name', 'unit_number': 'unit_number',
                   'floor': 'floor', 'contract_area': 'area', 'buyer_name': 'buyer_name', 'buyer_phone': 'buyer_phone',
                   'contract_number': 'contract_number', 'contract_date': 'contract_date_shamsi', 'total_price': 'contract_amount'}
        for target, source in mapping.items():
            if parsed.get(source) not in (None, '', 0): self.values[target].set(str(parsed[source]))
        for row in parsed.get('installments') or []:
            number = min(13, int(row.get('sequence_no') or 1))
            self.values[f'pay{number}_amount'].set(str(row.get('amount') or ''))
            self.values[f'pay{number}_date'].set(str(row.get('due_date_shamsi') or ''))
        self._show_step()

    def _price_per_m2(self):
        area = float(en(str(self.values['contract_area'].get() or '0')).replace(',', '') or 0)
        price = parse_money(self.values['total_price'].get())
        self.values['price_per_m2'].set(toman(price / area) if area > 0 else '')

    def _toggle_deal_mode(self):
        """In a meter-sale only contract area remains editable, as specified."""
        meter_sale = self.values['deal_mode'].get() == 'فروش متری'
        for widget in getattr(self, '_deal_widgets', []):
            widget.configure(state='disabled' if meter_sale else 'normal')

    @staticmethod
    def _payment_widget_container(widget):
        """Return the outer grid-managed field frame, not a packed date line."""
        current = widget
        while current.winfo_manager() == 'pack':
            current = current.master
        return current

    def _payment_installment_count(self) -> int:
        try:
            # The entered stage count includes cash, delivery and deed.  Only
            # the remaining stages are dynamic installment rows.
            return max(1, min(10, int(en(self.values['payment_count'].get())) - 3))
        except ValueError:
            return 1

    def _payment_count_changed(self, _event=None):
        """Keep cash and the two fixed final stages intact when count changes."""
        old_count = getattr(self, '_payment_visible_installments', self._payment_installment_count())
        new_count = self._payment_installment_count()
        if old_count != new_count:
            old_delivery, old_deed = old_count + 2, old_count + 3
            new_delivery, new_deed = new_count + 2, new_count + 3
            for suffix in ('amount', 'date'):
                delivery_value = self.values[f'pay{old_delivery}_{suffix}'].get()
                deed_value = self.values[f'pay{old_deed}_{suffix}'].get()
                if old_delivery != new_delivery:
                    self.values[f'pay{old_delivery}_{suffix}'].set('')
                    self.values[f'pay{new_delivery}_{suffix}'].set(delivery_value)
                if old_deed != new_deed:
                    self.values[f'pay{old_deed}_{suffix}'].set('')
                    self.values[f'pay{new_deed}_{suffix}'].set(deed_value)
        self._payment_visible_installments = new_count
        self._show_step()

    def _toggle_payment_rows(self):
        count = self._payment_installment_count() + 3
        for index, widgets in enumerate(getattr(self, '_payment_widgets', []), 1):
            for widget in widgets:
                if index <= count:
                    widget.grid()
                else:
                    widget.grid_remove()

    def _summary_text(self):
        labels = {
            'project_name': 'نام پروژه', 'block_name': 'بلوک', 'deal_mode': 'روش معامله', 'sale_kind': 'نوع فروش',
            'unit_number': 'شماره واحد', 'floor': 'طبقه', 'map_area': 'متراژ طبق نقشه', 'contract_area': 'متراژ قرارداد',
            'contract_number': 'شماره قرارداد', 'contract_date': 'تاریخ قرارداد', 'delivery_date': 'تاریخ تحویل',
            'deed_date': 'تاریخ سند', 'buyer_name': 'نام و نام خانوادگی', 'father_name': 'نام پدر',
            'national_id': 'شماره ملی', 'buyer_phone': 'شماره تماس', 'total_price': 'قیمت کل (تومان)',
            'price_per_m2': 'قیمت هر مترمربع (تومان)', 'broker_fee': 'کمیسیون بنگاه (تومان)',
            'payment_count': 'تعداد مراحل پرداخت', 'notary_name': 'نام دفترخانه', 'witness_1': 'شاهد ۱',
            'witness_2': 'شاهد ۲', 'buyer_representative': 'نماینده خریدار', 'seller_representative': 'نماینده فروشنده',
        }
        lines = [f"{labels.get(key, 'اطلاعات')}: {persian_digits(value.get() or '-')}"
                 for key, value in self.values.items() if not key.startswith('pay')]
        for index, title in enumerate(payment_labels_for_installments(self._payment_installment_count()), 1):
            amount = self.values[f'pay{index}_amount'].get() or '-'; due = self.values[f'pay{index}_date'].get() or '-'
            lines.append(f'{title}: مبلغ {persian_digits(amount)} تومان | تاریخ {persian_digits(due)}')
        return '\n'.join(lines)

    def _back(self):
        self.step = max(0, self.step - 1); self._show_step()

    def _next(self):
        if self.step < len(CONTRACT_STEPS) - 1:
            self.step += 1; self._show_step(); return
        self._save()

    def _save(self):
        raw = {key: value.get().strip() for key, value in self.values.items()}
        try:
            raw['payment_count'] = str(self._payment_installment_count())
            plan = payment_plan(raw)
            total = parse_money(raw['total_price'])
            if total <= 0 or total != sum(int(row['amount']) for row in plan):
                raise ValueError('جمع مراحل پرداخت باید دقیقاً با قیمت کل قرارداد برابر باشد.')
            contract_date = normalize_shamsi(raw['contract_date'])
            unit_number = raw['unit_number']
            if raw['deal_mode'] == 'فروش متری':
                unit_number = f"متری-{raw['contract_number'] or raw['buyer_name']}"
            if not contract_date or not raw['buyer_name'] or not unit_number:
                raise ValueError('شماره واحد، نام خریدار و تاریخ قرارداد الزامی است.')
            result = self.revenue.create_sale_bundle(
                project_name=raw['project_name'] or 'پروژه', block_name=raw['block_name'], unit_number=unit_number,
                floor=raw['floor'], area=float(en(raw['contract_area']).replace(',', '') or 0) or None,
                buyer_name=raw['buyer_name'], buyer_phone=raw['buyer_phone'], contract_number=raw['contract_number'],
                contract_date_shamsi=contract_date, contract_amount=total, installments=plan,
                notes=contract_note(raw), source='GUI_1871_CONTRACT', actor_id='GUI',
            )
            if self.document_path:
                AttachmentService(self.db, self.base_dir).add_local(entity_type='SALE', entity_id=result['sale_id'], source_path=self.document_path, document_type='CONTRACT')
        except Exception as exc:
            messagebox.showerror('ثبت قرارداد', str(exc), parent=self); return
        self.on_saved(); self.destroy()

    def _close_to_record(self):
        draft = {key: value.get().strip() for key, value in self.values.items()}
        if self.record_opener:
            self.record_opener(sale_id=self.sale_id, draft=draft, document_path=self.document_path)
        else:
            try: self.master.open_selected_record(sale_id=self.sale_id, draft=draft, document_path=self.document_path)
            except Exception: pass
        self.destroy()


class SalesRecord(tk.Toplevel):
    """Dedicated sales record: one editable view, separate from data-entry records."""
    _FIELDS = (
        ('project_name', 'نام پروژه'), ('block_name', 'بلوک'), ('deal_mode', 'روش معامله'),
        ('sale_kind', 'نوع فروش'), ('unit_number', 'شماره واحد'), ('floor', 'طبقه'),
        ('map_area', 'متراژ طبق نقشه'), ('contract_area', 'متراژ قرارداد'),
        ('contract_number', 'شماره قرارداد'), ('contract_date', 'تاریخ قرارداد'),
        ('delivery_date', 'تاریخ تحویل'), ('deed_date', 'تاریخ سند'),
        ('buyer_name', 'نام و نام خانوادگی'), ('father_name', 'نام پدر'),
        ('national_id', 'شماره ملی'), ('buyer_phone', 'شماره تماس'),
        ('total_price', 'قیمت کل (تومان)'), ('price_per_m2', 'قیمت هر مترمربع (تومان)'),
        ('broker_fee', 'کمیسیون بنگاه (تومان)'), ('payment_count', 'تعداد مراحل پرداخت'),
        ('notary_name', 'نام دفترخانه'), ('witness_1', 'شاهد ۱'),
        ('witness_2', 'شاهد ۲'), ('buyer_representative', 'نماینده خریدار'), ('seller_representative', 'نماینده فروشنده'),
    )
    _WIZARD_GROUPS = (
        ('۱. مورد معامله', ('project_name', 'block_name', 'deal_mode', 'sale_kind', 'unit_number', 'floor', 'map_area', 'contract_area')),
        ('۲. تاریخ‌ها', ('contract_number', 'contract_date', 'delivery_date', 'deed_date')),
        ('۳. خریدار', ('buyer_name', 'father_name', 'national_id', 'buyer_phone')),
        ('۴. مالی', ('total_price', 'price_per_m2', 'broker_fee')),
        ('۶. امضاکنندگان', ('notary_name', 'witness_1', 'witness_2', 'buyer_representative', 'seller_representative')),
    )
    _FIELD_LABELS = dict(_FIELDS)

    def __init__(self, parent, *, db, base_dir: Path, sale_id: str = '', draft: dict[str, Any] | None = None,
                 document_path: str = '', on_saved=None):
        super().__init__(parent)
        self.db, self.base_dir, self.sale_id, self.on_saved = db, Path(base_dir), str(sale_id or ''), on_saved
        self.revenue = RevenueService(db)
        self.document_path = document_path
        initial = self._initial_values(draft or {})
        self.values = {key: tk.StringVar(value=str(initial.get(key) or '')) for key, _label in self._FIELDS}
        if not self.values['payment_count'].get().strip():
            self.values['payment_count'].set('۴')
        for index in range(1, 14):
            self.values[f'pay{index}_amount'] = tk.StringVar(value=str(initial.get(f'pay{index}_amount') or ''))
            self.values[f'pay{index}_date'] = tk.StringVar(value=str(initial.get(f'pay{index}_date') or ''))
        self._payment_visible_installments = self._payment_installment_count()
        self._unit_records: dict[str, dict[str, Any]] = {}
        self._unit_combo = None
        self._digit_normalizing = False
        for key, variable in self.values.items():
            variable.trace_add('write', lambda *_args, current=variable, field=key: self._format_record_variable(field, current))
            self._persianize_variable(variable)
        self.values['project_name'].trace_add('write', lambda *_: self._refresh_unit_choices())
        self.values['block_name'].trace_add('write', lambda *_: self._refresh_unit_choices())
        self.values['total_price'].trace_add('write', lambda *_: self._record_price_per_m2())
        self.values['contract_area'].trace_add('write', lambda *_: self._record_price_per_m2())
        self._record_price_per_m2()
        self.title('پنجره رکورد فروش')
        # Keep the normal Windows title bar so native minimize/maximize/close
        # controls remain available in the white title area.
        self.geometry('1120x720'); self.minsize(860, 560); self.resizable(True, True)
        self._preview_image = None
        self._preview_pil_image = None
        self._build()

    def _initial_values(self, draft: dict[str, Any]) -> dict[str, Any]:
        values = dict(draft)
        if not self.sale_id:
            return values
        row = self.revenue.get_sale(self.sale_id) or {}
        try: metadata = json.loads(str(row.get('notes') or '{}'))
        except (ValueError, TypeError, json.JSONDecodeError): metadata = {}
        values = {**metadata, **values}
        values.update({
            'project_name': values.get('project_name') or row.get('project_name'),
            'block_name': values.get('block_name') or row.get('block_name'),
            'unit_number': values.get('unit_number') or row.get('unit_number'),
            'floor': values.get('floor') or row.get('floor'),
            'map_area': values.get('map_area') or row.get('area'),
            'contract_area': values.get('contract_area') or row.get('area'),
            'contract_number': values.get('contract_number') or row.get('contract_number'),
            'contract_date': values.get('contract_date') or row.get('contract_date_shamsi'),
            'buyer_name': values.get('buyer_name') or row.get('buyer_name'),
            'buyer_phone': values.get('buyer_phone') or row.get('buyer_phone'),
            'total_price': values.get('total_price') or row.get('contract_amount'),
        })
        return values

    def _persianize_variable(self, variable):
        if self._digit_normalizing:
            return
        value = variable.get(); display = persian_digits(value)
        if value != display:
            self._digit_normalizing = True
            try: variable.set(display)
            finally: self._digit_normalizing = False

    def _format_record_variable(self, key, variable):
        self._persianize_variable(variable)
        if key not in {'total_price', 'broker_fee', 'price_per_m2'} and not key.startswith('pay'):
            return
        value = formatted_amount_input(variable.get())
        if value != variable.get() and not self._digit_normalizing:
            self._digit_normalizing = True
            try: variable.set(value)
            finally: self._digit_normalizing = False

    def _build(self):
        bg, surface, panel, line, text = '#15202B', '#1E2C3A', '#20303F', '#324657', '#DDE7F0'
        self.configure(bg=bg)
        header = tk.Frame(self, bg=surface, highlightbackground=line, highlightthickness=1)
        header.pack(fill='x', padx=10, pady=8)
        tk.Label(header, text='پنجره رکورد فروش', bg=surface, fg=text,
                 font=('Segoe UI', 13, 'bold'), anchor='e').pack(side='right', padx=14, pady=9)
        tk.Label(header, text='اطلاعات اختصاصی قرارداد و فروش', bg=surface, fg='#9FB3C8',
                 anchor='w').pack(side='left', padx=14, pady=9)
        body = tk.PanedWindow(self, orient='horizontal', bg=line, sashwidth=6, relief='flat', bd=0)
        body.pack(fill='both', expand=True, padx=10, pady=(0, 8))
        preview = tk.Frame(body, bg=panel, highlightbackground=line, highlightthickness=1)
        form_area = tk.Frame(body, bg=panel, highlightbackground=line, highlightthickness=1)
        body.add(preview, minsize=260, stretch='always'); body.add(form_area, minsize=500, stretch='always')
        tk.Label(preview, text='پیش‌نمایش قرارداد', bg=panel, fg=text,
                 font=('Segoe UI', 11, 'bold'), anchor='e').pack(fill='x', padx=10, pady=8)
        preview_tools = tk.Frame(preview, bg=panel); preview_tools.pack(fill='x', padx=8)
        tk.Button(preview_tools, text='+', command=lambda: self._zoom_preview(1), bg=surface, fg=text, relief='flat', width=3).pack(side='left', padx=2)
        tk.Button(preview_tools, text='−', command=lambda: self._zoom_preview(-1), bg=surface, fg=text, relief='flat', width=3).pack(side='left', padx=2)
        self.preview_zoom = 1
        preview_area = tk.Frame(preview, bg='#0B1119'); preview_area.pack(fill='both', expand=True, padx=8, pady=8)
        self.preview_canvas = tk.Canvas(preview_area, bg='#0B1119', highlightthickness=0)
        preview_xscroll = ttk.Scrollbar(preview_area, orient='horizontal', command=self.preview_canvas.xview)
        preview_yscroll = ttk.Scrollbar(preview_area, orient='vertical', command=self.preview_canvas.yview)
        self.preview_canvas.configure(xscrollcommand=preview_xscroll.set, yscrollcommand=preview_yscroll.set)
        preview_yscroll.pack(side='right', fill='y'); preview_xscroll.pack(side='bottom', fill='x'); self.preview_canvas.pack(side='left', fill='both', expand=True)
        self.preview_canvas.bind('<Configure>', lambda _event: self._render_preview_image(), add='+')
        self.preview_canvas.bind('<MouseWheel>', self._preview_mouse_zoom, add='+')
        self._preview_source_image = None
        tk.Button(preview, text='نمایش فایل قرارداد', command=self._open_document, bg='#315D86', fg='white',
                  relief='flat', padx=8, pady=7).pack(fill='x', padx=8, pady=(0, 8))
        self._show_preview()
        tk.Label(form_area, text='مشخصات فروش', bg=panel, fg=text,
                 font=('Segoe UI', 11, 'bold'), anchor='e').pack(fill='x', padx=10, pady=8)
        canvas = tk.Canvas(form_area, highlightthickness=0)
        scroll = ttk.Scrollbar(form_area, orient='vertical', command=canvas.yview)
        canvas.configure(yscrollcommand=scroll.set)
        scroll.pack(side='right', fill='y'); canvas.pack(side='left', fill='both', expand=True)
        form = ttk.Frame(canvas); window = canvas.create_window((0, 0), window=form, anchor='nw')
        form.bind('<Configure>', lambda _e: canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.bind('<Configure>', lambda event: canvas.itemconfigure(window, width=event.width))
        form.grid_columnconfigure(0, weight=1)
        self._record_form = form
        upload = ttk.LabelFrame(form, text='۰. بارگذاری قرارداد', labelanchor='ne', padding=8)
        upload.grid(row=0, column=0, sticky='ew', padx=4, pady=4)
        ttk.Label(upload, text=self.document_path or 'فایلی انتخاب نشده است', anchor='e').pack(fill='x')
        tk.Button(upload, text='بارگذاری و تحلیل قرارداد', command=self._load_document, bg='#315D86', fg='white',
                  relief='flat', padx=14, pady=7).pack(anchor='e', pady=(7, 0))
        group_row = 1; signer_groups = []
        for title, keys in self._WIZARD_GROUPS:
            if title.startswith('۶.'):
                signer_groups.append((title, keys)); continue
            self._record_group(form, title, [(key, self._FIELD_LABELS[key]) for key in keys], group_row)
            group_row += 1
        self._payment_group_row = group_row
        self._render_payment_group()
        group_row += 1
        for title, keys in signer_groups:
            self._record_group(form, title, [(key, self._FIELD_LABELS[key]) for key in keys], group_row)
            group_row += 1
        final_group = self._record_group(form, '۷. تأیید نهایی', [], group_row)
        ttk.Label(final_group, text='اطلاعات این رکورد از ویزارد اختصاصی فروش قابل بازبینی و ثبت است.', anchor='e').pack(fill='x', padx=6, pady=4)
        actions = tk.Frame(self, bg=surface, highlightbackground=line, highlightthickness=1)
        actions.pack(fill='x', padx=10, pady=(0, 10))
        tk.Button(actions, text='ذخیرهٔ اطلاعات مجاز', command=self._save_metadata, bg='#2E6F4E', fg='white',
                  relief='flat', padx=16, pady=8).pack(side='right', padx=4, pady=5)
        tk.Button(actions, text='بستن', command=self.destroy, bg='#4A5568', fg='white',
                  relief='flat', padx=16, pady=8).pack(side='left', padx=4, pady=5)

    def _record_group(self, parent, title: str, fields: list[tuple[str, str]], row: int):
        group = ttk.LabelFrame(parent, text=title, labelanchor='ne', padding=8)
        group.grid(row=row, column=0, sticky='ew', padx=4, pady=4)
        for index, (key, label) in enumerate(fields):
            field_row, original_column = divmod(index, 2)
            field_column = 1 - original_column
            self._record_field(group, key, label, field_row, field_column)
        return group

    def _record_field(self, parent, key, label, row, column):
        box = ttk.Frame(parent); box.grid(row=row, column=column, sticky='ew', padx=6, pady=4)
        parent.grid_columnconfigure(column, weight=1)
        ttk.Label(box, text=label, anchor='e', justify='right').pack(fill='x')
        if key.endswith('_date'):
            line = ttk.Frame(box); line.pack(fill='x')
            ttk.Entry(line, textvariable=self.values[key], justify='right').pack(side='right', fill='x', expand=True)
            tk.Button(line, text='تقویم', command=lambda current=self.values[key]: ShamsiDatePicker(self, current),
                      bg='#324657', fg='#DDE7F0', relief='flat', padx=7).pack(side='left', padx=(5, 0))
        elif key == 'deal_mode':
            ttk.Combobox(box, textvariable=self.values[key], values=('واحد کامل', 'فروش متری'), state='readonly', justify='right').pack(fill='x')
        elif key == 'sale_kind':
            ttk.Combobox(box, textvariable=self.values[key], values=('پیش‌فروش', 'آماده'), state='readonly', justify='right').pack(fill='x')
        elif key == 'unit_number':
            self._unit_records = self._available_units()
            self._unit_combo = ttk.Combobox(box, textvariable=self.values[key], values=tuple(persian_digits(value) for value in self._unit_records), justify='right')
            self._unit_combo.bind('<<ComboboxSelected>>', self._unit_selected)
            self._unit_combo.pack(fill='x')
        elif key == 'floor':
            floors = tuple(dict.fromkeys(persian_digits(item.get('floor') or '') for item in self._unit_records.values() if item.get('floor') not in (None, '')))
            ttk.Combobox(box, textvariable=self.values[key], values=floors, justify='right').pack(fill='x')
        elif key == 'payment_count':
            widget = ttk.Entry(box, textvariable=self.values[key], justify='right')
            widget.bind('<FocusOut>', self._payment_count_changed)
            widget.bind('<Return>', self._payment_count_changed)
            widget.pack(fill='x')
        elif key == 'price_per_m2':
            ttk.Entry(box, textvariable=self.values[key], justify='right', state='readonly').pack(fill='x')
        else:
            ttk.Entry(box, textvariable=self.values[key], justify='right').pack(fill='x')

    def _payment_installment_count(self) -> int:
        try:
            # The entered stage count includes cash, delivery and deed.
            return max(1, min(10, int(en(self.values['payment_count'].get())) - 3))
        except ValueError:
            return 1

    def _render_payment_group(self):
        old_group = getattr(self, '_payment_group', None)
        if old_group is not None and old_group.winfo_exists():
            old_group.destroy()
        labels = payment_labels_for_installments(self._payment_installment_count())
        payment_fields = [('payment_count', 'تعداد اقساط')]
        self._payment_group = self._record_group(self._record_form, '۵. شرایط پرداخت', payment_fields, self._payment_group_row)
        for index, label in enumerate(labels, 1):
            self._record_field(self._payment_group, f'pay{index}_amount', f'{label} (تومان)', index, 1)
            self._record_field(self._payment_group, f'pay{index}_date', f'تاریخ {label}', index, 0)

    def _payment_count_changed(self, _event=None):
        old_count = getattr(self, '_payment_visible_installments', self._payment_installment_count())
        new_count = self._payment_installment_count()
        if old_count != new_count:
            old_delivery, old_deed = old_count + 2, old_count + 3
            new_delivery, new_deed = new_count + 2, new_count + 3
            for suffix in ('amount', 'date'):
                delivery_value = self.values[f'pay{old_delivery}_{suffix}'].get()
                deed_value = self.values[f'pay{old_deed}_{suffix}'].get()
                self.values[f'pay{old_delivery}_{suffix}'].set('')
                self.values[f'pay{old_deed}_{suffix}'].set('')
                self.values[f'pay{new_delivery}_{suffix}'].set(delivery_value)
                self.values[f'pay{new_deed}_{suffix}'].set(deed_value)
        self._payment_visible_installments = new_count
        self._render_payment_group()

    def _available_units(self) -> dict[str, dict[str, Any]]:
        project = self.values['project_name'].get().strip()
        block = self.values['block_name'].get().strip()
        try:
            with self.db.read() as con:
                rows = con.execute(
                    """SELECT unit_number, floor, area FROM units
                       WHERE status!='CANCELLED'
                         AND (?='' OR COALESCE(project_name,'')=?)
                         AND (?='' OR COALESCE(block_name,'')=?)
                       ORDER BY unit_number""",
                    (project, project, block, block),
                ).fetchall()
        except Exception:
            rows = []
        return {str(row['unit_number']): {'floor': row['floor'], 'area': row['area']} for row in rows if row['unit_number'] not in (None, '')}

    def _refresh_unit_choices(self):
        combo = self._unit_combo
        if combo is None:
            return
        self._unit_records = self._available_units()
        current = self.values['unit_number'].get().strip()
        choices = list(self._unit_records)
        if current and current not in choices:
            choices.insert(0, current)
        combo.configure(values=tuple(persian_digits(value) for value in choices))

    def _unit_selected(self, _event=None):
        selected = self.values['unit_number'].get().strip()
        record = self._unit_records.get(selected) or self._unit_records.get(en(selected))
        if not record:
            return
        self.values['floor'].set(persian_digits(record.get('floor') or ''))
        self.values['map_area'].set(persian_digits(record.get('area') or ''))

    def _record_price_per_m2(self):
        try:
            area = float(en(self.values['contract_area'].get()).replace(',', '') or 0)
        except ValueError:
            area = 0
        price = parse_money(self.values['total_price'].get())
        self.values['price_per_m2'].set(toman(price / area) if area > 0 else '')

    def _show_preview(self):
        path = Path(self.document_path) if self.document_path else None
        label = str(path) if path else 'فایل قرارداد انتخاب نشده است'
        self.preview_canvas.delete('all'); self._preview_source_image = None; self._preview_pil_image = None
        if path and path.suffix.lower() in {'.png', '.gif', '.jpg', '.jpeg', '.webp', '.bmp'} and path.exists():
            try:
                from PIL import Image
                self._preview_pil_image = Image.open(path).convert('RGB')
                self._render_preview_image()
            except (OSError, ValueError, tk.TclError):
                self.preview_canvas.create_text(10, 10, text=label, fill='#DDE7F0', anchor='nw')
        else:
            self.preview_canvas.create_text(10, 10, text=label, fill='#DDE7F0', anchor='nw')

    def _render_preview_image(self):
        self.preview_canvas.delete('all')
        if self._preview_pil_image is not None:
            from PIL import ImageTk
            source = self._preview_pil_image
            canvas_width = max(1, self.preview_canvas.winfo_width() - 8)
            canvas_height = max(1, self.preview_canvas.winfo_height() - 8)
            fit_scale = min(canvas_width / max(1, source.width), canvas_height / max(1, source.height))
            scale = max(0.05, min(1.0, fit_scale)) * self.preview_zoom
            size = (max(1, int(source.width * scale)), max(1, int(source.height * scale)))
            image = ImageTk.PhotoImage(source.resize(size))
            self._preview_image = image
            self.preview_canvas.create_image(0, 0, image=image, anchor='nw')
            self.preview_canvas.configure(scrollregion=(0, 0, image.width(), image.height()))
            return
        if self._preview_source_image is None:
            return
        image = self._preview_source_image
        if self.preview_zoom > 1:
            image = image.zoom(self.preview_zoom, self.preview_zoom)
        elif self.preview_zoom < 1:
            image = image.subsample(2, 2)
        self._preview_image = image
        self.preview_canvas.create_image(0, 0, image=image, anchor='nw')
        self.preview_canvas.configure(scrollregion=(0, 0, image.width(), image.height()))

    def _zoom_preview(self, delta):
        self.preview_zoom = max(0.2, min(4, self.preview_zoom + delta))
        self._render_preview_image()

    def _preview_mouse_zoom(self, event):
        self._zoom_preview(0.2 if event.delta > 0 else -0.2)
        return 'break'

    def _open_document(self):
        path = Path(self.document_path) if self.document_path else None
        if not path or not path.exists():
            return messagebox.showinfo('قرارداد', 'فایلی برای باز کردن انتخاب نشده است.', parent=self)
        try:
            import os
            os.startfile(str(path))  # type: ignore[attr-defined]
        except Exception as exc:
            messagebox.showerror('قرارداد', str(exc), parent=self)

    def _load_document(self):
        path = filedialog.askopenfilename(parent=self, title='قرارداد', filetypes=[('Documents', '*.pdf *.jpg *.jpeg *.png *.doc *.docx *.txt'), ('All files', '*.*')])
        if not path:
            return
        self.document_path = path; text = ''
        try:
            text = Path(path).read_text(encoding='utf-8') if Path(path).suffix.lower() == '.txt' else ''
        except (OSError, UnicodeDecodeError):
            text = ''
        parsed = parse_sale_form(text)
        for target, source in {'buyer_name': 'buyer_name', 'unit_number': 'unit_number', 'floor': 'floor', 'contract_area': 'area', 'contract_number': 'contract_number', 'contract_date': 'contract_date_shamsi', 'total_price': 'contract_amount'}.items():
            if parsed.get(source) not in (None, '', 0): self.values[target].set(str(parsed[source]))
        self._show_preview()

    def _save_metadata(self):
        if not self.sale_id:
            return messagebox.showinfo('رکورد فروش', 'این پیش‌نویس هنوز قرارداد ثبت‌شده ندارد؛ ثبت نهایی فقط از ویزارد قرارداد انجام می‌شود.', parent=self)
        raw = {key: value.get().strip() for key, value in self.values.items()}
        try:
            self.revenue.update_sale_metadata(
                self.sale_id, buyer_name=raw['buyer_name'], buyer_phone=raw['buyer_phone'],
                contract_number=raw['contract_number'], contract_date_shamsi=raw['contract_date'],
                notes=contract_note(raw), source='GUI_1871_SALES_RECORD', actor_id='GUI',
            )
        except Exception as exc:
            return messagebox.showerror('رکورد فروش', str(exc), parent=self)
        if self.on_saved: self.on_saved()
        messagebox.showinfo('رکورد فروش', 'اطلاعات غیرمالی ذخیره شد؛ مبالغ و وصول‌ها فقط از مسیر حسابرسی اصلاح می‌شوند.', parent=self)
        self.destroy()


def contract_export_preview_lines(specs, payload, headers, document_paths=None) -> list[tuple[str, str]]:
    """Build a visible, resilient preview even when optional export fields are incomplete."""
    lines: list[tuple[str, str]] = [('section', 'مشخصات قرارداد')]
    current_group = None
    for spec in specs or ():
        if isinstance(spec, (tuple, list)) and len(spec) >= 3:
            group, label, value = spec[:3]
        else:
            group, label, value = 'مشخصات قرارداد', 'اطلاعات', spec
        group, label, value = str(group or 'مشخصات قرارداد'), str(label or 'اطلاعات'), str(value or '-')
        if group != current_group:
            lines.append(('section', f'[{group}]'))
            current_group = group
        lines.append(('rtl', f'{label}: {value}'))
    lines.append(('section', 'جدول پرداخت‌ها'))
    lines.append(('table_header', ' | '.join(str(header) for header in headers or ())))
    for row in payload or ():
        values = row if isinstance(row, dict) else dict(zip(headers or (), row))
        lines.append(('rtl', ' | '.join(str(values.get(header, '')) for header in headers or ())))
    if not specs and not payload:
        lines.append(('rtl', 'داده‌ای برای پیش‌نمایش انتخاب نشده است.'))
    paths = [str(path) for path in (document_paths or ()) if str(path).strip()]
    if paths:
        lines.append(('section', 'فیش‌های واریزی انتخاب‌شده:'))
        lines.extend(('rtl', Path(path).name) for path in paths)
    return lines


class ContractExportPreview(tk.Toplevel):
    """Preview contract specs/payments before opening the native Save As dialog."""
    def __init__(self, parent, *, export_format, headers, payload, specs, on_confirm, document_paths=None,
                 contract_document_paths=None,
                 window_title='پیش‌نمایش خروجی قرارداد'):
        super().__init__(parent)
        self.title(window_title); self.geometry('980x650'); self.minsize(700, 440); self.resizable(True, True)
        self.headers, self.payload, self.specs, self.on_confirm = headers, payload, specs, on_confirm
        self.document_paths = [str(path) for path in (document_paths or []) if str(path).strip()]
        self.contract_document_paths = [str(path) for path in (contract_document_paths or []) if str(path).strip()]
        self._preview_images = []
        bar = ttk.Frame(self); bar.pack(fill='x', padx=8, pady=8)
        ttk.Label(bar, text='فرمت خروجی').pack(side='right', padx=6)
        self.format_choice = tk.StringVar(value=export_format)
        self.format_combo = ttk.Combobox(bar, textvariable=self.format_choice, values=('Excel', 'PDF', 'پرینت'), state='readonly', justify='right', width=18)
        self.format_combo.pack(side='right'); self.format_combo.bind('<<ComboboxSelected>>', lambda _e: self._render())
        ttk.Label(bar, text='نحوه نمایش').pack(side='right', padx=(18, 6))
        self.orientation_choice = tk.StringVar(value='افقی')
        self.orientation_combo = ttk.Combobox(bar, textvariable=self.orientation_choice, values=('افقی', 'عمودی'), state='readonly', justify='right', width=12)
        self.orientation_combo.pack(side='right'); self.orientation_combo.bind('<<ComboboxSelected>>', lambda _e: self._render())
        ttk.Button(bar, text='تأیید و ذخیره', command=self._confirm).pack(side='left', padx=6)
        preview_area = ttk.Frame(self); preview_area.pack(fill='both', expand=True, padx=8, pady=(0, 8))
        # Canvas is already the stable visual primitive used by the active sales
        # windows. It avoids the blank native Text control seen in this dialog.
        self.preview = tk.Canvas(preview_area, bg='#0A1A2A', highlightthickness=0)
        preview_scroll = ttk.Scrollbar(preview_area, orient='vertical', command=self.preview.yview)
        preview_xscroll = ttk.Scrollbar(preview_area, orient='horizontal', command=self.preview.xview)
        self.preview.configure(yscrollcommand=preview_scroll.set, xscrollcommand=preview_xscroll.set)
        preview_scroll.pack(side='right', fill='y'); preview_xscroll.pack(side='bottom', fill='x'); self.preview.pack(side='left', fill='both', expand=True)
        self.preview.bind('<Configure>', lambda _event: self._render(), add='+')
        self.preview.bind('<MouseWheel>', self._preview_mouse_zoom, add='+')
        self.preview.bind('<ButtonPress-1>', lambda event: self.preview.scan_mark(event.x, event.y), add='+')
        self.preview.bind('<B1-Motion>', lambda event: self.preview.scan_dragto(event.x, event.y, gain=1), add='+')
        self.preview_zoom = 1.0
        self._render()

    @staticmethod
    def _group_specs(specs):
        groups: dict[str, list[tuple[str, str]]] = {}
        for spec in specs or ():
            if isinstance(spec, (tuple, list)) and len(spec) >= 3:
                group, label, value = spec[:3]
            else:
                group, label, value = 'مشخصات', 'اطلاعات', spec
            groups.setdefault(str(group or 'مشخصات'), []).append((str(label or 'اطلاعات'), str(value or '-')))
        return groups

    def _render(self):
        self._preview_images = []
        self.preview.delete('all')
        scale = self.preview_zoom
        portrait = self.orientation_choice.get() == 'عمودی'
        page_width = 780 if portrait else 1120
        page_height = 1100 if portrait else 780
        margin = 24; x = margin; y = margin; page_start = 0
        def rect(left, top, right, bottom, **kwargs):
            return self.preview.create_rectangle(left * scale, top * scale, right * scale, bottom * scale, **kwargs)
        def text(right, top, value, *, width, size=9, bold=False, fill='#17212B'):
            return self.preview.create_text(right * scale, top * scale, text=value, anchor='ne', justify='right',
                                            width=width * scale, fill=fill, font=('Segoe UI', max(7, int(size * scale)), 'bold' if bold else 'normal'))
        rect(0, 0, page_width, page_height, fill='#FFFFFF', outline='#FFFFFF')
        text(page_width - margin, y, self.title(), width=page_width - 2 * margin, size=15, bold=True, fill='#102A43'); y += 32
        groups = self._group_specs(self.specs)
        if groups:
            group_width = (page_width - 3 * margin) / 2
            group_items = list(groups.items())
            for index in range(0, len(group_items), 2):
                # The first RTL group is on the physical right; the second is
                # on the left, matching the PDF box grid exactly.
                row_groups = list(reversed(group_items[index:index + 2]))
                heights = [32 + 25 * max(1, (len(entries) + 1) // 2) for _group, entries in row_groups]
                row_height = max(heights)
                for column, (group, entries) in enumerate(row_groups):
                    left = margin + column * (group_width + margin); right = left + group_width
                    height = 32 + 25 * max(1, (len(entries) + 1) // 2)
                    rect(left, y, right, y + height, outline='#496176', width=1)
                    rect(left, y, right, y + 24, fill='#DCE8F2', outline='#496176')
                    text(right - 7, y + 5, group, width=group_width - 14, size=10, bold=True, fill='#102A43')
                    for entry_index, (label, value) in enumerate(entries):
                        row, column_in_row = divmod(entry_index, 2)
                        cell_width = (group_width - 16) / 2
                        cell_right = right - 8 - column_in_row * cell_width
                        top = y + 30 + row * 25
                        text(cell_right, top, f'{label}: {value}', width=cell_width - 6, size=8)
                y += row_height + 10
        if not self.specs:
            text(page_width - margin, y, 'داده‌ای برای پیش‌نمایش انتخاب نشده است.', width=page_width - 2 * margin, size=10); y += 30
        table_width = page_width - 2 * margin
        text(page_width - margin, y, 'جدول پرداخت‌ها', width=table_width, size=11, bold=True, fill='#102A43'); y += 25
        document_headers = [header for header in self.headers if header in {'نمایش سند', 'فیش واریزی'}]
        document_width = 220 if document_headers else 0
        other_width = (table_width - document_width) / max(1, len(self.headers) - len(document_headers))
        columns = []
        right = page_width - margin
        for header in self.headers:
            width = document_width if header in document_headers else other_width
            columns.append((header, right - width, right)); right -= width
        def draw_payment_header(top):
            bottom = top + 27
            for header, left, right in columns:
                rect(left, top, right, bottom, fill='#153B5B', outline='#B8C2CC')
                text(right - 4, top + 6, header, width=right - left - 8, size=8, bold=True, fill='#FFFFFF')
            return bottom

        y = draw_payment_header(y)
        for row in self.payload:
            values = row if isinstance(row, dict) else dict(zip(self.headers, row))
            row_paths = []
            for header in document_headers:
                row_paths.extend(export_document_paths(values.get(header, '')))
            # Stack every receipt vertically in its own table row.  This keeps
            # second-page receipts visible instead of drawing them past a cell.
            receipt_preview_height = 118
            row_height = max(32, 12 + receipt_preview_height * len(row_paths))
            if y + row_height > page_start + page_height - margin:
                page_start += page_height + 18
                rect(0, page_start, page_width, page_start + page_height, fill='#FFFFFF', outline='#FFFFFF')
                text(page_width - margin, page_start + margin, 'ادامه جدول پرداخت‌ها', width=table_width, size=11, bold=True, fill='#102A43')
                y = draw_payment_header(page_start + margin + 25)
            for header, left, right in columns:
                rect(left, y, right, y + row_height, outline='#B8C2CC')
                if header not in document_headers:
                    text(right - 4, y + 7, str(values.get(header, '-') or '-'), width=right - left - 8, size=8)
            if document_headers:
                doc_left, doc_right = next((left, right) for header, left, right in columns if header in document_headers)
                image_x, image_y = doc_right - 6, y + 6
                for path in row_paths:
                    try:
                        from PIL import Image, ImageTk
                        image = Image.open(path).convert('RGB'); image.thumbnail((int((doc_right - doc_left - 12) * scale), int(receipt_preview_height * scale)))
                        photo = ImageTk.PhotoImage(image); self._preview_images.append(photo)
                        self.preview.create_image(image_x * scale, image_y * scale, image=photo, anchor='ne')
                        image_y += image.height() / scale + 6
                    except (OSError, ValueError, tk.TclError):
                        continue
                if not row_paths:
                    text(doc_right - 4, y + 7, 'سند ثبت نشده است', width=doc_right - doc_left - 8, size=8)
            y += row_height
        if self.contract_document_paths:
            y += 18
            text(page_width - margin, y, 'تصویر قرارداد', width=table_width, size=11, bold=True, fill='#102A43'); y += 28
            for path in self.contract_document_paths:
                try:
                    from PIL import Image, ImageTk
                    image = Image.open(path).convert('RGB'); image.thumbnail((int(360 * scale), int(510 * scale)))
                    photo = ImageTk.PhotoImage(image); self._preview_images.append(photo)
                    self.preview.create_image((page_width - margin) * scale, y * scale, image=photo, anchor='ne')
                    y += image.height() / scale + 18
                except (OSError, ValueError, tk.TclError):
                    text(page_width - margin, y, Path(path).name, width=table_width, size=8); y += 25
        self.preview.configure(scrollregion=(0, 0, page_width * scale + 20, y * scale + 30))
        self.preview.yview_moveto(0)

    def _preview_mouse_zoom(self, event):
        self.preview_zoom = max(.5, min(2.5, self.preview_zoom + (.1 if event.delta > 0 else -.1)))
        self._render()
        return 'break'

    def _confirm(self):
        # Keep the populated preview open if Save As is cancelled or fails.
        saved = self.on_confirm(self.format_choice.get(), self.orientation_choice.get())
        if saved:
            self.destroy()


class ContractDetails(tk.Toplevel):
    """Read-only contract and payment detail window over the existing revenue ledger."""
    _WIZARD_BOXES = (
        ('۱. مورد معامله', ('project_name', 'block_name', 'deal_mode', 'sale_kind', 'unit_number', 'floor', 'map_area', 'contract_area')),
        ('۲. تاریخ‌ها', ('contract_number', 'contract_date', 'delivery_date', 'deed_date')),
        ('۳. خریدار', ('buyer_name', 'father_name', 'national_id', 'buyer_phone')),
        ('۴. مالی', ('total_price', 'price_per_m2', 'broker_fee')),
        ('۶. امضاکنندگان', ('notary_name', 'witness_1', 'witness_2', 'buyer_representative', 'seller_representative')),
    )
    _FIELD_LABELS = dict(SalesRecord._FIELDS)
    def __init__(self, parent, *, db, base_dir: Path, sale_id: str):
        super().__init__(parent)
        self.db, self.base_dir, self.sale_id = db, Path(base_dir), str(sale_id)
        self.revenue = RevenueService(db); self.attachments = AttachmentService(db, self.base_dir)
        self.include_documents = tk.BooleanVar(value=True); self._documents: dict[str, list[dict[str, Any]]] = {}
        # The contract grid is stage-oriented, whereas receipt exports must
        # preserve the finance ledger's one-row-per-received-payment evidence.
        self._receipt_rows_by_installment: dict[str, list[dict[str, Any]]] = {}
        # Keep a normal overlapped Toplevel so Windows supplies native
        # minimize/maximize/restore/close controls in its title bar.
        self.title('جزئیات قرارداد فروش'); self.geometry('1280x720'); self.minsize(950, 560); self.resizable(True, True)
        self._build(); self._refresh()

    def _build(self):
        tools = ttk.Frame(self); tools.pack(fill='x', padx=8, pady=8)
        self.export_choice = tk.StringVar(value=EXPORT_PLACEHOLDER)
        ttk.Label(tools, text='خروجی').pack(side='right', padx=(8, 2))
        export_combo = ttk.Combobox(tools, textvariable=self.export_choice, values=(EXPORT_PLACEHOLDER, 'Excel', 'PDF', 'پرینت'), state='readonly', justify='right', width=18)
        export_combo.pack(side='right', padx=3)
        export_combo.bind('<<ComboboxSelected>>', self._export_selected)
        ttk.Checkbutton(tools, text='نمایش سندهای واریزی در خروجی', variable=self.include_documents).pack(side='right', padx=8)
        self.summary = tk.StringVar(value='')
        ttk.Label(self, textvariable=self.summary, anchor='e').pack(fill='x', padx=10, pady=(0, 8))
        cards_area = ttk.Frame(self); cards_area.pack(fill='both', expand=True, padx=10, pady=(0, 8))
        self.cards_canvas = tk.Canvas(cards_area, height=290, highlightthickness=0)
        cards_scroll = ttk.Scrollbar(cards_area, orient='vertical', command=self.cards_canvas.yview)
        self.cards_canvas.configure(yscrollcommand=cards_scroll.set)
        cards_scroll.pack(side='right', fill='y'); self.cards_canvas.pack(side='left', fill='both', expand=True)
        self.wizard_cards = ttk.Frame(self.cards_canvas)
        cards_window = self.cards_canvas.create_window((0, 0), window=self.wizard_cards, anchor='nw')
        self.wizard_cards.bind('<Configure>', lambda _e: self.cards_canvas.configure(scrollregion=self.cards_canvas.bbox('all')))
        self.cards_canvas.bind('<Configure>', lambda event: self.cards_canvas.itemconfigure(cards_window, width=event.width))
        table_area = ttk.Frame(self); table_area.pack(fill='both', expand=True, padx=10, pady=(0, 10))
        self.tree = ttk.Treeview(table_area, columns=tuple(key for key, _label, _width in PAYMENT_DETAIL_COLUMNS), show='headings')
        for key, label, width in PAYMENT_DETAIL_COLUMNS:
            self.tree.heading(key, text=label, anchor='e'); self.tree.column(key, width=width, anchor='e', stretch=True)
        self.tree.configure(displaycolumns=tuple(key for key, _label, _width in reversed(PAYMENT_DETAIL_COLUMNS)))
        table_xscroll = ttk.Scrollbar(table_area, orient='horizontal', command=self.tree.xview)
        table_yscroll = ttk.Scrollbar(table_area, orient='vertical', command=self.tree.yview)
        self.tree.configure(xscrollcommand=table_xscroll.set, yscrollcommand=table_yscroll.set)
        table_yscroll.pack(side='right', fill='y'); table_xscroll.pack(side='bottom', fill='x'); self.tree.pack(side='left', fill='both', expand=True)
        self.tree.bind('<Double-1>', lambda _event: self._open_document())

    def _toggle_maximize(self):
        try:
            self.state('normal' if self.state() == 'zoomed' else 'zoomed')
        except tk.TclError:
            self.attributes('-zoomed', self.state() != 'zoomed')

    def _refresh(self):
        sale = self.revenue.get_sale(self.sale_id) or {}; summary = self.revenue.sale_summary(self.sale_id)
        self.summary.set(f"مبلغ کل: {toman(sale.get('contract_amount'))} تومان | کل پرداختی: {toman(summary.cash_received)} تومان | معوق: {toman(summary.overdue_amount)} تومان | مانده: {toman(summary.receivable)} تومان")
        self._set_wizard_fields(sale)
        for iid in self.tree.get_children(): self.tree.delete(iid)
        self._documents = {}; self._receipt_rows_by_installment = {}
        with self.db.read() as con:
            rows = con.execute("""SELECT i.*, COALESCE(SUM(a.amount),0) paid, MAX(r.received_date_shamsi) received_date,
                                      MAX(r.transaction_id) transaction_id
                               FROM installments i
                               LEFT JOIN payment_allocations a ON a.installment_id=i.installment_id
                               LEFT JOIN received_payments r ON r.received_payment_id=a.received_payment_id AND r.status!='CANCELLED'
                               WHERE i.sale_id=? AND i.status!='CANCELLED'
                               GROUP BY i.installment_id ORDER BY COALESCE(i.sequence_no,999999),i.due_date""", (self.sale_id,)).fetchall()
            payment_links = con.execute("""SELECT a.installment_id, r.transaction_id, r.received_payment_id,
                                                  r.amount, r.received_date_shamsi
                                           FROM payment_allocations a
                                           JOIN received_payments r ON r.received_payment_id=a.received_payment_id
                                           WHERE a.installment_id IN (SELECT installment_id FROM installments WHERE sale_id=?)""", (self.sale_id,)).fetchall()
        links_by_installment: dict[str, list[dict[str, Any]]] = {}
        for link in payment_links:
            links_by_installment.setdefault(str(link['installment_id']), []).append(dict(link))
        today = date.today()
        for number, source in enumerate(rows, 1):
            item = dict(source); amount, paid = int(item.get('amount') or 0), int(item.get('paid') or 0)
            try: late = max(0, (today - date.fromisoformat(str(item.get('due_date')))).days) if paid < amount else 0
            except ValueError: late = 0
            documents, seen_documents, receipt_rows = [], set(), []
            for link in links_by_installment.get(str(item['installment_id']), []):
                entities = []
                if link.get('transaction_id'):
                    entities.append(('TRANSACTION', str(link['transaction_id'])))
                if link.get('received_payment_id'):
                    entities.extend([('RECEIVED_PAYMENT', str(link['received_payment_id'])), ('SALE_PAYMENT', str(link['received_payment_id']))])
                receipt_documents, seen_receipt_documents = [], set()
                for document in self.attachments.list_for_entities(entities):
                    marker = str(document.get('attachment_id') or document.get('local_path') or '')
                    if marker and marker not in seen_receipt_documents:
                        receipt_documents.append(document); seen_receipt_documents.add(marker)
                    if marker and marker not in seen_documents:
                        documents.append(document); seen_documents.add(marker)
                receipt_rows.append({
                    'received_payment_id': str(link.get('received_payment_id') or ''),
                    'amount': int(link.get('amount') or 0),
                    'received_date_shamsi': str(link.get('received_date_shamsi') or '-'),
                    'documents': receipt_documents,
                })
            self._documents[str(item['installment_id'])] = documents
            self._receipt_rows_by_installment[str(item['installment_id'])] = receipt_rows
            debt = amount - paid; balance = f"{'بدهکار' if debt >= 0 else 'بستانکار'}: {toman(abs(debt))}"
            self.tree.insert('', 'end', iid=str(item['installment_id']), values=(
                selectable_row(number), item.get('description') or f'قسط {number}', persian_digits(item.get('due_date_shamsi') or '-'),
                toman(amount), persian_digits(item.get('due_date_shamsi') or '-'), persian_digits(item.get('received_date') or '-'), toman(paid), fa(late), balance,
                'نمایش سند' if documents else '-',
            ))

    def _set_wizard_fields(self, sale: dict[str, Any]):
        try: metadata = json.loads(str(sale.get('notes') or '{}'))
        except (ValueError, TypeError, json.JSONDecodeError): metadata = {}
        fallback = {
            'project_name': sale.get('project_name'), 'block_name': sale.get('block_name'),
            'unit_number': sale.get('unit_number'), 'floor': sale.get('floor'), 'map_area': sale.get('area'),
            'contract_area': sale.get('area'), 'contract_number': sale.get('contract_number'),
            'contract_date': sale.get('contract_date_shamsi'), 'buyer_name': sale.get('buyer_name'),
            'buyer_phone': sale.get('buyer_phone'), 'total_price': toman(sale.get('contract_amount')),
        }
        values = {**fallback, **metadata}
        documents = self.attachments.list_for_entities([('SALE', self.sale_id)])
        self._contract_document_paths = [str(item.get('local_path') or '') for item in documents if str(item.get('local_path') or '').strip()]
        export_fields = []
        signer_boxes = []
        for title, keys in self._WIZARD_BOXES:
            if title.startswith('۶.'):
                signer_boxes.append((title, keys)); continue
            fields = [(self._FIELD_LABELS[key], str(values.get(key) or '-')) for key in keys]
            export_fields.extend((title, label, value) for (label, value) in fields)
            self._build_wizard_box(title, fields)
        # The ledger is authoritative for an existing contract.  Do not render
        # a fixed legacy ten-installment list when this sale actually has fewer
        # stages (for example cash + four agreed installments).
        recorded_installments = self.revenue.list_installments(self.sale_id)
        if recorded_installments:
            payments = [
                (
                    str(item.get('description') or f"مرحله {item.get('sequence_no') or '-'}"),
                    f"{toman(item.get('amount') or 0)} تومان | تاریخ: {item.get('due_date_shamsi') or '-'}",
                )
                for item in recorded_installments
            ]
        else:
            try:
                installment_count = max(1, min(10, int(en(str(values.get('payment_count') or '1')))))
            except ValueError:
                installment_count = 1
            labels = payment_labels_for_installments(installment_count)
            payments = [
                (label, f"{values.get(f'pay{index}_amount') or '-'} تومان | تاریخ: {values.get(f'pay{index}_date') or '-'}")
                for index, label in enumerate(labels, 1)
            ]
        export_fields.extend(('۵. شرایط پرداخت', label, value) for label, value in payments)
        self._build_wizard_box('۵. شرایط پرداخت', payments)
        for title, keys in signer_boxes:
            fields = [(self._FIELD_LABELS[key], str(values.get(key) or '-')) for key in keys]
            export_fields.extend((title, label, value) for (label, value) in fields)
            self._build_wizard_box(title, fields)
        self._wizard_export_fields = export_fields

    def _build_wizard_box(self, title: str, fields: list[tuple[str, str]]):
        box = ttk.LabelFrame(self.wizard_cards, text=title, labelanchor='ne', padding=8)
        box.pack(fill='x', expand=True, pady=4)
        for index, (label, value) in enumerate(fields):
            row, original_column = divmod(index, 2); column = 1 - original_column; box.grid_columnconfigure(column, weight=1)
            item = ttk.Frame(box); item.grid(row=row, column=column, sticky='ew', padx=5, pady=3)
            ttk.Label(item, text=label, anchor='e').pack(fill='x')
            ttk.Label(item, text=persian_digits(value), anchor='e', justify='right').pack(fill='x')

    def _open_document(self):
        selected = self.tree.selection()
        if not selected: return
        documents = self._documents.get(selected[0], [])
        if not documents:
            return messagebox.showinfo('سند پرداخت', 'برای این ردیف سندی ثبت نشده است.', parent=self)
        values = self.tree.item(selected[0], 'values')
        try:
            PaymentDocuments(
                self, db=self.db, base_dir=self.base_dir,
                payments=[{
                    'stage': str(values[1] or 'مرحله پرداخت'),
                    'date': str(values[5] or '-'),
                    'amount': parse_money(str(values[6] or '0')),
                    'documents': documents,
                }],
            )
        except Exception as exc:
            messagebox.showerror('سند پرداخت', str(exc), parent=self)

    def _export_selected(self, _event=None):
        export_format = {'Excel': 'xlsx', 'PDF': 'pdf', 'پرینت': 'print'}.get(self.export_choice.get())
        if export_format:
            self._export(export_format)

    def _export(self, export_format: str):
        # Keep contract metadata from the stage grid, but build its receipt
        # rows from the same individual payment allocations used by Finance.
        selected = list(self.tree.get_children())
        headers = [str(self.tree.heading(key, 'text')) for key in self.tree['columns']]
        row_index = {header: index for index, header in enumerate(headers)}
        stage_index = row_index.get('مرحله'); date_index = row_index.get('تاریخ واریز')
        amount_index = row_index.get('مبلغ واریز (تومان)')
        export_rows, export_ids, documents_by_export_row = [], [], {}
        for installment_id in selected:
            source = list(self.tree.item(installment_id, 'values'))
            receipt_rows = list(self._receipt_rows_by_installment.get(str(installment_id), []))
            if not receipt_rows:
                export_id = f'{installment_id}:stage'
                export_rows.append(tuple(source)); export_ids.append(export_id)
                documents_by_export_row[export_id] = list(self._documents.get(str(installment_id), []))
                continue
            multiple_receipts = len(receipt_rows) > 1
            for number, receipt in enumerate(receipt_rows, 1):
                row = list(source)
                if multiple_receipts and stage_index is not None:
                    row[stage_index] = f"{row[stage_index]}-{persian_digits(number)}"
                if date_index is not None:
                    row[date_index] = persian_digits(receipt.get('received_date_shamsi') or '-')
                if amount_index is not None:
                    row[amount_index] = toman(receipt.get('amount') or 0)
                export_id = f"{installment_id}:{receipt.get('received_payment_id') or number}"
                export_rows.append(tuple(row)); export_ids.append(export_id)
                documents_by_export_row[export_id] = list(receipt.get('documents') or [])
        payload = payment_detail_export_payload(headers, export_rows, export_ids,
                                                     documents_by_export_row, self.include_documents.get())
        specs = list(getattr(self, '_wizard_export_fields', []))
        choices = {'xlsx': 'Excel', 'pdf': 'PDF', 'print': 'پرینت'}
        document_paths = export_payload_document_paths(payload, 'نمایش سند')
        contract_document_paths = list(getattr(self, '_contract_document_paths', []))
        return ContractExportPreview(self, export_format=choices.get(export_format, 'PDF'), headers=headers,
                                        payload=payload, specs=specs,
                                        document_paths=document_paths,
                                        contract_document_paths=contract_document_paths,
                                        on_confirm=lambda choice, orientation: self._save_export(choice, headers, payload, specs, document_paths, orientation, contract_document_paths))

    def _save_export(self, choice: str, headers, payload, specs, document_paths=None, orientation='افقی', contract_document_paths=None):
        export_format = {'Excel': 'xlsx', 'PDF': 'pdf', 'پرینت': 'print'}.get(choice, 'pdf')
        actual_format = 'pdf' if export_format == 'print' else export_format
        extension = f'.{actual_format}'
        path = filedialog.asksaveasfilename(parent=self, title=f'خروجی {actual_format.upper()} جزئیات قرارداد', defaultextension=extension, filetypes=[('Excel', '*.xlsx')] if actual_format == 'xlsx' else [('PDF', '*.pdf')])
        if not path:
            return False
        try:
            destination = Path(path)
            destination.parent.mkdir(parents=True, exist_ok=True)
            if actual_format == 'xlsx':
                from openpyxl import Workbook
                book = Workbook(); sheet = book.active; sheet.title = 'مشخصات قرارداد'
                sheet.append(('کادر ویزارد', 'عنوان', 'مقدار'))
                for group, label, value in getattr(self, '_wizard_export_fields', []):
                    sheet.append((group, label, value))
                sheet.append(())
                sheet.append(('جدول پرداخت‌ها',))
                sheet.append(headers)
                for row in payload:
                    sheet.append([row.get(header, '') for header in headers])
                    document_path = str(row.get('نمایش سند') or '')
                    if document_path and Path(document_path).exists() and Path(document_path).suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.bmp'}:
                        try:
                            from openpyxl.drawing.image import Image as ExcelImage
                            image = ExcelImage(document_path); image.width, image.height = 160, 110
                            image.anchor = f'{chr(65 + max(0, len(headers) - 1))}{sheet.max_row}'
                            sheet.add_image(image)
                        except (OSError, ImportError, ValueError):
                            pass
                book.save(destination)
            else:
                summary = [f'{group} — {label}: {value}' for group, label, value in specs]
                export_receipts_pdf(destination, headers, [[row.get(header, '') for header in headers] for row in payload],
                                        'خروجی جزئیات قرارداد', summary, list(document_paths or []), orientation=orientation,
                                        contract_document_paths=list(contract_document_paths or []))
                if export_format == 'print' and hasattr(os, 'startfile'):
                    os.startfile(str(destination), 'print')
        except Exception as exc:
            messagebox.showerror('خروجی', str(exc), parent=self)
            return False
        messagebox.showinfo('خروجی', 'خروجی جزئیات قرارداد آماده شد.', parent=self)
        return True


class SalesClaimsPanel(ttk.Frame):
    """Right-to-left contracts and finance workspace backed by the existing ledger."""
    def __init__(self, parent, *, db, base_dir, on_changed=None):
        super().__init__(parent, padding=6)
        self.db, self.base_dir, self.on_changed = db, Path(base_dir), on_changed
        self.revenue = RevenueService(db)
        self.attachments = AttachmentService(db, self.base_dir)
        self.selected_sale_id = ''
        self.nb = ttk.Notebook(self); self.nb.pack(fill='both', expand=True)
        self.contracts = ttk.Frame(self.nb, padding=5); self.finance = ttk.Frame(self.nb, padding=5)
        self.nb.add(self.contracts, text='قراردادها'); self.nb.add(self.finance, text='مالی')
        style = ttk.Style(self)
        surface = style.lookup('TFrame', 'background') or '#F4F6F8'
        text = style.lookup('TLabel', 'foreground') or '#17212B'
        style.configure('SalesClaims.Treeview', background=surface, fieldbackground=surface,
                        foreground=text, borderwidth=0, relief='flat', rowheight=28)
        style.configure('SalesClaims.Treeview.Heading', background=surface, foreground=text,
                        borderwidth=0, relief='flat')
        style.configure('SalesClaims.Summary.TLabel', background=surface, foreground=text)
        tab_line = '#40566B' if surface.startswith('#') and int(surface[1:3], 16) < 100 else '#C7D2DE'
        self._rtl_tabs = SalesTabStrip(
            self.nb, bg=surface, line=tab_line, text=text, muted=text,
            active_bg='#315D86', active_fg='#FFFFFF',
        )
        self._build_contracts(); self._build_finance(); self.refresh()
        # The inherited global currency pass treats words such as «قرارداد» and
        # «فروش» as monetary.  Keep these three sales-table headings semantic.
        self.after(80, self._restore_contract_non_money_headings)
        self._price_refresh_job = self.after(60000, self._auto_refresh_prices)

    def _toolbar(self, parent, actions):
        bar = ttk.Frame(parent); bar.pack(fill='x', pady=(0, 5))
        for icon, title, command in actions:
            ttk.Button(bar, text=f'{icon} {title}', command=command).pack(side='right', padx=2)
        return bar

    def _tree(self, parent, columns):
        table_area = ttk.Frame(parent)
        table_area.pack(fill='both', expand=True)
        tree = ttk.Treeview(table_area, columns=tuple(key for key, _label, _width in columns), show='headings', selectmode='extended', style='SalesClaims.Treeview')
        for key, label, width in columns:
            tree.heading(key, text=label, anchor='e'); tree.column(key, width=width, anchor='e', stretch=True)
            if key == 'select':
                tree.heading(key, text='انتخاب ☐', command=lambda target=tree: self._toggle_all_checkboxes(target))
        tree.configure(displaycolumns=tuple(key for key, _label, _width in reversed(columns)))
        xscroll = ttk.Scrollbar(table_area, orient='horizontal', command=tree.xview)
        yscroll = ttk.Scrollbar(table_area, orient='vertical', command=tree.yview)
        tree.configure(xscrollcommand=xscroll.set, yscrollcommand=yscroll.set)
        yscroll.pack(side='right', fill='y')
        xscroll.pack(side='bottom', fill='x')
        tree.pack(side='left', fill='both', expand=True)
        return tree

    def _install_row_checkbox(self, tree):
        tree.bind('<Button-1>', lambda event: self._toggle_row_checkbox(event, tree), add='+')

    def _toggle_row_checkbox(self, event, tree):
        """Toggle the visible row checkbox while retaining normal Treeview selection."""
        iid = tree.identify_row(event.y)
        if tree.identify_region(event.x, event.y) != 'cell' or not iid:
            return
        visible_columns = tuple(tree['displaycolumns'] or tree['columns'])
        column_id = tree.identify_column(event.x)
        try:
            visible_key = visible_columns[int(column_id[1:]) - 1]
        except (ValueError, IndexError):
            return
        if visible_key != 'select':
            return
        values = list(tree.item(iid, 'values'))
        column_index = list(tree['columns']).index('select')
        checked = not str(values[column_index]).startswith('☑')
        values[column_index] = selection_cell(checked)
        tree.item(iid, values=values)
        # The dedicated checkbox is the explicit multi-selection control.
        # Do not replace earlier checked rows while the user ticks another one.
        if checked:
            tree.selection_add(iid)
        else:
            tree.selection_remove(iid)
        tree.focus(iid)
        tree.see(iid)
        return 'break'

    def _toggle_all_checkboxes(self, tree):
        """Toggle every visible row from the dedicated selection-column header."""
        item_ids = list(tree.get_children())
        if not item_ids:
            return
        column_index = list(tree['columns']).index('select')
        checked = not all(str(tree.item(iid, 'values')[column_index]).startswith('☑') for iid in item_ids)
        for iid in item_ids:
            values = list(tree.item(iid, 'values')); values[column_index] = selection_cell(checked)
            tree.item(iid, values=values)
        if checked:
            tree.selection_set(item_ids)
        else:
            tree.selection_remove(item_ids)
        tree.heading('select', text=f"انتخاب {'☑' if checked else '☐'}")

    def _checked_or_selected_rows(self, tree):
        """Return every checked row, falling back to normal Treeview selection."""
        try:
            select_index = list(tree['columns']).index('select')
            checked = [iid for iid in tree.get_children()
                       if str(tree.item(iid, 'values')[select_index]).startswith('☑')]
        except (ValueError, IndexError, tk.TclError):
            checked = []
        return checked or list(tree.selection())

    def _export_tree(self, tree, title: str, export_format: str | None = None):
        """Export selected rows, or the visible table when none is selected, without mutation."""
        item_ids = tree.selection() or tree.get_children()
        headers = [str(tree.heading(key, 'text')) for key in tree['columns']]
        rows = [tuple(tree.item(iid, 'values')) for iid in item_ids]
        payload = finance_export_rows(headers, rows) if tree is getattr(self, 'finance_tree', None) else visible_rows_payload(headers, rows)
        if not payload:
            messagebox.showinfo('خروجی', 'داده‌ای برای خروجی وجود ندارد.', parent=self)
            return
        formats = {'json': ('JSON', '.json', [('JSON', '*.json')]),
                   'csv': ('CSV', '.csv', [('CSV', '*.csv')]),
                   'txt': ('TXT', '.txt', [('Text', '*.txt')]),
                   'xlsx': ('Excel', '.xlsx', [('Excel', '*.xlsx')]),
                   'pdf': ('PDF', '.pdf', [('PDF', '*.pdf')]),
                   'print': ('پرینت', '.pdf', [('PDF', '*.pdf')])}
        selected = formats.get(str(export_format or '').lower())
        path = filedialog.asksaveasfilename(
            parent=self, title=f'خروجی {selected[0] if selected else "JSON/CSV/TXT"} {title}',
            defaultextension=selected[1] if selected else '.json',
            filetypes=selected[2] if selected else [item[2][0] for item in formats.values()],
        )
        if not path:
            return
        destination = Path(path)
        try:
            suffix = selected[1] if selected else destination.suffix.lower()
            if suffix == '.xlsx':
                from openpyxl import Workbook
                book = Workbook(); sheet = book.active; sheet.append(headers)
                for row in payload: sheet.append([row.get(header, '') for header in headers])
                book.save(destination)
            elif suffix == '.pdf':
                export_contract_pdf(destination, headers, payload, f'خروجی {title}')
                if export_format == 'print' and hasattr(os, 'startfile'):
                    os.startfile(str(destination), 'print')
            elif suffix == '.csv':
                with destination.open('w', encoding='utf-8-sig', newline='') as handle:
                    writer = csv.DictWriter(handle, fieldnames=headers)
                    writer.writeheader(); writer.writerows(payload)
            elif suffix == '.txt':
                with destination.open('w', encoding='utf-8') as handle:
                    handle.write('\n'.join('\t'.join(row.get(header, '') for header in headers) for row in payload))
            else:
                destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
        except OSError as exc:
            messagebox.showerror('خروجی', str(exc), parent=self)
            return
        messagebox.showinfo('خروجی', f'فقط دادهٔ نمایش‌داده‌شده صادر شد:\n{destination.name}', parent=self)

    def _export_contract_selected(self, _event=None):
        export_label = {'Excel': 'Excel', 'PDF': 'PDF', 'پرینت': 'پرینت'}.get(self.contract_export_choice.get())
        if not export_label:
            return
        selected = list(self.contract_tree.selection())
        if not selected:
            return messagebox.showwarning('خروجی قرارداد', 'ابتدا یک قرارداد را انتخاب کنید.', parent=self)
        # The main-tab selector opens the relevant contract only.  The actual
        # format is selected inside that window, where its populated preview
        # can be reviewed before Save As is opened.
        ContractDetails(self.winfo_toplevel(), db=self.db, base_dir=self.base_dir, sale_id=str(selected[0]))
        self.contract_export_choice.set(EXPORT_PLACEHOLDER)

    def _build_contracts(self):
        self._toolbar(self.contracts, [('➕', 'ثبت قرارداد', self.open_wizard), ('✏️', 'ویرایش', self.open_edit_wizard), ('🔄', 'بروزرسانی قیمت‌ها', self.refresh)])
        export_bar = ttk.Frame(self.contracts); export_bar.pack(fill='x', pady=(0, 5))
        self.contract_export_choice = tk.StringVar(value=EXPORT_PLACEHOLDER)
        ttk.Label(export_bar, text='خروجی').pack(side='right', padx=(8, 2))
        contract_export_combo = ttk.Combobox(export_bar, textvariable=self.contract_export_choice, values=(EXPORT_PLACEHOLDER, 'Excel', 'PDF', 'پرینت'), state='readonly', justify='right', width=18)
        contract_export_combo.pack(side='right', padx=3)
        contract_export_combo.bind('<<ComboboxSelected>>', self._export_contract_selected)
        search = ttk.Frame(self.contracts); search.pack(fill='x', pady=(0, 5))
        self.contract_search_kind = tk.StringVar(value='همه'); self.contract_search_value = tk.StringVar(value='همه')
        ttk.Label(search, text='نمایش بر اساس').pack(side='right', padx=4)
        kind = ttk.Combobox(search, textvariable=self.contract_search_kind, values=('همه', 'نام خریدار', 'شماره واحد', 'شماره قرارداد'), state='readonly', justify='right', width=18)
        kind.pack(side='right', padx=4); kind.bind('<<ComboboxSelected>>', lambda _e: self._refresh_contract_search_values())
        self.contract_search_combo = ttk.Combobox(search, textvariable=self.contract_search_value, state='readonly', justify='right', width=28)
        self.contract_search_combo.pack(side='right', padx=4); self.contract_search_combo.bind('<<ComboboxSelected>>', lambda _e: self.refresh())
        self._refresh_contract_search_values()
        self.contract_tree = self._tree(self.contracts, CONTRACT_TABLE_COLUMNS)
        self.contract_tree._barkav_non_currency_columns_186 = CONTRACT_NON_MONEY_COLUMNS
        self._install_row_checkbox(self.contract_tree)
        self.contract_tree.bind('<<TreeviewSelect>>', lambda _e: self._select_contract())
        self.contract_tree.bind('<Button-3>', lambda e: self._context(e, self.contract_tree))
        self.contract_tree.bind('<Double-1>', lambda _e: self._open_contract_details())

    def _restore_contract_non_money_headings(self):
        """Undo inherited currency decoration on non-monetary contract headings."""
        try:
            labels = {key: label for key, label, _width in CONTRACT_TABLE_COLUMNS}
            for key in CONTRACT_NON_MONEY_COLUMNS:
                self.contract_tree.heading(key, text=labels[key])
        except tk.TclError:
            return

    def _build_finance(self):
        self._toolbar(self.finance, [('➕', 'اضافه', self._finance_add), ('✏️', 'ویرایش', self._finance_edit), ('🗑', 'حذف', self._finance_delete), ('🔄', 'بروزرسانی', self._refresh_finance)])
        export_bar = ttk.Frame(self.finance); export_bar.pack(fill='x', pady=(0, 5))
        self.finance_export_choice = tk.StringVar(value=EXPORT_PLACEHOLDER)
        ttk.Label(export_bar, text='خروجی').pack(side='right', padx=(8, 2))
        finance_export_combo = ttk.Combobox(export_bar, textvariable=self.finance_export_choice, values=(EXPORT_PLACEHOLDER, 'Excel', 'PDF', 'پرینت'), state='readonly', justify='right', width=18)
        finance_export_combo.pack(side='right', padx=3)
        finance_export_combo.bind('<<ComboboxSelected>>', self._export_finance_selected)
        self.finance_summary = tk.StringVar(value='قراردادی انتخاب نشده است')
        summary_card = ttk.Frame(self.finance, padding=5); summary_card.pack(fill='x', padx=5, pady=5)
        ttk.Label(summary_card, textvariable=self.finance_summary, anchor='e', justify='right', style='SalesClaims.Summary.TLabel').pack(fill='x')
        search = ttk.Frame(self.finance); search.pack(fill='x', pady=(0, 5))
        self.finance_search_kind = tk.StringVar(value='همه'); self.finance_search_value = tk.StringVar(value='همه')
        ttk.Label(search, text='نمایش بر اساس').pack(side='right', padx=4)
        kind = ttk.Combobox(search, textvariable=self.finance_search_kind,
                            values=('همه', 'نام خریدار', 'شماره واحد', 'شماره قرارداد'),
                            state='readonly', justify='right', width=18)
        kind.pack(side='right', padx=4); kind.bind('<<ComboboxSelected>>', lambda _e: self._refresh_finance_search_values())
        self.finance_search_combo = ttk.Combobox(search, textvariable=self.finance_search_value,
                                                  state='readonly', justify='right', width=28)
        self.finance_search_combo.pack(side='right', padx=4)
        self.finance_search_combo.bind('<<ComboboxSelected>>', lambda _e: self._select_finance_search())
        self._refresh_finance_search_values()
        self.finance_tree = self._tree(self.finance, [('row', 'ردیف', 80), ('stage', 'مرحله پرداخت', 130), ('due', 'تاریخ پرداخت', 110), ('amount', 'مبلغ پرداخت (تومان)', 150), ('late', 'دیرکرد (روز)', 100), ('paid', 'کل مبلغ واریزی مرحله (تومان)', 180), ('balance', 'بدهکار/بستانکار تا این مرحله (تومان)', 180), ('final', 'بدهکار تا پایان قرارداد (تومان)', 180), ('document', 'نمایش سند', 110), ('select', 'انتخاب', 75)])
        self.finance_tree._barkav_non_currency_columns_186 = ('row', 'stage', 'due', 'late', 'document')
        self.finance_tree.heading('row', text='ردیف', anchor='e')
        # Keep non-monetary finance headings free of the money unit even if a
        # shared tree/theme layer supplies legacy heading text.
        self.finance_tree.heading('stage', text='مرحله پرداخت', anchor='e')
        self.finance_tree.heading('due', text='تاریخ پرداخت', anchor='e')
        self._install_row_checkbox(self.finance_tree)
        self.finance_tree.bind('<Double-1>', self._double_click_finance_payment)
        self.finance_tree.bind('<Button-1>', self._finance_document_click, add='+')
        self.finance_tree.bind('<Button-3>', lambda e: self._context(e, self.finance_tree))

    def _export_finance_selected(self, _event=None):
        export_label = {'Excel': 'Excel', 'PDF': 'PDF', 'پرینت': 'پرینت'}.get(self.finance_export_choice.get())
        if not export_label:
            return
        # Do not create a preview behind the receipt window.  The user chooses
        # the format again inside "اسناد واریزی", then receives its preview.
        self._open_selected_documents()
        self.finance_export_choice.set(EXPORT_PLACEHOLDER)

    def _save_finance_export(self, choice, headers, payload, item_ids):
        export_format = {'Excel': 'xlsx', 'PDF': 'pdf', 'پرینت': 'print'}.get(choice, 'pdf')
        actual_format = 'pdf' if export_format == 'print' else export_format
        path = filedialog.asksaveasfilename(parent=self, title=f'خروجی {actual_format.upper()} مالی',
                                            defaultextension=f'.{actual_format}',
                                            filetypes=[('Excel', '*.xlsx')] if actual_format == 'xlsx' else [('PDF', '*.pdf')])
        if not path:
            return
        destination = Path(path)
        try:
            if actual_format == 'xlsx':
                from openpyxl import Workbook
                book = Workbook(); sheet = book.active; sheet.append(headers)
                for row in payload: sheet.append([row.get(header, '') for header in headers])
                book.save(destination)
            else:
                documents = self._finance_export_documents(item_ids)
                export_receipts_pdf(destination, headers, [[row.get(header, '') for header in headers] for row in payload],
                                        'خروجی مالی', [], documents)
                if export_format == 'print' and hasattr(os, 'startfile'):
                    os.startfile(str(destination), 'print')
        except (OSError, ImportError) as exc:
            return messagebox.showerror('خروجی مالی', str(exc), parent=self)
        messagebox.showinfo('خروجی مالی', 'خروجی آماده شد.', parent=self)

    def _finance_export_documents(self, installment_ids):
        paths, seen = [], set()
        for installment_id in installment_ids:
            receipt = self._finance_receipt_context(str(installment_id))
            if not receipt:
                continue
            entities = [('RECEIVED_PAYMENT', str(receipt.get('received_payment_id') or '')),
                        ('SALE_PAYMENT', str(receipt.get('received_payment_id') or '')),
                        ('TRANSACTION', str(receipt.get('transaction_id') or ''))]
            for document in self.attachments.list_for_entities(entities):
                path = str(document.get('local_path') or '')
                if path and path not in seen:
                    paths.append(path); seen.add(path)
        return paths

    def _finance_document_click(self, event):
        """Open the selected row's receipts when the leftmost document cell is clicked."""
        iid = self.finance_tree.identify_row(event.y)
        if not iid or self.finance_tree.identify_region(event.x, event.y) != 'cell':
            return
        visible = tuple(self.finance_tree['displaycolumns'] or self.finance_tree['columns'])
        try:
            key = visible[int(self.finance_tree.identify_column(event.x)[1:]) - 1]
        except (ValueError, IndexError):
            return
        if key != 'document':
            return
        self.finance_tree.selection_set(iid); self.finance_tree.focus(iid)
        self._open_selected_documents()
        return 'break'

    def open_wizard(self, initial: dict[str, Any] | None = None, document_path: str = '', sale_id: str = ''):
        ContractWizard(self.winfo_toplevel(), db=self.db, base_dir=self.base_dir, on_saved=self._changed,
                          record_opener=self.open_sales_record, initial=initial, document_path=document_path, sale_id=sale_id)

    def open_edit_wizard(self):
        if not self.selected_sale_id:
            return messagebox.showwarning('ویرایش قرارداد', 'ابتدا یک قرارداد را انتخاب کنید.', parent=self)
        row = self.revenue.get_sale(self.selected_sale_id) or {}
        metadata = self._meta(row)
        initial = {
            **metadata,
            'project_name': metadata.get('project_name') or row.get('project_name'),
            'block_name': metadata.get('block_name') or row.get('block_name'),
            'unit_number': metadata.get('unit_number') or row.get('unit_number'),
            'floor': metadata.get('floor') or row.get('floor'),
            'map_area': metadata.get('map_area') or row.get('area'),
            'contract_area': metadata.get('contract_area') or row.get('area'),
            'contract_number': metadata.get('contract_number') or row.get('contract_number'),
            'contract_date': metadata.get('contract_date') or row.get('contract_date_shamsi'),
            'buyer_name': metadata.get('buyer_name') or row.get('buyer_name'),
            'buyer_phone': metadata.get('buyer_phone') or row.get('buyer_phone'),
            'total_price': metadata.get('total_price') or row.get('contract_amount'),
        }
        documents = self.attachments.list_for_entities([('SALE', self.selected_sale_id)])
        document_path = str((documents[0] if documents else {}).get('local_path') or '')
        self.open_wizard(initial=initial, document_path=document_path, sale_id=self.selected_sale_id)

    def open_sales_record(self, sale_id: str = '', draft: dict[str, Any] | None = None, document_path: str = ''):
        # A new wizard draft must never inherit the previously selected sale.
        # Keep the selected-sale fallback only for opening an existing record.
        sale_id = str(sale_id or (self.selected_sale_id if draft is None else '') or '')
        SalesRecord(self.winfo_toplevel(), db=self.db, base_dir=self.base_dir, sale_id=sale_id,
                       draft=draft, document_path=document_path, on_saved=self._changed)

    def _auto_refresh_prices(self):
        """The table price is derived from the ledger on every refresh; no price is posted."""
        try:
            if self.winfo_exists():
                self.refresh(); self._price_refresh_job = self.after(60000, self._auto_refresh_prices)
        except tk.TclError:
            pass

    def _refresh_contract_search_values(self):
        kind = getattr(self, 'contract_search_kind', tk.StringVar(value='همه')).get()
        if kind == 'همه':
            self.contract_search_value.set('همه'); self.contract_search_combo.configure(values=('همه',)); return
        key = {'نام خریدار': 'buyer_name', 'شماره واحد': 'unit_number', 'شماره قرارداد': 'contract_number'}[kind]
        values = sorted({str(row.get(key) or '') for row in self.revenue.list_sales(active_only=True, limit=500) if row.get(key)})
        self.contract_search_combo.configure(values=('همه', *values)); self.contract_search_value.set('همه')

    def _refresh_finance_search_values(self):
        kind = self.finance_search_kind.get()
        if kind == 'همه':
            self.finance_search_value.set('همه'); self.finance_search_combo.configure(values=('همه',)); return
        key = {'نام خریدار': 'buyer_name', 'شماره واحد': 'unit_number', 'شماره قرارداد': 'contract_number'}[kind]
        values = sorted({str(row.get(key) or '') for row in self.revenue.list_sales(active_only=True, limit=500) if row.get(key)})
        self.finance_search_combo.configure(values=('همه', *values)); self.finance_search_value.set('همه')

    def _select_finance_search(self):
        kind, value = self.finance_search_kind.get(), self.finance_search_value.get()
        if kind == 'همه' or value == 'همه':
            return
        key = {'نام خریدار': 'buyer_name', 'شماره واحد': 'unit_number', 'شماره قرارداد': 'contract_number'}[kind]
        match = next((row for row in self.revenue.list_sales(active_only=True, limit=500) if str(row.get(key) or '') == value), None)
        if not match:
            return
        self.selected_sale_id = str(match['sale_id'])
        if self.contract_tree.exists(self.selected_sale_id):
            self.contract_tree.selection_set(self.selected_sale_id)
        self._select_contract()

    def refresh(self):
        previous = self.selected_sale_id
        for iid in self.contract_tree.get_children(): self.contract_tree.delete(iid)
        rows = filter_contract_rows(self.revenue.list_sales(active_only=True, limit=500), self.contract_search_kind.get(), self.contract_search_value.get())
        for number, row in enumerate(rows, 1):
            detail = self.revenue.sale_summary(row['sale_id']); meta = self._meta(row)
            self.contract_tree.insert('', 'end', iid=row['sale_id'], values=contract_table_values(number, row, meta, detail))
        self._restore_contract_non_money_headings()
        if previous and self.contract_tree.exists(previous): self.contract_tree.selection_set(previous); self._select_contract()

    def _meta(self, row):
        try: return json.loads(str(row.get('notes') or '{}'))
        except (TypeError, ValueError, json.JSONDecodeError): return {}

    def _select_contract(self, open_docs=False):
        selected = self.contract_tree.selection()
        if not selected: return
        self.selected_sale_id = selected[0]
        row = self.revenue.get_sale(self.selected_sale_id) or {}; detail = self.revenue.sale_summary(self.selected_sale_id)
        self.finance_summary.set(f"قرارداد {row.get('contract_number') or '-'} | مبلغ کل: {toman(row.get('contract_amount'))} تومان | کل پرداختی: {toman(detail.cash_received)} تومان | معوق: {toman(detail.overdue_amount)} تومان | مانده: {toman(detail.receivable)} تومان")
        self._refresh_finance()

    def _open_contract_details(self):
        selected = self.contract_tree.selection()
        if not selected:
            return
        self.selected_sale_id = selected[0]
        ContractDetails(self.winfo_toplevel(), db=self.db, base_dir=self.base_dir, sale_id=self.selected_sale_id)

    def _refresh_finance(self):
        for iid in self.finance_tree.get_children(): self.finance_tree.delete(iid)
        if not self.selected_sale_id: return
        installments = self.revenue.list_installments(self.selected_sale_id)
        running = 0; total = sum(int(item.get('amount') or 0) for item in installments)
        today = date.today()
        for number, item in enumerate(installments, 1):
            amount, paid = int(item.get('amount') or 0), int(item.get('paid_amount') or 0)
            running += max(0, amount - paid)
            try: late = max(0, (today - date.fromisoformat(str(item.get('due_date')))).days) if paid < amount else 0
            except ValueError: late = 0
            self.finance_tree.insert('', 'end', iid=item['installment_id'], values=(fa(number), item.get('description') or f'قسط {number}', persian_digits(item.get('due_date_shamsi') or '-'), toman(amount), fa(late), toman(paid), f'بدهکار: {toman(max(0, running))}', toman(max(0, total - sum(int(x.get('paid_amount') or 0) for x in installments[:number]))), 'نمایش سند', selection_cell()))

    def _context(self, event, tree):
        iid = tree.identify_row(event.y)
        if iid: tree.selection_set(iid); tree.focus(iid)
        menu = tk.Menu(self, tearoff=0)
        if tree is self.finance_tree:
            menu.add_command(label='➕ اضافه', command=self._finance_add)
            menu.add_command(label='✏️ ویرایش', command=self._finance_edit)
            menu.add_command(label='🗑 حذف', command=self._finance_delete)
        else:
            menu.add_command(label='➕ ثبت قرارداد', command=self.open_wizard)
            menu.add_command(label='✏️ ویرایش', command=self._edit_hint)
            menu.add_command(label='🗑 لغو/حذف', command=self._cancel)
            menu.add_separator(); menu.add_command(label='📋 کپی', command=self._copy)
            menu.add_command(label='📥 پیست', command=lambda: messagebox.showinfo('پیست', 'پیست فقط برای فیلترهای متنی فعال است.', parent=self))
        try: menu.tk_popup(event.x_root, event.y_root)
        finally: menu.grab_release()

    def _copy(self):
        tree = self.contract_tree if self.nb.index(self.nb.select()) == 0 else self.finance_tree
        lines = ['\t'.join(map(str, tree.item(iid, 'values'))) for iid in tree.selection()]
        if lines: self.clipboard_clear(); self.clipboard_append('\n'.join(lines))

    def _edit_hint(self):
        self.open_edit_wizard()

    def _cancel(self):
        if not self.selected_sale_id: return
        if not messagebox.askyesno('لغو قرارداد', 'قرارداد انتخاب‌شده لغو شود؟ واریزی ثبت‌شده مانع لغو است.', parent=self): return
        try: self.revenue.cancel_sale(self.selected_sale_id, source='GUI_1871_CONTRACT', actor_id='GUI')
        except Exception as exc: messagebox.showerror('لغو قرارداد', str(exc), parent=self); return
        self.selected_sale_id = ''; self._changed()

    def _open_document(self):
        selected = self.doc_tree.selection()
        if not selected: return
        try: self.attachments.open(str(self.doc_tree.item(selected[0], 'values')[-1]))
        except Exception as exc: messagebox.showerror('اسناد', str(exc), parent=self)

    def _document_delete_hint(self):
        messagebox.showinfo('اسناد', 'حذف سند از این صفحه انجام نمی‌شود تا سابقهٔ مالی و پیوست‌ها حفظ شود.', parent=self)

    def _finance_delete_hint(self):
        messagebox.showinfo('مالی', 'ابطال واریزی فقط از مسیر بازبینی حسابرسی در تب ثبت تراکنش انجام می‌شود.', parent=self)

    def _finance_add(self):
        self._create_finance_pending('ADD')

    def _finance_delete(self):
        """Create a Pending deletion request for the selected finance row."""
        self._create_finance_pending('DELETE')

    def _create_finance_pending(self, action):
        dashboard = self.winfo_toplevel(); center = getattr(dashboard, 'approval_center', None)
        if center is None or not callable(getattr(center, 'open_wizard', None)):
            return messagebox.showwarning('عملیات مالی', 'پنجره ثبت دیتا در دسترس نیست.', parent=self)
        if not self.selected_sale_id:
            return messagebox.showwarning('عملیات مالی', 'ابتدا قرارداد را انتخاب کنید.', parent=self)
        sale = self.revenue.get_sale(self.selected_sale_id) or {}
        selected = self.finance_tree.selection()
        installment = self.revenue.get_installment(str(selected[0])) if selected else None
        receipt = self._finance_receipt_context(str(selected[0])) if selected else {}
        if action == 'EDIT' and not receipt:
            return messagebox.showinfo('ویرایش مالی', 'برای این مرحله پرداخت ثبت‌شده‌ای وجود ندارد.', parent=self)
        now_id = int(datetime.now().timestamp() * 1_000_000)
        payload = sales_finance_pending_payload(action=action, sale={**sale, 'sale_id': self.selected_sale_id},
                                                     installment=installment, receipt=receipt)
        try:
            self.db.register_telegram_message(chat_id=-187, message_id=now_id, raw_text='', content_type='sales_finance',
                                              message_date=sales_finance_pending_message_date(payload), state='RECEIVED')
            pending_id = self.db.create_or_replace_pending(telegram_chat_id=-187, telegram_message_id=now_id,
                                                            detected_type='TRANSACTION', payload=payload, confidence=1.0,
                                                            status='PENDING_REVIEW')
            center.refresh()
            row = next((item for item in getattr(center, 'rows', []) if str(item.get('pending_id')) == str(pending_id)), None)
            if row:
                notebook = getattr(dashboard, 'notebook', None); approval_tab = getattr(dashboard, 'approval_tab', None)
                if notebook is not None and approval_tab is not None: notebook.select(approval_tab)
                center.open_wizard(row)
        except Exception as exc:
            messagebox.showerror('عملیات مالی', str(exc), parent=self)

    def _finance_edit(self):
        """Create an audit-safe Pending edit request for the selected receipt."""
        self._create_finance_pending('EDIT')

    def _finance_receipt_context(self, installment_id):
        """Read an existing receipt reference only; no financial row is mutated here."""
        with self.db.read() as con:
            row = con.execute(
                """SELECT r.received_payment_id, r.transaction_id, r.amount, r.received_date_shamsi
                   FROM payment_allocations a
                   JOIN received_payments r ON r.received_payment_id=a.received_payment_id
                   WHERE a.installment_id=? ORDER BY r.received_date_shamsi DESC LIMIT 1""",
                (str(installment_id),),
            ).fetchone()
        return dict(row) if row else {}

    def _open_pending_editor_for_finance(self, title):
        dashboard = self.winfo_toplevel()
        center = getattr(dashboard, 'approval_center', None)
        if center is None:
            messagebox.showwarning(title, 'پنجره ثبت دیتا در دسترس نیست.', parent=self); return False
        rows = list(getattr(center, 'rows', []) or [])
        candidates = []
        for row in rows:
            payload = row.get('payload') or {}
            if str(payload.get('sale_id') or '') == str(self.selected_sale_id):
                candidates.append(row)
        if len(candidates) != 1:
            messagebox.showinfo(title, 'برای این قرارداد ردیف متناظر یکتا در Pending پیدا نشد.', parent=self); return False
        row = candidates[0]
        try:
            center.current = row
            center.open_editor(row)
            notebook = getattr(dashboard, 'notebook', None)
            approval_tab = getattr(dashboard, 'approval_tab', None)
            if notebook is not None and approval_tab is not None:
                notebook.select(approval_tab)
            return True
        except Exception as exc:
            messagebox.showerror(title, str(exc), parent=self); return False

    def _open_record(self):
        self.open_sales_record()

    def _show_payment_documents(self):
        messagebox.showinfo('اسناد پرداخت', 'اسناد پرداخت از تب ثبت تراکنش به همین قرارداد متصل می‌شوند. برای نمایش فیش، رکورد واریزی متناظر را باز کنید.', parent=self)

    def _double_click_finance_payment(self, event):
        iid = self.finance_tree.identify_row(event.y)
        if iid:
            self.finance_tree.selection_set(iid); self.finance_tree.focus(iid)
            self._open_selected_documents()

    def _open_selected_documents(self):
        selected = self._checked_or_selected_rows(self.finance_tree)
        if not selected:
            return messagebox.showwarning('اسناد واریزی', 'ابتدا ردیف‌های موردنظر را انتخاب کنید.', parent=self)
        payments = []
        with self.db.read() as con:
            for installment_id in selected:
                row = [dict(item) for item in con.execute(
                    """SELECT i.description, i.sequence_no, r.received_payment_id, r.transaction_id,
                              r.amount, r.received_date_shamsi
                       FROM installments i
                       LEFT JOIN payment_allocations a ON a.installment_id=i.installment_id
                       LEFT JOIN received_payments r ON r.received_payment_id=a.received_payment_id
                       WHERE i.installment_id=? ORDER BY r.received_date_shamsi""", (str(installment_id),)
                ).fetchall()]
                if not row:
                    continue
                first = row[0]; stage = str(first.get('description') or f"مرحله {first.get('sequence_no') or '-'}")
                receipts, known_receipts = [], set()
                for item in row:
                    receipt_id = str(item.get('received_payment_id') or '')
                    if not receipt_id or receipt_id in known_receipts:
                        continue
                    receipts.append(item); known_receipts.add(receipt_id)
                # An installment may receive several payments.  Preserve each
                # receipt as its own export/card row, rather than summing it
                # into one stage and losing the one-to-one receipt evidence.
                if not receipts:
                    payments.append({'stage': stage, 'date': '-', 'amount': 0, 'documents': []})
                    continue
                multiple_receipts = len(receipts) > 1
                for number, receipt in enumerate(receipts, 1):
                    entities = [('RECEIVED_PAYMENT', str(receipt['received_payment_id'])),
                                ('SALE_PAYMENT', str(receipt['received_payment_id']))]
                    transaction_id = str(receipt.get('transaction_id') or '')
                    if transaction_id:
                        entities.append(('TRANSACTION', transaction_id))
                    documents, known = [], set()
                    for document in self.attachments.list_for_entities(entities):
                        attachment_id = str(document.get('attachment_id') or '')
                        if attachment_id and attachment_id in known:
                            continue
                        documents.append(document)
                        if attachment_id:
                            known.add(attachment_id)
                    receipt_stage = f'{stage}-{persian_digits(number)}' if multiple_receipts else stage
                    payments.append({
                        'stage': receipt_stage,
                        'date': receipt.get('received_date_shamsi') or '-',
                        'amount': int(receipt.get('amount') or 0),
                        'documents': documents,
                    })
        sale = self.revenue.get_sale(self.selected_sale_id) or {}
        return PaymentDocuments(self.winfo_toplevel(), db=self.db, base_dir=self.base_dir,
                                   payments=payments, contract_total=int(sale.get('contract_amount') or 0))

    def _changed(self):
        self.refresh()
        if self.on_changed: self.on_changed()
