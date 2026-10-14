from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def insert_before(target, template, text: str):
    paragraph = target.insert_paragraph_before()
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if template._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return paragraph


def update(source: Path, destination: Path) -> None:
    document = Document(source)
    paragraphs = document.paragraphs
    replace_text(paragraphs[1], 'نسخه 18.5 Feature Freeze آماده تثبیت و پذیرش نهایی ویندوز')
    replace_text(paragraphs[6], '• قابلیت‌های برنامه تا نسخه 18.5 تکمیل و Feature Freeze انجام شد؛ تا عبور از گیت تثبیت قابلیت جدید اضافه نشود.')

    step_two = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith('گام ۲ —'))
    heading_template = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith('گام ۱ —'))
    bullet_template = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith('• تکمیل') or paragraph.text.startswith('• قابلیت'))
    insert_before(step_two, heading_template, 'وضعیت تحویل نسخه 18.5 — Feature Freeze مورخ 2026-08-20')
    insert_before(step_two, bullet_template, '• نسخه 18.5 با اصلاحات پایداری، ظاهر Dark/Light، مشخصات پایه، ساختار پویا و قابل‌ویرایش طبقات/واحدها/انباری/پارکینگ، مسئول پیگیری اداری، فراخوانی واقعی بکاپ چت و چیدمان افقی شاخص‌های مالی ساخته شد.')
    insert_before(step_two, bullet_template, '• نتیجه آزمون خودکار: 209 تست موفق و 30 تست GUI اختیاری Skip؛ هیچ شکست در Suite پیش‌فرض ثبت نشد.')
    insert_before(step_two, bullet_template, '• آزمون GUI اجباری در محیط ساخت به علت نبود Tcl/Tk (init.tcl) قابل اجرا نبود؛ پذیرش دستی روی Windows با DPI واقعی، Telegram/Gemini، TGJU، Printer/PDF و داده کپی‌شده پروژه همچنان الزامی است.')
    insert_before(step_two, bullet_template, '• گام ۱ هنوز تکمیل‌شده محسوب نمی‌شود؛ پس از Smoke Test ویندوز، آزمون Backup/Restore و تأیید نبود باگ Critical/High می‌توان وضعیت را «آماده عبور» کرد.')

    status_cell = document.tables[0].cell(0, 0)
    status_lines = [
        'وضعیت فعلی پروژه: گام ۱ — تثبیت نسخه نهایی دسکتاپ',
        '• نسخه فعلی: 18.5 — Feature Freeze / Release Candidate',
        '• نسخه هدف برای Freeze قابلیت‌ها: 18.5 — محقق شد',
        '• تصمیم فعلی: تا پایان گام ۱ قابلیت جدید اضافه نشود؛ تمرکز روی Smoke Test ویندوز، رفع ایراد و ایمنی داده باشد.',
        '• نتیجه آزمون خودکار نسخه: 209 Passed / 30 GUI-Skipped / 0 Failed',
        '• این سند باید همراه هر نسخه جدید به‌روزرسانی شود و «گام فعلی» در ابتدای آن مشخص باشد.',
    ]
    while len(status_cell.paragraphs) < len(status_lines):
        new_paragraph = status_cell.add_paragraph()
        template = status_cell.paragraphs[-2]
        if template._p.pPr is not None:
            if new_paragraph._p.pPr is not None:
                new_paragraph._p.remove(new_paragraph._p.pPr)
            new_paragraph._p.insert(0, deepcopy(template._p.pPr))
        if template.runs and template.runs[0]._r.rPr is not None:
            run = new_paragraph.add_run('')
            run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    for paragraph, text in zip(status_cell.paragraphs, status_lines):
        replace_text(paragraph, text)

    summary_table = document.tables[1]
    replace_text(summary_table.cell(1, 3).paragraphs[0], 'Feature Freeze انجام شد؛ Acceptance ویندوز باقی است')

    version_table = document.tables[2]
    values = [
        '18.5 — Feature Freeze / Release Candidate',
        'گام ۱ — تثبیت محصول دسکتاپ',
        'Feature Freeze انجام شده؛ تثبیت و پذیرش ویندوز در حال انجام',
        'اصلاحات عمومی و گرافیکی، طبقات/واحدها، ثبت دیتا، گزارش مالی، لانچر امن و 209 تست موفق',
        'Smoke Test ویندوز، Backup/Restore، بستن Critical/High؛ سپس تصمیم عبور به گام ۲',
    ]
    for row, value in zip(version_table.rows, values):
        replace_text(row.cells[1].paragraphs[0], value)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


if __name__ == '__main__':
    update(Path(sys.argv[1]), Path(sys.argv[2]))
