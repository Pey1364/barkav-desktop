"""One immutable report document for the non-calendar preview and every output."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from dataclasses import dataclass

DOCUMENT_HEADERS = {'نمایش سند', 'فیش واریزی', 'رسید', 'سند'}


def fa_number(value):
    return str(value).translate(str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹'))


def selected_tree_items(tree):
    """Keep visual order and restrict output to selected rows when present."""
    selected = set(getattr(tree, 'selection', lambda: ())())
    return [iid for iid in tree.get_children() if not selected or iid in selected]


def selected_table_rows(table):
    if hasattr(table, 'sorted_rows'):
        rows = table.sorted_rows()
        if not hasattr(table, 'tree'):
            return [list(row) for row in rows]
        return [list(rows[table.tree.index(iid)]) for iid in selected_tree_items(table.tree)]
    return [list(table.tree.item(iid, 'values')) for iid in selected_tree_items(table.tree)]


def resolve_report_document_paths(db, row):
    """Include pending evidence and evidence relinked by financial approval."""
    row = dict(row or {})
    paths = []
    def append(value):
        value = str(value or '').strip()
        if value and value not in paths:
            paths.append(value)
    for key in ('local_path', 'receipt_path', 'attachment_path', 'document_path'):
        append(row.get(key))
    if db is None:
        return paths
    pending_ids = {str(row.get(key) or '').strip() for key in
                   ('pending_id', 'source_pending_id', 'origin_pending_id')} - {''}
    entities = set()
    for pending_id in pending_ids:
        if hasattr(db, 'list_pending_attachments'):
            for attachment in db.list_pending_attachments(pending_id) or ():
                append(attachment.get('local_path'))
        if hasattr(db, 'get_financial_posting'):
            posting = db.get_financial_posting(pending_id) or {}
            for key, kind in (('expense_id', 'EXPENSE'), ('transaction_id', 'TRANSACTION'), ('cheque_id', 'CHEQUE')):
                if posting.get(key):
                    entities.add((kind, str(posting[key])))
    for key, kind in (('expense_id', 'EXPENSE'), ('transaction_id', 'TRANSACTION'), ('cheque_id', 'CHEQUE'), ('loan_id', 'LOAN')):
        if row.get(key):
            entities.add((kind, str(row[key])))
    if hasattr(db, 'list_entity_attachments'):
        for kind, identity in entities:
            for attachment in db.list_entity_attachments(kind, identity) or ():
                append(attachment.get('local_path'))
    return paths


@dataclass(frozen=True)
class ReportDocument187:
    pdf: bytes
    headers: tuple
    rows: tuple
    page_count: int
    title: str

    def page_images(self, scale=1.0):
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(self.pdf) as document:
            for index in range(len(document)):
                page = document[index]
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil().copy()
                bitmap.close()
                page.close()
                yield image

    def save(self, path, output_format):
        path = Path(path)
        if output_format != 'Excel':
            path.write_bytes(self.pdf)
            return path
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        book.remove(book.active)
        streams = []  # Image streams must remain open until workbook.save().
        for index, page in enumerate(self.page_images(scale=2), 1):
            sheet = book.create_sheet(f'صفحه {fa_number(index)}')
            sheet.sheet_view.rightToLeft = True
            sheet.sheet_view.showGridLines = False
            stream = BytesIO()
            page.save(stream, format='PNG')
            stream.seek(0)
            streams.append(stream)
            image = Image(stream)
            image.width, image.height = page.width / 2, page.height / 2
            sheet.add_image(image, 'A1')
            # Excel caps row heights at 409 points. Cover the image with
            # ordinary rows so tall portrait pages are never cropped.
            sheet.column_dimensions['A'].width = max(1, (image.width - 5) / 7)
            import math
            row_count = max(1, math.ceil(image.height / 20))
            for row_index in range(1, row_count + 1):
                sheet.row_dimensions[row_index].height = min(20, image.height - (row_index - 1) * 20) * .75
            sheet.print_area = f'A1:A{row_count}'
            sheet.page_setup.orientation = 'landscape' if page.width > page.height else 'portrait'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
            sheet.page_margins.left = sheet.page_margins.right = 0
            sheet.page_margins.top = sheet.page_margins.bottom = 0
            sheet.page_margins.header = sheet.page_margins.footer = 0
        data = book.create_sheet('داده‌های گزارش')
        data.sheet_view.rightToLeft = True
        data.append(list(self.headers))
        for row in self.rows:
            data.append(list(row))
        for row in data:
            for cell in row:
                cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
                cell.font = Font(name='Tahoma', size=11, bold=cell.row == 1)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for column in data.columns:
            data.column_dimensions[column[0].column_letter].width = min(42, max(12, max(len(str(c.value or '')) for c in column) + 2))
        data.freeze_panes = 'A2'
        book.active = 0
        book.save(path)
        for stream in streams:
            stream.close()
        return path


def build_report_document(headers, payload, specs, title, *, orientation='افقی',
                          audience='سایر', include_documents=True,
                          document_paths=(), contract_document_paths=(),
                          progress=None, project_name=''):
    """Build once; preview rasterizes these exact bytes and Save copies them."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import SimpleDocTemplate, Flowable, Table, TableStyle, Spacer, Image, PageBreak
    from reportlab.lib.utils import ImageReader
    from export_utils import _pdf_font, _rtl, _pdf_column_widths
    from sales_claims_stage18_7 import export_document_paths_187, export_project_name_187, _export_logo_image_187
    font = _pdf_font()
    page_size = A4 if orientation == 'عمودی' else landscape(A4)
    width, height = page_size
    margin = 24
    available = width - 2 * margin
    style = ParagraphStyle('report187', fontName=font, fontSize=8, leading=12, alignment=TA_RIGHT)
    heading = ParagraphStyle('report187heading', parent=style, fontSize=11, leading=15)
    class RTLText(Flowable):
        """Wrap logical text first, then shape each line for correct RTL order."""
        def __init__(self, value, selected_style, fixed_lines=None):
            super().__init__()
            self.value = str(value if value is not None else '')
            self.style = selected_style
            self.fixed_lines = fixed_lines
        def wrap(self, available_width, available_height):
            from reportlab.pdfbase.pdfmetrics import stringWidth
            self.width = max(1, available_width)
            if self.fixed_lines is not None:
                self.lines = list(self.fixed_lines)
            else:
                self.lines = []
                for logical_line in self.value.splitlines() or ['']:
                    line = ''
                    for word in logical_line.split():
                        candidate = f'{line} {word}'.strip()
                        if line and stringWidth(_rtl(candidate), font, self.style.fontSize) > self.width:
                            self.lines.append(line)
                            line = ''
                        if stringWidth(_rtl(word), font, self.style.fontSize) > self.width:
                            for character in word:
                                candidate = line + character
                                if line and stringWidth(_rtl(candidate), font, self.style.fontSize) > self.width:
                                    self.lines.append(line)
                                    line = character
                                else:
                                    line = candidate
                        else:
                            line = f'{line} {word}'.strip()
                    self.lines.append(line)
            self.height = max(1, len(self.lines)) * self.style.leading
            return self.width, self.height
        def split(self, available_width, available_height):
            self.wrap(available_width, available_height)
            count = int(available_height // self.style.leading)
            if count <= 0:
                return []
            if count >= len(self.lines):
                return [self]
            return [RTLText('', self.style, self.lines[:count]),
                    RTLText('', self.style, self.lines[count:])]
        def draw(self):
            self.canv.setFont(font, self.style.fontSize)
            self.canv.setFillColor(self.style.textColor)
            for index, line in enumerate(self.lines):
                self.canv.drawRightString(self.width,
                    self.height - self.style.fontSize - index * self.style.leading, _rtl(line))
    def paragraph(value, selected_style=style):
        return RTLText(value, selected_style)
    def table_style(header=False):
        commands = [('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#B8C2CC')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
                    ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]
        if header:
            commands.append(('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DCE8F2')))
        return TableStyle(commands)
    story = []
    grouped = {}
    for spec in specs or ():
        if isinstance(spec, (tuple, list)) and len(spec) >= 3:
            group, label, value = spec[:3]
        else:
            group, label, value = 'مشخصات', 'اطلاعات', spec
        grouped.setdefault(str(group), []).append((label, value))
    boxes = []
    for group, entries in grouped.items():
        cells = [[paragraph(group, heading), '']] + [[paragraph(value), paragraph(label)] for label, value in entries]
        box = Table(cells, colWidths=[available / 4] * 2, repeatRows=1, splitInRow=1, hAlign='RIGHT')
        box.setStyle(table_style(True))
        box.setStyle(TableStyle([('SPAN', (0, 0), (-1, 0))]))
        boxes.append(box)
    for index in range(0, len(boxes), 2):
        pair = boxes[index:index + 2]
        if any(box.wrap(available / 2, height)[1] > height - 150 for box in pair):
            for box in pair:
                story.extend([box, Spacer(1, 8)])
            continue
        paired = Table([[boxes[index + 1] if index + 1 < len(boxes) else '', boxes[index]]], colWidths=[available / 2] * 2)
        paired.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
        story.extend([paired, Spacer(1, 8)])
    headers = list(headers or ['گزارش'])
    records = [dict(row) if isinstance(row, dict) else {**dict(zip(headers, row)), '__values_187': list(row)} for row in payload]
    document_columns = [key for key in headers if key in DOCUMENT_HEADERS]
    display_headers = [key for key in headers if include_documents or key not in DOCUMENT_HEADERS]
    if not display_headers:
        headers.append('ردیف')
        display_headers = ['ردیف']
        for index, record in enumerate(records, 1):
            record['ردیف'] = fa_number(index)
    def display_value(row, key, index=None):
        raw = row.get('__values_187')
        value = raw[index] if index is not None and raw is not None and index < len(raw) else row.get(key, '')
        if key in DOCUMENT_HEADERS:
            return '\n'.join(Path(path).name for path in export_document_paths_187(value) if Path(path).suffix) if include_documents else ''
        return value
    plain_rows = [tuple(display_value(row, key, index) for index, key in enumerate(headers)) for row in records]
    physical_indices = list(reversed([index for index, key in enumerate(headers) if key in display_headers]))
    physical_headers = [headers[index] for index in physical_indices]
    widths = _pdf_column_widths(physical_headers, [[display_value(row, headers[index], index) for index in physical_indices] for row in records], available)
    table_heading = heading
    if 'شرح گزارش' in headers:
        # Keep each date intact and use compact column headings in calendar reports.
        from reportlab.pdfbase.pdfmetrics import stringWidth
        table_heading = ParagraphStyle('calendar187heading', parent=style, fontSize=8, leading=12)
        minimums = [max(24, max((stringWidth(_rtl(word), font, 8) for word in key.split()), default=0) + 12)
                    for key in physical_headers]
        for index, key in enumerate(physical_headers):
            if key in {'تاریخ', 'روز هفته'}:
                minimums[index] = max(minimums[index], stringWidth(_rtl('۱۴۰۴/۰۷/۰۱'), font, 8) + 12)
        weights = [5 if key == 'شرح گزارش' else 2 if key == 'تاریخ' else
                   .3 if key == 'ردیف' else 1 for key in physical_headers]
        remaining = max(0, available - sum(minimums))
        if sum(minimums) <= available:
            widths = [minimum + remaining * weight / sum(weights)
                      for minimum, weight in zip(minimums, weights)]
    head = [paragraph(key, table_heading) for key in physical_headers]
    pending_images = []
    def document_flowables(raw_paths):
        result = []
        for raw in raw_paths:
            candidate = Path(str(raw))
            if not candidate.is_file():
                result.append(paragraph(f'سند در دسترس نیست: {candidate.name}'))
                continue
            try:
                if candidate.suffix.lower() == '.pdf':
                    import pypdfium2 as pdfium
                    with pdfium.PdfDocument(str(candidate)) as source:
                        for doc_page in source:
                            bitmap = doc_page.render(scale=1.5)
                            picture = bitmap.to_pil().copy()
                            bitmap.close()
                            doc_page.close()
                            buffer = BytesIO()
                            picture.save(buffer, format='PNG')
                            buffer.seek(0)
                            pending_images.append(buffer)
                            result.append(Image(buffer, width=min(available, 250), height=min(height - 180, 300), kind='proportional'))
                else:
                    result.append(Image(str(candidate), width=min(available, 250), height=min(height - 180, 300), kind='proportional'))
            except (OSError, ValueError):
                result.append(paragraph(f'سند: {candidate.name}'))
        return result
    any_row_docs = False
    if not include_documents:
        document_columns = []
    paths_by_row = []
    for row in records:
        paths = []
        if include_documents:
            for key in document_columns + ['__document_paths_187']:
                for path in export_document_paths_187(row.get(key, '')):
                    if path not in paths and (Path(path).suffix.lower() in {'.pdf', '.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'} or Path(path).is_file()):
                        paths.append(path)
        paths_by_row.append(paths)
        any_row_docs = any_row_docs or bool(paths)
    if any_row_docs:
        for index, row in enumerate(records):
            table = Table([head, [paragraph(display_value(row, headers[column], column) or '-') for column in physical_indices]], colWidths=widths, repeatRows=1, splitInRow=1, hAlign='RIGHT')
            table.setStyle(table_style(True))
            story.append(table)
            if paths_by_row[index]:
                story.extend([Spacer(1, 5), paragraph(f'اسناد ردیف {fa_number(index + 1)}', heading)])
                for item in document_flowables(paths_by_row[index]):
                    story.extend([item, Spacer(1, 8)])
    else:
        table = Table([head] + [[paragraph(display_value(row, headers[column], column) or '-') for column in physical_indices] for row in records], colWidths=widths, repeatRows=1, splitInRow=1, hAlign='RIGHT')
        table.setStyle(table_style(True))
        story.append(table)
    if include_documents:
        extra_paths = [path for path in document_paths if not any(path in group for group in paths_by_row)]
        for label, paths in [('اسناد گزارش', extra_paths), ('تصویر قرارداد', contract_document_paths)]:
            if paths:
                story.extend([PageBreak(), paragraph(label, heading), Spacer(1, 8)])
                for item in document_flowables(paths):
                    story.extend([item, Spacer(1, 8)])
    if progress is not None:
        from sales_claims_stage18_7 import _PDF_TABLE_HEADER_BG, _PDF_TABLE_GRID_LINE
        overall = max(0.0, min(100.0, float(progress)))
        class ProjectProgressBar187(Flowable):
            def __init__(self):
                super().__init__()
                self.width, self.height = available, 36
            def wrap(self, available_width, available_height):
                self.width = min(available, available_width)
                return self.width, self.height
            def draw(self):
                canvas = self.canv
                canvas.saveState()
                canvas.setFillColor(colors.HexColor(_PDF_TABLE_GRID_LINE))
                canvas.roundRect(0, 0, self.width, 18, 9, stroke=0, fill=1)
                filled = self.width * overall / 100
                if filled > 0:
                    canvas.setFillColor(colors.HexColor(_PDF_TABLE_HEADER_BG))
                    canvas.roundRect(self.width - filled, 0, filled, 18,
                                     min(9, filled / 2), stroke=0, fill=1)
                canvas.setFillColor(colors.HexColor(_PDF_TABLE_HEADER_BG))
                canvas.setFont(font, 9)
                canvas.drawCentredString(self.width / 2, 25,
                    _rtl(f'پیشرفت کلی پروژه: {fa_number(f"{overall:.1f}")}٪'))
                canvas.restoreState()
        story.extend([Spacer(1, 12), ProjectProgressBar187()])
    project_name = str(project_name or export_project_name_187(specs))
    logo = _export_logo_image_187(background='#153B5B', height=180)
    logo_reader = ImageReader(logo) if logo is not None else None
    page_count = []
    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.saved_pages = []
        def showPage(self):
            self.saved_pages.append(dict(self.__dict__))
            self._startPage()
        def save(self):
            count = len(self.saved_pages)
            page_count.append(count)
            for saved in self.saved_pages:
                self.__dict__.update(saved)
                self.saveState()
                self.setFont(font, 8)
                self.setFillColor(colors.HexColor('#17212B'))
                self.drawRightString(width - margin, height - 14, _rtl(f'تعداد صفحات گزارش: {fa_number(count)}'))
                self.setFillColor(colors.HexColor('#F7FAFC'))
                self.setStrokeColor(colors.HexColor('#496176'))
                self.rect(margin, height - 100, available, 76, fill=1)
                section = available / 3
                for divider in (margin + section, margin + 2 * section):
                    self.line(divider, height - 100, divider, height - 24)
                self.setFillColor(colors.HexColor('#17212B'))
                self.setFont(font, 11)
                self.drawCentredString(margin + section * 1.5, height - 47, _rtl('شرکت مهندسی ساخت بارکاو'))
                self.setFont(font, 10)
                self.drawCentredString(margin + section * 1.5, height - 72, _rtl(f'پروژه {project_name}' if project_name != 'پروژه' else 'پروژه'))
                self.drawCentredString(margin + section * 2.5, height - 44, _rtl(title))
                self.setFont(font, 8)
                self.drawRightString(width - margin - 8, height - 62, _rtl('جهت ارائه به:'))
                self.drawCentredString(margin + section * 2.5, height - 82, _rtl(audience))
                if logo_reader:
                    self.setFillColor(colors.HexColor('#153B5B'))
                    self.rect(margin, height - 100, section, 76, fill=1)
                    self.drawImage(logo_reader, margin + 8, height - 87, width=section - 16, height=50, preserveAspectRatio=True, anchor='c', mask='auto')
                else:
                    self.drawCentredString(margin + section * .5, height - 62, 'BARKAV')
                self.drawCentredString(width / 2, 14, _rtl(f'{fa_number(self._pageNumber)} از {fa_number(count)}'))
                self.restoreState()
                super().showPage()
            super().save()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=page_size, leftMargin=margin, rightMargin=margin, topMargin=110, bottomMargin=28, title=title)
    document.build(story, canvasmaker=NumberedCanvas)
    pdf = output.getvalue()
    count = page_count[0]
    return ReportDocument187(pdf, tuple(headers), tuple(plain_rows), count, title)
