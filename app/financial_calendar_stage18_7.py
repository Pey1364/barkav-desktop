from __future__ import annotations

from dataclasses import dataclass
import os
from datetime import date, timedelta
import json
import tkinter as tk
from tkinter import filedialog, font as tkfont, messagebox, ttk
import logging
import hashlib
import re
import tempfile
import threading
from typing import Any
from pathlib import Path

import jdatetime

from ui_stage18_3 import en, fa
from barkav_ui import BarkavWindow
from barkav_ui.theme import font as kit_font, metrics as kit_metrics
from barkav_ui.table import apply_tree_rtl
from barkav_themes.tokens import mix
from barkav_themes.engine import stripe_tree
from export_utils import export_pdf, export_xlsx, print_file
from sales_claims_stage18_7 import (
    ShamsiDatePicker187, show_export_preview_187, export_receipts_pdf_187,
    style_xlsx_workbook,
)
from ui_stage18_7 import palette_187, theme_selection_187
from v18_store import load_ui, save_ui
from project_report_ai_187 import (
    build_project_prompt, generate_project_text, parse_invoice_response,
    parse_project_report_rows_187,
)


from calendar_report_tables_187 import (
    REPORT_FIELDS, ReportTable187, build_period_rows, phase_fields,
    prepare_answer_records, parse_grounded_answer, counting_days,
    clean_report_description, report_table, present_answer_table,
    is_report_fragment,
)

from calendar_semantic_report_187 import generate_semantic_report, semantic_records


from ai_section_models187 import MODEL_LABELS as AI_REPORT_ENGINE_LABELS_187, AIModelPicker187, section_choice
AI_REPORT_ENGINE_VALUES_187 = {label: key for key, label in AI_REPORT_ENGINE_LABELS_187.items()}


def _project_report_profile_187(service) -> dict[str, Any]:
    db_path = getattr(getattr(service, 'db', None), 'db_path', '')
    if not db_path:
        return {}
    try:
        path = Path(db_path).resolve().parent / 'project_profile_v18.json'
        value = json.loads(path.read_text(encoding='utf-8')) if path.is_file() else {}
        return value if isinstance(value, dict) else {}
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return {}


def _project_report_project_187(service) -> dict[str, Any]:
    db = getattr(service, 'db', None)
    if db is None:
        return {}
    try:
        with db.read() as con:
            row = con.execute('SELECT name, address FROM projects ORDER BY created_at LIMIT 1').fetchone()
        return dict(row) if row else {}
    except Exception:
        return {}


def _project_report_name_187(service) -> str:
    project = _project_report_project_187(service)
    return str(project.get('name') or os.environ.get('PEYMAAN_ACTIVE_PROJECT_NAME') or 'پروژه')


def _project_report_specs_187(label: str, start: str, end: str) -> list[tuple[str, str, str]]:
    first = '' if str(start or '').strip() in {'', '—'} else str(start).strip()
    last = '' if str(end or '').strip() in {'', '—'} else str(end).strip()
    report_date = first if first and first == last else f'{first} تا {last}'.strip(' تا')
    return [
        ('گزارش', 'عنوان گزارش', f'گزارش {label}'),
        ('گزارش', 'تاریخ گزارش', report_date),
    ]


def _project_report_audience_187(service, choice: str) -> str:
    profile = _project_report_profile_187(service)
    if choice == 'مجری':
        team = profile.get('team') or {}
        group = team.get('executor') or {}
        if isinstance(group, dict) and not group.get('name'):
            group = next(iter(group.values()), {})
        return f"مجری: {group.get('name') or 'نام ثبت نشده'}" if isinstance(group, dict) else 'مجری: نام ثبت نشده'
    if choice == 'ناظرها':
        output = []
        supervisors = (profile.get('team') or {}).get('supervisor') or {}
        if isinstance(supervisors, dict):
            for discipline, value in supervisors.items():
                if isinstance(value, dict) and str(value.get('name') or '').strip():
                    output.append(f'{discipline}: {value["name"]}')
        return 'ناظرها: ' + ('، '.join(output) if output else 'نامی ثبت نشده')
    # The addressee is a role label in this report; employer contact details
    # are stored in project factors and must not be repeated in the output.
    return 'کارفرما'


def _project_report_employer_specs_187(service, choice: str) -> list[tuple[str, str, str]]:
    if choice != 'کارفرما':
        return []
    employer = _project_report_profile_187(service).get('employer') or {}
    if not isinstance(employer, dict):
        return []
    return [('کارفرما', label, str(employer.get(key) or '—')) for key, label in (
        ('name', 'نام و نام خانوادگی / نام شرکت'), ('phone', 'شماره تماس'), ('address', 'آدرس'),
    )]


def _project_report_context_187(events, start='', end='') -> list[dict[str, Any]]:
    first, last = _normalize_shamsi(start), _normalize_shamsi(end)
    result = []
    for event in events or ():
        report_date = _report_date_from_text_187(event.title, event.details) if event.kind == 'DAILY_REPORT' else ''
        event_shamsi = _normalize_shamsi(report_date or event.due_date_shamsi)
        if first and event_shamsi and event_shamsi < first:
            continue
        if last and event_shamsi and event_shamsi > last:
            continue
        weekday = ''
        try:
            actual_date = _gregorian_from_shamsi_187(event_shamsi) if event_shamsi else event.due_date[:10]
            weekday = ('دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه', 'شنبه', 'یکشنبه')[
                date.fromisoformat(actual_date).weekday()
            ]
        except (TypeError, ValueError):
            pass
        metadata = event.metadata or {}
        source_fields = phase_fields(metadata, event.phase)
        for field, key in (('تعداد عوامل پیمانکار', 'report_contractor_count'),
                           ('تعداد عوامل روزمزد', 'report_daily_wage_count'),
                           ('ماشین‌آلات فعال', 'report_machinery'), ('وضعیت آب‌وهوا', 'report_weather')):
            source_fields[field] = metadata.get(key, metadata.get(field))
        result.append({
            'تاریخ': event_shamsi or event.due_date_shamsi,
            'روز هفته': weekday,
            'تاریخ ثبت': event.registered_date_shamsi,
            'دسته‌بندی': EVENT_LABELS.get(event.kind, event.kind),
            'شرح': event.title,
            'جزئیات': event.details,
            'مبلغ ثبت‌شده': event.amount if event.amount else None,
            'فاز': event.phase,
            **{key: value for key, value in metadata.items() if key in REPORT_FIELDS and key not in {'تاریخ', 'روز هفته'}},
            **source_fields,
        })
    return result


def _project_report_cost_context_187(service, start: str, end: str) -> list[dict[str, Any]]:
    """Read approved expenses for the requested dates as phase clues, never as report facts."""
    db = getattr(service, 'db', None)
    first, last = _normalize_shamsi(start), _normalize_shamsi(end)
    if db is None or not first or not last:
        return []
    try:
        with db.read() as con:
            rows = con.execute(
                "SELECT shamsi_date, description, construction_stage, category, final_amount "
                "FROM expenses WHERE shamsi_date>=? AND shamsi_date<=? "
                "AND COALESCE(status,'APPROVED') NOT IN ('CANCELLED','DRAFT') "
                "AND COALESCE(deleted_at,'')='' ORDER BY shamsi_date, created_at",
                (first, last),
            ).fetchall()
        return [{
            'تاریخ': str(row['shamsi_date'] or ''),
            'دسته‌بندی': 'هزینه ثبت‌شده — فقط قرینه برای فاز',
            'شرح': str(row['description'] or ''),
            'فاز ثبت‌شده': str(row['construction_stage'] or ''),
            'مبلغ': row['final_amount'],
            'منبع': 'حواله/هزینه ثبت‌شده؛ به‌تنهایی اثبات‌کننده انجام فعالیت نیست',
        } for row in rows]
    except Exception:
        logging.getLogger(__name__).exception('Project report cost context could not be read')
        return []


def _daily_report_image_paths_187(owner, source_id: str) -> list[str]:
    service = getattr(owner, 'document_service', None) or getattr(owner, 'service', None)
    paths = []
    provider = getattr(service, 'list_documents', None)
    if callable(provider):
        try:
            for row in provider() or ():
                if str(row.get('document_id') or '') == source_id and row.get('local_path'):
                    paths.append(str(row['local_path']))
        except Exception:
            logging.getLogger(__name__).exception('Daily report attachments could not be listed')
    db = getattr(getattr(owner, 'service', None), 'db', None)
    if db is not None:
        try:
            with db.read() as con:
                rows = con.execute(
                    "SELECT local_path FROM attachments WHERE entity_type='DOCUMENT' AND entity_id=?",
                    (source_id,),
                ).fetchall()
            paths.extend(str(row['local_path'] or '') for row in rows)
        except Exception:
            logging.getLogger(__name__).exception('Daily report attachment links could not be read')
    roots = []
    db_path = str(getattr(db, 'db_path', '') or '')
    if db_path:
        root = Path(db_path).resolve().parent
        roots.extend((root, root / 'receipt_archive'))
    resolved = []
    for raw in paths:
        path = Path(raw)
        if path.is_file():
            resolved.append(str(path)); continue
        for root in roots:
            candidate = root / path.name
            if candidate.is_file():
                resolved.append(str(candidate)); break
    image_exts = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.tif', '.tiff'}
    return list(dict.fromkeys(path for path in resolved if Path(path).suffix.lower() in image_exts))


def _image_signature_187(path: str) -> bytes:
    try:
        from PIL import Image
        with Image.open(path) as source:
            image = source.convert('L').resize((8, 8))
            pixels = list(image.getdata())
        average = sum(pixels) / max(1, len(pixels))
        return bytes(1 if value >= average else 0 for value in pixels)
    except (ImportError, OSError, ValueError):
        try:
            return hashlib.sha256(Path(path).read_bytes()).digest()
        except OSError:
            return path.encode('utf-8', errors='ignore')


def _visually_duplicate_187(left: bytes, right: bytes) -> bool:
    if len(left) != len(right):
        return left == right
    return sum(a != b for a, b in zip(left, right)) <= 5


@dataclass(frozen=True)
class FinancialEvent187:
    due_date: str
    due_date_shamsi: str
    kind: str
    title: str
    amount: int
    direction: str
    details: str = ''
    registered_date_shamsi: str = ''
    alert_enabled: bool = True
    source_id: str = ''
    phase: str = ''
    metadata: dict[str, Any] | None = None


EVENT_LABELS = {
    'CHEQUE': 'سررسید چک هزینه',
    'RECEIVABLE': 'سررسید مطالبات',
    'LOAN': 'قسط تسهیلات',
    'WORKSHOP_CALENDAR': 'تقویم کارگاهی',
    'DAILY_REPORT': 'گزارش روزانه',
    'DOCUMENT': 'داکیومنت',
    'REMINDER': 'یادآوری',
    'NOTE': 'یادداشت',
    'INFO': 'اطلاع‌رسانی',
    # A note about an invoice - distinct from the AI's own document_type=
    # INVOICE, which classifies an attached photo/scan, not this calendar
    # event's own kind.
    'INVOICE_NOTE': 'فاکتور',
}

CALENDAR_EVENT_FILTERS_187 = {
    'همه': None,
    'چک': 'CHEQUE',
    'مطالبات': 'RECEIVABLE',
    'تسهیلات': 'LOAN',
    'گزارش روزانه': 'DAILY_REPORT',
    'تقویم کارگاهی': 'WORKSHOP_CALENDAR',
    'یادداشت': 'NOTE',
    'داکیومنت': 'DOCUMENT',
    'اطلاع‌رسانی': 'INFO',
    'یادآوری': 'REMINDER',
    'فاکتور': 'INVOICE_NOTE',
}


def filter_calendar_events_187(events, kind_label='همه', from_shamsi='', to_shamsi=''):
    """Apply the project-calendar kind and inclusive Solar-Hijri date range."""
    kind = CALENDAR_EVENT_FILTERS_187.get(str(kind_label or 'همه'), None)
    start = _normalize_shamsi(from_shamsi)
    end = _normalize_shamsi(to_shamsi)
    if start and end and start > end:
        start, end = end, start
    return [event for event in events if (
        (kind is None or event.kind == kind)
        and (not start or event.due_date_shamsi >= start)
        and (not end or event.due_date_shamsi <= end)
    )]

ALERT_PREFS_KEY_187 = 'financial_message_alerts_187'
ALERT_DEFAULTS_187 = {'CHEQUE': True, 'RECEIVABLE': True, 'LOAN': True, 'REMINDER': True}


def _alert_preferences_187() -> dict[str, bool]:
    try:
        raw = load_ui().get(ALERT_PREFS_KEY_187) or {}
        return {str(key): bool(value) for key, value in dict(raw).items()}
    except (AttributeError, TypeError, ValueError):
        return {}


def _alert_enabled_187(kind: str, preferences: dict[str, bool] | None = None) -> bool:
    prefs = preferences if preferences is not None else _alert_preferences_187()
    return bool(prefs.get(kind, ALERT_DEFAULTS_187.get(kind, False)))


def is_startup_alert_due_187(event: FinancialEvent187, today: date | None = None, preferences: dict[str, bool] | None = None,
                             hidden_events: set[str] | None = None) -> bool:
    """Return whether an event may open the startup alert window.

    Reminders wait for their selected reminder date; other due alerts retain
    their established startup behaviour.
    """
    if not _event_alert_enabled_187(event, preferences):
        return False
    hidden = startup_hidden_events_187() if hidden_events is None else hidden_events
    if _message_state_key_187(event) in hidden:
        return False
    if event.kind != 'REMINDER':
        return True
    current = today or date.today()
    try:
        return date.fromisoformat(str(event.due_date)[:10]) == current
    except (TypeError, ValueError):
        return False

ROW_ALERT_PREFS_KEY_187 = 'financial_message_row_alerts_187'
STARTUP_HIDDEN_EVENTS_KEY_187 = 'financial_message_startup_hidden_187'
READ_EVENTS_KEY_187 = 'financial_message_read_187'
_DUE_DATE_EVENT_KINDS_187 = frozenset({'CHEQUE', 'RECEIVABLE', 'LOAN', 'REMINDER'})


def _event_alert_key_187(event: FinancialEvent187) -> str:
    return '\x1f'.join((str(event.kind or ''), str(event.due_date or ''), str(event.title or ''), str(event.details or ''), str(event.registered_date_shamsi or '')))


def _message_state_key_187(event: FinancialEvent187) -> str:
    """Stable per-message key for startup hiding and the footer unread badge."""
    return '\x1f'.join((
        str(event.kind or ''), str(event.source_id or ''), str(event.due_date or ''),
        str(event.title or ''), str(event.details or ''), str(event.registered_date_shamsi or ''),
        str(event.amount or 0), str(event.direction or ''), str(event.phase or ''),
    ))


def _saved_message_keys_187(setting_key: str) -> set[str]:
    try:
        raw = load_ui().get(setting_key) or {}
        return {str(key) for key, value in dict(raw).items() if bool(value)}
    except (AttributeError, TypeError, ValueError):
        return set()


def _save_message_keys_187(setting_key: str, keys) -> None:
    settings = dict(load_ui())
    settings[setting_key] = {str(key): True for key in set(keys)}
    save_ui(settings)


def startup_hidden_events_187() -> set[str]:
    return _saved_message_keys_187(STARTUP_HIDDEN_EVENTS_KEY_187)


def hide_startup_events_187(events) -> None:
    keys = startup_hidden_events_187()
    keys.update(_message_state_key_187(event) for event in events)
    _save_message_keys_187(STARTUP_HIDDEN_EVENTS_KEY_187, keys)


def unread_message_events_187(events) -> list[FinancialEvent187]:
    """Return alertable messages not yet opened in this project."""
    hidden = startup_hidden_events_187()
    read = _saved_message_keys_187(READ_EVENTS_KEY_187)
    return [
        event for event in events
        if _event_alert_enabled_187(event)
        and _message_state_key_187(event) not in hidden
        and _message_state_key_187(event) not in read
    ]


def mark_messages_read_187(events) -> None:
    keys = _saved_message_keys_187(READ_EVENTS_KEY_187)
    keys.update(_message_state_key_187(event) for event in events)
    _save_message_keys_187(READ_EVENTS_KEY_187, keys)


def _row_alert_preferences_187() -> dict[str, bool]:
    try:
        raw = load_ui().get(ROW_ALERT_PREFS_KEY_187) or {}
        return {str(key): bool(value) for key, value in dict(raw).items()}
    except (AttributeError, TypeError, ValueError):
        return {}


def _event_alert_enabled_187(event: FinancialEvent187, preferences: dict[str, bool] | None = None) -> bool:
    prefs = preferences if preferences is not None else _row_alert_preferences_187()
    # Installments are actionable financial reminders on first load. A saved
    # row preference still wins after the user explicitly toggles it.
    default = event.kind in {'RECEIVABLE', 'LOAN'}
    return bool(prefs.get(_event_alert_key_187(event), default))


def _set_event_alert_187(event: FinancialEvent187, enabled: bool) -> None:
    settings = dict(load_ui())
    prefs = _row_alert_preferences_187()
    prefs[_event_alert_key_187(event)] = bool(enabled)
    settings[ROW_ALERT_PREFS_KEY_187] = prefs
    save_ui(settings)


def _display_due_date_187(event: FinancialEvent187) -> str:
    return fa(event.due_date_shamsi) if event.kind in _DUE_DATE_EVENT_KINDS_187 else '—'

def _set_alert_preference_187(kind: str, enabled: bool) -> None:
    settings = dict(load_ui())
    prefs = _alert_preferences_187()
    prefs[str(kind)] = bool(enabled)
    settings[ALERT_PREFS_KEY_187] = prefs
    save_ui(settings)


def _registered_shamsi_187(value: Any, fallback: str = '') -> str:
    raw = str(value or '').strip()
    if '/' in raw and len(raw.split('/')) >= 3:
        return _normalize_shamsi(raw.split()[0])
    return _shamsi_from_gregorian(raw[:10]) or _normalize_shamsi(fallback)

# Stable Solar Hijri public holidays.  Fridays are included as the recurring
# official weekly holiday; lunar holidays are intentionally not guessed.
_FIXED_OFFICIAL_HOLIDAYS_187 = frozenset({
    (1, 1), (1, 2), (1, 3), (1, 4), (1, 12), (1, 13),
    (3, 14), (3, 15), (11, 22), (12, 29),
})


def _is_official_holiday_187(year: int, month: int, day: int) -> bool:
    if (int(month), int(day)) in _FIXED_OFFICIAL_HOLIDAYS_187:
        return True
    try:
        # Python weekday 4 is Friday, the weekly official holiday.
        return jdatetime.date(int(year), int(month), int(day)).togregorian().weekday() == 4
    except (TypeError, ValueError):
        return False


def _shamsi_from_gregorian(value: str) -> str:
    try:
        converted = jdatetime.date.fromgregorian(date=date.fromisoformat(str(value)[:10]))
        return f'{converted.year:04d}/{converted.month:02d}/{converted.day:02d}'
    except Exception:
        return ''


def _normalize_shamsi(value: Any, fallback_gregorian: str = '') -> str:
    raw = en(str(value or '')).replace('-', '/').strip()
    parts = raw.split('/')
    if len(parts) == 3:
        try:
            return f'{int(parts[0]):04d}/{int(parts[1]):02d}/{int(parts[2]):02d}'
        except Exception:
            pass
    return _shamsi_from_gregorian(fallback_gregorian)


def _report_date_keys_187(value: Any) -> list[str]:
    normalized = en(str(value or '')).replace('-', '/')
    return list(dict.fromkeys(
        _normalize_shamsi(match.group(0))
        for match in re.finditer(r'(?<!\d)\d{4}/\d{1,2}/\d{1,2}(?!\d)', normalized)
    ))


def _attach_report_documents_187(ai_rows, source_rows) -> list[dict[str, Any]]:
    by_date: dict[str, list[str]] = {}
    for source in source_rows or ():
        paths = source.get('_document_paths') or ()
        for report_date in _report_date_keys_187(source.get('تاریخ')):
            existing = by_date.setdefault(report_date, [])
            existing.extend(path for path in paths if path not in existing)
    used: set[str] = set()
    output = []
    for row in ai_rows or ():
        item = dict(row)
        paths = []
        for report_date in _report_date_keys_187(item.get('تاریخ')):
            for path in by_date.get(report_date, ()):
                if path not in used:
                    paths.append(path)
                    used.add(path)
        item['_document_paths'] = paths
        output.append(item)
    return output


def _gregorian_from_shamsi_187(value: str) -> str:
    try:
        parts = [int(part) for part in _normalize_shamsi(value).split('/')]
        if len(parts) != 3:
            return ''
        return jdatetime.date(*parts).togregorian().isoformat()
    except (TypeError, ValueError):
        return ''


def _period_bounds_187(year: int, month: int, day: int, period: str) -> tuple[date, date]:
    """Return Gregorian bounds for a selected Persian day/week/month."""
    current = jdatetime.date(int(year), int(month), int(day)).togregorian()
    if period == 'month':
        start = current.replace(day=1)
        if month == 12:
            next_month = jdatetime.date(int(year) + 1, 1, 1).togregorian()
        else:
            next_month = jdatetime.date(int(year), int(month) + 1, 1).togregorian()
        return start, next_month - timedelta(days=1)
    # The grid starts on Saturday.  Convert Python's Monday-based weekday to
    # the same Saturday-first week used by the Persian calendar UI.
    start = current - timedelta(days=(current.weekday() + 2) % 7)
    return start, start + timedelta(days=6)


_REPORT_STATUS_RE_187 = re.compile(r'\s*(?:شروع|آغاز|ادامه(?:\s+یافت)?|در\s+حال\s+اجرا|تکمیل(?:\s+عملیات)?|کامل(?:\s+شد)?|اتمام)\s*', re.I)


def _report_activity_key_187(text: str) -> str:
    cleaned = _REPORT_STATUS_RE_187.sub(' ', str(text or ''))
    return ' '.join(cleaned.split()).strip(' -،:')


def summarize_daily_reports_187(
    events: list[FinancialEvent187], start: date, end: date,
) -> list[str]:
    """Summarize daily-report text once per activity in an inclusive range.

    Completion wins over start/continued entries for the same activity.  This
    keeps weekly/monthly reports concise while preserving the user's wording.
    """
    grouped: dict[str, dict[str, Any]] = {}
    for event in events:
        if event.kind != 'DAILY_REPORT':
            continue
        try:
            event_date = date.fromisoformat(event.due_date[:10])
        except (TypeError, ValueError):
            continue
        if not start <= event_date <= end:
            continue
        text = ' '.join(str(event.details or event.title or '').split())
        if not text:
            continue
        key = _report_activity_key_187(text) or text
        record = grouped.setdefault(key, {'text': text, 'completed': False})
        if re.search(r'تکمیل|کامل|اتمام', text, re.I):
            record['completed'] = True
            record['text'] = text
    output: list[str] = []
    for key in sorted(grouped):
        record = grouped[key]
        if record['completed']:
            output.append(f'تکمیل عملیات {key}')
        else:
            output.append(str(record['text']))
    return output




_REPORT_DATE_NUMERIC_RE_187 = re.compile(r'(?<![\d۰-۹])([1۱][3۳4۴][\d۰-۹]{2})\s*[/\-.]\s*([\d۰-۹]{1,2})\s*[/\-.]\s*([\d۰-۹]{1,2})(?![\d۰-۹])')
_REPORT_DATE_WORD_RE_187 = re.compile(
    r'(?<![\d۰-۹])([\d۰-۹]{1,2})\s+(فروردین|اردیبهشت|خرداد|تیر|مرداد|شهریور|مهر|آبان|آذر|دی|بهمن|اسفند)\s+([1۱][3۳4۴][\d۰-۹]{2})(?![\d۰-۹])'
)
_REPORT_MONTHS_187 = {name: index for index, name in enumerate((
    '', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور',
    'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند',
)) if name}


def _report_date_from_text_187(*values: Any) -> str:
    """Find a Solar-Hijri date written in a daily-report sentence."""
    source = ' '.join(str(value or '') for value in values)
    source = en(source).replace('٫', '.').replace('／', '/')
    match = _REPORT_DATE_NUMERIC_RE_187.search(source)
    if match:
        return _normalize_shamsi('/'.join(match.groups()))
    match = _REPORT_DATE_WORD_RE_187.search(source)
    if match:
        day, month, year = match.groups()
        return _normalize_shamsi(f'{year}/{_REPORT_MONTHS_187[month]}/{day}')
    return ''


def _strip_report_date_187(value: Any) -> str:
    return clean_report_description(value)


def _phase_from_values_187(*values: Any) -> str:
    source = ' '.join(str(value or '') for value in values)
    match = re.search(r'(?:فاز|مرحله(?:\s+ساخت)?)\s*[:：\-]?\s*([^\n،؛|]+)', source)
    return ' '.join(match.group(1).split()).strip() if match else ''


def _report_phase_187(event: FinancialEvent187) -> str:
    """Extract the user-entered construction phase without inventing one."""
    return str(event.phase or _phase_from_values_187(event.title, event.details) or '—')


def _daily_report_entries_187(events, start, end):
    """Keep each source activity's own execution day for semantic grouping."""
    entries = []
    for event in events:
        if event.kind != 'DAILY_REPORT':
            continue
        entered = _report_date_from_text_187(event.title, event.details)
        try:
            actual_date = _gregorian_from_shamsi_187(entered) if entered else str(event.due_date)[:10]
            event_day = date.fromisoformat(actual_date)
        except (TypeError, ValueError):
            continue
        if not start <= event_day <= end:
            continue
        meta = event.metadata or {}
        original = next((cleaned for value in (meta.get('row_description'), event.title, event.details)
                         if (cleaned := _strip_report_date_187(value))), '')
        if not original:
            continue
        activities = [
            part.strip(' •-') for part in re.split(r'[\n؛;،]+', original)
            if part.strip(' •-') and not is_report_fragment(part)
        ]
        for activity in activities:
            values = {
                '_day': event_day, 'activity': activity,
                'activity_key': _report_activity_key_187(activity),
                '_legacy_phase': str(meta.get('main_phase') or _report_phase_187(event)),
                'details': _strip_report_date_187(event.details),
                '_source_id': str(event.source_id or ''),
                **phase_fields(meta, _report_phase_187(event)),
            }
            for field, key in (
                ('تعداد عوامل پیمانکار', 'report_contractor_count'),
                ('تعداد عوامل روزمزد', 'report_daily_wage_count'),
                ('ماشین‌آلات فعال', 'report_machinery'), ('وضعیت آب‌وهوا', 'report_weather'),
            ):
                value = meta.get(key)
                values[field] = '' if value is None else str(value)
            entries.append(values)
    return entries


def daily_report_rows_187(
    events: list[FinancialEvent187], start: date, end: date,
    *, deduplicate: bool = True,
) -> list[dict[str, Any]]:
    return build_period_rows(_daily_report_entries_187(events, start, end), start, end,
                             daily=not deduplicate)


def _message_document_kind_187(value: Any, title: Any = '') -> str:
    """Normalize current and legacy document type labels for message filters."""
    raw = str(value or '').strip().upper()
    aliases = {
        'NOTES': 'NOTE', 'یادداشت': 'NOTE',
        'REMIND': 'REMINDER', 'یادآوری': 'REMINDER', 'یاداوری': 'REMINDER',
        'NOTIFICATION': 'INFO', 'اطلاع رسانی': 'INFO', 'اطلاع‌رسانی': 'INFO',
        'WORKSHOP': 'WORKSHOP_CALENDAR', 'تقویم کارگاهی': 'WORKSHOP_CALENDAR',
        'DAILY': 'DAILY_REPORT', 'گزارش روزانه': 'DAILY_REPORT',
        'DOC': 'DOCUMENT', 'داکیومنت': 'DOCUMENT',
        'OTHER': 'DOCUMENT',
        'فاکتور': 'INVOICE_NOTE',
    }
    kind = aliases.get(raw, raw)
    text = str(title or '')
    if 'یادداشت' in text:
        return 'NOTE'
    if 'یاداوری' in text or 'یادآوری' in text:
        return 'REMINDER'
    if 'اطلاع' in text:
        return 'INFO'
    if kind in EVENT_LABELS:
        return kind
    return 'DOCUMENT'


def _document_events_from_rows_187(rows) -> list[FinancialEvent187]:
    """Convert the shared non-financial-document rows into calendar events."""
    events = []
    for raw_row in rows or ():
        try:
            row = dict(raw_row)
        except (TypeError, ValueError):
            continue
        if str(row.get('status') or 'ACTIVE').upper() == 'DELETED':
            continue
        created_at = row.get('created_at') or ''
        registered = _registered_shamsi_187(created_at, '')
        shamsi = _normalize_shamsi(row.get('event_date_shamsi'))
        shamsi = shamsi or registered
        due = _gregorian_from_shamsi_187(shamsi)
        if not due:
            due = date.today().isoformat()
            shamsi = shamsi or _shamsi_from_gregorian(due)

        raw_kind = row.get('document_type') or ''
        raw_title = row.get('description') or ''
        kind = _message_document_kind_187(raw_kind, raw_title)
        details = str(row.get('raw_text') or '')
        title = str(row.get('description') or EVENT_LABELS[kind])
        phase = str(
            row.get('construction_stage') or row.get('main_phase')
            or row.get('phase_name') or _phase_from_values_187(title, details)
        ).strip()
        try:
            amount = int(row.get('amount') or 0)
        except (TypeError, ValueError):
            amount = 0
        report_rows = row.get('daily_report_rows') or row.get('report_rows') or []
        if isinstance(report_rows, str):
            try:
                report_rows = json.loads(report_rows)
            except (TypeError, ValueError, json.JSONDecodeError):
                report_rows = []
        if kind == 'DAILY_REPORT' and isinstance(report_rows, list) and report_rows:
            for index, report_row in enumerate(report_rows, 1):
                if isinstance(report_row, dict):
                    row_title = str(report_row.get('description') or report_row.get('شرح') or '').strip()
                else:
                    row_title = str(report_row or '').strip()
                if not row_title:
                    continue
                events.append(FinancialEvent187(
                    due, shamsi, kind, row_title, amount, 'INFO', details,
                    registered, _alert_enabled_187(kind),
                    str(row.get('document_id') or ''),
                    str(report_row.get('main_phase') or report_row.get('فاز اصلی') or phase),
                    {
                        'main_phase': str(report_row.get('main_phase') or report_row.get('فاز اصلی') or phase),
                        'sub_phase': str(report_row.get('sub_phase') or report_row.get('فاز فرعی') or ''),
                        'main_phase_no': report_row.get('main_phase_no') or report_row.get('شماره فاز اصلی'),
                        'sub_phase_no': report_row.get('sub_phase_no') or report_row.get('شماره فاز فرعی'),
                        'report_contractor_count': row.get('report_contractor_count'),
                        'report_daily_wage_count': row.get('report_daily_wage_count'),
                        'report_machinery': row.get('report_machinery'),
                        'report_weather': row.get('report_weather'),
                        'row_description': row_title,
                    },
                ))
        else:
            events.append(FinancialEvent187(
                due, shamsi, kind, title, amount, 'INFO', details,
                registered, _alert_enabled_187(kind),
                str(row.get('document_id') or ''), phase,
            ))
    return events


def _document_events_187(service) -> list[FinancialEvent187]:
    """Read every non-deleted document, including legacy document types.

    The messages panel uses the same document provider as the non-financial
    documents tab when one is supplied. This keeps both screens on the same
    project/database and preserves rows that the tab already displays.
    """
    provider = getattr(service, 'list_documents', None)
    if callable(provider):
        try:
            listed = provider()
        except Exception:
            listed = None
        if listed:
            db = getattr(service, 'db', None)
            if db is not None and hasattr(db, 'get_pending_item'):
                enriched = []
                for item in listed:
                    row = dict(item)
                    pending_id = str(row.get('pending_id') or '').strip()
                    if pending_id:
                        try:
                            payload = dict((db.get_pending_item(pending_id) or {}).get('payload') or {})
                            row.setdefault('construction_stage', payload.get('construction_stage'))
                            row.setdefault('main_phase', payload.get('main_phase'))
                            row.setdefault('phase_name', payload.get('phase_name'))
                            row.setdefault('daily_report_rows', payload.get('daily_report_rows') or payload.get('report_rows'))
                            for key in ('report_contractor_count', 'report_daily_wage_count', 'report_machinery', 'report_weather'):
                                row.setdefault(key, payload.get(key))
                        except Exception:
                            pass
                    enriched.append(row)
                listed = enriched
            converted = _document_events_from_rows_187(listed)
            if converted:
                return converted

    db = getattr(service, 'db', None)
    if db is None:
        return []
    queries = (
        "SELECT document_id,document_type,event_date_shamsi,description,amount,created_at,raw_text "
        "FROM financial_documents WHERE COALESCE(status,'ACTIVE') != 'DELETED'",
        "SELECT document_id,document_type,event_date_shamsi,description,amount,created_at "
        "FROM financial_documents WHERE COALESCE(status,'ACTIVE') != 'DELETED'",
        "SELECT document_type,event_date_shamsi,description,amount "
        "FROM financial_documents WHERE COALESCE(status,'ACTIVE') != 'DELETED'",
        "SELECT document_type,event_date_shamsi,description,amount "
        "FROM financial_documents",
    )
    rows = None
    try:
        with db.read() as con:
            for query in queries:
                try:
                    rows = con.execute(query).fetchall()
                    break
                except Exception:
                    continue
    except Exception:
        return []
    if rows is None:
        return []
    return _document_events_from_rows_187(rows)


def collect_financial_events_187(service, days: int = 370, *, document_service=None) -> list[FinancialEvent187]:
    """Build one read-only financial event stream from existing SQLite services."""
    events: list[FinancialEvent187] = []
    preferences = _alert_preferences_187()
    for row in service.upcoming_cheques(days):
        due = str(row.get('due_date') or '')[:10]
        events.append(FinancialEvent187(
            due, _normalize_shamsi(row.get('due_date_shamsi'), due), 'CHEQUE',
            str(row.get('description') or row.get('expense_description') or 'سررسید چک'),
            int(row.get('amount') or 0), 'OUT',
            str(row.get('receiver_name') or row.get('issuer_bank') or ''),
            _registered_shamsi_187(row.get('created_at'), row.get('due_date_shamsi') or due),
            _alert_enabled_187('CHEQUE', preferences),
        ))
    buyer_rows = list(service.upcoming_buyer_installments(days))
    try:
        buyer_rows.extend(service.overdue_buyer_installments())
    except Exception:
        pass
    seen_buyer: set[str] = set()
    for row in buyer_rows:
        row_id = str(row.get('installment_id') or '')
        if row_id and row_id in seen_buyer:
            continue
        seen_buyer.add(row_id)
        due = str(row.get('due_date') or '')[:10]
        events.append(FinancialEvent187(
            due, _normalize_shamsi(row.get('due_date_shamsi'), due), 'RECEIVABLE',
            f"وصول واحد {row.get('unit_number') or '-'}",
            int(row.get('outstanding') or 0), 'IN',
            str(row.get('buyer_name') or ''),
            _registered_shamsi_187(row.get('created_at'), row.get('due_date_shamsi') or due),
            _alert_enabled_187('RECEIVABLE', preferences),
        ))
    for row in service.upcoming_loan_installments(days):
        due = str(row.get('due_date') or '')[:10]
        events.append(FinancialEvent187(
            due, _normalize_shamsi(row.get('due_date_shamsi'), due), 'LOAN',
            str(row.get('loan_title') or 'قسط تسهیلات'),
            int(row.get('outstanding') or 0), 'OUT',
            str(row.get('lender_name') or ''),
            _registered_shamsi_187(row.get('created_at'), row.get('due_date_shamsi') or due),
            _alert_enabled_187('LOAN', preferences),
        ))
    sources = [service if document_service is None else document_service]
    if document_service is not None and document_service is not service:
        sources.append(service)
    for source in sources:
        events.extend(_document_events_187(source))
    unique = {}
    for event in events:
        key = (event.due_date, event.kind, event.title, event.source_id)
        unique[key] = event
    events = list(unique.values())
    return sorted(events, key=lambda item: (item.due_date, item.kind, item.title))


def upcoming_events_187(service, days: int = 30) -> list[FinancialEvent187]:
    today = date.today()
    output = []
    for event in collect_financial_events_187(service, max(31, days)):
        try:
            delta = (date.fromisoformat(event.due_date) - today).days
        except Exception:
            continue
        if 0 <= delta <= days:
            output.append(event)
    return output


def message_events_187(
    service, days: int = 30, *, document_service=None,
) -> list[FinancialEvent187]:
    """Return financial due items plus documents from the shared document tab.

    service remains the financial-event source. When provided,
    document_service is the same RichBudgetService used by the
    non-financial documents table, so the messages window cannot drift to a
    different or incomplete document source.
    """
    document_sources = [service if document_service is None else document_service]
    if document_service is not None and document_service is not service:
        document_sources.append(service)
    # Keep document-backed messages visible when a financial provider is unavailable.
    # This path is read-only; each document source is still loaded independently below.
    try:
        events = [*upcoming_events_187(service, days)]
    except Exception:
        events = []
    for source in document_sources:
        events.extend(_document_events_187(source))
    unique: dict[tuple[str, str, str, str, str], FinancialEvent187] = {}
    for event in events:
        key = (event.due_date, event.kind, event.title, event.details, event.registered_date_shamsi)
        unique[key] = event
    return sorted(unique.values(), key=lambda item: (item.due_date, item.kind, item.title))


def _format_amount(value: int) -> str:
    return fa(f'{int(value or 0):,}')


def message_export_indices_187(visible_indices, selected_indices) -> list[int]:
    """Export checked visible rows, or every visible row when nothing is checked."""
    visible = [int(index) for index in visible_indices]
    selected = {int(index) for index in selected_indices}
    checked = [index for index in visible if index in selected]
    return checked or visible


class A6DocumentPreview187(BarkavWindow):
    """Document preview constrained to an A6 portrait page."""

    size_class = 'small'

    def __init__(self, parent, path: str | Path):
        super().__init__(parent, 'نمایش سند', subtitle='پیش‌نمایش A6', show_footer=False)
        self.geometry('420x595')
        self.minsize(360, 500)
        self._image = None
        from PIL import Image, ImageTk

        image = Image.open(path).convert('RGB')
        image.thumbnail((380, 525))
        self._image = ImageTk.PhotoImage(image)
        label = tk.Label(self.body, image=self._image, bg=palette_187()['bg'])
        label.pack(fill='both', expand=True, padx=12, pady=12)


class StartupMessagesPopup187(BarkavWindow):
    """Small blocking notice shown before the dashboard when alerts are due."""

    size_class = 'small'

    def __init__(self, parent, events: list[FinancialEvent187], *, on_hidden=None):
        self.events = list(events or ())
        self.hidden_indices: set[int] = set()
        self._on_hidden_187 = on_hidden
        super().__init__(parent, 'پیام‌های جدید', subtitle='پیام‌های پروژه', show_footer=False)
        self.transient(parent)
        self.geometry('520x330')
        self.minsize(420, 260)
        self.resizable(True, True)
        p = palette_187()
        close_bar = tk.Frame(self.body, bg=p['surface'])
        close_bar.pack(fill='x', padx=10, pady=(8, 0))
        # Keep an explicit close action in the content as well as the native
        # caption control.  This remains available when Windows dims or hides
        # the title-bar controls while the modal owns focus.
        tk.Button(
            close_bar, text='✕ بستن', command=self._close_modal_187,
            bg=p['panel2'], fg=p['text'], activebackground=p['blue'],
            activeforeground=p['text'], relief='flat', cursor='hand2',
            padx=12, pady=4,
        ).pack(side='right')
        tk.Button(
            close_bar, text='عدم نمایش مجدد', command=self._hide_selected_187,
            bg=p['panel2'], fg=p['text'], activebackground=p['blue'],
            activeforeground=p['text'], relief='flat', cursor='hand2',
            padx=12, pady=4,
        ).pack(side='right', padx=(0, 6))
        tk.Label(
            self.body, text=f'{fa(len(self.events))} پیام جدید برای بررسی وجود دارد.',
            bg=p['surface'], fg=p['text'], anchor='e', justify='right',
            font=kit_font('body', bold=True), padx=12, pady=10,
        ).pack(fill='x')
        table_host = tk.Frame(self.body, bg=p['panel'])
        table_host.pack(fill='both', expand=True, padx=10, pady=(0, 8))
        self.tree = ttk.Treeview(
            table_host, columns=('hide', 'kind', 'title', 'date'), show='headings',
            height=7, style='Barkav.Data.Treeview',
        )
        for key, label, width in (
            ('hide', 'عدم نمایش مجدد', 135), ('kind', 'نوع', 130),
            ('title', 'شرح', 260), ('date', 'تاریخ', 105),
        ):
            self.tree.heading(key, text=label, anchor='e')
            self.tree.column(key, width=width, minwidth=70, anchor='e', stretch=True)
        apply_tree_rtl(self.tree)
        scroll = ttk.Scrollbar(table_host, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side='right', fill='y')
        self.tree.pack(side='left', fill='both', expand=True)
        self.tree.bind('<Button-1>', self._toggle_hide_187, add='+')
        self._render_events_187()
        ttk.Button(
            self.body, text='بستن و ورود به داشبورد', style='Primary.TButton',
            command=self._close_modal_187,
        ).pack(anchor='e', padx=10, pady=(0, 10))
        self.protocol('WM_DELETE_WINDOW', self._close_modal_187)
        self.after_idle(self._activate_modal_187)

    def _activate_modal_187(self):
        try:
            self.grab_set()
            self.focus_force()
        except tk.TclError:
            pass

    def _close_modal_187(self):
        try:
            self.grab_release()
        except tk.TclError:
            pass
        self.destroy()

    def _render_events_187(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for index, event in enumerate(self.events):
            self.tree.insert('', 'end', iid=str(index), values=(
                '☑' if index in self.hidden_indices else '☐',
                EVENT_LABELS.get(event.kind, event.kind), event.title or '—',
                fa(event.due_date_shamsi or event.registered_date_shamsi or '—'),
            ))

    def _toggle_hide_187(self, event):
        row = self.tree.identify_row(event.y)
        column = self.tree.identify_column(event.x)
        try:
            position = int(column[1:]) - 1
            key = tuple(self.tree.cget('displaycolumns'))[position]
        except (IndexError, TypeError, ValueError, tk.TclError):
            key = ''
        if not row or key != 'hide':
            return
        try:
            index = int(row)
        except ValueError:
            return 'break'
        if index in self.hidden_indices:
            self.hidden_indices.remove(index)
        else:
            self.hidden_indices.add(index)
        self._render_events_187()
        return 'break'

    def _hide_selected_187(self):
        selected = [self.events[index] for index in sorted(self.hidden_indices)]
        if not selected:
            messagebox.showinfo('عدم نمایش مجدد', 'ابتدا پیام‌های موردنظر را تیک بزنید.', parent=self)
            return
        hide_startup_events_187(selected)
        if callable(self._on_hidden_187):
            self._on_hidden_187()
        self._close_modal_187()


class FinancialMessagesDialog187(BarkavWindow):
    size_class = 'medium'

    def __init__(self, parent, service, *, startup: bool = False,
                 document_service=None, on_open_calendar=None, on_messages_opened=None):
        self._calendar_parent_187 = parent
        self.service = service
        self.document_service = document_service
        self._on_open_calendar_187 = on_open_calendar
        self._on_messages_opened_187 = on_messages_opened
        self.events = message_events_187(
            service, 30, document_service=document_service,
        )
        mark_messages_read_187(self.events)
        if callable(self._on_messages_opened_187):
            self._on_messages_opened_187()
        self.alert_preferences = _row_alert_preferences_187()
        self.selected_indices: set[int] = set()
        self.visible_indices: list[int] = []
        self._document_file_index_187 = self._build_document_file_index_187()
        self.document_images = self._document_image_paths_187()
        title = 'یادآوری شروع برنامه' if startup else 'پیام‌ها'
        super().__init__(parent, title, subtitle='رویدادهای پروژه', show_footer=False)
        self.search_placeholder187 = 'جستجو'
        self.search_var = tk.StringVar(value=self.search_placeholder187)
        self.kind_filter_var = tk.StringVar(value='همه')
        self.alert_filter_var = tk.StringVar(value='همه')
        self.due_filter_var = tk.StringVar(value='همه')
        self.export_var = tk.StringVar(value='انتخاب فرمت خروجی')
        p = palette_187()
        m = kit_metrics()
        # Flat data tables use an explicit row checkbox glyph.  Remove the
        # native tree-item indicator so a second checkbox is not rendered at
        # the opposite edge after RTL display columns are applied.
        table_style = ttk.Style(self)
        table_style.layout(
            'Barkav.Data.Treeview.Item',
            [('Treeitem.padding', {
                'sticky': 'nswe',
                'children': [('Treeitem.text', {'sticky': 'nswe'})],
            })],
        )
        head = tk.Frame(self.body, bg=p['surface'],
                        highlightbackground=p['border_subtle'], highlightthickness=1)
        head.pack(fill='x')
        head.grid_columnconfigure(1, weight=1)
        right = tk.Frame(head, bg=p['surface'])
        right.grid(row=0, column=2, sticky='e', padx=m['gap'], pady=m['gap'])
        ttk.Button(right, text='تقویم پروژه', style='Primary.TButton',
                   command=self._open_calendar_from_messages_187,
                   ).pack(side='right')

        center = tk.Frame(head, bg=p['surface'])
        center.grid(row=0, column=1, sticky='ew', padx=12, pady=m['gap'])
        search_host = tk.Frame(center, bg=p['surface'], width=290, height=46,
                               highlightthickness=0)
        search_host.pack(anchor='center')
        search_host.pack_propagate(False)
        def round_search_host():
            try:
                from barkav_ui.skin import rounded_backdrop
                rounded_backdrop(search_host, fill=p['input'], page=p['surface'], radius=18,
                                 border=p['border_subtle'], master=self)
            except (ImportError, tk.TclError):
                try:
                    search_host.configure(highlightbackground=p['border_subtle'], highlightthickness=1)
                except tk.TclError:
                    pass
        self.after_idle(round_search_host)
        ttk.Style(self).configure(
            'Messages.Search.TEntry', fieldbackground=p['input'],
            foreground=p['muted'], borderwidth=0, padding=0,
            lightcolor=p['input'], darkcolor=p['input'], bordercolor=p['input'],
        )
        self.search_entry = ttk.Entry(
            search_host, textvariable=self.search_var, justify='right', width=28,
            style='Messages.Search.TEntry',
        )
        self.search_entry.pack(fill='both', expand=True, padx=16, pady=7)
        self.search_entry.configure(foreground=p['muted'])
        self.search_entry.bind('<FocusIn>', self._activate_search_187, add='+')
        self.search_entry.bind('<FocusOut>', self._restore_search_placeholder_187, add='+')
        self.search_entry.bind('<Return>', lambda _event: self._render_table(), add='+')
        self.search_var.trace_add('write', lambda *_args: self._render_table())

        left = tk.Frame(head, bg=p['surface'])
        left.grid(row=0, column=0, sticky='w', padx=m['gap'], pady=m['gap'])
        export_combo = ttk.Combobox(left, textvariable=self.export_var,
                     values=('انتخاب فرمت خروجی', 'Excel', 'PDF', 'پرینت'),
                     state='readonly', width=18, justify='right',
                     style='Output.TCombobox')
        export_combo.pack(side='left', padx=2)
        export_combo.bind('<<ComboboxSelected>>', lambda _event: self.export_current(), add='+')
        ttk.Button(left, text='خروجی', style='Toolbar.TButton',
                   command=self.export_current).pack(side='left', padx=2)

        filters = tk.Frame(self.body, bg=p['panel2'],
                           highlightbackground=p['border_subtle'], highlightthickness=1)
        filters.pack(fill='x', pady=(m['gap'], 0))
        kind_values = ['همه'] + [label for _kind, label in EVENT_LABELS.items()]
        ttk.Label(filters, text='نوع', style='Caption.TLabel').pack(side='right', padx=(6, 2))
        self.kind_filter_box = ttk.Combobox(
            filters, textvariable=self.kind_filter_var, values=kind_values,
            state='readonly', width=20, justify='right', style='TCombobox',
        )
        self.kind_filter_box.pack(side='right', padx=3, pady=4)
        ttk.Label(filters, text='وضعیت اعلان', style='Caption.TLabel').pack(side='right', padx=(10, 2))
        self.alert_filter_box = ttk.Combobox(
            filters, textvariable=self.alert_filter_var,
            values=('همه', 'روشن', 'خاموش'), state='readonly', width=10,
            justify='right', style='TCombobox',
        )
        self.alert_filter_box.pack(side='right', padx=3, pady=4)
        ttk.Label(filters, text='تاریخ سر رسید', style='Caption.TLabel').pack(side='right', padx=(10, 2))
        self.due_filter_box = ttk.Combobox(
            filters, textvariable=self.due_filter_var,
            values=('همه', 'امروز', '۷ روز آینده', '۳۰ روز آینده'),
            state='readonly', width=14, justify='right', style='TCombobox',
        )
        self.due_filter_box.pack(side='right', padx=3, pady=4)
        for variable in (self.kind_filter_var, self.alert_filter_var, self.due_filter_var):
            variable.trace_add('write', lambda *_args: self._render_table())

        self.summary_var = tk.StringVar(value='')
        tree_host = tk.Frame(self.body, bg=p['panel'])
        tree_host.pack(fill='both', expand=True, pady=(kit_metrics()['gap'], 0))
        tree_host.grid_rowconfigure(0, weight=1)
        tree_host.grid_columnconfigure(0, weight=1)
        message_columns = ('row', 'registered', 'kind', 'title', 'amount', 'due', 'phase', 'document', 'alert')
        self.tree = ttk.Treeview(
            tree_host,
            columns=message_columns, displaycolumns=message_columns,
            selectmode='extended',
            show='headings', style='Barkav.Data.Treeview',
        )
        # This table owns its checkbox glyph in the explicit ``row`` column;
        # the inherited 18.5 visual pass must not add a second native #0 box.
        self.tree._barkav_kit_table_187 = True
        headings = [
            ('row', 'ردیف', 68), ('registered', 'تاریخ ثبت', 125), ('kind', 'نوع', 180),
            ('title', 'شرح', 300), ('amount', 'مبلغ (تومان)', 150),
            ('due', 'تاریخ سر رسید', 130), ('phase', 'مرحله ساخت', 150),
            ('document', 'نمایش سند', 110), ('alert', 'وضعیت اعلان', 125),
        ]
        for key, label, width in headings:
            heading_options = {'text': label, 'anchor': 'e'}
            if key == 'row':
                heading_options['command'] = self._toggle_all_visible_187
            self.tree.heading(key, **heading_options)
            self.tree.column(key, width=width, minwidth=52, anchor='e', stretch=True)
        apply_tree_rtl(self.tree)
        # The legacy tree column can retain its native indicator after the RTL
        # display order is rebuilt. The row column already owns the checkbox.
        self.tree.column('#0', width=0, minwidth=0, stretch=False)
        stripe_tree(self.tree)
        self.tree.grid(row=0, column=0, sticky='nsew')
        vbar = ttk.Scrollbar(tree_host, orient='vertical', command=self.tree.yview)
        hbar = ttk.Scrollbar(tree_host, orient='horizontal', command=self.tree.xview)
        vbar.grid(row=0, column=1, sticky='ns')
        hbar.grid(row=1, column=0, sticky='ew')
        self.tree.configure(yscrollcommand=vbar.set, xscrollcommand=hbar.set)
        self.tree.bind('<Button-1>', self._on_tree_click, add='+')
        self.tree.bind('<Double-1>', self._on_tree_double_click, add='+')
        self._column_keys = message_columns
        self._render_table()
        # Reload once the window is mapped so late-loaded shared documents appear.
        self.after_idle(self._reload_events_187)

    def _open_calendar_from_messages_187(self):
        """Open the dashboard-owned calendar when a host callback is provided."""
        if self._on_open_calendar_187 is not None:
            return self._on_open_calendar_187()
        # Keep standalone/legacy callers working when no host callback exists.
        return FinancialCalendarDialog187(self._calendar_parent_187, self.service)

    def _build_document_file_index_187(self) -> dict[str, str]:
        """Index migrated attachments by name when their old absolute path changed.

        The transferred SQLite keeps the original machine path in
        ``attachments.local_path``.  The files themselves live in this
        project's receipt archive, so resolving by basename keeps the source
        database read-only while making previews portable on D:.
        """
        roots: list[Path] = []
        db_path = str(getattr(getattr(self.service, 'db', None), 'db_path', '') or '')
        if db_path:
            project_root = Path(db_path).parent
            roots.extend((project_root, project_root / 'receipt_archive'))
        release_root = Path(__file__).resolve().parent.parent
        roots.append(release_root / 'data')
        index: dict[str, str] = {}
        for root in roots:
            if not root.is_dir():
                continue
            try:
                candidates = root.rglob('*')
            except OSError:
                continue
            for candidate in candidates:
                if candidate.is_file():
                    index.setdefault(candidate.name.casefold(), str(candidate))
        return index

    def _resolve_document_path_187(self, raw_path: str) -> str:
        path = Path(str(raw_path or '').strip())
        if path.is_file():
            return str(path)
        return self._document_file_index_187.get(path.name.casefold(), '')

    def _document_image_paths_187(self) -> dict[int, list[str]]:
        ids = {event.source_id for event in self.events if event.source_id}
        if not ids:
            return {}
        by_id: dict[str, list[str]] = {}
        # Non-financial documents keep their attachment path on the shared
        # project-document row rather than in the financial attachments table.
        provider = getattr(self.document_service, 'list_documents', None)
        if callable(provider):
            try:
                for row in provider() or ():
                    document_id = str(row.get('document_id') or '')
                    path = self._resolve_document_path_187(str(row.get('local_path') or ''))
                    if document_id and path and Path(path).suffix.lower() in {
                        '.png', '.jpg', '.jpeg', '.bmp', '.gif', '.webp', '.tif', '.tiff',
                    }:
                        by_id.setdefault(document_id, []).append(path)
            except Exception:
                logging.getLogger(__name__).exception('Project document paths could not be indexed')
        try:
            placeholders = ','.join('?' for _ in ids)
            with self.service.db.read() as con:
                rows = con.execute(
                    f"SELECT entity_id, local_path, mime_type FROM attachments "
                    f"WHERE entity_type='DOCUMENT' AND entity_id IN ({placeholders})", tuple(ids),
                ).fetchall()
        except Exception:
            return {index: by_id.get(event.source_id, []) for index, event in enumerate(self.events) if event.source_id}
        image_suffixes = {'.png', '.jpg', '.jpeg', '.bmp', '.gif', '.webp', '.tif', '.tiff'}
        for row in rows:
            path = self._resolve_document_path_187(str(row['local_path'] or ''))
            mime = str(row['mime_type'] or '').lower()
            if path and (mime.startswith('image/') or Path(path).suffix.lower() in image_suffixes):
                by_id.setdefault(str(row['entity_id'] or ''), []).append(path)
        return {index: by_id.get(event.source_id, []) for index, event in enumerate(self.events) if event.source_id}

    def _show_document_187(self, index: int):
        paths = self.document_images.get(index, [])
        if not paths:
            return
        path = Path(paths[0])
        if not path.is_file():
            messagebox.showinfo('نمایش سند', 'فایل تصویر سند در دسترس نیست.', parent=self); return
        try:
            A6DocumentPreview187(self, path)
        except Exception:
            messagebox.showinfo('نمایش سند', 'تصویر سند قابل نمایش نیست.', parent=self)
    def _reload_events_187(self) -> None:
        """Reload project events after the dialog is mapped.

        This retries the shared document source once the window is visible,
        covering providers that finish loading after the initial constructor.
        """
        self.events = message_events_187(
            self.service, 30, document_service=self.document_service,
        )
        self.alert_preferences = _row_alert_preferences_187()
        self._document_file_index_187 = self._build_document_file_index_187()
        self.document_images = self._document_image_paths_187()
        self._render_table()

    def _summary_text(self) -> str:
        active_alerts = sum(1 for event in self.events if self._alert_state(event))
        return f"۳۰ روز آینده: {fa(len(self.events))} پیام — اعلان روشن: {fa(active_alerts)}"

    def _alert_state(self, event: FinancialEvent187) -> bool:
        return _event_alert_enabled_187(event, self.alert_preferences)

    def _kind_filter(self) -> str | None:
        selected = self.kind_filter_var.get()
        if selected in ('', 'همه'):
            return None
        for kind, label in EVENT_LABELS.items():
            if label == selected:
                return kind
        return selected

    def _visible_event_indices(self) -> list[int]:
        kind_filter = self._kind_filter()
        alert_filter = self.alert_filter_var.get()
        due_filter = self.due_filter_var.get()
        query = self._search_query_187()
        today = date.today()
        visible = []
        for index, event in enumerate(self.events):
            if kind_filter and event.kind != kind_filter:
                continue
            alert = self._alert_state(event)
            if alert_filter == 'روشن' and not alert:
                continue
            if alert_filter == 'خاموش' and alert:
                continue
            try:
                delta = (date.fromisoformat(event.due_date[:10]) - today).days
            except (TypeError, ValueError):
                delta = 99999
            if due_filter == 'امروز' and delta != 0:
                continue
            if due_filter == '۷ روز آینده' and not 0 <= delta <= 7:
                continue
            if due_filter == '۳۰ روز آینده' and not 0 <= delta <= 30:
                continue
            if query:
                values = (
                    event.registered_date_shamsi, event.due_date_shamsi,
                    EVENT_LABELS.get(event.kind, event.kind), event.title, event.details,
                    _format_amount(event.amount), _display_due_date_187(event), event.phase,
                    'روشن' if alert else 'خاموش', event.source_id,
                )
                haystack = ' '.join(str(value or '') for value in values)
                if query not in haystack and en(query) not in en(haystack):
                    continue
            visible.append(index)
        return visible

    def _search_query_187(self) -> str:
        """Return the normalized query, excluding the visible placeholder."""
        query = ' '.join(str(self.search_var.get() or '').split())
        if query.casefold() == self.search_placeholder187.casefold():
            return ''
        return query

    def _row_values(self, index: int, row_number: int) -> list[str]:
        event = self.events[index]
        return [
            f"{'☑' if index in self.selected_indices else '☐'} {fa(row_number)}",
            fa(event.registered_date_shamsi or event.due_date_shamsi),
            EVENT_LABELS.get(event.kind, event.kind), event.title or '-',
            _format_amount(event.amount), _display_due_date_187(event),
            event.phase or '—',
            ('نمایش' if self.document_images.get(index) else '—'),
            ('روشن' if self._alert_state(event) else 'خاموش') + ' ▾',
        ]
    def _render_table(self):
        tree = getattr(self, 'tree', None)
        if tree is None or not tree.winfo_exists():
            return
        for item in tree.get_children():
            tree.delete(item)
        self.visible_indices = self._visible_event_indices()
        for row_number, index in enumerate(self.visible_indices, 1):
            event = self.events[index]
            stripe = 'even' if row_number % 2 else 'odd'
            tree.insert('', 'end', iid=str(index), tags=(stripe, event.kind),
                        values=self._row_values(index, row_number))
        checked_rows = [str(index) for index in self.visible_indices if index in self.selected_indices]
        if checked_rows:
            tree.selection_set(checked_rows)
            tree.focus(checked_rows[-1])
        all_checked = bool(self.visible_indices) and all(
            index in self.selected_indices for index in self.visible_indices
        )
        tree.heading('row', text=f"ردیف {'☑' if all_checked else '☐'}",
                     anchor='e', command=self._toggle_all_visible_187)
        self._auto_size_columns()
        if hasattr(self, 'summary_var'):
            active_alerts = sum(1 for index in self.visible_indices if self._alert_state(self.events[index]))
            self.summary_var.set(
                f"{fa(len(self.visible_indices))} پیام مطابق فیلتر — اعلان روشن: {fa(active_alerts)}"
            )

    def _auto_size_columns(self):
        try:
            measure = tkfont.Font(root=self, font=kit_font('body')).measure
        except (tk.TclError, ValueError):
            measure = lambda value: len(str(value)) * 9
        bounds = {
            'row': (60, 70), 'registered': (104, 140), 'kind': (125, 190),
            'title': (160, 560), 'amount': (115, 180), 'due': (105, 140),
            'phase': (110, 190), 'document': (84, 115), 'alert': (98, 135),
        }
        for position, key in enumerate(self._column_keys):
            samples = [self.tree.heading(key, 'text')]
            for item in self.tree.get_children():
                values = self.tree.item(item, 'values')
                if position < len(values):
                    samples.append(values[position])
            low, high = bounds[key]
            width = max(low, min(high, max(measure(value) for value in samples) + 28))
            self.tree.column(key, width=width)

    def _display_key(self, event) -> str | None:
        row = self.tree.identify_row(event.y)
        column = self.tree.identify_column(event.x)
        if not row or not column.startswith('#'):
            return None
        try:
            position = int(column[1:]) - 1
            return tuple(self.tree.cget('displaycolumns'))[position]
        except (IndexError, TypeError, ValueError, tk.TclError):
            return None

    def _on_tree_click(self, event):
        row = self.tree.identify_row(event.y)
        if not row:
            return
        try:
            index = int(row)
        except ValueError:
            return
        key = self._display_key(event)
        self.tree.selection_set(row)
        self.tree.focus(row)
        if key == 'alert':
            self._open_alert_menu_187(index, event)
            return 'break'
        if key == 'document':
            self._show_document_187(index)
            return 'break'
        if key != 'row':
            return 'break'
        if index in self.selected_indices:
            self.selected_indices.remove(index)
        else:
            self.selected_indices.add(index)
        self._render_table()
        self.tree.selection_set(row)
        self.tree.focus(row)
        return 'break'

    def _toggle_all_visible_187(self) -> None:
        visible = set(self.visible_indices)
        if not visible:
            return
        if visible.issubset(self.selected_indices):
            self.selected_indices.difference_update(visible)
        else:
            self.selected_indices.update(visible)
        self._render_table()

    def _on_tree_double_click(self, event):
        row = self.tree.identify_row(event.y)
        if row and self._display_key(event) == 'alert':
            try:
                self._toggle_alert(int(row))
            except ValueError:
                pass

    def _open_alert_menu_187(self, index: int, event) -> None:
        if not 0 <= index < len(self.events):
            return
        menu = tk.Menu(self, tearoff=0)
        for enabled, label in ((False, 'خاموش'), (True, 'روشن')):
            menu.add_command(label=label, command=lambda value=enabled: self._set_alert_187(index, value))
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    def _toggle_alert(self, index: int):
        if 0 <= index < len(self.events):
            self._set_alert_187(index, not self._alert_state(self.events[index]))

    def _set_alert_187(self, index: int, enabled: bool) -> None:
        if not 0 <= index < len(self.events):
            return
        current = self.events[index]
        self.alert_preferences[_event_alert_key_187(current)] = bool(enabled)
        _set_event_alert_187(current, bool(enabled))
        self._render_table()

    def _activate_search_187(self, _event=None):
        if self.search_var.get() == self.search_placeholder187:
            self.search_var.set('')
            self.search_entry.configure(foreground=palette_187()['text'])

    def _restore_search_placeholder_187(self, _event=None):
        if not str(self.search_var.get() or '').strip():
            self.search_var.set(self.search_placeholder187)
            self.search_entry.configure(foreground=palette_187()['muted'])

    def _export_rows(self) -> list[dict[str, str]]:
        headers = ('تاریخ ثبت', 'نوع', 'شرح', 'مبلغ (تومان)', 'تاریخ سر رسید', 'مرحله ساخت', 'نمایش سند', 'وضعیت اعلان')
        rows = []
        export_indices = message_export_indices_187(self.visible_indices, self.selected_indices)
        for row_number, index in enumerate(export_indices, 1):
            values = self._row_values(index, row_number)[1:]
            values[headers.index('نمایش سند')] = '\n'.join(self.document_images.get(index, []))
            rows.append(dict(zip(headers, values)))
        return rows
    def _save_export(self, choice: str, headers, rows, *, orientation='افقی'):
        export_rows = []
        for row in rows:
            if isinstance(row, dict):
                export_rows.append([row.get(header, '') for header in headers])
                continue
            values = list(row or ())
            export_rows.append(values[:len(headers)] + [''] * max(0, len(headers) - len(values)))
        is_print = choice == 'پرینت'
        if is_print:
            handle, destination = tempfile.mkstemp(suffix='.pdf')
            import os
            os.close(handle)
        else:
            extension = '.xlsx' if choice == 'Excel' else '.pdf'
            destination = filedialog.asksaveasfilename(
                parent=self, title=f'خروجی {choice} پیام‌ها', defaultextension=extension,
                filetypes=[('Excel', '*.xlsx')] if choice == 'Excel' else [('PDF', '*.pdf')],
            )
            if not destination:
                return False
        try:
            if choice == 'Excel':
                document_column = next((index for index, header in enumerate(headers)
                                        if header == 'نمایش سند'), None)
                if document_column is None:
                    export_xlsx(destination, headers, export_rows, title='پیام‌ها')
                else:
                    from openpyxl import Workbook
                    from openpyxl.drawing.image import Image as ExcelImage
                    book = Workbook()
                    sheet = book.active
                    sheet.title = 'پیام‌ها'
                    sheet.append(list(headers))
                    for row in export_rows:
                        sheet.append(list(row))
                        document_path = str(row[document_column] or '').splitlines()[0].strip()
                        if document_path and Path(document_path).is_file() and Path(document_path).suffix.lower() in {
                            '.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif',
                        }:
                            try:
                                image = ExcelImage(document_path)
                                image.width, image.height = 160, 110
                                image.anchor = f'{chr(65 + document_column)}{sheet.max_row}'
                                sheet.add_image(image)
                            except (OSError, ImportError, ValueError):
                                pass
                    style_xlsx_workbook(book)
                    book.save(destination)
            else:
                document_paths = [
                    path for row in export_rows
                    for path in str(row[headers.index('نمایش سند')]).splitlines()
                    if path.strip()
                ] if 'نمایش سند' in headers else []
                export_receipts_pdf_187(
                    destination, headers, export_rows, 'پیام‌ها',
                    [self.summary_var.get()], document_paths,
                    orientation=orientation if orientation in {'افقی', 'عمودی'} else 'افقی',
                )
                if is_print:
                    print_file(destination)
        except Exception as exc:
            messagebox.showerror('خروجی پیام‌ها', 'خروجی انجام نشد: ' + str(exc), parent=self)
            return False
        if not is_print:
            messagebox.showinfo('خروجی پیام‌ها', 'فایل آماده شد: ' + str(destination), parent=self)
        return True

    def export_current(self):
        choice = str(self.export_var.get() or '')
        if choice == 'انتخاب فرمت خروجی':
            messagebox.showinfo('خروجی پیام‌ها', 'ابتدا فرمت خروجی را انتخاب کنید.', parent=self)
            return
        self.export_var.set('انتخاب فرمت خروجی')
        headers = ['تاریخ ثبت', 'نوع', 'شرح', 'مبلغ (تومان)', 'تاریخ سر رسید', 'مرحله ساخت', 'نمایش سند', 'وضعیت اعلان']
        rows = self._export_rows()
        if not rows:
            messagebox.showinfo('خروجی پیام‌ها', 'برای خروجی داده‌ای وجود ندارد.', parent=self)
            return
        return show_export_preview_187(
            self,
            export_format=choice,
            headers=headers,
            rows=rows,
            specs=[('خلاصه', 'وضعیت', self.summary_var.get())],
            document_paths=[path for row in rows for path in str(row.get('نمایش سند') or '').splitlines() if path],
            title='پیش‌نمایش پیام‌ها',
            table_caption='پیام‌ها',
            row_label='ردیف',
            on_confirm=lambda selected, orientation: self._save_export(
                selected, headers, rows, orientation=orientation,
            ),
        )



class DailyReportSummaryDialog187(BarkavWindow):
    """Report-page view for a selected daily/weekly/monthly date range."""

    size_class = 'full'

    def __init__(self, parent, *, owner, label, start, end, rows):
        super().__init__(
            parent, f'گزارش {label} پروژه',
            subtitle='BARKAV 18.7.1', size_class='full', show_footer=False,
        )
        self.owner = owner
        self.label = str(label)
        self.start, self.end = str(start), str(end)
        self.rows = [dict(row) for row in rows]
        self._attach_daily_images_187()
        self.checked: set[str] = set()
        self._build()
        self._refresh_tree()

    def _attach_daily_images_187(self):
        used = set()
        for row in self.rows:
            identifiers = row.get('_source_ids') or [row.get('_source_id')]
            signatures, unique = [], []
            for identifier in identifiers:
                if not identifier:
                    continue
                for path in _daily_report_image_paths_187(self.owner, str(identifier)):
                    if path in used:
                        continue
                    signature = _image_signature_187(path)
                    if any(_visually_duplicate_187(signature, existing) for existing in signatures):
                        continue
                    signatures.append(signature); unique.append(path); used.add(path)
                    if len(unique) >= 4:
                        break
                if len(unique) >= 4:
                    break
            row['_document_paths'] = unique

    def _build(self):
        p = palette_187()
        toolbar = tk.Frame(
            self.body, bg=p['surface'],
            highlightbackground=p['border_subtle'], highlightthickness=1,
        )
        toolbar.pack(fill='x', padx=6, pady=(0, 6))
        actions = tk.Frame(toolbar, bg=p['surface'])
        actions.pack(side='right', padx=6, pady=5)
        for text, command in (
            ('افزودن', self._add_187),
            ('ویرایش', self._edit_187),
            ('حذف', self._delete_187),
            ('کپی', self._copy_187),
        ):
            ttk.Button(
                actions, text=text, style='Toolbar.TButton', command=command,
            ).pack(side='right', padx=2)

        output = tk.Frame(toolbar, bg=p['surface'])
        output.pack(side='left', padx=6, pady=5)
        self.export_choice = tk.StringVar(value='انتخاب فرمت خروجی')
        self.export_box = ttk.Combobox(
            output, textvariable=self.export_choice,
            values=('انتخاب فرمت خروجی', 'Excel', 'PDF', 'پرینت'),
            state='readonly', justify='right', width=18,
            style='Output.TCombobox', font=kit_font('body'),
        )
        self.export_box.pack(side='left', padx=2)
        self.export_box.bind(
            '<<ComboboxSelected>>', lambda _event: self.export_current(),
            add='+',
        )
        ttk.Button(
            output, text='خروجی', style='Toolbar.TButton',
            command=self.export_current,
        ).pack(side='left', padx=2)

        summary = tk.Label(
            self.body,
            text=f'بازه: {fa(self.start)} تا {fa(self.end)} — گزارش {self.label}',
            bg=p['panel2'], fg=p['text'], anchor='e',
            font=kit_font('body', bold=True),
        )
        summary.pack(fill='x', padx=6, pady=(0, 6))
        progress = _project_report_overall_progress_187(getattr(self.owner, 'service', None))
        progress_host = ttk.Frame(self.body)
        progress_host.pack(fill='x', padx=6, pady=(0, 6))
        ttk.Label(progress_host, text=f'پیشرفت کلی پروژه: {fa(f"{progress:.1f}")}٪',
                  anchor='e').pack(side='right', padx=6)
        self.project_progress187 = ttk.Progressbar(progress_host, orient='horizontal',
                                                   mode='determinate', maximum=100, value=progress)
        self.project_progress187.pack(side='right', fill='x', expand=True)

        ai_controls = tk.Frame(self.body, bg=p['surface'])
        ai_controls.pack(fill='x', padx=6, pady=(0, 6))
        self.ai_detail187 = tk.StringVar(value='کامل' if self.label == 'روزانه' else 'خلاصه')
        self.ai_audience187 = tk.StringVar(value='کارفرما')
        ui = load_ui()
        current_engine = section_choice('calendar', ui)['mode']
        engine_label = AI_REPORT_ENGINE_LABELS_187.get(current_engine, AI_REPORT_ENGINE_LABELS_187['AUTO'])
        self.ai_engine187 = tk.StringVar(value=engine_label)
        ttk.Label(ai_controls, text='نوع گزارش').pack(side='right', padx=(8, 3))
        ttk.Combobox(ai_controls, textvariable=self.ai_detail187, values=('خلاصه', 'کامل'),
                     state='readonly', width=9, justify='right').pack(side='right', padx=3)
        ttk.Label(ai_controls, text='مخاطب').pack(side='right', padx=(8, 3))
        ttk.Combobox(ai_controls, textvariable=self.ai_audience187,
                     values=('کارفرما', 'مجری', 'ناظرها'), state='readonly', width=12,
                     justify='right').pack(side='right', padx=3)
        ttk.Label(ai_controls, text='مدل هوش مصنوعی').pack(side='right', padx=(8, 3))
        section = getattr(self.owner, '_ai_section187', 'calendar')
        self.ai_models187 = AIModelPicker187(ai_controls, section=section,
            base_dir=Path(__file__).resolve().parent,
            on_change=lambda choice: self.ai_engine187.set(AI_REPORT_ENGINE_LABELS_187.get(choice['mode'],
                AI_REPORT_ENGINE_LABELS_187.get((choice['models'] or ['AUTO'])[0]))))
        self.ai_models187.pack(side='right', padx=3)
        self.ai_report_button187 = ttk.Button(ai_controls, text='تهیه گزارش با AI', style='Toolbar.TButton',
                   command=self._generate_ai_report187)
        self.ai_report_button187.pack(side='left', padx=3)
        self.ai_status187 = tk.StringVar(value='')
        ttk.Label(ai_controls, textvariable=self.ai_status187).pack(side='left', padx=3)

        table_area = ttk.Frame(self.body)
        table_area.pack(fill='both', expand=True, padx=6, pady=(0, 6))
        self.report_headers187, _ = report_table([], periodic=self.label != 'روزانه', documents=False)
        self.tree = ttk.Treeview(table_area, columns=tuple(self.report_headers187),
                                 show='headings', selectmode='extended')
        for key in self.report_headers187:
            width = 330 if key == 'شرح گزارش' else 220 if key == 'تاریخ' else 90 if key in {'ردیف', 'تعداد روز'} else 160
            self.tree.heading(key, text=key, anchor='e')
            self.tree.column(key, width=width, minwidth=80, anchor='e', stretch=True)
        apply_tree_rtl(self.tree)
        stripe_tree(self.tree)
        yscroll = ttk.Scrollbar(table_area, orient='vertical', command=self.tree.yview)
        xscroll = ttk.Scrollbar(table_area, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        yscroll.pack(side='right', fill='y')
        xscroll.pack(side='bottom', fill='x')
        self.tree.pack(side='left', fill='both', expand=True)
        self.tree.bind('<Button-1>', self._row_click_187, add='+')
        self.tree.bind('<Double-1>', lambda _event: self._edit_187(), add='+')

    def _refresh_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        _, display_rows = report_table(self.rows, periodic=self.label != 'روزانه', documents=False)
        for index, row in enumerate(display_rows):
            iid = str(index)
            checked = '☑' if iid in self.checked else '☐'
            values = [f'{checked} {row[key]}' if key == 'ردیف' else row.get(key, '')
                      for key in self.report_headers187]
            self.tree.insert(
                '', 'end', iid=iid, values=values,
                tags=('even' if index % 2 == 0 else 'odd',),
            )

    def _row_click_187(self, event):
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        display = tuple(self.tree.cget('displaycolumns'))
        column = self.tree.identify_column(event.x)
        if column and display[int(column[1:]) - 1] == 'ردیف':
            if iid in self.checked:
                self.checked.remove(iid)
            else:
                self.checked.add(iid)
            self._refresh_tree()
            self.tree.selection_set(iid)
            self.tree.focus(iid)
            return 'break'
        self.tree.selection_set(iid)
        self.tree.focus(iid)

    def _selected_rows_187(self):
        checked = [int(iid) for iid in self.checked if iid.isdigit() and int(iid) < len(self.rows)]
        indexes = sorted(set(checked))
        if not indexes:
            indexes = [
                int(iid) for iid in self.tree.selection()
                if iid.isdigit() and int(iid) < len(self.rows)
            ]
        return [self.rows[index] for index in sorted(set(indexes))]

    def _dashboard_187(self):
        return getattr(self.owner, '_parent', None)

    def _add_187(self):
        handler = getattr(self._dashboard_187(), '_open_document_wizard_187', None)
        if callable(handler):
            self.destroy()
            return handler()
        messagebox.showwarning('گزارش روزانه', 'ویرایشگر ثبت دیتا در دسترس نیست.', parent=self)

    def _source_row_187(self, row):
        if len(row.get('_source_ids') or []) > 1:
            return None
        source_id = str(row.get('_source_id') or '').strip()
        service = getattr(self.owner, 'service', None)
        db = getattr(service, 'db', None)
        if not source_id or db is None:
            return None
        try:
            with db.read() as con:
                found = con.execute(
                    'SELECT document_id,pending_id,document_type,event_date_shamsi,'
                    'description,amount FROM financial_documents '
                    'WHERE document_id=? AND status=\'ACTIVE\' LIMIT 1',
                    (source_id,),
                ).fetchone()
            return dict(found) if found else None
        except Exception:
            return None

    def _edit_187(self):
        selected = self._selected_rows_187()
        if len(selected) != 1:
            return messagebox.showinfo(
                'ویرایش گزارش', 'یک ردیف گزارش را انتخاب کنید.', parent=self,
            )
        handler = getattr(self._dashboard_187(), '_open_document_edit_wizard_187', None)
        source = self._source_row_187(selected[0])
        if callable(handler) and source:
            self.destroy()
            return handler(source)
        messagebox.showwarning(
            'ویرایش گزارش',
            'این ردیف ممکن است شامل چند ثبت باشد؛ ثبت موردنظر را از بخش ثبت دیتا ویرایش کنید.',
            parent=self,
        )

    def _delete_187(self):
        selected = self._selected_rows_187()
        if not selected:
            return messagebox.showinfo(
                'حذف از گزارش', 'حداقل یک ردیف را انتخاب کنید.', parent=self,
            )
        sources=[self._source_row_187(row) for row in selected]
        dashboard=self._dashboard_187()
        handler=getattr(dashboard,'_delete_lifecycle_rows187',None)
        if not callable(handler) or any(source is None for source in sources):
            return messagebox.showwarning('تعیین تکلیف','مبدا بعضی ردیف‌های گزارش مشخص نیست؛ حذف انجام نشد.',parent=self)
        handler(sources)
        self.owner._reload_events_187()
        self.rows = daily_report_rows_187(
            self.owner.events, date.fromisoformat(_gregorian_from_shamsi_187(self.start)),
            date.fromisoformat(_gregorian_from_shamsi_187(self.end)),
            deduplicate=self.label != 'روزانه',
        )
        self._attach_daily_images_187()
        self.checked.clear(); self._refresh_tree()

    def _copy_187(self):
        selected = self._selected_rows_187()
        if not selected:
            return messagebox.showinfo(
                'کپی', 'حداقل یک ردیف را انتخاب کنید.', parent=self,
            )
        headers, rows = report_table(selected, periodic=self.label != 'روزانه', documents=False)
        text = '\n'.join('\t'.join(str(row.get(header, '')).replace('\n', '؛ ') for header in headers)
                         for row in [dict(zip(headers, headers)), *rows])
        self.clipboard_clear()
        self.clipboard_append(text)

    def _save_export_187(self, choice, headers, rows, *, orientation='افقی', include_documents=True):
        is_print = choice == 'پرینت'
        extension = '.xlsx' if choice == 'Excel' else '.pdf'
        if is_print:
            handle, destination = tempfile.mkstemp(suffix='.pdf')
            import os
            os.close(handle)
        else:
            destination = filedialog.asksaveasfilename(
                parent=self, title=f'خروجی گزارش {self.label}',
                defaultextension=extension,
                filetypes=[('Excel', '*.xlsx')] if choice == 'Excel' else [('PDF', '*.pdf')],
            )
            if not destination:
                return False
        try:
            normalized_rows = [
                [row.get(header, '') for header in headers] if isinstance(row, dict) else list(row)
                for row in rows
            ]
            if not include_documents:
                for row in normalized_rows:
                    for index, header in enumerate(headers):
                        if header in {'نمایش سند', 'سند', 'رسید'} and index < len(row):
                            row[index] = ''
            if choice == 'Excel':
                export_xlsx(destination, headers, normalized_rows, title=f'گزارش {self.label}')
            else:
                service = getattr(self.owner, 'service', None)
                project_name = _project_report_name_187(service)
                audience = _project_report_audience_187(service, self.ai_audience187.get())
                document_paths = [
                    path for row in normalized_rows
                    for index, header in enumerate(headers)
                    if header in {'نمایش سند', 'سند', 'رسید'}
                    for path in str(row[index] or '').splitlines() if path
                ]
                export_receipts_pdf_187(
                    destination, headers, normalized_rows, f'گزارش {self.label}',
                    [f'پروژه: {project_name}',
                     *[f'{group} — {label}: {value}' for group, label, value
                       in _project_report_specs_187(self.label, self.start, self.end)]],
                    list(dict.fromkeys(document_paths)),
                    orientation=orientation if orientation in {'افقی', 'عمودی'} else 'افقی',
                    documents_below_rows=True,
                    report_audience=audience,
                    report_progress=_project_report_overall_progress_187(service),
                    report_project_name=project_name,
                    blank_missing=True,
                )
                if is_print:
                    print_file(destination)
        except Exception as exc:
            messagebox.showerror('خروجی گزارش', 'خروجی انجام نشد: ' + str(exc), parent=self)
            return False
        if not is_print:
            messagebox.showinfo('خروجی گزارش', 'فایل آماده شد: ' + str(destination), parent=self)
        return True

    def export_current(self):
        choice = str(self.export_choice.get() or '')
        if choice == 'انتخاب فرمت خروجی':
            return messagebox.showinfo(
                'خروجی گزارش', 'ابتدا فرمت خروجی را انتخاب کنید.', parent=self,
            )
        if getattr(self, '_semantic_cache_key187', None) != self._semantic_key187():
            return self._run_semantic_report187(self.export_current)
        self.export_choice.set('انتخاب فرمت خروجی')
        headers, rows = report_table(self.rows, periodic=self.label != 'روزانه')
        if not rows:
            return messagebox.showinfo(
                'خروجی گزارش', 'برای خروجی داده‌ای وجود ندارد.', parent=self,
            )
        service = getattr(self.owner, 'service', None)
        audience = _project_report_audience_187(service, self.ai_audience187.get())
        project_name = _project_report_name_187(service)
        preview = show_export_preview_187(
            self,
            export_format=choice,
            headers=headers,
            rows=rows,
            specs=_project_report_specs_187(self.label, self.start, self.end),
            report_audience=audience,
            report_progress=_project_report_overall_progress_187(service),
            report_project_name=project_name,
            documents_below_rows=True,
            on_confirm_with_documents=True,
            title=f'پیش‌نمایش گزارش {self.label}',
            table_caption=f'گزارش {self.label}',
            row_label='ردیف', include_addressee=True,
            on_confirm=lambda selected, orientation, include_documents: self._save_export_187(
                selected, headers, rows, orientation=orientation,
                include_documents=include_documents,
            ),
        )
        if preview is not None and hasattr(preview, 'include_documents'):
            preview.include_documents.set(getattr(self, 'default_include_documents187', True))
            if not preview.include_documents.get() and hasattr(preview, '_render'):
                preview._render()
        return preview

    def _semantic_entries187(self):
        return _daily_report_entries_187(self.owner.events,
            date.fromisoformat(_gregorian_from_shamsi_187(self.start)),
            date.fromisoformat(_gregorian_from_shamsi_187(self.end)))

    def _semantic_key187(self):
        entries = self._semantic_entries187()
        return json.dumps([semantic_records(entries), [row['_source_id'] for row in entries], self.label,
                           self.ai_detail187.get(), self.ai_engine187.get(),
                           self.ai_models187.choice if hasattr(self, 'ai_models187') else None], ensure_ascii=False, sort_keys=True)

    def _run_semantic_report187(self, on_ready):
        if getattr(self, '_ai_running187', False):
            return
        self._ai_running187 = True
        self.ai_status187.set('در حال تحلیل و ادغام فعالیت‌ها...')
        self.ai_report_button187.configure(state='disabled')
        entries = self._semantic_entries187()
        cache_key = self._semantic_key187()
        start = date.fromisoformat(_gregorian_from_shamsi_187(self.start))
        end = date.fromisoformat(_gregorian_from_shamsi_187(self.end))
        daily = self.label == 'روزانه'
        mode = self.ai_detail187.get()
        engine = AI_REPORT_ENGINE_VALUES_187.get(self.ai_engine187.get(), 'AUTO')
        settings = load_ui()
        selection = dict(self.ai_models187.choice) if hasattr(self, 'ai_models187') else None
        section = getattr(self.owner, '_ai_section187', 'calendar')
        def scoped_generator(prompt, **kwargs):
            return generate_project_text(prompt, **kwargs, selection=selection, section=section)


        def job():
            try:
                model, ai_rows = generate_semantic_report(entries, start, end, daily=daily,
                    mode=mode, engine=engine, base_dir=Path(__file__).resolve().parent,
                    ui_settings=settings, generator=scoped_generator)
                error_text = ''
            except Exception as exc:
                model, ai_rows, error_text = '', None, str(exc)

            def finish():
                if not self.winfo_exists():
                    return
                self._ai_running187 = False
                self.ai_status187.set('')
                self.ai_report_button187.configure(state='normal')
                if error_text:
                    messagebox.showerror('گزارش هوشمند', error_text, parent=self)
                    return
                self.rows = ai_rows
                self.checked.clear()
                self._attach_daily_images_187()
                self._refresh_tree()
                self._semantic_cache_key187 = cache_key
                self._semantic_model187 = model
                on_ready()
            try:
                self.after(0, finish)
            except tk.TclError:
                pass
        threading.Thread(target=job, daemon=True).start()

    def _generate_ai_report187(self):
        def show_result():
            service = getattr(self.owner, 'service', None)
            ProjectAIResultDialog187(
                self, service=service, model=getattr(self, '_semantic_model187', ''), text='',
                audience=_project_report_audience_187(service, self.ai_audience187.get()),
                start=self.start, end=self.end, label=self.label, detail=self.ai_detail187.get(),
                report_rows=self.rows, ai_report_rows=self.rows)
        if getattr(self, '_semantic_cache_key187', None) == self._semantic_key187():
            return show_result()
        return self._run_semantic_report187(show_result)


def _project_report_overall_progress_187(service) -> float:
    """Return the same estimate-weighted overall percentage shown on the dashboard."""
    try:
        with service.db.read() as con:
            items = con.execute(
                "SELECT pi.phase_no, pi.qty_estimate, pi.factor, pi.local_estimate_rate, "
                "p.estimate_rate FROM rich_phase_items pi LEFT JOIN rich_price_items p "
                "ON p.price_id=pi.price_id WHERE pi.active=1"
            ).fetchall()
            recorded_rows = con.execute(
                'SELECT phase_no, physical_progress FROM phase_progress_v163'
            ).fetchall()
        phase_estimates: dict[int, int] = {}
        for item in items:
            phase_no = int(item['phase_no'] or 0)
            factor = float(item['factor'] or 1)
            quantity = float(item['qty_estimate'] or 0)
            rate = int(item['estimate_rate'] or item['local_estimate_rate'] or 0)
            phase_estimates[phase_no] = phase_estimates.get(phase_no, 0) + int(round(quantity * factor * rate))
        if not phase_estimates:
            return 0.0
        progress_by_phase = {
            int(row['phase_no']): float(row['physical_progress'] or 0)
            for row in recorded_rows
        }
        weights = {phase_no: max(1, estimate) for phase_no, estimate in phase_estimates.items()}
        total_weight = sum(weights.values())
        overall = sum(weight * progress_by_phase.get(phase_no, 0.0)
                      for phase_no, weight in weights.items()) / total_weight
        return max(0.0, min(100.0, overall))
    except Exception:
        logging.getLogger(__name__).exception('Project report progress could not be read')
        return 0.0


class ProjectAIResultDialog187(BarkavWindow):
    """Display the generated Persian answer and open the shared branded preview."""
    size_class = 'large'

    def __init__(self, parent, *, service, model, text, audience, start, end,
                 label='درخواست', detail='خلاصه', report_rows=None,
                 invoice_rows=None, invoice_total=0, unread_invoice_amounts=0,
                 ai_report_rows=None, answer_table=None):
        super().__init__(parent, f'گزارش هوشمند {label}', subtitle=str(model), show_footer=False)
        self.service, self.model, self.report_text = service, str(model), str(text)
        self.answer_table187 = answer_table
        self.audience, self.start, self.end = str(audience), str(start), str(end)
        self.label, self.detail = str(label), str(detail)
        self.report_rows = [dict(row) for row in (report_rows or [])]
        self.ai_report_rows = None if ai_report_rows is None else [dict(row) for row in ai_report_rows]
        self.is_invoice_report = invoice_rows is not None
        self.invoice_rows = [dict(row) for row in (invoice_rows or [])]
        self.invoice_total = int(invoice_total or 0)
        self.unread_invoice_amounts = int(unread_invoice_amounts or 0)
        if self.is_invoice_report:
            self.report_text = (
                f"تعداد فاکتورهای بررسی‌شده: {fa(len(self.invoice_rows))}\n"
                f"جمع مبالغ خوانا: {fa(f'{self.invoice_total:,}')} تومان\n"
                f"مبالغ ناخوانا: {fa(self.unread_invoice_amounts)} مورد"
            )
        actions = ttk.Frame(self.body); actions.pack(fill='x', padx=8, pady=8)
        ttk.Button(actions, text='پیش‌نمایش پیش‌نویس', command=self._preview187).pack(side='left', padx=4)
        headers, rows = self._table_data187()
        self.output = ReportTable187(self.body, headers, rows)
        self.output.pack(fill='both', expand=True, padx=8, pady=(0, 8))

    def _table_data187(self):
        if self.answer_table187 is not None:
            return self.answer_table187
        if self.is_invoice_report:
            headers = ['ردیف', 'تاریخ', 'شرح کالا/مصالح', 'وزن', 'فروشنده', 'مبلغ (تومان)']
            rows = [dict(row) for row in self.invoice_rows]
            rows.append({'ردیف': 'جمع', 'شرح کالا/مصالح': 'جمع مبالغ خوانا',
                         'مبلغ (تومان)': fa(f'{self.invoice_total:,}')})
            if self.unread_invoice_amounts:
                rows.append({'ردیف': '', 'شرح کالا/مصالح': 'تعداد مبالغ خوانده‌نشده',
                             'مبلغ (تومان)': fa(self.unread_invoice_amounts)})
            return present_answer_table(headers, rows)
        source = self.ai_report_rows if self.ai_report_rows is not None else self.report_rows or []
        if source:
            return report_table(source, periodic=self.label != 'روزانه')
        return ['ردیف', 'تاریخ', 'شرح'], [{'ردیف': '۱', 'تاریخ': '', 'شرح': 'در اطلاعات ثبت‌شده یافت نشد'}]

    def _report_specs187(self):
        return _project_report_specs_187(self.label, self.start, self.end)

    def _preview187(self):
        headers, rows = self._table_data187()
        specs = self._report_specs187()
        progress = _project_report_overall_progress_187(self.service)
        def save(selected, orientation, include_documents=True):
            is_print = selected == 'پرینت'
            extension = '.xlsx' if selected == 'Excel' else '.pdf'
            if is_print:
                handle, destination = tempfile.mkstemp(suffix='.pdf')
                import os as _os
                _os.close(handle)
            else:
                destination = filedialog.asksaveasfilename(
                    parent=self, title='خروجی پیش‌نویس گزارش', defaultextension=extension,
                    filetypes=[('Excel', '*.xlsx')] if selected == 'Excel' else [('PDF', '*.pdf')],
                )
                if not destination:
                    return False
            try:
                if selected == 'Excel':
                    normalized = [[row.get(header, '') for header in headers] if isinstance(row, dict) else list(row) for row in rows]
                    if not include_documents:
                        for row in normalized:
                            for index, header in enumerate(headers):
                                if header in {'نمایش سند', 'سند', 'رسید'} and index < len(row):
                                    row[index] = ''
                    export_xlsx(destination, headers, normalized, title=f'گزارش {self.label}')
                else:
                    normalized = [[row.get(header, '') for header in headers] if isinstance(row, dict) else list(row) for row in rows]
                    if not include_documents:
                        for row in normalized:
                            for index, header in enumerate(headers):
                                if header in {'نمایش سند', 'سند', 'رسید'} and index < len(row):
                                    row[index] = ''
                    document_paths = [
                        path for row in normalized
                        for index, header in enumerate(headers)
                        if header in {'نمایش سند', 'سند', 'رسید'}
                        for path in str(row[index] or '').splitlines() if path
                    ]
                    project_name = _project_report_name_187(self.service)
                    export_receipts_pdf_187(
                        destination, headers, normalized, f'گزارش {self.label}',
                        [f'پروژه: {project_name}',
                         *[f'{group} — {label}: {value}' for group, label, value in specs]],
                        list(dict.fromkeys(document_paths)),
                        orientation=orientation if orientation in {'افقی', 'عمودی'} else 'عمودی',
                        documents_below_rows=bool(self.report_rows),
                        report_audience=self.audience, report_progress=progress,
                        report_project_name=project_name, blank_missing=True,
                    )
                    if is_print:
                        print_file(destination)
            except Exception as exc:
                messagebox.showerror('خروجی گزارش', 'خروجی انجام نشد: ' + str(exc), parent=self)
                return False
            if not is_print:
                messagebox.showinfo('خروجی گزارش', 'فایل آماده شد: ' + str(destination), parent=self)
            return True
        return show_export_preview_187(
            self, export_format='PDF', headers=headers, rows=rows, specs=specs,
            report_audience=self.audience, report_progress=progress,
            report_project_name=_project_report_name_187(self.service),
            documents_below_rows=bool(self.report_rows),
            on_confirm_with_documents=bool(self.report_rows),
            title=f'پیش‌نمایش گزارش {self.label}', table_caption=f'گزارش {self.label}',
            row_label='ردیف', include_addressee=True, on_confirm=save,
        )

class ProjectAIRequestDialog187(BarkavWindow):
    """Ask a question across the project's calendar and non-financial records."""
    size_class = 'large'

    def __init__(self, owner):
        super().__init__(owner, 'درخواست از هوش مصنوعی', subtitle='تقویم و اطلاعات پروژه', show_footer=False)
        self.owner = owner
        self._running187 = False
        p = palette_187()
        controls = tk.Frame(self.body, bg=p['surface'])
        controls.pack(fill='x', padx=10, pady=8)
        self.start = tk.StringVar(); self.end = tk.StringVar()
        for label, variable in (('از تاریخ', self.start), ('تا تاریخ', self.end)):
            cell = tk.Frame(controls, bg=p['surface']); cell.pack(side='right', padx=4)
            tk.Label(cell, text=label, bg=p['surface'], fg=p['text'], anchor='e').pack(fill='x')
            ttk.Entry(cell, textvariable=variable, justify='right', width=14).pack(fill='x')
        ui = load_ui()
        selected = section_choice('calendar', ui)['mode']
        self.engine = tk.StringVar(value=AI_REPORT_ENGINE_LABELS_187.get(selected, AI_REPORT_ENGINE_LABELS_187['AUTO']))
        ttk.Label(controls, text='مدل هوش مصنوعی').pack(side='right', padx=(8, 3))
        self.ai_models187 = AIModelPicker187(controls, section='calendar', base_dir=Path(__file__).resolve().parent)
        self.ai_models187.pack(side='right', padx=3)
        self.audience = tk.StringVar(value='کارفرما')
        ttk.Label(controls, text='مخاطب').pack(side='right', padx=(8, 3))
        ttk.Combobox(controls, textvariable=self.audience, values=('کارفرما', 'مجری', 'ناظرها'),
                     state='readonly', width=12, justify='right').pack(side='right', padx=3)
        tk.Label(self.body, text='پرسش یا درخواست گزارش', bg=p['surface'], fg=p['text'],
                 anchor='e', font=kit_font('body', bold=True)).pack(fill='x', padx=12)
        self.question = tk.Text(self.body, height=6, wrap='word', font=kit_font('body'),
            bg=p.get('input', p['surface']), fg=p['text'], insertbackground=p['text'],
            padx=10, pady=8, relief='flat', highlightthickness=1, highlightbackground=p['border_subtle'])
        self.question.pack(fill='x', padx=12, pady=6)
        self.question.tag_configure('rtl', justify='right', spacing1=3, spacing3=5)
        self.question.bind('<<Modified>>', self._question_rtl187, add='+')
        self.status = tk.StringVar(value='')
        tk.Label(self.body, textvariable=self.status, bg=p['surface'], fg=p['muted'], anchor='e').pack(fill='x', padx=12)
        ttk.Button(self.body, text='جست‌وجو و تهیه پاسخ', style='Toolbar.TButton',
                   command=self._ask187).pack(anchor='e', padx=12, pady=8)

    def _question_rtl187(self, _event=None):
        if self.question.edit_modified():
            self.question.tag_add('rtl', '1.0', 'end')
            self.question.edit_modified(False)

    def _ask187(self):
        if self._running187:
            return
        query = self.question.get('1.0', 'end').strip()
        if not query:
            return messagebox.showinfo('درخواست از AI', 'پرسش را وارد کنید.', parent=self)
        start, end = _normalize_shamsi(self.start.get()), _normalize_shamsi(self.end.get())
        if bool(start) != bool(end):
            return messagebox.showwarning('بازه تاریخ', 'برای فیلتر تاریخی هر دو تاریخ را وارد کنید.', parent=self)
        if start and start > end:
            return messagebox.showwarning('بازه تاریخ', 'تاریخ شروع پس از تاریخ پایان است.', parent=self)
        service = self.owner.service
        audience = _project_report_audience_187(service, self.audience.get())
        normalized_query = query.replace('ي', 'ی').replace('ك', 'ک')
        invoice_query = 'فاکت' in normalized_query
        relevant_events = [event for event in self.owner.events
                           if not invoice_query or event.kind == 'INVOICE_NOTE']
        records = _project_report_context_187(relevant_events, start, end)
        records = prepare_answer_records(records)
        if not records:
            return ProjectAIResultDialog187(self, service=service, model='اطلاعات ثبت‌شده', text='',
                audience=audience, start=start or '—', end=end or '—', label='پاسخ',
                answer_table=parse_grounded_answer(
                    '{"columns":["تاریخ","شرح"],"record_ids":[]}', [], count_days=counting_days(query)))
        prompt = build_project_prompt(query=query, events=records, start=start, end=end,
                                      mode='خلاصه', audience=audience, grounded_table=True)
        engine = 'AUTO'
        selection = dict(self.ai_models187.choice)
        settings = load_ui()
        self._running187 = True
        self.status.set('در حال بررسی رکوردهای پروژه…')
        def job():
            answer_table = None
            try:
                model, result = generate_project_text(prompt, engine=engine,
                    base_dir=Path(__file__).resolve().parent, ui_settings=settings, selection=selection, section='calendar')
                answer_table = parse_grounded_answer(result, records, count_days=counting_days(query))
                result = ''
                error_text = ''
            except Exception as exc:
                model, result, error_text = '', '', str(exc)
            def finish():
                self._running187 = False
                self.status.set('')
                if error_text:
                    messagebox.showerror('درخواست از AI', error_text, parent=self)
                    return
                ProjectAIResultDialog187(self, service=service, model=model, text=result,
                    audience=audience, start=start or '—', end=end or '—',
                    label='فاکتورها' if invoice_query else 'پاسخ', detail='خلاصه',
                    answer_table=answer_table)
            try:
                self.after(0, finish)
            except tk.TclError:
                pass
        threading.Thread(target=job, daemon=True).start()


class FinancialCalendarDialog187(BarkavWindow):
    DAYS_FA = ('شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنجشنبه', 'جمعه')
    MONTHS_FA = (
        '', 'فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور',
        'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند',
    )

    size_class = 'full'

    def __init__(self, parent, service):
        super().__init__(parent, 'تقویم پروژه', subtitle='BARKAV 18.7.1', show_footer=False)
        self.service = service
        self.document_service = getattr(parent, 'rich', None)
        self.events = collect_financial_events_187(service, 370, document_service=self.document_service)
        # Keep the lower table in lockstep with the rows shown by the پیام‌ها
        # dialog.  The month grid keeps its longer planning horizon.
        try:
            self.message_events = message_events_187(service, 30)
        except Exception:
            logging.getLogger(__name__).exception('Project calendar message stream could not be loaded')
            self.message_events = list(self.events)
        self.by_day: dict[str, list[FinancialEvent187]] = {}
        for event in self.events:
            self.by_day.setdefault(event.due_date_shamsi, []).append(event)
        icon_style = ttk.Style(self)
        toolbar_options = icon_style.configure('Toolbar.TButton') or {}
        toolbar_options.update(font=kit_font('title'), padding=(5, 0))
        icon_style.configure('CalendarIcon.TButton', **toolbar_options)
        today = jdatetime.date.today()
        self.year, self.month = today.year, today.month
        self.search = tk.StringVar(value='جستجو')
        self._search_placeholder_active_187 = True
        self.event_type_filter = tk.StringVar(value='همه')
        self.from_date_filter = tk.StringVar()
        self.to_date_filter = tk.StringVar()
        self.from_date_filter.trace_add('write', lambda *_: self._refresh_selected_rows_187())
        self.to_date_filter.trace_add('write', lambda *_: self._refresh_selected_rows_187())
        self.year_var = tk.StringVar(value=fa(today.year))
        self.month_var = tk.StringVar(value=self.MONTHS_FA[today.month])
        self.month_title = tk.StringVar()
        self.selected_title = tk.StringVar(value='برای نمایش جزئیات، یک روز را انتخاب کنید.')
        self.summary_text = tk.StringVar(value='گزارش هفتگی یا ماهانه را از پایین انتخاب کنید.')
        self.selected_shamsi: str | None = None
        self._export_rows: list[list[str]] = []
        self._day_buttons: dict[str, tk.Button] = {}
        self._build()
        self.after_idle(self._reload_events_187)
        self.render_month()

    def _reload_events_187(self) -> None:
        """Refresh the calendar stream after the window has been mapped.

        The old constructor scheduled this callback even though the method
        only existed on the messages dialog.  Tk then reported an asynchronous
        AttributeError, leaving the calendar shell partly painted and its
        details table empty.  Keep the refresh local to this dialog and retain
        the last good stream if a provider is temporarily unavailable.
        """
        try:
            events = collect_financial_events_187(
                self.service, 370, document_service=self.document_service,
            )
        except Exception:
            logging.getLogger(__name__).exception('Project calendar data reload failed')
            try:
                messagebox.showerror(
                    'تقویم پروژه',
                    'داده‌های تقویم بارگذاری نشد. جزئیات خطا در لاگ ثبت شد.',
                    parent=self,
                )
            except tk.TclError:
                pass
            return
        self.events = events
        try:
            self.message_events = message_events_187(self.service, 30)
        except Exception:
            logging.getLogger(__name__).exception('Project calendar message stream refresh failed')
            self.message_events = list(self.events)
        self.by_day = {}
        for event in events:
            self.by_day.setdefault(event.due_date_shamsi, []).append(event)
        self.render_month()
        if self.selected_shamsi:
            try:
                self.select_day(int(self.selected_shamsi.rsplit('/', 1)[-1]))
                return
            except (TypeError, ValueError):
                pass
        # Show useful rows immediately when the current month contains an
        # event; the operator no longer has to guess which day to click.
        current_month_days = [
            int(event.due_date_shamsi.rsplit('/', 1)[-1])
            for event in events
            if event.due_date_shamsi.startswith(f'{self.year:04d}/{self.month:02d}/')
        ]
        if current_month_days:
            self.select_day(min(current_month_days))
            return
        # If the current month is quiet, open on the nearest project event so
        # the report rows are visible immediately instead of presenting an
        # apparently empty details panel.
        for event in sorted(events, key=lambda item: item.due_date_shamsi or ''):
            try:
                year, month, day = (int(part) for part in event.due_date_shamsi.split('/'))
            except (AttributeError, TypeError, ValueError):
                continue
            self.year, self.month = year, month
            self._sync_selectors_187()
            self.render_month()
            self.select_day(day)
            break

    def _build(self):
        p = palette_187()
        ttk.Style(self).layout(
            'Barkav.Data.Treeview.Item',
            [('Treeitem.padding', {
                'sticky': 'nswe',
                'children': [('Treeitem.text', {'sticky': 'nswe'})],
            })],
        )
        nav = tk.Frame(self.body, bg=p['surface'], highlightbackground=p['border_subtle'], highlightthickness=1)
        nav.pack(fill='x')
        left = tk.Frame(nav, bg=p['surface'])
        left.pack(side='left', padx=6, pady=5)
        for label, delta in (('ماه بعد  ←', 1), ('→  ماه قبل', -1)):
            ttk.Button(left, text=label, style='Toolbar.TButton',
                       command=lambda value=delta: self.move_month(value)).pack(side='left', padx=2)

        center = tk.Frame(nav, bg=p['surface'])
        center.pack(side='left', fill='x', expand=True, padx=12, pady=5)
        search_host = tk.Frame(center, bg=p['surface'], width=260, height=36)
        search_host.pack(anchor='center')
        search_host.pack_propagate(False)
        def round_search_host():
            try:
                from barkav_ui.skin import rounded_backdrop
                rounded_backdrop(search_host, fill=p['input'], page=p['surface'], radius=18,
                                 border=p['border_subtle'], master=self)
            except (ImportError, tk.TclError):
                try:
                    search_host.configure(highlightbackground=p['border_subtle'], highlightthickness=1)
                except tk.TclError:
                    pass
        self.after_idle(round_search_host)
        entry = ttk.Entry(search_host, textvariable=self.search, justify='right', width=28)
        entry.pack(fill='both', expand=True, padx=10, pady=4)
        entry.configure(foreground=p['muted'])
        entry.bind('<FocusIn>', self._activate_search_187, add='+')
        entry.bind('<FocusOut>', self._restore_search_placeholder_187, add='+')
        entry.bind('<Return>', lambda _event: self.search_date(), add='+')
        self.search_entry = entry

        right = tk.Frame(nav, bg=p['surface'])
        right.pack(side='right', padx=6, pady=5)
        ttk.Button(right, text='امروز', style='Toolbar.TButton', command=self.go_today).pack(side='right', padx=2)
        self.year_box = ttk.Combobox(
            right, textvariable=self.year_var,
            values=tuple(fa(year) for year in range(1300, 1501)),
            state='readonly', width=7, justify='right', style='TCombobox',
        )
        self.year_box.pack(side='right', padx=2)
        self.year_box.bind('<<ComboboxSelected>>', self._select_year_187, add='+')
        self.month_box = ttk.Combobox(
            right, textvariable=self.month_var,
            values=self.MONTHS_FA[1:], state='readonly', width=10,
            justify='right', style='TCombobox',
        )
        self.month_box.pack(side='right', padx=2)
        self.month_box.bind('<<ComboboxSelected>>', self._select_month_187, add='+')
        tk.Label(right, textvariable=self.month_title, bg=p['surface'], fg=p['text'],
                 font=kit_font('title')).pack(side='right', padx=8)

        self.grid_host = tk.Frame(self.body, bg=p['bg'], padx=8, pady=8)
        self.grid_host.pack(fill='both', expand=True, pady=kit_metrics()['gap'])
        for column in range(7):
            self.grid_host.grid_columnconfigure(column, weight=1, uniform='calendar187')
        for row in range(7):
            self.grid_host.grid_rowconfigure(row, weight=1, uniform='calendar187row', minsize=54)

        details = tk.Frame(self.body, bg=p['panel'], highlightbackground=p['border_subtle'], highlightthickness=1)
        details.pack(fill='x')
        details_head = tk.Frame(details, bg=p['panel2'])
        details_head.pack(fill='x')
        reports = tk.Frame(details_head, bg=p['panel2'])
        reports.pack(side='right', padx=6, pady=4)
        ttk.Button(reports, text='گزارش روزانه', style='Toolbar.TButton',
                   command=lambda: self.show_period_summary('day')).pack(side='right', padx=2)
        ttk.Button(reports, text='گزارش هفتگی', style='Toolbar.TButton',
                   command=lambda: self.show_period_summary('week')).pack(side='right', padx=2)
        ttk.Button(reports, text='گزارش ماهیانه', style='Toolbar.TButton',
                   command=lambda: self.show_period_summary('month')).pack(side='right', padx=2)
        ttk.Button(reports, text='درخواست از AI', style='Toolbar.TButton',
                   command=lambda: ProjectAIRequestDialog187(self)).pack(side='right', padx=2)
        actions = tk.Frame(details_head, bg=p['panel2'])
        actions.pack(side='left', padx=5, pady=4)
        self.export_var = tk.StringVar(value='انتخاب فرمت خروجی')
        export_combo = ttk.Combobox(actions, textvariable=self.export_var,
                     values=('انتخاب فرمت خروجی', 'Excel', 'PDF', 'پرینت'),
                     state='readonly', width=18, justify='right',
                     style='Output.TCombobox')
        export_combo.pack(side='left', padx=2)
        export_combo.bind('<<ComboboxSelected>>', lambda _event: self.export_current(), add='+')


        filters = tk.Frame(details_head, bg=p['panel2'])
        filters.pack(side='right', padx=4, pady=4)
        tk.Label(
            filters, text='نمایش براساس', bg=p['panel2'], fg=p['text'],
            anchor='e', font=kit_font('body', bold=True),
        ).pack(side='right', padx=(4, 7))
        type_box = ttk.Combobox(
            filters, textvariable=self.event_type_filter,
            values=tuple(CALENDAR_EVENT_FILTERS_187), state='readonly',
            width=16, justify='right', style='Output.TCombobox',
            font=kit_font('body'),
        )
        type_box.pack(side='right', padx=3)
        type_box.bind(
            '<<ComboboxSelected>>',
            lambda _event: self._refresh_selected_rows_187(), add='+',
        )
        self._calendar_filter_date_field_187(filters, 'از تاریخ', self.from_date_filter)
        self._calendar_filter_date_field_187(filters, 'تا تاریخ', self.to_date_filter)
        details_table = tk.Frame(details, bg=p['panel'])
        details_table.pack(fill='x', padx=6, pady=6)
        details_table.grid_columnconfigure(0, weight=1)
        details_table.grid_rowconfigure(0, weight=1)
        self.details_tree = ttk.Treeview(
            details_table, columns=('row', 'kind', 'title', 'details', 'phase', 'amount', 'direction'),
            show='headings', height=8, style='Barkav.Data.Treeview',
        )
        self.details_tree._barkav_kit_table_187 = True
        columns = [
            ('row', 'ردیف', 60), ('kind', 'نوع', 170), ('title', 'شرح', 300),
            ('details', 'جزییات', 250), ('phase', 'مرحله ساخت', 150),
            ('amount', 'مبلغ', 140), ('direction', 'جریان', 100),
        ]
        for key, label, width in columns:
            self.details_tree.heading(key, text=label, anchor='e')
            self.details_tree.column(key, width=width, anchor='e')
        apply_tree_rtl(self.details_tree)
        self.details_tree.column('#0', width=0, minwidth=0, stretch=False)
        self.details_tree.grid(row=0, column=0, sticky='nsew')
        details_vbar = ttk.Scrollbar(details_table, orient='vertical', command=self.details_tree.yview)
        details_vbar.grid(row=0, column=1, sticky='ns')
        details_hbar = ttk.Scrollbar(details_table, orient='horizontal', command=self.details_tree.xview)
        details_hbar.grid(row=1, column=0, sticky='ew')
        self.details_tree.configure(
            yscrollcommand=details_vbar.set, xscrollcommand=details_hbar.set,
        )

    @staticmethod
    def _month_days(year: int, month: int) -> int:
        if month <= 6:
            return 31
        if month <= 11:
            return 30
        return 30 if jdatetime.date(year, 1, 1).isleap() else 29

    def _sync_selectors_187(self):
        self.year_var.set(fa(self.year))
        self.month_var.set(self.MONTHS_FA[self.month])

    def _select_year_187(self, _event=None):
        try:
            year = int(en(self.year_var.get()))
            jdatetime.date(year, self.month, 1)
        except (TypeError, ValueError):
            return
        self.year = year
        self.selected_shamsi = None
        self.render_month()

    def _select_month_187(self, _event=None):
        try:
            self.month = self.MONTHS_FA[1:].index(self.month_var.get()) + 1
        except ValueError:
            return
        self.selected_shamsi = None
        self.render_month()

    def render_month(self, selected_day: int | None = None):
        p = palette_187()
        self._day_buttons = {}
        backdrop = getattr(self.grid_host, '_barkav_backdrop', None)
        for child in self.grid_host.winfo_children():
            if child is not backdrop:
                child.destroy()
        if selected_day is not None:
            self.selected_shamsi = f'{self.year:04d}/{self.month:02d}/{int(selected_day):02d}'
        self._sync_selectors_187()
        self.month_title.set(f'{self.MONTHS_FA[self.month]} {fa(self.year)}')
        for index, label in enumerate(self.DAYS_FA):
            column = 6 - index
            tk.Label(self.grid_host, text=label, bg=p['panel2'], fg=p['text'],
                     font=kit_font('small', bold=True), pady=6).grid(
                         row=0, column=column, sticky='nsew', padx=2, pady=2
                     )
        first_g = jdatetime.date(self.year, self.month, 1).togregorian()
        saturday_index = (first_g.weekday() + 2) % 7
        for day in range(1, self._month_days(self.year, self.month) + 1):
            slot = saturday_index + day - 1
            week, day_index = divmod(slot, 7)
            column = 6 - day_index
            sh = f'{self.year:04d}/{self.month:02d}/{day:02d}'
            events = self.by_day.get(sh, [])
            kinds = {event.kind for event in events}
            if 'CHEQUE' in kinds:
                accent = p['red']
            elif 'LOAN' in kinds:
                accent = p['orange']
            elif 'RECEIVABLE' in kinds:
                accent = p['green']
            else:
                accent = p['muted']
            today = jdatetime.date.today()
            is_selected = self.selected_shamsi == sh
            is_today = (today.year, today.month, today.day) == (self.year, self.month, day)
            holiday = _is_official_holiday_187(self.year, self.month, day)
            holiday_accent = mix(p['red'], p['bg'], 0.42)
            if is_selected:
                background = mix(p['panel'], p['selection'], 0.42)
            elif is_today:
                background = mix(p['panel'], p['cyan'], 0.20)
            elif holiday:
                background = mix(p['panel'], holiday_accent, 0.22)
            else:
                background = p['panel']
            if holiday:
                accent = holiday_accent
            # Keep every day-cell equal in size; event presence is conveyed by its colour.
            text = fa(day)

            cell = tk.Frame(
                self.grid_host, bg=p['bg'], width=100, height=54,
            )
            cell.grid(row=week + 1, column=column, sticky='nsew', padx=2, pady=2)
            cell.grid_propagate(False)
            cell_border = (
                holiday_accent if holiday else
                (p['selection'] if is_selected else
                 (p['cyan'] if is_today else (accent if events else p['line'])))
            )
            button = tk.Button(
                cell, text=text, command=lambda value=day: self.select_day(value),
                bg=background,
                fg=accent if events else p['text'], activebackground=p['glass'],
                activeforeground=p['text'], relief='flat', bd=0,
                highlightbackground=cell_border,
                highlightthickness=2 if (holiday or is_selected or is_today or events) else 1,
                font=kit_font('body', bold=True), justify='center', cursor='hand2',
            )
            button.pack(fill='both', expand=True, padx=5, pady=5)
            button._calendar_base_background = background
            button._calendar_base_foreground = accent if events else p['text']
            button._calendar_base_border = cell_border
            button._calendar_base_thickness = 2 if (holiday or is_selected or is_today or events) else 1
            self._day_buttons[sh] = button
        self.after_idle(self._round_calendar_grid_187)

    def _round_calendar_grid_187(self):
        """Paint one cached rounded plate behind the month grid.

        Rebuilding a PIL rounded image for every day after each click caused
        the visible selection lag.  The grid needs one backdrop; day buttons
        already carry their own border and selection colors.
        """
        try:
            if getattr(self.grid_host, '_barkav_backdrop', None):
                return
            from barkav_ui.skin import rounded_backdrop
            rounded_backdrop(
                self.grid_host, fill=palette_187()['surface'], page=palette_187()['bg'],
                radius=22, border=palette_187()['border_subtle'], master=self,
            )
        except (ImportError, tk.TclError):
            return

    def _set_selected_day_visual_187(self, previous: str | None, current: str):
        palette = palette_187()
        old_button = self._day_buttons.get(previous or '')
        if old_button is not None:
            old_button.configure(
                bg=old_button._calendar_base_background,
                fg=old_button._calendar_base_foreground,
                highlightbackground=old_button._calendar_base_border,
                highlightthickness=old_button._calendar_base_thickness,
            )
        new_button = self._day_buttons.get(current)
        if new_button is not None:
            new_button.configure(
                bg=mix(palette['panel'], palette['selection'], 0.42),
                highlightbackground=palette['selection'], highlightthickness=2,
            )

    def select_day(self, day: int):
        sh = f'{self.year:04d}/{self.month:02d}/{day:02d}'
        previous = self.selected_shamsi
        self.selected_shamsi = sh
        if sh in self._day_buttons:
            self._set_selected_day_visual_187(previous, sh)
        else:
            self.render_month()
        rows = self._filtered_rows_187(self.message_events)
        self.selected_title.set(f'پیام‌های پروژه — {fa(len(rows))} ردیف')
        self._export_rows = []
        for item in self.details_tree.get_children():
            self.details_tree.delete(item)
        for index, event in enumerate(rows):
            flow = 'اطلاع‌رسانی' if event.direction == 'INFO' else ('ورودی' if event.direction == 'IN' else 'خروجی')
            values = [fa(index + 1), EVENT_LABELS.get(event.kind, event.kind), event.title,
                      event.details or '-', event.phase or '—', _format_amount(event.amount), flow]
            self._export_rows.append(values)
            self.details_tree.insert('', 'end', iid=str(index), values=(
                *values,
            ))
        self.summary_text.set('گزارش هفتگی یا ماهانه را از پایین انتخاب کنید.')


    def _calendar_filter_date_field_187(self, parent, label, variable):
        p = palette_187()
        tk.Label(
            parent, text=label, bg=p['panel2'], fg=p['text'],
            anchor='e', font=kit_font('body', bold=True),
        ).pack(side='right', padx=(7, 2))
        host = tk.Frame(parent, bg=p['panel2'])
        host.pack(side='right', padx=2)
        entry = ttk.Entry(
            host, textvariable=variable, justify='right', width=12,
            font=kit_font('body'),
        )
        entry.pack(side='right')
        entry.bind(
            '<FocusOut>', lambda _event: self._refresh_selected_rows_187(),
            add='+',
        )
        ttk.Button(
            host, text='▣', width=4, style='CalendarIcon.TButton',
            takefocus=False,
            command=lambda: ShamsiDatePicker187(self, variable),
        ).pack(side='left', padx=(2, 0))

    def _filtered_rows_187(self, rows):
        return filter_calendar_events_187(
            rows, self.event_type_filter.get(), self.from_date_filter.get(), self.to_date_filter.get(),
        )

    def _refresh_selected_rows_187(self):
        if not self.selected_shamsi:
            return
        try:
            self.select_day(int(self.selected_shamsi.rsplit('/', 1)[-1]))
        except (TypeError, ValueError):
            return

    def _activate_search_187(self, _event=None):
        if self._search_placeholder_active_187:
            self.search.set('')
            self.search_entry.configure(foreground=palette_187()['text'])
            self._search_placeholder_active_187 = False

    def _restore_search_placeholder_187(self, _event=None):
        if not self.search.get().strip():
            self.search.set('جستجو')
            self.search_entry.configure(foreground=palette_187()['muted'])
            self._search_placeholder_active_187 = True


    def _show_summary_result_187(self, label, start, end, rows):
        return DailyReportSummaryDialog187(
            self, owner=self, label=label, start=start, end=end, rows=rows,
        )

    def _period_summary_dialog_187(self, period):
        label = {'day': 'روزانه', 'month': 'ماهیانه'}.get(period, 'هفتگی')
        default = self.selected_shamsi or f'{self.year:04d}/{self.month:02d}/01'
        try:
            year, month, day = (int(part) for part in default.split('/'))
            if period == 'day':
                initial_start = initial_end = default
            else:
                initial_start, initial_end = _period_bounds_187(year, month, day, period)
                initial_start = _normalize_shamsi(
                    _shamsi_from_gregorian(initial_start.isoformat()),
                )
                initial_end = _normalize_shamsi(
                    _shamsi_from_gregorian(initial_end.isoformat()),
                )
        except (TypeError, ValueError):
            initial_start = initial_end = default

        dialog = BarkavWindow(
            self, f'بازه گزارش {label}', subtitle='تقویم پروژه',
            size_class='small', show_footer=False,
        )
        dialog.grab_set()
        dialog.geometry('620x250')
        dialog.minsize(560, 230)
        dialog.resizable(True, False)
        content = dialog.body
        tk.Label(
            content, text=f'بازه گزارش {label}',
            bg=palette_187()['surface'], fg=palette_187()['text'],
            font=kit_font('title'), anchor='e',
        ).pack(fill='x', padx=18, pady=(14, 7))

        fields = tk.Frame(content, bg=palette_187()['surface'])
        fields.pack(fill='x', padx=18, pady=6)
        fields.columnconfigure(0, weight=1)
        fields.columnconfigure(1, weight=1)
        fields.grid_anchor('e')
        start_var = tk.StringVar(value=fa(initial_start))
        end_var = tk.StringVar(value=fa(initial_end))
        for column, (caption, variable) in enumerate(
            (('تا تاریخ', end_var), ('از تاریخ', start_var)),
        ):
            field = tk.Frame(fields, bg=palette_187()['surface'])
            field.grid(row=0, column=column, sticky='ew', padx=5)
            tk.Label(
                field, text=caption, bg=palette_187()['surface'],
                fg=palette_187()['text'], anchor='e',
                font=kit_font('body', bold=True),
            ).pack(fill='x')
            line = tk.Frame(field, bg=palette_187()['surface'])
            line.pack(fill='x', pady=(3, 0))
            ttk.Entry(
                line, textvariable=variable, justify='right',
                font=kit_font('body'),
            ).pack(side='right', fill='x', expand=True)
            ttk.Button(
                line, text='▣', width=4, style='CalendarIcon.TButton',
                command=lambda current=variable: ShamsiDatePicker187(dialog, current),
            ).pack(side='left', padx=(3, 0))

        def submit():
            start, end = _normalize_shamsi(start_var.get()), _normalize_shamsi(end_var.get())
            if not start or not end:
                messagebox.showwarning(
                    'گزارش', 'هر دو تاریخ را با تقویم انتخاب کنید.', parent=dialog,
                )
                return
            if start > end:
                messagebox.showwarning(
                    'گزارش', 'تاریخ «از» نمی‌تواند بعد از تاریخ «تا» باشد.',
                    parent=dialog,
                )
                return
            try:
                first = _gregorian_from_shamsi_187(start)
                last = _gregorian_from_shamsi_187(end)
                rows = daily_report_rows_187(
                    self.events,
                    date.fromisoformat(first),
                    date.fromisoformat(last),
                    deduplicate=period != 'day',
                )
            except (TypeError, ValueError):
                messagebox.showwarning(
                    'گزارش', 'بازه تاریخ معتبر نیست.', parent=dialog,
                )
                return
            dialog.destroy()
            self._show_summary_result_187(label, start, end, rows)

        ttk.Button(
            content, text='نمایش گزارش', style='Toolbar.TButton',
            command=submit,
        ).pack(anchor='e', padx=18, pady=14)

    def _rows_for_month_187(self) -> list[list[str]]:
        rows: list[list[str]] = []
        month_prefix = f'{self.year:04d}/{self.month:02d}/'
        for index, event in enumerate(
            item for item in self.events if item.due_date_shamsi.startswith(month_prefix)
        ):
            flow = 'اطلاع‌رسانی' if event.direction == 'INFO' else ('ورودی' if event.direction == 'IN' else 'خروجی')
            rows.append([
                fa(index + 1), EVENT_LABELS.get(event.kind, event.kind), event.title,
                event.details or '-', event.phase or '—', _format_amount(event.amount), flow,
            ])
        return rows

    def _save_export(self, choice: str, headers, rows, *, orientation='افقی'):
        is_print = choice == 'پرینت'
        extension = '.xlsx' if choice == 'Excel' else '.pdf'
        if is_print:
            handle, destination = tempfile.mkstemp(suffix='.pdf')
            import os
            os.close(handle)
        else:
            destination = filedialog.asksaveasfilename(
                parent=self, title=f'خروجی {choice} تقویم پروژه',
                defaultextension=extension,
                filetypes=[('Excel', '*.xlsx')] if choice == 'Excel' else [('PDF', '*.pdf')],
            )
            if not destination:
                return False
        try:
            if choice == 'Excel':
                export_xlsx(destination, headers, rows, title='تقویم پروژه')
            else:
                export_pdf(
                    destination, headers, rows, title='تقویم پروژه',
                    summary_lines=[self.selected_title.get()],
                    orientation=orientation if orientation in {'افقی', 'عمودی'} else 'افقی',
                )
                if is_print:
                    print_file(destination)
        except Exception as exc:
            messagebox.showerror('خروجی تقویم', 'خروجی انجام نشد: ' + str(exc), parent=self)
            return False
        if not is_print:
            messagebox.showinfo('خروجی تقویم', 'فایل آماده شد: ' + str(destination), parent=self)
        return True

    def export_current(self):
        choice = str(self.export_var.get() or '')
        if choice == 'انتخاب فرمت خروجی':
            messagebox.showinfo('خروجی تقویم', 'ابتدا فرمت خروجی را انتخاب کنید.', parent=self)
            return
        self.export_var.set('انتخاب فرمت خروجی')
        headers = ['ردیف', 'نوع', 'شرح', 'جزییات', 'مرحله ساخت', 'مبلغ', 'جریان']
        rows = self._export_rows or self._rows_for_month_187()
        if not rows:
            messagebox.showinfo('خروجی تقویم', 'برای خروجی داده‌ای وجود ندارد.', parent=self)
            return
        return show_export_preview_187(
            self,
            export_format=choice,
            headers=headers,
            rows=rows,
            specs=[('خلاصه', 'بازه', self.selected_title.get())],
            title='پیش‌نمایش تقویم پروژه',
            table_caption='رویدادهای تقویم',
            row_label='ردیف',
            on_confirm=lambda selected, orientation: self._save_export(
                selected, headers, rows, orientation=orientation,
            ),
        )

    def show_period_summary(self, period: str):
        self._period_summary_dialog_187(period)

    def move_month(self, delta: int):
        value = self.year * 12 + self.month - 1 + int(delta)
        self.year, zero_month = divmod(value, 12)
        self.month = zero_month + 1
        self.selected_shamsi = None
        self.render_month()

    def go_today(self):
        today = jdatetime.date.today()
        self.year, self.month = today.year, today.month
        self.selected_shamsi = f'{today.year:04d}/{today.month:02d}/{today.day:02d}'
        self.render_month()
        self.select_day(today.day)

    def search_date(self):
        raw = en(self.search.get()).replace('-', '/').strip()
        parts = [part for part in raw.split('/') if part]
        try:
            if len(parts) not in (2, 3):
                raise ValueError
            year, month = int(parts[0]), int(parts[1])
            day = int(parts[2]) if len(parts) == 3 else None
            if not 1 <= month <= 12 or (day is not None and not 1 <= day <= self._month_days(year, month)):
                raise ValueError
        except (TypeError, ValueError):
            messagebox.showwarning('جست‌وجوی تقویم', 'تاریخ را به شکل ۱۴۰۵/۰۵ یا ۱۴۰۵/۰۵/۲۰ وارد کنید.', parent=self)
            return
        self.year, self.month = year, month
        self.selected_shamsi = f'{year:04d}/{month:02d}/{day:02d}' if day is not None else None
        self.render_month()
        if day is not None:
            self.select_day(day)
